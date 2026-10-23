"""R36: STAWN developer purge of TRIAL / TEST data with every linked record."""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase/migrations/20261022100000_qcms_r36_dev_purge.sql"

PO_ID = "aaaaaaaa-0000-0000-0000-000000000001"
PO_ROWS = [
    {"id": PO_ID, "tenant_id": "00000000-0000-0000-0000-000000000001", "po_number": "PD-TRIAL-001", "order_date": "2026-10-20",
     "supplier_id": None, "status": "DRAFT", "created_by": None, "created_at": "2026-10-20T05:00:00+00:00"},
    {"id": "aaaaaaaa-0000-0000-0000-000000000002", "tenant_id": "00000000-0000-0000-0000-000000000001", "po_number": "PD-TRIAL-002",
     "order_date": "2026-10-21", "supplier_id": None, "status": "APPROVED", "created_by": None, "created_at": "2026-10-21T05:00:00+00:00"},
]
PREVIEW = {
    "root_table": "supply_purchase_orders", "root_count": 1, "total_rows": 4, "tenant_id": "00000000-0000-0000-0000-000000000001",
    "tables": [
        {"table": "supply_purchase_orders", "count": 1, "min_depth": 0, "max_depth": 0, "is_root_table": True,
         "samples": [{"id": PO_ID, "label": "PD-TRIAL-001", "status": "DRAFT", "depth": 0}], "rows": [PO_ROWS[0]]},
        {"table": "supply_purchase_order_items", "count": 2, "min_depth": 1, "max_depth": 1, "is_root_table": False,
         "samples": [{"id": "b1", "label": "001", "depth": 1}, {"id": "b2", "label": "002", "depth": 1}],
         "rows": [{"id": "b1", "purchase_order_id": PO_ID, "meta": {"a": 1}}, {"id": "b2", "purchase_order_id": PO_ID}]},
        {"table": "document_attachments", "count": 1, "min_depth": 1, "max_depth": 1, "is_root_table": False,
         "samples": [{"id": "c1", "label": "po.pdf", "depth": 1}], "rows": [{"id": "c1", "object_path": "po/po.pdf"}]},
    ],
    "nullify": [{"table": "part_supplier_price_history", "column": "source_purchase_order_item_id", "references": "supply_purchase_order_items", "count": 1}],
    "blockers": [], "storage_paths": ["po/po.pdf"], "depth_limit_reached": False,
}
CALLS: list[tuple[str, dict]] = []


def _status(dev: bool) -> dict:
    return {"valid": True, "status": "ACTIVE", "plan": "ENTERPRISE", "is_stawn_developer": dev, "is_company_admin": True}


def _fake_rpc_factory(dev: bool):
    def _fake_rpc(self, name, params=None):
        CALLS.append((name, dict(params or {})))
        if name == "qcms_licence_status":
            return _status(dev)
        if name == "qcms_dev_purge_preview":
            return dict(PREVIEW)
        if name == "qcms_dev_purge_execute":
            return {"run_id": "r1", "deleted_counts": {"supply_purchase_orders": 1, "supply_purchase_order_items": 2, "document_attachments": 1},
                    "total_deleted": 4, "storage_paths": ["po/po.pdf"]}
        if name == "qcms_dev_licence_overview":
            return []
        raise RuntimeError("RPC functions are unavailable in controlled preview mode.")
    return _fake_rpc


def _fake_select(self, table, **kwargs):
    if table == "supply_purchase_orders":
        return [dict(r) for r in PO_ROWS]
    return []


# ----------------------------------------------------------------------------------------------- static SQL checks
def _fn_body(sql: str, name: str) -> str:
    start = sql.index(f"create or replace function public.{name}(")
    return sql[start:sql.index("end $fn$;", start)]


def test_migration_functions_are_developer_only_and_safe():
    sql = MIGRATION.read_text()
    low = sql.lower()
    for token in ("create table if not exists public.qcms_purge_runs", "create table if not exists public.qcms_purge_archive",
                  "alter table public.qcms_purge_runs enable row level security",
                  "alter table public.qcms_purge_archive enable row level security",
                  "revoke insert, update, delete, truncate on public.qcms_purge_runs from authenticated",
                  "revoke insert, update, delete, truncate on public.qcms_purge_archive from authenticated",
                  "grant execute on function public.qcms_dev_purge_preview(text, uuid[], boolean) to authenticated",
                  "grant execute on function public.qcms_dev_purge_execute(text, uuid[], text, text) to authenticated",
                  "revoke all on function public.qcms_dev_purge_execute(text, uuid[], text, text) from public, anon"):
        assert token in sql, token
    assert "drop table" not in low and "drop column" not in low and "drop function" not in low
    # the migration itself never deletes or updates data outside function bodies
    outside = re.sub(r"\$fn\$.*?\$fn\$|\$\$.*?\$\$|\$patch\$.*?\$patch\$|\$pol\$.*?\$pol\$", "", sql, flags=re.S).lower()
    assert not re.search(r"\bdelete\s+from\b", outside) and not re.search(r"\bupdate\s+public\.", outside)
    # developer check in both RPCs (and the shared collector)
    for name in ("qcms_dev_purge_preview", "qcms_dev_purge_execute"):
        body = _fn_body(sql, name)
        assert "if not public.qcms_is_stawn_developer() then" in body, name
        assert "security definer" in body and "set search_path = public, pg_temp" in body, name
    assert sql.count("if not public.qcms_is_stawn_developer() then") >= 4


def test_execute_confirms_archives_before_delete_and_bypass_is_scoped():
    sql = MIGRATION.read_text()
    body = _fn_body(sql, "qcms_dev_purge_execute")
    assert "coalesce(p_confirm, '') <> 'DELETE TRIAL DATA'" in body
    assert "length(btrim(coalesce(p_reason, ''))) < 5" in body
    archive = body.index("insert into public.qcms_purge_archive")
    assert archive < body.index("update public.%I c set %I = null") < body.index("delete from public.%I")
    assert "order by s.depth desc" in body  # leaves first
    assert "perform set_config('qcms.dev_purge', 'on', true)" in body
    assert "insert into public.qcms_purge_runs" in body
    # trigger bypass only inside the purge and only for the developer
    assert "current_setting(''qcms.dev_purge'', true) = ''on'' and public.qcms_is_stawn_developer() then return old" in sql
    assert "tg_op = ''UPDATE'' and current_setting(''qcms.dev_purge'', true)" in sql
    assert "'qcms_guard_approved_delete'" in sql and "pg_get_functiondef" in sql


def test_masters_never_roots_and_service_matches_sql():
    from core import dev_purge_service as purge

    sql = MIGRATION.read_text()
    roots_sql = sql[sql.index("qcms_dev_purge_root_tables()\nreturns"):]
    roots_sql = roots_sql[:roots_sql.index("$$;", roots_sql.index("$$") + 2)]
    sql_roots = set(re.findall(r"'([a-z_]+)'", roots_sql))
    assert sql_roots == set(purge.ROOT_TABLES)
    protected = sql[sql.index("function qcms_private.qcms_dev_purge_protected"):]
    for master in ("parts", "parties", "material_grades", "processes", "employees", "profiles", "tenants", "inspection_plans"):
        assert master not in purge.ROOT_TABLES and master not in sql_roots, master
        assert f"'{master}'" in protected[:2000], master
    assert "p_table like 'part\\_%'" in protected and "p_table like 'qcms\\_%'" in protected
    assert not (purge.ROOT_TABLES & purge.MASTER_TABLES)
    for module, spec in purge.MODULES.items():
        assert not spec[0].startswith(("qcms_", "part_")), module


def test_backup_workbook_has_one_sheet_per_table():
    import openpyxl
    from io import BytesIO
    from core import dev_purge_service as purge

    data = purge.backup_workbook(PREVIEW)
    book = openpyxl.load_workbook(BytesIO(data))
    assert book.sheetnames == ["Summary", "supply_purchase_orders", "supply_purchase_order_items", "document_attachments"]
    assert book["supply_purchase_order_items"].max_row == 3


def test_licence_page_wires_purge_before_rights():
    src = (ROOT / "app_pages/licence_company.py").read_text()
    assert "    from app_pages import dev_purge\n    dev_purge.render_section()\n    section_bar(\"STAWN DEVELOPER RIGHTS\")" in src


# ----------------------------------------------------------------------------------------------- UI
_SCRIPT = '''
import streamlit as st
from core.auth import PREVIEW_PROFILE
st.session_state["_qsms_preview"] = True
st.session_state["profile"] = PREVIEW_PROFILE.copy()
from app_pages import dev_purge
dev_purge.render_section()
'''


def test_non_developer_sees_only_error(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from core.repository import Repository
    monkeypatch.setattr(Repository, "rpc", _fake_rpc_factory(False))
    monkeypatch.setattr(Repository, "select", _fake_select)
    app = AppTest.from_string(_SCRIPT).run(timeout=60)
    assert not app.exception
    assert len(app.error) == 1 and "only available to the STAWN developer" in app.error[0].value
    assert not app.selectbox and not app.button and not app.text_input


def test_developer_preview_and_execute(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import core.auth
    from core.repository import Repository
    CALLS.clear()
    monkeypatch.setattr(Repository, "rpc", _fake_rpc_factory(True))
    monkeypatch.setattr(Repository, "select", _fake_select)
    monkeypatch.setattr(core.auth, "verify_current_password", lambda password: None)
    app = AppTest.from_string(_SCRIPT).run(timeout=60)
    assert not app.exception
    assert not app.error
    app.selectbox(key="qcms_dev_purge_module").set_value("Purchase Orders").run(timeout=60)
    pick = app.multiselect(key="qcms_dev_purge_pick_supply_purchase_orders")
    first = [o for o in pick.options if "PD-TRIAL-001" in o]
    assert first
    pick.set_value(first).run(timeout=60)
    app.button(key="qcms_dev_purge_preview_btn").click().run(timeout=60)
    assert not app.exception
    text = " ".join(str(m.value) for m in app.markdown)
    assert "4 row(s) in 3 table(s) will be deleted" in text
    assert any(name == "qcms_dev_purge_preview" and params.get("p_ids") == [PO_ID] for name, params in CALLS)
    assert app.button(key="qcms_dev_purge_execute_btn").disabled
    app.text_input(key="qcms_dev_purge_confirm").set_value("DELETE TRIAL DATA")
    app.text_area(key="qcms_dev_purge_reason").set_value("Trial PO created during training")
    app.text_input(key="qcms_dev_purge_password").set_value("secret")
    app.checkbox(key="qcms_dev_purge_agree").check()
    app.run(timeout=60)
    assert not app.button(key="qcms_dev_purge_execute_btn").disabled
    app.button(key="qcms_dev_purge_execute_btn").click().run(timeout=60)
    assert not app.exception
    executed = [p for n, p in CALLS if n == "qcms_dev_purge_execute"]
    assert executed and executed[0]["p_confirm"] == "DELETE TRIAL DATA" and executed[0]["p_ids"] == [PO_ID]
    assert any("Last purge removed 4 row(s)" in str(s.value) for s in app.success)
