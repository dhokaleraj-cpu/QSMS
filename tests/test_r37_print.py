"""R37: print-layout PDF (preview / download) for every record, search result and report."""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]

DB: dict[str, list[dict]] = {
    "parts": [{"id": "p1", "part_number": "40256626", "fsi_part_number": "FSI-777", "part_name": "Pinion Gear", "customer_id": "c1"}],
    "parties": [{"id": "c1", "party_code": "CUST01", "party_name": "Acme Tractors"},
                {"id": "s1", "party_code": "SUP01", "party_name": "Bharat Forge"},
                {"id": "v1", "party_code": "OSP01", "party_name": "Heat Treat Works"}],
    "employees": [{"id": "e1", "employee_code": "E100", "first_name": "Rajesh", "last_name": "Dhokale"}],
    "processes": [{"id": "pr1", "process_code": "HT", "process_name": "Heat Treatment"}],
    "material_grades": [{"id": "g1", "grade_code": "20MnCr5", "standard": "EN"}],
    "company_branches": [{"id": "b1", "branch_code": "D9", "branch_name": "Four Star Industries D9", "address_line1": "Plot D9, Chakan MIDC",
                          "city": "Pune", "gstin": "27AAGCF3769A1ZP", "status": "ACTIVE"}],
    "supply_customer_orders": [{"id": "o1", "master_reference_no": "MR-0001", "customer_order_no": "CPO-555", "customer_id": "c1",
                                "part_id": "p1", "order_type": "SCHEDULE", "order_qty": 100, "status": "OPEN", "order_date": "2026-10-01",
                                "responsible_employee_id": "e1"}],
    "supply_downstream_events": [{"id": "d1", "customer_order_id": "o1", "event_type": "DISPATCH", "reference_no": "DSP-9", "qty_pcs": 50}],
    "supply_rm_receipts": [{"id": "r1", "receipt_number": "RMR-0042", "heat_number": "H12345", "supplier_challan": "CH-7",
                            "receipt_date": "2026-10-05", "received_qty_kg": 1200.5, "status": "RECEIVED"}],
    "supply_forging_orders": [{"id": "f1", "supplier_order_no": "FO-0077", "forging_supplier_id": "s1", "order_date": "2026-10-02",
                               "ordered_qty_pcs": 500, "status": "OPEN"}],
    "supply_forging_receipts": [{"id": "fr1", "forging_order_id": "f1", "receipt_number": "FR-1", "received_qty_pcs": 200}],
    "osp_jobs": [{"id": "j1", "osp_job_number": "OSP-0100", "vendor_id": "v1", "process_id": "pr1", "part_id": "p1",
                  "dispatch_date": "2026-10-06", "quantity_dispatched": 300, "status": "AT_VENDOR"}],
    "osp_receipts": [{"id": "or1", "osp_job_id": "j1", "receipt_number": "OSPR-1", "quantity_received": 120}],
    "npd_orders": [{"id": "n1", "order_number": "NPD-0005", "part_id": "p1", "customer_id": "c1", "order_qty": 20, "status": "IN_PROGRESS"}],
    "npd_order_steps": [{"id": "ns1", "npd_order_id": "n1", "sequence_no": 1, "step_name": "Drawing Review", "status": "DONE"},
                        {"id": "ns2", "npd_order_id": "n1", "sequence_no": 2, "step_name": "Trial Forging", "status": "OPEN"}],
    "qc_calculation_records": [{"id": "q1", "calculation_number": "QC-0001", "calculation_type": "HARDNESS_CONVERSION", "primary_value": 30,
                                "primary_unit": "HRC", "result_value": 286, "conversion_unit": "HB", "status": "SAVED", "result_payload": {}}],
}


import pytest


@pytest.fixture(autouse=True)
def _generic_osp_print(monkeypatch):
    # these tests check the generic print; the OSP vendor challan (agent A module) is covered in test_r37_universal_pdf
    import sys
    monkeypatch.setitem(sys.modules, "core.osp_challan_pdf", None)


class FakeRepo:
    def select(self, table, eq=None, in_=None, limit=None, order_by=None, desc=False, **_):
        rows = list(DB.get(table, []))
        return [dict(r) for r in rows if all(str(r.get(k)) == str(v) for k, v in (eq or {}).items())]

    def get(self, table, record_id):
        rows = self.select(table, eq={"id": record_id})
        return rows[0] if rows else None


def _text(pdf: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)


def test_generic_print_for_five_transaction_types():
    from core.record_print import print_pdf
    cases = [
        ("supply_customer_orders", "o1", "CUSTOMER ORDER / SCHEDULE", "MR-0001", ["Part", "40256626", "Pinion Gear", "Responsible Employee", "Rajesh Dhokale", "MACHINING / FG / DISPATCH", "DSP-9"]),
        ("supply_rm_receipts", "r1", "RM RECEIPT", "RMR-0042", ["Heat Number", "H12345", "Supplier Challan", "Received Qty kg"]),
        ("supply_forging_orders", "f1", "FORGING ORDER", "FO-0077", ["Forging Supplier", "Bharat Forge", "FORGING RECEIPTS", "FR-1"]),
        ("osp_jobs", "j1", "OSP MATERIAL OUT", "OSP-0100", ["Vendor", "Heat Treat Works", "Process", "Heat Treatment", "OSP INWARD RECEIPTS", "OSPR-1"]),
        ("npd_orders", "n1", "NPD ORDER / STATUS", "NPD-0005", ["NPD STEPS", "Drawing Review", "Trial Forging", "Customer", "Acme Tractors"]),
    ]
    for table, rid, title, number, labels in cases:
        name, pdf = print_pdf(table, rid, FakeRepo())
        assert pdf[:4] == b"%PDF" and name.endswith(".pdf") and number in name, (table, name)
        text = _text(pdf)
        assert title in text, (table, text[:400])
        assert number in text and "DOC NO." in text and "DATE" in text
        assert "Generated by STAWN QCMS" in text and "Page 1 of 1" in text
        assert "Four Star Industries D9" in text and "GSTIN: 27AAGCF3769A1ZP" in text
        for label in labels:
            assert label in text, (table, label)
        assert "c1" not in text.split() and "p1" not in text.split()  # raw ids are never printed


def test_customer_safe_copy_hides_customer():
    from core.record_print import print_pdf, record_print_data
    name, pdf = print_pdf("supply_customer_orders", "o1", FakeRepo(), external=True)
    text = _text(pdf)
    assert "customer_safe" in name
    for secret in ("Acme Tractors", "CUST01", "CPO-555", "MR-0001"):
        assert secret not in text, secret
    assert "40256626" in text and "Customer-safe copy" in text
    internal = _text(print_pdf("supply_customer_orders", "o1", FakeRepo())[1])
    assert "Acme Tractors" in internal and "CPO-555" in internal
    data = record_print_data("npd_orders", "n1", FakeRepo(), external=True)
    assert all("Acme" not in str(v) for _k, v in data["fields"])


def test_dedicated_layouts_are_reused_and_failures_fall_back():
    from core.record_print import print_pdf
    name, pdf = print_pdf("qc_calculation_records", "q1", FakeRepo())
    assert name.startswith("QC-0001") and "QC CALCULATION RECORD" in _text(pdf)
    # A controlled print that cannot be built (no payload in the fake repo) never breaks printing.
    DB["inward_lots"] = [{"id": "i1", "inward_number": "INW-0009", "part_id": "p1", "heat_number": "H1", "status": "RELEASED"}]
    try:
        name, pdf = print_pdf("inward_lots", "i1", FakeRepo())
        assert pdf[:4] == b"%PDF" and "INW-0009" in _text(pdf)
    finally:
        DB.pop("inward_lots", None)
    src = (ROOT / "core/record_print.py").read_text()
    for token in ("_generated_report", "qc_calculation_pdf_bytes", "_calibration_pdf", "_standard_room_pdf", "strip_customer",
                  "transaction_record_pdf_bytes", "supply_purchase_order_items"):
        assert token in src, token


def test_register_pdf_and_record_excel():
    from openpyxl import load_workbook
    from core.record_print import record_excel_bytes, register_pdf
    frame = pd.DataFrame([{"Heat Number": "H1", "Qty kg": 10.5}, {"Heat Number": "H2", "Qty kg": 4.5}])
    text = _text(register_pdf("Heat Register", frame, "Heat: All"))
    assert "HEAT REGISTER" in text and "TOTAL" in text and "Heat: All" in text
    name, data = record_excel_bytes("npd_orders", "n1", FakeRepo())
    book = load_workbook(io.BytesIO(data))
    assert name.endswith(".xlsx") and book.sheetnames[0] == "Record" and len(book.sheetnames) >= 2
    values = [c.value for row in book["Record"].iter_rows() for c in row]
    assert "NPD-0005" in values


def test_print_panel_renders_preview_and_downloads():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string('''
import streamlit as st
import sys
sys.path.insert(0, %r)
from tests.test_r37_print import FakeRepo
from core.record_print import print_panel
print_panel("supply_rm_receipts", "r1", key="t", repo=FakeRepo())
''' % str(ROOT)).run(timeout=60)
    assert not app.exception
    assert any(b.label == "Preview print PDF" for b in app.button)
    app.button(key="_rp_t_supply_rm_receipts_r1_preview").click().run(timeout=60)
    assert not app.exception


def test_global_search_records_centre_and_reports_are_wired():
    gs = (ROOT / "app_pages/global_search.py").read_text()
    viewer = gs[gs.index("def _record_viewer"):gs.index("def _open_record")]
    assert "print_panel(" in viewer and "PRINT LAYOUT PDF" in viewer
    assert "_result_print_area(repo, items, key, clicked)" in gs and "def _print_target" in gs and "_resolve_target(repo, table, row)" in gs
    assert 'key="gs_classic"' in gs
    rc = (ROOT / "app_pages/records_center.py").read_text()
    assert "selectable_table(" in rc and "sources=[(\"rmtc_approvals\"" in rc and "sources=[(\"inward_lots\"" in rc
    assert "print_panel(table, record_id" in rc and "register_print_panel(" in rc and "print_panel(table, selected_id" in rc
    rp = (ROOT / "app_pages/reports.py").read_text()
    assert rp.count("preview_toggle(") >= 4 and "register_pdf(" in rp
    for page in (ROOT / "app_pages").glob("*.py"):
        assert "st.dataframe(" not in page.read_text(), page.name
    from app_pages import global_search, records_center, reports  # noqa: F401 - import check
    from core.reporting import transaction_record_pdf_bytes  # noqa: F401
