from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]


def _sample(status: str = "APPROVED"):
    header = {
        "id": "po-test-1",
        "po_number": "PD9210900001",
        "po_type": "FORGING",
        "approval_status": status,
        "status": "OPEN",
        "order_date": "2026-09-21",
        "currency": "INR",
        "supplier_snapshot": {"party_name": "SUPPLIER TEST"},
    }
    item = {
        "item_no": "FSI-ITEM-001",
        "fsi_part_number_snapshot": "FSI-ITEM-001",
        "item_description": "Machined component",
        "part_description_master": "Differential Shaft Pinion",
        "quantity": 100,
        "uom": "PCS",
        "unit_price": 25,
        "gst_percent": 18,
        "gst_amount": 450,
        "line_total": 2950,
        "customer_source_rows": [
            {
                "customer_po_number": "CUST-PO-SECRET-55",
                "po_position": "10",
                "part_number": "CUSTOMER-PART-SECRET-99",
                "part_description": "Differential Shaft Pinion",
                "quantity": 100,
                "uom": "PCS",
                "customer_delivery_date": "2026-10-15",
            }
        ],
    }
    return header, [item]


def test_supplier_pdf_omits_customer_part_and_delivery_and_watermarks_every_page():
    from core.purchase_order_reporting import purchase_order_pdf_bytes
    header, items = _sample("APPROVED")
    pdf = purchase_order_pdf_bytes(header, items, terms_path=ROOT / "templates" / "FSI_STANDARD_PO_TERMS_2023.pdf")
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) >= 2
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "CUST-PO-SECRET-55" in text
    assert "CUSTOMER-PART-SECRET-99" not in text
    assert "15-10-2026" not in text
    first_text = reader.pages[0].extract_text() or ""
    source_section = first_text.split("PO SOURCE REFERENCE", 1)[1].split("RAW MATERIAL / FORGING PARAMETERS", 1)[0]
    assert "PART NUMBER" not in source_section
    assert "DELIVERY" not in source_section
    assert all("APPROVED" in (page.extract_text() or "") for page in reader.pages)


def test_pending_pdf_watermarks_every_page():
    from core.purchase_order_reporting import purchase_order_pdf_bytes
    header, items = _sample("PENDING_APPROVAL")
    pdf = purchase_order_pdf_bytes(header, items, terms_path=ROOT / "templates" / "FSI_STANDARD_PO_TERMS_2023.pdf")
    reader = PdfReader(BytesIO(pdf))
    assert all("PENDING APPROVAL" in (page.extract_text() or "") for page in reader.pages)


def test_excel_print_watermark_and_supplier_safe_source_reference():
    from core.purchase_order_reporting import purchase_order_excel_bytes
    header, items = _sample("APPROVED")
    data = purchase_order_excel_bytes(header, items)
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb["Purchase Order"]
    assert "APPROVED" in str(ws.oddHeader.center.text)
    values = "\n".join(str(cell.value or "") for row in ws.iter_rows() for cell in row)
    assert "CUST-PO-SECRET-55" in values
    assert "CUSTOMER-PART-SECRET-99" not in values
    assert "2026-10-15" not in values


def test_batch_download_is_one_click_zip_of_individual_pdfs():
    from core.purchase_order_reporting import purchase_order_pdf_files, purchase_order_files_zip_bytes
    h1, i1 = _sample("APPROVED")
    h2, i2 = _sample("PENDING_APPROVAL")
    h2 = dict(h2); h2["id"] = "po-test-2"; h2["po_number"] = "PD9210900002"
    files = purchase_order_pdf_files([(h1, i1), (h2, i2)], terms_path=ROOT / "templates" / "FSI_STANDARD_PO_TERMS_2023.pdf")
    assert len(files) == 2
    packed = purchase_order_files_zip_bytes(files)
    with ZipFile(BytesIO(packed)) as zf:
        names = zf.namelist()
        assert len(names) == 2
        assert all(name.lower().endswith(".pdf") for name in names)
    supply = (ROOT / "app_pages" / "supply_chain.py").read_text()
    assert "Download All Selected POs" in supply


def test_reminder_guard_mobile_collapse_icon_and_ios_package_are_present():
    migration = (ROOT / "supabase" / "migrations" / "20260921133000_qcms_v41435_po_confirmation_reminder_guard.sql").read_text()
    edge = (ROOT / "supabase" / "functions" / "qcms-po-confirmation-reminder" / "index.ts").read_text()
    assert "qcms_po_reminder_allowed" in migration
    assert "qcms_claim_notification_for_send" in migration
    assert "qcms_notification_send_is_current" in migration
    assert "SUPPLIER_PO_CONFIRMATION" in migration
    assert "qcms_po_reminder_allowed" in edge
    assert "qcms_notification_send_is_current" in edge

    android = (ROOT / "mobile" / "android_qcms" / "app" / "src" / "main" / "java" / "com" / "fourstar" / "qcms" / "MainActivity.java").read_text()
    assert ("__qcmsMobileNavOpen" in android or ("openDrawer" in android and "closeDrawer" in android))
    assert ("navButton" in android or "drawerPanel" in android)
    assert ("QCMSMobile/0.1.3" in android or "QCMSMobile/0.1.4" in android or "QCMSMobile/0.1.5" in android or "QCMSMobile/0.1.6" in android or "QCMSMobile/0.1.7" in android or "QCMSMobile/0.1.8" in android or "QCMSMobile/0.1.9" in android or "QCMSMobile/0.2.0" in android or "QCMSMobile/0.2.1" in android or "QCMSMobile/0.2.2" in android or "QCMSMobile/0.2.3" in android or "QCMSMobile/0.2.4" in android)
    assert (ROOT / "mobile" / "android_qcms" / "app" / "src" / "main" / "res" / "drawable-nodpi" / "stawn_icon.png").exists()

    ios = ROOT / "mobile" / "ios_qcms"
    assert (ios / "QCMSMobileIOS.xcodeproj" / "project.pbxproj").exists()
    assert (ios / "QCMSMobileIOS" / "QCMSWebView.swift").exists()
    assert (ios / "BUILD_IPA_ON_MAC.command").exists()
    assert (ios / "QCMSMobileIOS" / "Assets.xcassets" / "AppIcon.appiconset" / "AppIcon-1024.png").exists()
