"""R16: Report Builder, Master List Excel, Bend layouts at any stage, company-mailbox email as the user."""
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from core import mail_service as ms
from core import report_builder as rb

ROOT = Path(__file__).resolve().parents[1]
ROWS = [
    {"id": "1", "tenant_id": "t", "part_id": "p1", "qty": "10", "status": "OPEN", "order_date": "2026-09-01", "tags": ["a", "b"]},
    {"id": "2", "tenant_id": "t", "part_id": "p2", "qty": 5, "status": "CLOSED", "order_date": "2026-10-05"},
    {"id": "3", "tenant_id": "t", "part_id": "p1", "qty": 7, "status": "OPEN", "order_date": "2026-10-06"},
]
LOAD = lambda t: [{"id": "p1", "part_number": "40256626", "part_name": "Yoke"}, {"id": "p2", "part_number": "X2"}] if t == "parts" else []


def test_readable_frame_labels_and_hides_system_columns():
    f = rb.readable_frame(ROWS, LOAD)
    assert "tenant_id" not in f.columns and "id" not in f.columns
    assert f.loc[0, "part_id"] == "40256626 · Yoke" and f.loc[0, "tags"] == "a, b"


def test_build_report_filters_group_sort_limit():
    f = rb.readable_frame(ROWS, LOAD)
    spec = rb.ReportSpec(dataset="x", group_by=["part_id"], measures=[{"column": "qty", "agg": "sum"}], sort_by="Sum of Qty", descending=True)
    out = rb.build_report(f, spec)
    assert list(out.columns) == ["Part", "Records", "Sum of Qty"] and out.loc[0, "Sum of Qty"] == 17
    dated = rb.build_report(f, rb.ReportSpec(dataset="x", filters=[{"column": "order_date", "op": "between", "value": "2026-10-01", "value2": "2026-10-31"}]))
    assert len(dated) == 2
    eq = rb.build_report(f, rb.ReportSpec(dataset="x", columns=["qty"], filters=[{"column": "status", "op": "equals", "value": "open"}], sort_by="qty", limit=1))
    assert len(eq) == 1 and str(eq.iloc[0, 0]) == "7"
    with pytest.raises(ValueError):
        rb.apply_filter(f, "nope", "contains", "x")
    assert rb.ReportSpec.from_json(spec.to_json()).group_by == ["part_id"]


def test_excel_is_clean_with_title_header_and_filter():
    data = rb.excel_bytes([("Parts", pd.DataFrame({"Part": ["A", "B"], "Qty": [1, 2]})), ("Parts", pd.DataFrame())], title="Master List")
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Parts", "Parts 2"]
    ws = wb["Parts"]
    assert "Master List" in ws["A1"].value and ws["A4"].value == "Part" and ws["A4"].font.bold
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:B6"
    assert ws["A5"].fill.fgColor.rgb in ("00000000", None) or ws["A5"].fill.fill_type is None


def test_master_catalog_covers_all_masters():
    from app_pages.report_builder import master_catalog, master_frame
    keys = {i["key"] for i in master_catalog()}
    assert {"parts", "customers", "suppliers", "steel_mills", "osp_vendors", "material_grades", "processes", "employees"} <= keys
    customers = next(i for i in master_catalog() if i["key"] == "customers")
    rows = [{"id": "1", "party_code": "C1", "party_name": "Dana", "party_types": ["CUSTOMER"]}, {"id": "2", "party_code": "S1", "party_name": "JSW", "party_types": ["SUPPLIER"]}]
    frame = master_frame(customers, rows, lambda t: [])
    assert len(frame) == 1 and "Name" in frame.columns


def test_routes_registered():
    app = (ROOT / "streamlit_app.py").read_text()
    for token in ('"report-builder"', '"master-lists"', "report_builder.render_master_lists", '"Report Builder": "report-builder"'):
        assert token in app


def test_company_outbox_rows_show_user_as_sender_and_split_bcc():
    to = [ms.Recipient("a@fsi.in", "Asha", "Employee")]
    bcc = [ms.Recipient(f"u{i}@fsi.in", "", "Dept") for i in range(150)]
    rows = ms.company_outbox_rows(subject="Hi", html_body="<p>x</p>", text_body="x", to=to, bcc=bcc, sender_name="Rajesh Dhokale", sender_email="rajesh@fsi.in")
    assert len(rows) == 2 and rows[0]["context"]["sender_email"] == "rajesh@fsi.in" and rows[0]["context"]["to_emails"] == ["a@fsi.in"]
    assert rows[1]["recipient_email"] == "rajesh@fsi.in" and rows[1]["context"]["to_emails"] == []
    assert sum(len(r["bcc_emails"]) for r in rows) == 150 and all(r["event_key"] == "USER_EMAIL" for r in rows)
    with pytest.raises(ValueError):
        ms.company_outbox_rows(subject="Hi", html_body="", text_body="", to=to, sender_name="X", sender_email="not-an-email")


def test_edge_function_sends_as_user_with_fallback():
    ts = (ROOT / "supabase/functions/qcms-send-email/index.ts").read_text()
    for token in ("qcms_claim_notification_for_send", "qcms_notification_send_is_current", "ctx.sender_email", "replyTo: senderEmail", "via QCMS", "sendAsRefused", "535 5.7.139"):
        assert token in ts


def test_bend_layout_any_stage_and_all_bend_layouts_selectable():
    layouts = (ROOT / "app_pages/inspection_layouts.py").read_text()
    assert "Any stage (Bend Test)" in layouts
    from core.inspection_service import InspectionService
    svc = InspectionService.__new__(InspectionService)
    svc.plans = lambda *a, **k: [{"id": "b1", "inspection_method": "BEND_TEST", "inward_type": "OSP_PROCESS"}, {"id": "m1", "inspection_method": "GENERAL"}, {"id": "b2", "inspection_method": "GENERAL"}]
    svc.plan_inspection_method = lambda pid: "BEND_TEST" if pid == "b2" else "GENERAL_METLAB"
    assert [r["id"] for r in svc.bend_plans("p")] == ["b1", "b2"]
    from app_pages.metlab_report import _filter_plans_by_method
    assert [r["id"] for r in _filter_plans_by_method(svc, [], "BEND_TEST", "p")] == ["b1", "b2"]
