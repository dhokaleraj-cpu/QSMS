"""R38 · Part Processes & Drawings (Part Master), ECN / PCN (format FSI/806/F05) and per-Part Sections applicable."""
from __future__ import annotations

import json
import re
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase/migrations/20261028120000_qcms_r38_part_process_ecn.sql"
MIGRATION_2 = ROOT / "supabase/migrations/20261028130000_qcms_r38_part_process_ecn.sql"
FULL = {k: True for k in ("can_view", "can_create", "can_edit", "can_archive", "can_validate", "can_approve")}


class FakeRepo:
    def __init__(self, tables=None, rpc_numbers=False):
        self.tables = deepcopy(tables or {})
        self.calls = []
        self.rpc_numbers = rpc_numbers
        self.tenant_id = "t1"

    def select(self, table, eq=None, in_=None, order_by=None, desc=False, limit=None, **kwargs):
        rows = [r for r in self.tables.get(table, []) if all(str(r.get(k)) == str(v) for k, v in (eq or {}).items())]
        if order_by:
            rows = sorted(rows, key=lambda r: str(r.get(order_by) or ""), reverse=bool(desc))
        return deepcopy(rows[:limit] if limit else rows)

    def get(self, table, rid):
        return next((deepcopy(r) for r in self.tables.get(table, []) if str(r.get("id")) == str(rid)), None)

    def find_one(self, table, *, eq):
        rows = self.select(table, eq=eq, limit=1)
        return rows[0] if rows else None

    def insert(self, table, payload):
        row = {"id": f"{table}-{len(self.tables.get(table, [])) + 1}", "created_at": "2026-10-28T10:00:00+00:00", **deepcopy(payload)}
        self.tables.setdefault(table, []).append(row)
        self.calls.append(("insert", table, dict(payload)))
        return deepcopy(row)

    def update(self, table, rid, payload):
        for row in self.tables.get(table, []):
            if str(row.get("id")) == str(rid):
                row.update(deepcopy(payload))
        self.calls.append(("update", table, rid, dict(payload)))
        return self.get(table, rid)

    def delete(self, table, rid):
        self.tables[table] = [r for r in self.tables.get(table, []) if str(r.get("id")) != str(rid)]

    def upsert_by(self, table, payload, *, natural_key):
        for row in self.tables.get(table, []):
            if all(str(row.get(k)) == str(v) for k, v in natural_key.items()):
                row.update(payload)
                return deepcopy(row), "UPDATED"
        return self.insert(table, payload), "CREATED"

    def count(self, table, **kwargs):
        return len(self.select(table, **kwargs))

    def rpc(self, name, params=None):
        self.calls.append(("rpc", name, dict(params or {})))
        if name == "qcms_next_ecn_number" and self.rpc_numbers:
            from core.ecn_service import next_from_existing
            return next_from_existing((r.get("ecn_number") for r in self.tables.get("part_ecns", [])), params["p_change_type"],
                                      params["p_plant_code"], date.fromisoformat(params["p_date"]))
        raise RuntimeError(f"RPC {name} unavailable in the fake repository")


BASE = {
    "parts": [{"id": "p1", "part_number": "7568", "part_name": "Shaft Diff. Pinion", "drawing_number": "40227568", "drawing_revision": "A",
               "status": "ACTIVE", "section_applicability": {"jominy": False, "osp_metlab": "no"}},
              {"id": "p2", "part_number": "7553", "part_name": "Shaft", "status": "ACTIVE"}],
    "employees": [{"id": "e1", "first_name": "Rajesh", "last_name": "Dhokale", "email": "rajesh@fourstarindustries.com", "status": "ACTIVE", "profile_id": "u1"},
                  {"id": "e2", "first_name": "Nitin", "last_name": "Gawde", "email": "nitin@fourstarindustries.com", "status": "ACTIVE"}],
    "processes": [{"id": "pr1", "process_code": "P10", "process_name": "Finish Machining", "process_type": "IN_HOUSE", "status": "ACTIVE"},
                  {"id": "pr2", "process_code": "P20", "process_name": "Plating - ENP", "process_type": "OUTSOURCED", "status": "ACTIVE"}],
    "part_jominy_requirements": [{"id": "j1", "part_id": "p1", "status": "ACTIVE"}],
    "part_raw_material_details": [{"id": "r1", "part_id": "p1"}],
}


# ------------------------------------------------------------------------------------------- sections applicable
def test_section_keys_cover_part_master_and_rmtc_sections():
    from core.part_sections import PART_SECTIONS, SECTION_KEYS
    assert SECTION_KEYS == (
        "controlled_drawings", "customer_standards", "approved_sources", "raw_material", "supplier_technical_data", "price_history",
        "jominy", "osp_metlab", "final_metallurgical", "process_routing", "chemistry", "di_value", "heat_treatment", "microstructure",
    )
    assert {g for _, _, g, _ in PART_SECTIONS} == {"PART", "RMTC"}
    doc = (ROOT / "core/part_sections.py").read_text()
    for key in SECTION_KEYS:
        assert key in doc.split('"""')[1]  # every key is documented in the module contract


def test_applicability_defaults_true_and_reads_stored_values():
    from core.part_sections import applicable_from_row, applicable_sections, normalize_applicability, section_applicable, stored_applicability
    assert all(normalize_applicability(None).values())
    assert all(normalize_applicability("not json").values())
    repo = FakeRepo(BASE)
    values = applicable_sections("p1", repo)
    assert values["jominy"] is False and values["osp_metlab"] is False and values["raw_material"] is True
    assert section_applicable("p1", "jominy", repo) is False
    assert section_applicable("p1", "chemistry", repo) is True
    assert section_applicable("p1", "unknown_key", repo) is True
    assert section_applicable("missing-part", "jominy", repo) is True
    assert section_applicable("", "jominy", repo) is True
    assert applicable_from_row({"section_applicability": json.dumps({"chemistry": "false"})})["chemistry"] is False
    stored = stored_applicability({"jominy": False})
    assert set(stored) == set(normalize_applicability({})) and stored["jominy"] is False and stored["chemistry"] is True


def test_print_filter_and_completeness_skip_not_applicable_sections():
    from core.part_sections import filter_print_sections, not_applicable_labels, section_completeness
    mapping = {"jominy": False}
    printed = filter_print_sections({"Jominy Requirements": [1], "Raw Material Details": [2], "Something Else": [3]}, mapping)
    assert "Jominy Requirements" not in printed and "Raw Material Details" in printed and "Something Else" in printed
    assert not_applicable_labels(mapping) == ["Jominy Requirement / Jominy Results"]
    status = {row["key"]: row["status"] for row in section_completeness(FakeRepo(BASE), "p1", mapping)}
    assert status["jominy"] == "N/A"
    assert status["raw_material"] == "COMPLETE"
    assert status["customer_standards"] == "MISSING"
    assert status["chemistry"] == "OPTIONAL"




# ------------------------------------------------------------------------------------------- ECN numbering / rules
def test_ecn_number_series_per_type_plant_and_year():
    from core.ecn_service import format_ecn_number, next_ecn_number, next_from_existing
    assert format_ecn_number("ECN", "d-9", 7, date(2026, 10, 28)) == "ECN-D9-2026-0007"
    assert format_ecn_number("pcn", "", 1, date(2026, 1, 1)) == "PCN-FSI-2026-0001"
    numbers = ["ECN-D9-2026-0001", "ECN-D9-2026-0009", "PCN-D9-2026-0004", "ECN-G27-2026-0050", "ECN-D9-2025-0099"]
    assert next_from_existing(numbers, "ECN", "D9", date(2026, 5, 1)) == "ECN-D9-2026-0010"
    assert next_from_existing(numbers, "PCN", "D9", date(2026, 5, 1)) == "PCN-D9-2026-0005"
    assert next_from_existing(numbers, "ECN", "D9", date(2027, 1, 2)) == "ECN-D9-2027-0001"  # yearly restart
    repo = FakeRepo({"part_ecns": [{"id": "x", "ecn_number": "ECN-D9-2026-0003"}]})
    assert next_ecn_number(repo, "ECN", "D9", date(2026, 3, 1)) == "ECN-D9-2026-0004"  # RPC unavailable -> register fallback
    service_repo = FakeRepo(BASE, rpc_numbers=True)
    from core.ecn_service import ECNService
    svc = ECNService(service_repo)
    numbers = [svc.create(_full_form(change_type=t), plant_code="D9")["ecn_number"] for t in ("ECN", "ECN", "PCN")]
    assert numbers == ["ECN-D9-2026-0001", "ECN-D9-2026-0002", "PCN-D9-2026-0001"]


def test_status_flow_validation_and_permission_matrix():
    from core.ecn_service import (ECN_STATUSES, allowed_targets, can_transition, normalize_ecn_payload, normalize_status,
                                  status_label, validate_ecn)
    assert ECN_STATUSES == ("DRAFT", "PENDING_APPROVAL", "APPROVED", "DISAPPROVED", "COMPLETED", "CANCELLED")
    assert normalize_status("submitted") == "PENDING_APPROVAL" and normalize_status("REJECTED") == "DISAPPROVED"
    assert normalize_status("IMPLEMENTED") == "COMPLETED" and normalize_status(None) == "DRAFT"
    assert status_label("PENDING_APPROVAL") == "Pending Approval"
    draft = normalize_ecn_payload({"part_number": "7568", "revision_requested": "New PFD"})
    assert validate_ecn(draft) == []
    errors = validate_ecn(draft, submitting=True)
    assert any("FSI Representative" in e for e in errors) and any("Drawing No." in e for e in errors)
    assert validate_ecn(normalize_ecn_payload({})) and "Revision Requested" in " ".join(validate_ecn({"part_number": "1"}))
    editor = {"can_edit": True}
    approver = {"can_edit": True, "can_approve": True}
    assert can_transition("DRAFT", "PENDING_APPROVAL", editor)
    assert not can_transition("PENDING_APPROVAL", "APPROVED", editor)
    assert can_transition("PENDING_APPROVAL", "APPROVED", approver) and can_transition("PENDING_APPROVAL", "DISAPPROVED", approver)
    assert can_transition("SUBMITTED", "APPROVED", approver)  # legacy name of the first cut
    assert not can_transition("DRAFT", "APPROVED", approver)
    assert can_transition("APPROVED", "COMPLETED", editor) and not can_transition("COMPLETED", "DRAFT", approver)
    assert can_transition("DISAPPROVED", "DRAFT", editor) and not can_transition("DISAPPROVED", "COMPLETED", approver)
    assert allowed_targets("PENDING_APPROVAL", editor) == ["DRAFT"]
    assert allowed_targets("PENDING_APPROVAL", approver) == ["APPROVED", "DISAPPROVED", "DRAFT"]
    payload = normalize_ecn_payload({"inventory_affected": False, "inventory_qty": 50, "demands_affected": True, "demands_qty": "12",
                                     "supporting_documents": {"pfd": 1, "bogus": True}, "fsi_decision": "maybe"})
    assert payload["inventory_qty"] is None and payload["demands_qty"] == 12.0
    assert payload["supporting_documents"]["pfd"] is True and "bogus" not in payload["supporting_documents"]
    assert payload["fsi_decision"] is None


def _full_form(**extra):
    form = {"part_id": "p1", "part_number": "7568", "part_name": "Shaft Diff. Pinion", "drawing_number": "40227568",
            "previous_description": "Old route G27", "revision_requested": "New route D9", "fsi_representative_ids": ["e1", "e2"],
            "to_contact": {"name": "Prasanjit Samanta", "phone": "9892320872"}, "change_type": "PCN", "ecn_date": date(2026, 10, 28)}
    form.update(extra)
    return form


def test_service_create_save_transition_flow():
    from core.ecn_service import ECNService
    repo = FakeRepo(BASE, rpc_numbers=True)
    service = ECNService(repo)
    ecn = service.create(_full_form(), status="PENDING_APPROVAL", plant_code="D9")
    assert ecn["ecn_number"] == "PCN-D9-2026-0001" and ecn["status"] == "PENDING_APPROVAL"
    with pytest.raises(PermissionError):
        service.transition(ecn["id"], "APPROVED", {"can_edit": True})
    approved = service.transition(ecn["id"], "APPROVED", FULL, remarks="OK for D9", approver_employee_id="e2")
    assert approved["status"] == "APPROVED" and approved["fsi_decision"] == "APPROVED" and "OK for D9" in approved["comments"]
    assert approved["fsi_approver_employee_id"] == "e2" and approved["fsi_approval_date"] == date.today().isoformat()
    # after approval only the FSI-use / completion block can change
    saved = service.save(ecn["id"], {"revision_requested": "changed!", "completed_date": date(2026, 11, 2)})
    assert saved["revision_requested"] == "New route D9" and saved["completed_date"] == "2026-11-02"
    done = service.transition(ecn["id"], "COMPLETED", FULL, approver_employee_id="e1")
    assert done["status"] == "COMPLETED" and done["completed_by_employee_id"] == "e1"
    second = service.create(_full_form(change_type="ECN"), plant_code="D9")
    assert second["ecn_number"] == "ECN-D9-2026-0001" and second["status"] == "DRAFT"
    third = service.create(_full_form(change_type="ECN"), status="PENDING_APPROVAL", plant_code="D9")
    rejected = service.transition(third["id"], "DISAPPROVED", FULL)
    assert rejected["status"] == "DISAPPROVED" and rejected["fsi_decision"] == "DISAPPROVED"
    assert service.save(third["id"], {"revision_requested": "Reworked"})["revision_requested"] == "Reworked"  # editable again
    with pytest.raises(ValueError):
        service.create({"part_number": ""}, status="DRAFT")
    with pytest.raises(ValueError):
        service.create(_full_form(), status="APPROVED")
    with pytest.raises(ValueError):
        service.transition(service.create({"part_number": "7568", "revision_requested": "x"}, plant_code="D9")["id"], "PENDING_APPROVAL", FULL)


def test_release_with_ecn_links_revision_and_cancels_on_failure():
    from core.ecn_service import ECNService, release_with_ecn, revision_defaults
    repo = FakeRepo(BASE, rpc_numbers=True)
    service = ECNService(repo)
    current = {"id": "att-old", "drawing_number": "PD-10", "revision": "01", "revision_date": "2026-01-01"}
    defaults = revision_defaults(repo.get("parts", "p1"), document_scope="PROCESS_DRAWING", drawing_number="PD-10", current=current,
                                 new_revision="02", new_revision_date=date(2026, 10, 28), process_text="Op 10 · Finish Machining")
    assert defaults["previous_attachment_id"] == "att-old" and defaults["part_number"] == "7568" and defaults["part_name"] == "Shaft Diff. Pinion"
    assert "Rev 01" in defaults["previous_description"] and "Rev 02" in defaults["revision_requested"]
    finish = revision_defaults(repo.get("parts", "p1"), document_scope="FINISH_DRAWING", drawing_number="40227568",
                               current={"revision": "A"}, new_revision="B")
    assert finish["change_type"] == "ECN" and finish["drawing_number"] == "40227568" and finish["previous_revision"] == "A" and finish["new_revision"] == "B"
    ecn, released = release_with_ecn(service, {**defaults, "plant_code": "D9"}, lambda: {"id": "att-new", "superseded_id": "att-old"})
    stored = repo.get("part_ecns", ecn["id"])
    assert stored["document_attachment_id"] == "att-new" and stored["previous_attachment_id"] == "att-old" and stored["status"] == "DRAFT"

    def boom():
        raise RuntimeError("storage offline")
    with pytest.raises(RuntimeError):
        release_with_ecn(service, {**defaults, "plant_code": "D9"}, boom)
    cancelled = [r for r in repo.tables["part_ecns"] if r["status"] == "CANCELLED"]
    assert len(cancelled) == 1 and "storage offline" in cancelled[0]["comments"]


def test_register_frame_email_text_and_default_representative():
    from core.ecn_service import ECNService, default_representative, ecn_email_text
    repo = FakeRepo(BASE, rpc_numbers=True)
    service = ECNService(repo)
    ecn = service.create(_full_form(), plant_code="D9")
    frame = service.register_frame(service.list(part_id="p1"), {"p1": repo.get("parts", "p1")})
    assert list(frame["ECN No."]) == [ecn["ecn_number"]] and frame.iloc[0]["Revision Linked"] == "No" and frame.iloc[0]["Status"] == "Draft"
    subject, body = ecn_email_text(ecn)
    assert subject.startswith("Engineering · PCN " + ecn["ecn_number"]) and "Process Change Note" in subject and "New route D9" in body
    from core.notification_service import department_subject
    assert department_subject(subject, "Engineering", "ECN_PCN") == subject  # not prefixed twice
    assert default_representative(BASE["employees"], {"id": "u1"}) == ["e1"]
    assert default_representative(BASE["employees"], {"email": "NITIN@fourstarindustries.com"}) == ["e2"]
    assert default_representative(BASE["employees"], {}) == []


# ------------------------------------------------------------------------------------------- part processes
def test_part_process_payloads_validate_and_default_source():
    from core.part_process_service import part_process_payloads
    procs = {r["id"]: r for r in BASE["processes"]}
    labels = {"P10 · Finish Machining": "pr1", "P20 · Plating - ENP": "pr2"}
    grid = pd.DataFrame([
        {"_id": None, "Op No.": 10, "Process": "P10 · Finish Machining", "Mandatory": True, "Source": "IN_HOUSE", "Remarks": "", "Status": "ACTIVE"},
        {"_id": "route-2", "Op No.": 20, "Process": "P20 · Plating - ENP", "Mandatory": False, "Source": None, "Remarks": "ENP", "Status": "ACTIVE"},
        {"_id": None, "Op No.": None, "Process": "", "Mandatory": None, "Source": None, "Remarks": None, "Status": None},
    ])
    payloads = part_process_payloads(grid, labels, procs)
    assert len(payloads) == 2
    assert payloads[0]["is_mandatory"] is True and payloads[0]["_id"] is None
    assert payloads[1]["source_type"] == "OSP" and payloads[1]["is_mandatory"] is False and payloads[1]["_id"] == "route-2"
    dup = pd.concat([grid.iloc[:1], grid.iloc[:1]])
    with pytest.raises(ValueError, match="Duplicate Op No"):
        part_process_payloads(dup, labels, procs)
    with pytest.raises(ValueError, match="valid Process"):
        part_process_payloads(pd.DataFrame([{"Op No.": 10, "Process": "nope"}]), labels, procs)
    inactive_dup = pd.concat([grid.iloc[:1], grid.iloc[:1].assign(Status="INACTIVE")])
    assert len(part_process_payloads(inactive_dup, labels, procs)) == 2  # an INACTIVE row may reuse the Op No.


ROUTES = [
    {"id": "route-1", "part_id": "p1", "operation_no": 10, "process_id": "pr1", "is_mandatory": True, "source_type": "IN_HOUSE", "status": "ACTIVE"},
    {"id": "route-2", "part_id": "p1", "operation_no": 20, "process_id": "pr2", "is_mandatory": False, "source_type": "OSP", "status": "ACTIVE", "remarks": "Vendor ABC"},
    {"id": "route-3", "part_id": "p1", "operation_no": 30, "process_id": "pr2", "is_mandatory": True, "source_type": "OSP", "status": "INACTIVE"},
    {"id": "route-4", "part_id": "p2", "operation_no": 10, "process_id": "pr1", "is_mandatory": True, "source_type": "IN_HOUSE", "status": "ACTIVE"},
]


def test_mandatory_processes_contract_and_route_text():
    from core.part_process_service import PartProcessService, mandatory_processes, process_change_defaults, route_text
    repo = FakeRepo({**BASE, "part_process_routes": ROUTES})
    rows = mandatory_processes("p1", repo)
    assert [(r["operation_no"], r["process_name"], r["route_id"]) for r in rows] == [(10, "Finish Machining", "route-1")]
    assert rows[0]["process_code"] == "P10" and rows[0]["source_type"] == "IN_HOUSE"
    assert mandatory_processes("", repo) == [] and mandatory_processes("p1", object()) == []  # never raises
    service = PartProcessService(repo)
    assert [r["id"] for r in service.active_routes("p1")] == ["route-1", "route-2"]
    text = route_text(service.active_routes("p1"), {p["id"]: p for p in BASE["processes"]})
    assert text.splitlines() == ["1) Op 10 P10 · Finish Machining (Mandatory, In-house)", "2) Op 20 P20 · Plating - ENP (Optional, OSP) - Vendor ABC"]
    pcn = process_change_defaults(repo.get("parts", "p1"), "1) Op 10 old", text)
    assert pcn["change_type"] == "PCN" and pcn["part_number"] == "7568" and "Op 10 old" in pcn["previous_description"]
    assert "Plating" in pcn["revision_requested"] and pcn["supporting_documents"]["pfd"] and pcn["supporting_documents"]["pfmea"]
    rows = service.print_rows("p1", {p["id"]: p for p in BASE["processes"]})
    assert [r["Mandatory"] for r in rows] == ["Yes", "No"] and rows[1]["Source"] == "OSP"
    assert service.save_routes("p2", [{"_id": "route-4", "operation_no": 15, "process_id": "pr1"}, {"_id": None, "operation_no": 25, "process_id": "pr2"}]) == 2
    assert [r["operation_no"] for r in service.routes("p2")] == [15, 25]


class _Bucket:
    def __init__(self, log):
        self.log = log

    def upload(self, path, content, options):
        self.log.append(("upload", path, len(content), options["upsert"]))

    def remove(self, paths):
        self.log.append(("remove", tuple(paths)))


class _Client:
    def __init__(self):
        self.log = []
        self.storage = self

    def from_(self, bucket):
        assert bucket == "quality-documents"
        return _Bucket(self.log)


class _RevisionRepo(FakeRepo):
    """Mimics qcms_activate_part_process_document_revision: new revision ACTIVE, previous same number INACTIVE."""

    def rpc(self, name, params=None):
        if name != "qcms_activate_part_process_document_revision":
            return super().rpc(name, params)
        p = dict(params)
        docs = self.tables.setdefault("document_attachments", [])
        if any(d["entity_id"] == p["p_process_route_id"] and d["drawing_number"] == p["p_document_number"] and d["revision"] == p["p_revision"] for d in docs):
            raise RuntimeError("This Drawing / Specification Number and Revision Number already exist in revision history")
        prev = next((d for d in docs if d["entity_id"] == p["p_process_route_id"] and d["document_type"] == p["p_document_type"]
                     and d["drawing_number"] == p["p_document_number"] and d["status"] == "ACTIVE"), None)
        if prev:
            prev["status"] = "INACTIVE"
        row = {"id": f"att-{len(docs) + 1}", "entity_type": "PART_PROCESS", "entity_id": p["p_process_route_id"], "document_type": p["p_document_type"],
               "drawing_number": p["p_document_number"], "revision": p["p_revision"], "revision_date": p["p_revision_date"],
               "file_name": p["p_file_name"], "object_path": p["p_object_path"], "status": "ACTIVE", "created_at": f"2026-10-28T10:0{len(docs)}:00"}
        docs.append(row)
        return [{**row, "superseded_id": prev["id"] if prev else None}]


def test_process_document_revision_new_rev_makes_old_inactive():
    from core.part_process_service import PartProcessService, StoredUpload, revision_status, save_process_document_revision
    repo = _RevisionRepo({**BASE, "part_process_routes": ROUTES})
    client = _Client()
    route = repo.get("part_process_routes", "route-1")
    first = save_process_document_revision(repo, route, "PROCESS_DRAWING", "PD-10", "Turning", "01", date(2026, 1, 1),
                                           StoredUpload("pd10_r01.pdf", b"%PDF-1", "application/pdf"), client=client)
    second = save_process_document_revision(repo, route, "PROCESS_DRAWING", "PD-10", "Turning", "02", date(2026, 10, 28),
                                            StoredUpload("pd10_r02.pdf", b"%PDF-2", "application/pdf"), client=client)
    assert second["superseded_id"] == first["id"]
    docs = {d["revision"]: d for d in PartProcessService(repo).documents("route-1")}
    assert revision_status(docs["01"]) == "OBSOLETE" and revision_status(docs["02"]) == "CURRENT"
    assert [d["revision"] for d in PartProcessService(repo).current_documents("route-1")] == ["02"]
    assert client.log[0][0] == "upload" and "/process_documents/route-1/process_drawing/PD-10/01_" in client.log[0][1] and client.log[0][3] == "false"
    with pytest.raises(RuntimeError):  # duplicate revision -> storage object removed again
        save_process_document_revision(repo, route, "PROCESS_DRAWING", "PD-10", "", "02", date(2026, 10, 28),
                                       StoredUpload("dup.pdf", b"x", None), client=client)
    assert client.log[-1][0] == "remove"
    with pytest.raises(ValueError, match="Revision Number"):
        save_process_document_revision(repo, route, "PROCESS_SPEC", "SP-1", "", "", date.today(), StoredUpload("a.pdf", b"x", None), client=client)
    with pytest.raises(ValueError, match="empty"):
        save_process_document_revision(repo, route, "PROCESS_SPEC", "SP-1", "", "A", date.today(), StoredUpload("a.pdf", b"", None), client=client)


# ------------------------------------------------------------------------------------------- print
def _print_ecn(**extra):
    from core.ecn_service import normalize_ecn_payload
    return {**normalize_ecn_payload(_full_form(supporting_documents={"pfd": True, "ppap": True}, inventory_affected=True, inventory_qty=1200, **extra)),
            "ecn_number": "PCN-D9-2026-0001", "status": "PENDING_APPROVAL", "benefit_of_change": "Productivity & Product quality Improvements"}


def test_ecn_pdf_follows_fsi_806_f05_layout():
    pypdf = pytest.importorskip("pypdf")
    from core.ecn_pdf import build_ecn_pdf, watermark_label
    ecn = _print_ecn()
    pdf = build_ecn_pdf(ecn, employees={"e1": BASE["employees"][0]})
    assert pdf.startswith(b"%PDF")
    text = " ".join(page.extract_text() for page in pypdf.PdfReader(BytesIO(pdf)).pages)
    for token in ("PROCESS CHANGE NOTE", "FSI/806/F05", "25.09.2024", "PCN-D9-2026-0001", "FSI REPRESENTATIVE", "Rajesh Dhokale",
                  "IDENTIFICATION", "Prasanjit Samanta", "DRAWING CRITERIA", "40227568", "7568", "Shaft Diff. Pinion",
                  "PREVIOUS DESCRIPTION", "Old route G27", "REVISION REQUESTED", "New route D9",
                  "SUPPORTING DOCUMENTS REQUIRED", "Purchase Orders Affected", "Current Inventory Affected", "1,200",
                  "Current Demands Affected", "FSI USE ONLY", "Drawing change to occur", "Approval Signatures", "Completed By",
                  "Benefit to this Change", "PENDING APPROVAL"):
        assert token in text, token
    eng = build_ecn_pdf({**ecn, "change_type": "ECN"})
    eng_text = pypdf.PdfReader(BytesIO(eng)).pages[0].extract_text()
    assert "ENGINEERING CHANGE NOTE" in eng_text and "SUPPORTING DOCUMENTS ATTACHED" in eng_text
    assert watermark_label({"status": "PENDING_APPROVAL"}) == "PENDING APPROVAL" and watermark_label({}) == "DRAFT"
    assert watermark_label({"status": "SUBMITTED"}) == "PENDING APPROVAL" and watermark_label({"status": "COMPLETED"}) == "COMPLETED"


def test_ecn_excel_uses_the_workbook_layout():
    from openpyxl import load_workbook
    from core.ecn_pdf import build_ecn_xlsx
    pcn = load_workbook(BytesIO(build_ecn_xlsx(_print_ecn(), employees={"e1": BASE["employees"][0]})))
    ws = pcn["PCN"]
    values = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")}
    text = " | ".join(str(v) for v in values.values())
    assert values["A3"] == "PROCESS CHANGE NOTE" and values["AA1"] == "FSI/806/F05" and values["AA2"] == "25.09.2024"
    for token in ("FSI Representative(s)", "Rajesh Dhokale", "Contact Name", "Prasanjit Samanta", "Drawing No.", "40227568",
                  "Part No.", "7568", "Part Name", "Previous Description", "Old route G27", "Revision requested", "New route D9",
                  "Supporting Documents Required", "PPAP", "PFMEA", "MAINTAIN INVENTORY", "FSI Use ONLY", "Approval Signatures",
                  "Completed By:", "Comments:", "Benefit to this Change : Productivity"):
        assert token in text, token
    ecn = load_workbook(BytesIO(build_ecn_xlsx({**_print_ecn(), "change_type": "ECN"})))["ECN Form"]
    etext = " | ".join(str(c.value) for row in ecn.iter_rows() for c in row if c.value)
    for token in ("ENGINEERING CHANGE NOTE", "Purchase Orders Affected", "Current Inventory Affected", "Current Demands Affected", "1,200", "Qty."):
        assert token in etext, token


# ------------------------------------------------------------------------------------------- source / DB contracts
def test_part_master_wires_r38_sections_and_ecn_hooks():
    text = (ROOT / "app_pages/part_master.py").read_text()
    for token in ("SECTIONS APPLICABLE", "part_master_render_entry_i", "part_master_render_entry_j", "PART PROCESSES & DRAWINGS",
                  "ENGINEERING / PROCESS CHANGE NOTES (ECN / PCN)", "from app_pages.ecn import",
                  "_render_pending_ecn(repo, existing, \"PART_DRAWING\", perms)", "_queue_pending_ecn(part_id, {",
                  "filter_print_sections", "_require_section(applicable, \"jominy\")", "applicable[\"price_history\"]"):
        assert token in text, token
    assert "st.dataframe(" not in text and "def _ecn_form" not in text  # UI lives in app_pages/ecn.py
    for token in ("qcms_activate_part_drawing_revision", "Previous revision is now INACTIVE", "DRAWING REVISION HISTORY"):
        assert token in text  # existing controlled-drawing contract is kept
    ecn_page = (ROOT / "app_pages/ecn.py").read_text()
    for token in ("Save ECN as DRAFT & Release Revision", "Submit ECN for Approval & Release Revision", "Cancel Replacement",
                  "Raise PCN", "Copy processes from NPD Process Flow", "OBSOLETE", "Download ECN Excel", "SUPPORTING DOCUMENTS ATTACHED",
                  "queue_pending_ecn(part_id, {"):
        assert token in ecn_page, token
    assert "st.dataframe(" not in ecn_page


def test_route_registered_under_engineering_and_verified():
    app = (ROOT / "streamlit_app.py").read_text()
    assert 'st.Page(ecn_page.render, title="Engineering Change Notes"' in app and '"ecn-register":"PART_MASTER"' in app
    assert '"ecn-register": "Engineering"' in app and '(PAGE_BY_PATH["ecn-register"], "ECN / PCN", "Engineering"' in app
    assert '"Engineering": (' in app
    assert '"ecn-register"' in (ROOT / "scripts/verify_phase1.py").read_text()


def test_migrations_are_additive_tenant_scoped_and_listed():
    sql = MIGRATION.read_text()
    for token in ("add column if not exists section_applicability", "create table if not exists public.part_process_routes",
                  "create table if not exists public.part_ecns", "public.current_tenant_id()", "enable row level security",
                  "revoke all on public.%I from anon", "qcms_next_ecn_number", "qcms_activate_part_process_document_revision",
                  "pg_advisory_xact_lock", "ux_part_ecns_number", "'PART_PROCESS'", "qcms_effective_module_permission('PART_MASTER', 'approve')"):
        assert token in sql, token
    sql2 = MIGRATION_2.read_text()
    for token in ("'PENDING_APPROVAL'", "'DISAPPROVED'", "'COMPLETED'", "not valid", "Only an APPROVED ECN / PCN can be marked COMPLETED",
                  "qcms_effective_module_permission('PART_MASTER', 'approve')", "qcms_part_mandatory_processes",
                  "'PART_ECN'", "pg_get_functiondef", "execute 'select public.qcms_search_table_map()' into v_map",
                  "'part_ecns', jsonb_build_array('PART_MASTER'", "'part_process_routes', jsonb_build_array(",
                  "revoke all on function public.qcms_part_mandatory_processes(uuid) from public, anon"):
        assert token in sql2, token
    for text in (sql, sql2):
        for word in ("delete", "update", "drop", "truncate"):
            assert not re.search(rf"\b{word}\b", text, re.I), word
    files = json.loads((ROOT / "supabase/auto_apply.json").read_text())["files"]
    first = files.index("supabase/migrations/20261028120000_qcms_r38_part_process_ecn.sql")
    assert files.index("supabase/migrations/20261028130000_qcms_r38_part_process_ecn.sql") > first


# ------------------------------------------------------------------------------------------- AppTest
SCRIPT_HEAD = '''
import streamlit as st
st.data_editor = lambda frame, **kwargs: frame
from tests.test_r38_part_process_ecn import FakeRepo, BASE, ROUTES, _full_form
from core.ecn_service import ECNService
import app_pages.part_master as page
import app_pages.ecn as ecnp
repo = st.session_state.setdefault("r38_repo", FakeRepo(BASE, rpc_numbers=True))
if not repo.tables.get("part_ecns"):
    ECNService(repo).create(_full_form(), status="PENDING_APPROVAL", plant_code="D9")
FULL = {k: True for k in ("can_view", "can_create", "can_edit", "can_archive", "can_validate", "can_approve")}
page.Repository = lambda: repo
ecnp.Repository = lambda: repo
page.current_permissions = lambda key: dict(FULL)
ecnp.current_permissions = lambda key: dict(FULL)
ecnp._profile = lambda: {"id": "u1"}
ecnp.render_attachment_manager = lambda **kwargs: None
st.session_state["_qsms_pages"] = {k: "streamlit_app.py" for k in ("masters", "part-entry", "part-records", "process-entry", "dashboard", "standards-entry")}
page.subpage_navigation = lambda *a, **k: None
'''


def test_ecn_register_page_renders_with_actions():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string(SCRIPT_HEAD + "\necnp.render()\n").run(timeout=30)
    assert not app.exception, app.exception
    downloads = [b.label for b in app.get("download_button")]
    assert "Download ECN PDF" in downloads and "Download ECN Excel" in downloads
    labels = [b.label for b in app.button]
    assert "Approve" in labels and "Disapprove" in labels and "Preview ECN Print" in labels
    app.button(key=next(b.key for b in app.button if b.label == "Approve")).click().run(timeout=30)
    assert not app.exception, app.exception
    row = app.session_state["r38_repo"].tables["part_ecns"][0]
    assert row["status"] == "APPROVED" and row["fsi_approver_employee_id"] == "e1"  # signed-in user's employee
    assert "Mark Completed" in [b.label for b in app.button]


def test_part_entry_renders_part_processes_and_hides_not_applicable():
    from streamlit.testing.v1 import AppTest
    script = SCRIPT_HEAD + '''
repo.tables.setdefault("part_process_routes", [dict(r) for r in ROUTES])
st.page_link = lambda *a, **k: None
import core.record_history as rh
rh.record_history_panel = lambda *a, **k: None
page.current_profile = lambda: {"id": "u1", "role": "ADMIN"}
page.section_permissions = lambda *a, **k: {"can_view": True, "can_edit": True}
if "edit_part_id" not in st.session_state and not st.session_state.get("_r38_once"):
    st.session_state["edit_part_id"] = "p1"; st.session_state["_r38_once"] = True
page.render_entry()
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception, app.exception
    expanders = [e.label for e in app.expander]
    assert any("PART PROCESSES & DRAWINGS" in label for label in expanders)
    assert any("ENGINEERING / PROCESS CHANGE NOTES" in label for label in expanders)
    infos = " ".join(i.value for i in app.info)
    assert "Jominy Requirement / Jominy Results is marked NOT APPLICABLE" in infos
    labels = [b.label for b in app.button]
    for label in ("Save Sections Applicable", "Release Process Document Revision", "Save Part Processes", "Raise PCN"):
        assert label in labels, label
    app.button(key="raise_pcn_p1").click().run(timeout=30)
    assert not app.exception, app.exception
    assert "Submit PCN for Approval" in [b.label for b in app.button]


def test_replacing_a_current_revision_requires_an_ecn():
    """A replacement is only queued; the revision is released together with the ECN (or not at all on cancel)."""
    from streamlit.testing.v1 import AppTest
    script = SCRIPT_HEAD + '''
repo.tables.setdefault("part_process_routes", [dict(r) for r in ROUTES])
repo.tables.setdefault("document_attachments", [{"id": "att-old", "entity_type": "PART_PROCESS", "entity_id": "route-1", "document_type": "PROCESS_DRAWING",
    "drawing_number": "PD-10", "revision": "01", "revision_date": "2026-01-01", "status": "ACTIVE", "file_name": "pd10.pdf"}])
released = st.session_state.setdefault("released", [])
def fake_release(*args, **kwargs):
    released.append(args)
    return {"id": "att-new", "superseded_id": "att-old"}
ecnp.save_process_document_revision = fake_release
ecnp.AttachmentService = lambda repo: type("A", (), {"download": lambda self, row: b"%PDF"})()
st.page_link = lambda *a, **k: None
import core.record_history as rh
rh.record_history_panel = lambda *a, **k: None
page.current_profile = lambda: {"id": "u1", "role": "ADMIN"}
page.section_permissions = lambda *a, **k: {"can_view": True, "can_edit": True}
from core.part_process_service import StoredUpload
if not st.session_state.get("_queued"):
    st.session_state["_queued"] = True
    st.session_state["edit_part_id"] = "p1"
    ecnp.queue_pending_ecn("p1", {"kind": "PROCESS_DOC", "part_id": "p1", "document_type": "PROCESS_DRAWING", "document_number": "PD-10",
        "title": "Turning", "revision": "02", "revision_date": "2026-10-28", "current": repo.get("document_attachments", "att-old"),
        "route": {"id": "route-1", "part_id": "p1"}, "process_text": "Op 10 · Finish Machining"}, StoredUpload("pd10_r02.pdf", b"%PDF-2", "application/pdf"))
page.render_entry()
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception, app.exception
    assert app.session_state["released"] == []  # nothing released before the ECN
    labels = [b.label for b in app.button]
    assert "Submit ECN for Approval & Release Revision" in labels and "Cancel Replacement" in labels
    assert any("ECN REQUIRED" in str(m.value) for m in app.markdown)
    assert app.text_input(key="ecn_part_no_pending_PROCESS_DOC_p1").value == "7568"
    assert app.text_input(key="ecn_drawing_pending_PROCESS_DOC_p1").value == "PD-10"
    assert app.text_input(key="ecn_rev_from_pending_PROCESS_DOC_p1").value == "01"
    assert app.text_input(key="ecn_rev_to_pending_PROCESS_DOC_p1").value == "02"
    assert app.multiselect(key="ecn_reps_pending_PROCESS_DOC_p1").value == ["e1"]  # FSI Representative = signed-in user
    submit = next(b for b in app.button if b.label == "Submit ECN for Approval & Release Revision")
    submit.click().run(timeout=30)
    assert not app.exception, app.exception
    repo = app.session_state["r38_repo"]
    assert len(app.session_state["released"]) == 1
    ecn = [r for r in repo.tables["part_ecns"] if r.get("document_attachment_id") == "att-new"]
    assert len(ecn) == 1 and ecn[0]["status"] == "PENDING_APPROVAL" and ecn[0]["previous_attachment_id"] == "att-old"
    assert ecn[0]["ecn_number"].startswith("ECN-") and ecn[0]["drawing_number"] == "PD-10" and ecn[0]["new_revision"] == "02"
