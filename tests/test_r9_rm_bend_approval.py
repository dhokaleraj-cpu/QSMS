from copy import deepcopy
from io import BytesIO
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from test_r7_linked_raw_flow import service, po_payload, FINISHED, LINK
from test_r6_metlab_stage_isolation import service as inspection_service
from core.reporting import metlab_record_pdf_bytes, quality_record_excel_bytes


def rm_service(flow='FSI_RM_DIRECT_PRODUCTION'):
    svc=service()
    orders={'one':dict(id='one',part_id='finished',raw_material_detail_id='link',supply_flow=flow,required_rm_kg=100,rm_procurement_required=True),
            'two':dict(id='two',part_id='second',raw_material_detail_id='second-link',supply_flow='FSI_RM',required_rm_kg=200,rm_procurement_required=True)}
    svc.order=lambda oid: deepcopy(orders[oid])
    svc.totals=lambda oid: {'rm_ordered_kg':0,'forging_ordered_pcs':0}
    svc.repo.tables['parts'].append({**FINISHED,'id':'second','part_number':'SECOND','fsi_part_number':'SECOND'})
    svc.repo.tables['parts'][0]['fsi_part_number']='FIRST'
    svc.repo.tables['part_raw_material_details'].append({**LINK,'id':'second-link','part_id':'second'})
    svc.current_price=lambda *a,**k: 92
    svc._record_purchase_price=lambda **k: None
    svc.price_history_for_po=lambda *a,**k: []
    return svc

@pytest.mark.parametrize('flow',['FSI_RM','FSI_RM_DIRECT_PRODUCTION'])
def test_rm_po_allows_both_rm_flows_and_multiple_parts(flow):
    svc=rm_service(flow)
    saved=svc.create_purchase_order({**po_payload(),'po_type':'RAW_MATERIAL','customer_order_ids':['one','two'],'allocations':{'one':70,'two':160}})
    assert len(saved['items'])==2 and len(saved['stages'])==2
    assert {r['customer_order_id']:r['allocated_qty'] for r in svc.repo.tables['supply_purchase_order_sources']}=={'one':70,'two':160}
    assert {r['part_id'] for r in saved['items']}=={'finished','second'}
    assert sum(r['quantity'] for r in saved['items'])==230
    assert len(svc.repo.tables['supply_purchase_orders'])==1


def test_direct_forging_cannot_create_rm_po():
    svc=rm_service('DIRECT_FORGING')
    with pytest.raises(ValueError,match='Direct Forging'):
        svc.create_purchase_order({**po_payload(),'po_type':'RAW_MATERIAL','customer_order_ids':['one']})
    assert not svc.repo.writes


def test_forging_source_only_after_rm_dispatch_and_not_direct_production():
    svc=rm_service();svc.pending_direct_forging_orders=lambda: [];svc.forging_orders=lambda: []
    svc.rm_dispatches=lambda: []
    assert svc.pending_forging_po_sources()==[]
    svc.rm_dispatches=lambda:[{'id':'d1','customer_order_id':'one'},{'id':'d2','customer_order_id':'two'}]
    assert [r['_source_id'] for r in svc.pending_forging_po_sources()]==['d2']

@pytest.mark.parametrize('angle',[None,0,90,180,360])
def test_bend_details_survive_save_and_exports(angle):
    svc=inspection_service()
    saved=svc.save_metlab({'part_id':'40257237','layout_plan_id':'raw','inspection_scope':'RAW_MATERIAL_STAGE'},
        {'inspection_method':'BEND_TEST','bend_angle_degrees':angle,'conclusion_remark':'No crack','reference_documents':['BEND-SPEC'],'rows':[]})
    assert saved['results']['bend_angle_degrees']==angle
    assert saved['results']['inspection_method']=='BEND_TEST'
    assert saved['results']['conclusion_remark']=='No crack'
    assert saved['results']['reference_documents']==['BEND-SPEC']
    payload={'record':saved,'results':saved['results']}
    pdf=metlab_record_pdf_bytes(payload)
    text=' '.join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    assert 'Part Bend Angle:' in text and 'PART PHOTOGRAPHS' in text
    wb=load_workbook(BytesIO(quality_record_excel_bytes(payload,'METLAB')))
    rows=dict(wb['Report Summary'].iter_rows(min_row=2,values_only=True))
    assert rows['Part Bend Angle (degrees)']==angle

@pytest.mark.parametrize('angle',[float('nan'),float('inf'),-1,361,'invalid'])
def test_invalid_angle_cannot_write_report(angle):
    svc=inspection_service()
    with pytest.raises(ValueError,match='Bend angle'):
        svc.save_metlab({'inspection_scope':'FINAL_DISPATCH_STAGE'},{'bend_angle_degrees':angle})
    assert svc.repo.saved is None


def test_bend_photograph_is_embedded_in_pdf():
    from PIL import Image
    photo=BytesIO();Image.new('RGB',(120,80),'blue').save(photo,format='PNG')
    payload={'results':{'inspection_method':'BEND_TEST','bend_angle_degrees':120},'microstructure_images':[{'bytes':photo.getvalue(),'caption':'Part Bend Photograph'}]}
    reader=PdfReader(BytesIO(metlab_record_pdf_bytes(payload)))
    assert any(len(page.images)>0 for page in reader.pages)
    assert 'Part Bend Photograph' in ' '.join(page.extract_text() for page in reader.pages)


def test_email_confirmation_requires_review_again_after_recipient_change():
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_string('''
import streamlit as st
from core.notification_ui import notification_confirmation
class Notifier:
    def preview(self,*a,**k): return {'enabled':True,'recipient_email':'approver@example.com','cc_emails':[]}
pref=notification_confirmation(Notifier(),'METLAB_APPROVAL_PENDING',key='saved_draft',default_send=False)
st.button('Save and send for approval',disabled=pref['send'] and not pref['confirmed'])
''').run()
    assert not app.exception
    app.checkbox[0].check().run()
    assert app.button[-1].disabled
    app.session_state['saved_draft_confirmed_signature']='METLAB_APPROVAL_PENDING|approver@example.com'
    app.run()
    assert not app.button[-1].disabled
    app.text_input[0].set_value('different@example.com').run()
    assert app.button[-1].disabled
