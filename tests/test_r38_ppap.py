"""R38: PPAP document manager linked to APQP (module PPAP, routes ppap-home / ppap-package / ppap-templates):
tracker import, element templates, due digests (qcms-ppap-due-notifier), ZIP layout, customer submission, pages."""
from __future__ import annotations

import io
import re
import uuid
import zipfile
from datetime import date, timedelta
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
TRACKER = ROOT / "tests/fixtures/r38_ppap_tracker.xlsx"
MIGRATION = ROOT / "supabase/migrations/20261028120000_qcms_r38_ppap.sql"
FUNCTION = ROOT / "supabase/functions/qcms-ppap-due-notifier"


# ----------------------------------------------------------------------------- pure helpers
def test_status_and_date_mapping():
    from core.ppap_service import normalize_status, parse_date, status_from_label, status_label
    assert normalize_status("Completed") == "COMPLETED"
    assert normalize_status("NA") == normalize_status("N/A") == "NA"
    assert normalize_status("Pending ") == "NOT_STARTED" and status_label("NOT_STARTED") == "Pending"
    assert normalize_status("In progress") == "IN_PROGRESS" and status_label("IN_PROGRESS") == "In progress"
    assert normalize_status("") == "NOT_STARTED"
    assert normalize_status("Submitted") == "SUBMITTED" and normalize_status("rejected") == "REJECTED"
    assert [status_from_label(x) for x in ("Pending", "In progress", "Completed", "NA", "Rejected", "IN_PROGRESS")] == \
        ["NOT_STARTED", "IN_PROGRESS", "COMPLETED", "NA", "REJECTED", "IN_PROGRESS"]
    assert status_label("FULL_SUBMITTED") == "Full Submitted"
    assert parse_date("-") is None and parse_date("NA") is None and parse_date(None) is None
    assert parse_date("05-12-2025") == date(2025, 12, 5) and parse_date("2026-07-15") == date(2026, 7, 15)
    assert parse_date("15-Jul-2026") == date(2026, 7, 15)


def test_progress_excludes_na_and_due_logic():
    from core.ppap_service import due_items, element_state, is_due_item, package_progress, portfolio_kpis
    today = date(2026, 10, 28)
    rows = [
        {"id": "1", "package_id": "p", "element_no": 1, "status": "COMPLETED", "due_date": "2026-10-01"},
        {"id": "2", "package_id": "p", "element_no": 2, "status": "NA", "due_date": None},
        {"id": "3", "package_id": "p", "element_no": 3, "status": "IN_PROGRESS", "due_date": "2026-10-27"},   # overdue
        {"id": "4", "package_id": "p", "element_no": 4, "status": "NOT_STARTED", "due_date": "2026-10-31"},   # due in 3 days
        {"id": "5", "package_id": "p", "element_no": 5, "status": "SUBMITTED", "due_date": "2026-11-20"},     # later
        {"id": "6", "package_id": "p", "element_no": 6, "status": "NOT_STARTED", "due_date": None},           # not planned
        {"id": "7", "package_id": "p", "element_no": 7, "status": "REJECTED", "due_date": "2026-11-03"},      # due in 6 days
    ]
    progress = package_progress(rows)
    assert progress == {"total": 7, "na": 1, "applicable": 6, "completed": 1, "open": 5, "percent": round(100 / 6, 1)}
    assert package_progress([{"status": "NA"}])["percent"] == 100.0 and package_progress([])["percent"] == 0.0
    assert [element_state(r, today) for r in rows][:6] == ["completed", "na", "overdue", "due_soon", "submitted", "not_planned"]
    assert [r["id"] for r in due_items(rows, today)] == ["3", "4"]          # reminder: due within 3 days or overdue
    assert [r["id"] for r in due_items(rows, today, days_ahead=7)] == ["3", "4", "7"]
    assert not is_due_item({"status": "SUBMITTED", "due_date": "2026-10-01"}, today)        # with the customer, not the employee
    assert not is_due_item({"status": "COMPLETED", "due_date": "2020-01-01"}, today)
    kpi = portfolio_kpis([{"id": "p", "status": "IN_PROGRESS"}], rows, today)
    assert kpi["overdue"] == 1 and kpi["due_week"] == 2 and kpi["pending"] == 4
    assert portfolio_kpis([{"id": "p", "status": "APPROVED"}], rows, today)["overdue"] == 0


def test_due_digests_one_per_employee_with_manager_cc():
    from core.ppap_service import by_responsible, due_digests
    today = date(2026, 10, 28)
    employees = [
        {"id": "e1", "first_name": "Asha", "last_name": "Patil", "email": "asha@fsi.in", "department": "QUALITY", "reports_to_employee_id": "m1"},
        {"id": "e2", "first_name": "Ravi", "email": "ravi@fsi.in", "department": "engineering", "reports_to_employee_id": "m1"},
        {"id": "e3", "first_name": "No", "last_name": "Mail", "email": "", "department": "Production"},
        {"id": "m1", "first_name": "Manager", "email": "boss@fsi.in", "department": "Quality"},
    ]
    packages = [{"id": "p1", "package_no": "PPAP-69166-40286128", "customer_part_no": "40286128", "customer_ppap_no": "69166", "status": "IN_PROGRESS"},
                {"id": "p2", "package_no": "PPAP-OLD", "customer_part_no": "X", "status": "APPROVED"}]
    elements = [
        {"id": "a", "package_id": "p1", "element_no": 10, "element_name": "PFMEA", "status": "IN_PROGRESS", "due_date": "2026-10-25", "responsible_employee_id": "e1"},
        {"id": "b", "package_id": "p1", "element_no": 11, "element_name": "Process Control Plan", "status": "NOT_STARTED", "due_date": "2026-10-30", "responsible_employee_id": "e1"},
        {"id": "c", "package_id": "p1", "element_no": 7, "element_name": "Design Record", "status": "NOT_STARTED", "due_date": "2026-10-31", "responsible_employee_id": "e2"},
        {"id": "d", "package_id": "p1", "element_no": 8, "element_name": "Process Flow", "status": "NOT_STARTED", "due_date": "2026-11-15", "responsible_employee_id": "e2"},
        {"id": "e", "package_id": "p1", "element_no": 9, "element_name": "Floor Plan Layout", "status": "COMPLETED", "due_date": "2026-10-01", "responsible_employee_id": "e2"},
        {"id": "f", "package_id": "p1", "element_no": 12, "element_name": "PM Plan", "status": "NOT_STARTED", "due_date": "2026-10-20", "responsible_employee_id": "e3"},
        {"id": "g", "package_id": "p2", "element_no": 1, "element_name": "Supplier Timing", "status": "NOT_STARTED", "due_date": "2026-10-01", "responsible_employee_id": "e1"},
    ]
    digests = due_digests(packages, elements, employees, today)
    assert [d["email"] for d in digests] == ["asha@fsi.in", "ravi@fsi.in"]          # e3 has no e-mail; p2 approved
    asha, ravi = digests
    assert [r["element"] for r in asha["rows"]] == ["10 PFMEA", "11 Process Control Plan"]   # most overdue first
    assert asha["rows"][0] == {"package_id": "p1", "element_id": "a", "part": "40286128", "ppap_no": "69166", "package_no": "PPAP-69166-40286128",
                               "element": "10 PFMEA", "due_date": "2026-10-25", "days_overdue": 3, "status": "In progress"}
    assert asha["cc"] == ["boss@fsi.in"] and asha["overdue"] == 1                   # overdue -> reporting manager in CC
    assert asha["subject"] == "Quality · PPAP elements due · 2 item(s) · 28-10-2026"
    assert asha["dedupe_key"] == "PPAP_ELEMENT_DUE:2026-10-28:e1"
    assert ravi["cc"] == [] and [r["element"] for r in ravi["rows"]] == ["07 Design Record"]   # not overdue -> no CC
    assert ravi["subject"].startswith("Engineering · PPAP elements due")
    load = by_responsible(packages, elements, employees, today)
    assert [(r["Responsible"], r["Open"], r["Overdue"]) for r in load][:2] == [("Asha Patil", 2, 1), ("No Mail", 1, 1)]


def test_edge_function_mirrors_the_digest_rules():
    ts = (FUNCTION / "index.ts").read_text()
    assert 'const DUE_STATUSES = new Set(["NOT_STARTED", "IN_PROGRESS", "REJECTED"])' in ts
    assert 'const INACTIVE_PACKAGES = new Set(["APPROVED", "CANCELLED", "CLOSED"])' in ts
    assert "reports_to_employee_id" in ts and "cc.push(managerEmail)" in ts and "if (overdue)" in ts
    assert "`${label} · PPAP elements due · ${rows.length} item(s) · ${dmy(today)}`" in ts
    assert "`PPAP_ELEMENT_DUE:${today}:${employeeId}`" in ts and 'eq("dedupe_key", digest.dedupe_key)' in ts
    assert 'req.headers.get("X-QCMS-Scheduler")' in ts and "qcms_notification_scheduler_config" in ts and "sha256(schedulerToken)" in ts
    assert 'admin.rpc("qcms_ppap_due_notifications"' in ts and 'from "./reminder_kit.ts"' in ts
    kit = (FUNCTION / "reminder_kit.ts").read_text()
    assert kit == (ROOT / "supabase/functions/qcms-overdue-notifier/reminder_kit.ts").read_text()
    assert (FUNCTION / "deno.json").exists()


def test_revision_zip_storage_and_recipient_helpers():
    from core.ppap_service import (MAX_EMAIL_BYTES, build_zip, element_folder, foreign_party_recipients, next_revision, own_documents,
                                   package_number, split_emails, storage_path, submission_method)
    assert [next_revision(x) for x in ("", "A", "Z", "01", "B3")] == ["A", "B", "AA", "02", "B4"]
    assert element_folder(1, "Supplier Timing") == "01 Supplier Timing"
    assert element_folder(19, "Dimensional Results") == "19 Dimensional Results"
    assert element_folder(21, "Shipping / Packaging") == "21 Shipping - Packaging"
    assert storage_path("t1", "p1", {"element_no": 1, "element_name": "Supplier Timing"}, "ab_x.pdf") == "t1/ppap/p1/01_Supplier_Timing/ab_x.pdf"
    data = build_zip([
        {"element_no": 1, "element_name": "Supplier Timing", "file_name": "timing.pdf", "content": b"a"},
        {"element_no": 1, "element_name": "Supplier Timing", "file_name": "timing.pdf", "content": b"b"},
        {"element_no": 21, "element_name": "Shipping / Packaging", "file_name": "pack plan.xlsx", "content": b"c"},
    ], extra_files=[("00 PSW Index PPAP-1.pdf", b"%PDF")])
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert names == ["00 PSW Index PPAP-1.pdf", "01 Supplier Timing/timing.pdf", "01 Supplier Timing/timing (2).pdf",
                     "21 Shipping - Packaging/pack plan.xlsx"]
    assert submission_method([{"size_bytes": 1024}] * 3) == "EMAIL_ATTACHMENTS"
    assert MAX_EMAIL_BYTES == 10 * 1024 * 1024 and submission_method([{"size_bytes": MAX_EMAIL_BYTES + 1}]) == "SECURE_LINK"
    assert submission_method([{"size_bytes": 10}] * 25) == "SECURE_LINK"
    assert split_emails("a@x.com; bad, B@y.in a@x.com") == ["a@x.com", "B@y.in"]
    parties = [{"id": "c1", "email": "ppap@gm.com"}, {"id": "s1", "email": "sales@forge.in", "notification_emails": "qa@forge.in"}]
    assert foreign_party_recipients(["ppap@gm.com", "qa@forge.in", "new@gm.com"], parties, "c1") == ["qa@forge.in"]
    assert package_number("69166", "40286128 Tesla", ["PPAP-69166-40286128"]) == "PPAP-69166-40286128-TESLA"
    assert package_number("69166", "40286128", ["PPAP-69166-40286128"]) == "PPAP-69166-40286128-2"
    elements = [{"id": "e1", "package_id": "p1"}]
    docs = [{"id": "d1", "package_id": "p1", "package_element_id": "e1"}, {"id": "d2", "package_id": "p9", "package_element_id": "e9"},
            {"id": "d3", "package_id": "p9", "package_element_id": "e1"}]
    assert [d["id"] for d in own_documents(docs, elements)] == ["d1"]     # never another package's (customer's) file


def test_import_mapping_from_real_tracker_file():
    from core.ppap_service import note_due_date, note_targets, parse_tracker
    parsed = parse_tracker(TRACKER.read_bytes())
    assert parsed.sheet == "Sheet1" and not parsed.warnings
    assert [c.part_text for c in parsed.columns] == ["40256626 TRMF", "40257237", "40229265", "40102994", "40256626 CKMF",
                                                      "40286128 Tesla", "40269907", "40280398", "40280399", "40286128"]
    assert [c.ppap_no for c in parsed.columns] == ["", "", "", "", "66266", "69166", "68093", "68133", "69011", "69166"]
    assert len(parsed.element_names) == 32 and parsed.element_names[1] == "Supplier Timing" and parsed.element_names[32] == "IMDS"
    assert parsed.element_names[19] == "Dimensional Results" and parsed.element_names[21] == "Shipping / Packaging"
    first, last = parsed.columns[0], parsed.columns[-1]
    assert first.part_token == "40256626" and all(len(c.elements) == 32 for c in parsed.columns)
    assert sum(len(c.elements) for c in parsed.columns) == 320                                       # 10 parts x 32 elements
    assert first.elements[0] == {"element_no": 1, "element_name": "Supplier Timing", "status": "COMPLETED", "raw_status": "Completed", "due_date": "2025-12-05"}
    assert first.elements[4]["status"] == "NA" and first.elements[4]["due_date"] is None          # DFMEA NA / NA
    assert parsed.columns[1].elements[0]["due_date"] is None                                          # '-' due date
    statuses = [e["status"] for e in last.elements]
    assert statuses.count("NOT_STARTED") == 6 and statuses.count("NA") == 4 and statuses.count("COMPLETED") == 22   # Pending = NOT_STARTED
    assert {e["raw_status"] for e in last.elements if e["status"] == "NOT_STARTED"} == {"Pending"}
    assert last.elements[31]["due_date"] == "2026-05-26"                                             # IMDS W39
    assert len(parsed.notes) == 2 and parsed.notes[0].startswith("Part No. 6128")
    assert note_due_date(parsed.notes[0]) == date(2026, 7, 15)
    packages = [{"customer_part_no": c.part_text, "customer_ppap_no": c.ppap_no} for c in parsed.columns]
    targets = note_targets(parsed.notes, packages)
    assert targets[0][1] == [5, 9] and targets[1][1] == [5, 9]   # 6128 -> both 40286128 columns; follow-up note continues


def test_template_round_trip_and_cover_sheet():
    from openpyxl import load_workbook
    from core.ppap_service import package_cover_xlsx, parse_tracker, tracker_workbook_bytes
    blank = parse_tracker(tracker_workbook_bytes([]))
    assert blank.columns == [] and len(blank.element_names) == 32 and blank.notes == []
    data = tracker_workbook_bytes([], [{"id": "p1", "customer_part_no": "71.784.3", "customer_ppap_no": "1234"}],
                                  {"p1": [{"element_no": 10, "status": "IN_PROGRESS", "due_date": "2026-11-02"},
                                          {"element_no": 5, "status": "NA"}, {"element_no": 7, "status": "NOT_STARTED"}]},
                                  ["Part 71.784.3 waiting for gauge"])
    parsed = parse_tracker(data)
    column = parsed.columns[0]
    assert (column.part_text, column.ppap_no) == ("71.784.3", "1234")
    by_no = {e["element_no"]: e for e in column.elements}
    assert by_no[10]["status"] == "IN_PROGRESS" and by_no[10]["due_date"] == "2026-11-02" and by_no[5]["status"] == "NA"
    assert by_no[7]["raw_status"] == "Pending" and by_no[7]["status"] == "NOT_STARTED"
    assert by_no[1]["status"] == "NOT_STARTED" and parsed.notes == ["Part 71.784.3 waiting for gauge"]
    package = {"id": "p1", "package_no": "PPAP-1", "customer_part_no": "71.784.3", "customer_ppap_no": "1234", "full_submit_date": "2026-11-20"}
    elements = [{"id": "e1", "package_id": "p1", "element_no": 1, "element_name": "Supplier Timing", "status": "COMPLETED", "responsible_employee_id": "x"},
                {"id": "e2", "package_id": "p1", "element_no": 19, "element_name": "Dimensional Results", "status": "IN_PROGRESS", "due_date": "2026-11-05"}]
    cover = load_workbook(io.BytesIO(package_cover_xlsx(package, elements, [{"package_element_id": "e2", "file_name": "dim.pdf", "revision": "B", "is_latest": True}],
                                                         employee_labels={"x": "Asha Patil"})))
    assert cover.sheetnames[:2] == ["Package register", "Sheet1"]
    values = [[c.value for c in row] for row in cover["Package register"].iter_rows()]
    assert ["Full Submit", "20-11-2026"] == values[7][:2]
    assert any(r[:3] == [19, "Dimensional Results", "In progress"] and r[7] == "dim.pdf (Rev. B)" for r in values)
    assert any(r[:3] == [1, "Supplier Timing", "Completed"] and r[5] == "Asha Patil" for r in values)
    assert len(parse_tracker(package_cover_xlsx(package, elements)).columns) == 1    # Sheet1 = customer tracker layout


# ----------------------------------------------------------------------------- service with an in-memory repository
class FakeRepo:
    tenant_id = "00000000-0000-0000-0000-000000000001"

    def __init__(self, **tables):
        self.db = {k: [dict(r) for r in v] for k, v in tables.items()}
        self.rpcs: list[tuple[str, dict]] = []

    def select(self, table, *, eq=None, in_=None, order_by=None, desc=False, limit=None, **_):
        rows = [dict(r) for r in self.db.get(table, [])]
        for k, v in (eq or {}).items():
            rows = [r for r in rows if str(r.get(k)) == str(v)]
        for k, values in (in_ or {}).items():
            rows = [r for r in rows if str(r.get(k)) in {str(x) for x in values}]
        if order_by:
            rows.sort(key=lambda r: str(r.get(order_by) or ""), reverse=desc)
        return rows[:limit] if limit else rows

    def get(self, table, record_id):
        return next((r for r in self.select(table, eq={"id": record_id})), None)

    def insert(self, table, payload):
        row = {"id": str(uuid.uuid4()), **{k: (v.isoformat() if isinstance(v, date) else v) for k, v in payload.items()}}
        self.db.setdefault(table, []).append(row)
        return dict(row)

    def update(self, table, record_id, payload):
        for row in self.db.get(table, []):
            if str(row["id"]) == str(record_id):
                row.update({k: (v.isoformat() if isinstance(v, date) else v) for k, v in payload.items()})
                return dict(row)
        raise KeyError(record_id)

    def delete(self, table, record_id):
        self.db[table] = [r for r in self.db.get(table, []) if str(r["id"]) != str(record_id)]

    def rpc(self, name, params=None):
        self.rpcs.append((name, dict(params or {})))
        return None


def test_service_create_save_and_progress():
    from core.ppap_service import PpapService
    repo = FakeRepo()
    service = PpapService(repo)
    package = service.create_package({"customer_part_no": "40286128", "customer_ppap_no": "69166", "customer_id": "c1"},
                                     default_due=date(2026, 11, 30), responsible_employee_id="e1")
    assert package["package_no"] == "PPAP-69166-40286128"
    master = repo.db["ppap_element_master"]
    assert len(master) == 32 and {m["element_no"]: m["default_department"] for m in master}[21] == "Logistics"
    elements = service.elements(package["id"])
    assert len(elements) == 32
    by_no = {e["element_no"]: e for e in elements}
    assert by_no[5]["status"] == "NA" and by_no[5]["due_date"] is None and by_no[5]["responsible_employee_id"] is None   # DFMEA NA
    assert by_no[1]["status"] == "NOT_STARTED" and by_no[1]["due_date"] == "2026-11-30" and by_no[1]["responsible_employee_id"] == "e1"
    assert by_no[1]["responsible_department"] == "NPD" and by_no[10]["responsible_department"] == "Quality"
    rows = [{**e, "status": "Completed"} if e["element_no"] in (1, 2) else e for e in elements]     # grid labels are accepted
    rows[2] = {**rows[2], "responsible_employee_id": "e2"}
    assigned = service.save_elements(package["id"], rows)
    assert [a["employee_id"] for a in assigned] == ["e2"]
    saved = {e["element_no"]: e for e in service.elements(package["id"])}
    assert saved[1]["status"] == "COMPLETED" and saved[1]["completed_date"] == date.today().isoformat()
    assert repo.get("ppap_packages", package["id"])["completion_percent"] == round(2 * 100 / 29, 1)   # 3 default NA elements
    assert service.notify_assignments(package, assigned) == 1 and repo.rpcs[0][0] == "qcms_notify"
    assert repo.rpcs[0][1]["p_web_route"] == "ppap-package" and repo.rpcs[0][1]["p_module_key"] == "PPAP"
    assert repo.rpcs[0][1]["p_recipient_employee_id"] == "e2"


def test_customer_template_is_used_for_new_packages_and_can_be_reordered():
    from core.ppap_service import PpapService
    repo = FakeRepo()
    service = PpapService(repo)
    assert service.copy_template_to_customer("cust-gm") == 32
    with pytest.raises(ValueError):
        service.copy_template_to_customer("cust-gm")
    rows = service.customer_template("cust-gm")
    # rename 01, move 32 IMDS to the top, drop 14 Vision Systems, add a customer element 33 with 5 lead days
    edited = [{**r, "element_name": "Supplier Timing Plan"} if r["element_no"] == 1 else dict(r) for r in rows if r["element_no"] != 14]
    edited = [{**r, "sort_order": 0} if r["element_no"] == 32 else r for r in edited]
    edited.append({"element_no": 33, "element_name": "GM BIQS form", "default_required": True, "default_lead_days": 5,
                   "default_department": "Quality", "sort_order": 99})
    counts = service.save_template(edited, customer_id="cust-gm")
    assert counts == {"saved": 31, "added": 1, "removed": 1}
    assert len(repo.db["ppap_element_master"]) == 32                       # default template untouched
    package = service.create_package({"customer_part_no": "P-GM", "customer_id": "cust-gm"}, default_due=date(2026, 12, 1))
    elements = service.elements(package["id"])
    assert [e["element_no"] for e in elements][:2] == [32, 1] and elements[-1]["element_no"] == 33
    assert 14 not in {e["element_no"] for e in elements} and elements[1]["element_name"] == "Supplier Timing Plan"
    assert elements[-1]["due_date"] == (date.today() + timedelta(days=5)).isoformat()
    other = service.create_package({"customer_part_no": "P-OTHER", "customer_id": "cust-ford"})
    assert len(service.elements(other["id"])) == 32                          # other customers keep the default template
    with pytest.raises(ValueError):
        service.save_template([{"element_no": 1, "element_name": "A"}, {"element_no": 1, "element_name": "B"}])


def test_service_import_tracker_creates_packages_elements_and_notes():
    from core.ppap_service import PpapService, parse_tracker
    parts = [{"id": "part-1", "part_number": "40286128", "fsi_part_number": "6128/2702", "customer_id": "cust-1"},
             {"id": "part-2", "part_number": "40256626", "fsi_part_number": "6626", "customer_id": "cust-1"}]
    repo = FakeRepo()
    service = PpapService(repo)
    parsed = parse_tracker(TRACKER.read_bytes())
    summary = service.import_tracker(parsed, parts=parts, owner_employee_id="owner-1", author="Tester")
    assert summary["created"] == 10 and summary["updated"] == 0 and summary["elements"] == 320
    assert len(repo.db["ppap_package_elements"]) == 320
    packages = repo.db["ppap_packages"]
    tesla = next(p for p in packages if p["customer_part_no"] == "40286128 Tesla")
    assert tesla["part_id"] == "part-1" and tesla["customer_id"] == "cust-1" and tesla["customer_ppap_no"] == "69166"
    assert tesla["customer_due_date"] == "2026-07-15"                      # from the status note
    notes = [n for n in repo.db["ppap_status_notes"] if n["package_id"] == tesla["id"]]
    assert len(notes) == 2 and {n["source"] for n in notes} == {"IMPORT"}
    plain = next(p for p in packages if p["customer_part_no"] == "40286128")
    assert repo.get("ppap_packages", plain["id"])["completion_percent"] == round(22 * 100 / 28, 1)
    elements = service.elements(plain["id"])
    completed = next(e for e in elements if e["element_no"] == 1)
    assert completed["status"] == "COMPLETED" and completed["completed_date"] == completed["due_date"] == "2026-06-04"
    assert sum(e["status"] == "NOT_STARTED" for e in elements) == 6 and {e["responsible_department"] for e in elements} >= {"Quality", "Engineering"}
    # re-import updates the same packages (no duplicates) and does not duplicate notes
    again = service.import_tracker(parsed, parts=parts)
    assert again["created"] == 0 and again["updated"] == 10 and again["notes"] == 0 and len(repo.db["ppap_packages"]) == 10
    assert len(repo.db["ppap_package_elements"]) == 320


def test_upload_document_path_revision_and_status(monkeypatch):
    from core import ppap_service
    repo = FakeRepo()
    service = ppap_service.PpapService(repo)
    package = service.create_package({"customer_part_no": "P1", "customer_id": "c1"})
    element = next(e for e in service.elements(package["id"]) if e["element_no"] == 19)
    uploaded = []

    class Storage:
        def upload(self, path, data, options):
            uploaded.append(path)

    monkeypatch.setattr(ppap_service.PpapService, "_storage", lambda self: Storage())
    first = service.upload_document(package, element, file_name="dim report.pdf", content=b"%PDF-1", uploader_name="QA")
    second = service.upload_document(package, element, file_name="dim report.pdf", content=b"%PDF-2", uploader_name="QA")
    assert uploaded[0].startswith(f"{repo.tenant_id}/ppap/{package['id']}/19_Dimensional_Results/") and uploaded[0].endswith("_dim_report.pdf")
    assert (first["revision"], second["revision"]) == ("A", "B")
    docs = {d["id"]: d for d in repo.db["ppap_element_documents"]}
    assert docs[first["id"]]["is_latest"] is False and docs[second["id"]]["is_latest"] is True
    assert repo.get("ppap_package_elements", element["id"])["status"] == "IN_PROGRESS"     # Pending -> In progress on upload


def test_send_package_small_uses_attachments_and_logs(monkeypatch):
    from core import ppap_service
    from core.notification_service import NotificationService
    repo = FakeRepo()
    service = ppap_service.PpapService(repo)
    package = service.create_package({"customer_part_no": "P1", "customer_id": "c1", "status": "IN_PROGRESS"})
    element = service.elements(package["id"])[0]
    doc = repo.insert("ppap_element_documents", {"package_id": package["id"], "package_element_id": element["id"], "revision": "A",
                                                 "file_name": "timing.pdf", "object_path": "t/ppap/x/01/timing.pdf", "size_bytes": 2048,
                                                 "is_latest": True, "status": "ACTIVE"})
    captured = {}

    def enqueue(self, event_key, **kwargs):
        captured.update(kwargs, event=event_key, manifest=self.attachment_manifest(None, None))
        return {"id": "outbox-1"}

    monkeypatch.setattr(NotificationService, "enqueue", enqueue)
    monkeypatch.setattr(NotificationService, "dispatch", lambda self, rows: {"sent": 1})
    monkeypatch.setattr(NotificationService, "_persist_generated_pdf", lambda self, name, content: None)
    result = service.send_to_customer(package, [doc], service.elements(package["id"]), to=["ppap@customer.com"], cc=["qa@fsi.in"],
                                      sender="Tester", customer_name="Customer")
    assert result["method"] == "EMAIL_ATTACHMENTS" and result["status"] == "SENT"
    assert captured["event"] == "PPAP_PACKAGE_SUBMISSION" and captured["recipient_email"] == "ppap@customer.com"
    assert captured["context"]["to_emails"] == ["ppap@customer.com"] and not any(k.startswith("customer") for k in captured["context"])
    assert [m["object_path"] for m in captured["manifest"]] == ["t/ppap/x/01/timing.pdf"]
    log = repo.db["ppap_customer_submissions"][0]
    assert log["method"] == "EMAIL_ATTACHMENTS" and log["status"] == "SENT" and log["documents"][0]["revision"] == "A"
    assert log["outbox_id"] == "outbox-1" and repo.get("ppap_packages", package["id"])["status"] == "SUBMITTED"
    # a document of another package (another customer's PPAP) is refused
    with pytest.raises(ValueError):
        service.send_to_customer(package, [{**doc, "package_id": "someone-else"}], service.elements(package["id"]), to=["ppap@customer.com"])


def test_send_package_large_uses_signed_link(monkeypatch):
    from core import ppap_service
    from core.notification_service import NotificationService
    repo = FakeRepo()
    service = ppap_service.PpapService(repo)
    package = service.create_package({"customer_part_no": "P2", "customer_id": "c1"})
    elements = service.elements(package["id"])
    docs = [repo.insert("ppap_element_documents", {"package_id": package["id"], "package_element_id": elements[i]["id"], "revision": "A",
                                                   "file_name": f"f{i}.pdf", "object_path": f"t/{i}.pdf", "size_bytes": 6 * 1024 * 1024,
                                                   "is_latest": True}) for i in range(2)]

    class Storage:
        uploaded: list = []

        def upload(self, path, data, options):
            Storage.uploaded.append((path, data))

        def download(self, path):
            return b"x"

        def create_signed_url(self, path, seconds):
            assert seconds == 7 * 24 * 3600
            return {"signedURL": f"https://files/{path}?token=1"}

    monkeypatch.setattr(ppap_service.PpapService, "_storage", lambda self: Storage())
    body = {}
    monkeypatch.setattr(NotificationService, "enqueue", lambda self, ev, **kw: body.update(kw, manifest=self.attachment_manifest()) or {"id": "o2"})
    monkeypatch.setattr(NotificationService, "dispatch", lambda self, rows: {"sent": 0, "failed": 0})
    result = service.send_to_customer(package, docs, elements, to=["ppap@customer.com"])
    assert result["method"] == "SECURE_LINK" and result["link"].startswith("https://files/") and result["status"] == "QUEUED"
    assert body["manifest"] == [] and result["link"] in body["body_text"]
    zip_path, zip_bytes = Storage.uploaded[0]
    assert zip_path.endswith(".zip") and f"/ppap/{package['id']}/submissions/" in zip_path
    assert zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist() == ["01 Supplier Timing/f0.pdf", "02 Fixtures - Tooling - Gaging/f1.pdf"]
    assert repo.db["ppap_customer_submissions"][0]["method"] == "SECURE_LINK" and repo.db["ppap_customer_submissions"][0]["link_expires_at"]


def test_package_zip_with_cover_sheets(monkeypatch):
    from core import ppap_service
    repo = FakeRepo()
    service = ppap_service.PpapService(repo)
    package = service.create_package({"customer_part_no": "P3", "customer_id": "c1", "package_no": "PPAP-P3"})
    elements = service.elements(package["id"])
    docs = [{"id": "d1", "package_id": package["id"], "package_element_id": elements[18]["id"], "file_name": "dim.pdf", "object_path": "x"},
            {"id": "d2", "package_id": package["id"], "package_element_id": elements[20]["id"], "file_name": "pack.xlsx", "object_path": "y"}]
    monkeypatch.setattr(ppap_service.PpapService, "document_bytes", lambda self, doc: b"data")
    data = service.package_zip(package, docs, elements, index_pdf=b"%PDF", cover_xlsx=b"xlsx")
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["00 PSW Index PPAP-P3.pdf", "00 PPAP Tracker PPAP-P3.xlsx",
                                                            "19 Dimensional Results/dim.pdf", "21 Shipping - Packaging/pack.xlsx"]


def test_package_index_pdf_builds():
    from core.ppap_service import package_index_pdf
    pdf = package_index_pdf({"package_no": "PPAP-1", "psw_status": "NOT_SUBMITTED", "submission_level": 3, "interim_submit_date": "2026-10-01"},
                            [{"id": "e1", "element_no": 1, "element_name": "Supplier Timing", "status": "COMPLETED"}],
                            [{"package_element_id": "e1", "file_name": "t.pdf", "revision": "A", "size_bytes": 10}])
    assert pdf[:4] == b"%PDF" and len(pdf) > 1500


def test_permissions_fall_back_to_npd_apqp(monkeypatch):
    from app_pages import ppap
    calls = []
    monkeypatch.setattr(ppap, "module_permissions_with_source", lambda profile, key: ({"can_view": True, "can_create": False, "can_edit": False}, "LEGACY ROLE · X"))
    monkeypatch.setattr(ppap, "module_permissions", lambda profile, key: calls.append(key) or {"can_view": True, "can_create": True, "can_edit": True})
    assert ppap.ppap_permissions({"id": "u"})["can_create"] is True and calls == ["NPD_APQP"]
    monkeypatch.setattr(ppap, "module_permissions_with_source", lambda profile, key: ({"can_view": True, "can_create": False}, "USER OVERRIDE"))
    assert ppap.ppap_permissions({"id": "u"})["can_create"] is False
    from core.access import MODULES, _legacy_role_permissions
    assert ("PPAP", "PPAP") in MODULES
    assert _legacy_role_permissions({"role": "BUSINESS_DEVELOPMENT"}, "PPAP")["can_edit"] is True


def test_matrix_filter_modes():
    from app_pages.ppap import ppap_matrix_filter
    today = date(2026, 10, 28)
    data = {"p": [{"element_no": 1, "status": "COMPLETED", "responsible_employee_id": "a"},
                  {"element_no": 2, "status": "NOT_STARTED", "due_date": "2026-10-01", "responsible_employee_id": "a"},
                  {"element_no": 3, "status": "IN_PROGRESS", "due_date": "2026-11-30", "responsible_employee_id": "b"}]}
    assert [e["element_no"] for e in ppap_matrix_filter(data, mode="Open only", today=today)["p"]] == [2, 3]
    assert [e["element_no"] for e in ppap_matrix_filter(data, mode="Overdue only", today=today)["p"]] == [2]
    assert [e["element_no"] for e in ppap_matrix_filter(data, responsible_id="b", today=today)["p"]] == [3]


# ----------------------------------------------------------------------------- pages / AppTest (in-memory preview repository)
def _script(call: str) -> str:
    return f'''
import streamlit as st
from core.auth import PREVIEW_PROFILE
st.session_state["_qsms_preview"] = True
st.session_state["profile"] = PREVIEW_PROFILE.copy()
st.data_editor = lambda data, **kwargs: data  # pyarrow-free test runtime
from app_pages import ppap
ppap.{call}()
'''


APQP_SCRIPT = '''
import streamlit as st
from core.auth import PREVIEW_PROFILE
st.session_state["_qsms_preview"] = True
st.session_state["profile"] = PREVIEW_PROFILE.copy()
st.data_editor = lambda data, **kwargs: data
from app_pages import npd_apqp
npd_apqp.render_apqp()
'''


def _button(app, label):
    return next(b for b in app.button if b.label == label)


def test_ppap_home_and_templates_render():
    from streamlit.testing.v1 import AppTest
    home = AppTest.from_string(_script("render_ppap_home")).run(timeout=30)
    assert not home.exception, home.exception
    assert [t.label for t in home.tabs][:3] == ["PPAP Packages", "Tracker Matrix", "Import Tracker (Excel)"]
    assert "New PPAP Package" in [b.label for b in home.button]
    templates = AppTest.from_string(_script("render_ppap_templates")).run(timeout=30)
    assert not templates.exception, templates.exception
    assert len(templates.session_state["_qsms_demo_store"].get("ppap_element_master", [])) == 32
    assert {"Save Element Template", "Create"} <= {b.label for b in templates.button}


def test_ppap_package_page_creates_package_and_shows_sections():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string(_script("render_ppap_package")).run(timeout=30)
    assert not app.exception, app.exception
    store = app.session_state["_qsms_demo_store"]
    cust_id = next(str(p["id"]) for p in store["parties"] if "CUSTOMER" in (p.get("party_types") or []))
    next(s for s in app.selectbox if s.label == "Customer").set_value(cust_id).run(timeout=30)
    next(t for t in app.text_input if t.label == "Customer Part No.").input("40286128")
    next(t for t in app.text_input if t.label == "Customer PPAP No.").input("69166")
    _button(app, "Save PPAP Package").click().run(timeout=30)
    assert not app.exception, app.exception
    store = app.session_state["_qsms_demo_store"]
    packages = store.get("ppap_packages", [])
    assert len(packages) == 1 and packages[0]["package_no"] == "PPAP-69166-40286128"
    assert len([e for e in store["ppap_package_elements"] if e["package_id"] == packages[0]["id"]]) == 32
    assert app.selectbox(key="ppap_package_selector").value == str(packages[0]["id"])
    labels = [b.label for b in app.button]
    for label in ("Update PPAP Package", "Save PPAP Elements", "Add status note", "Download all (ZIP)", "Download selected (ZIP)",
                  "Generate PSW / index PDF", "Generate tracker Excel", "Send PPAP package to customer", "Delete PPAP Package"):
        assert label in labels, label
    next(t for t in app.text_area if t.label == "New status note").input("Waiting for customer to enable Final Submit")
    _button(app, "Add status note").click().run(timeout=30)
    assert not app.exception
    assert app.session_state["_qsms_demo_store"]["ppap_status_notes"][0]["note"] == "Waiting for customer to enable Final Submit"
    _button(app, "Generate PSW / index PDF").click().run(timeout=30)
    assert not app.exception, app.exception
    assert any(getattr(d, "label", "") == "Download PSW / index PDF" for d in app.get("download_button"))
    _button(app, "Generate tracker Excel").click().run(timeout=30)
    assert not app.exception, app.exception
    assert any(getattr(d, "label", "") == "Download tracker Excel" for d in app.get("download_button"))


def test_apqp_page_links_and_creates_ppap_package():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string(APQP_SCRIPT).run(timeout=30)
    assert not app.exception
    app.text_input[1].input("APQP-PPAP-1")
    _button(app, "Save APQP Project").click().run(timeout=30)
    assert not app.exception, app.exception
    assert "Create PPAP package from this APQP project" in [b.label for b in app.button]
    _button(app, "Create PPAP package from this APQP project").click().run(timeout=30)
    assert not app.exception, app.exception
    store = app.session_state["_qsms_demo_store"]
    project = next(p for p in store["ppap_projects"] if p["project_code"] == "APQP-PPAP-1")
    package = next(p for p in store["ppap_packages"])
    assert str(package["apqp_project_id"]) == str(project["id"]) and package["submission_level"] == 3
    assert any(b.label == "Open PPAP" for b in app.button)


# ----------------------------------------------------------------------------- wiring / migration safety
def test_routes_permissions_and_menu():
    app = (ROOT / "streamlit_app.py").read_text()
    for route, fn, title in (("ppap-home", "render_ppap_home", "PPAP"), ("ppap-package", "render_ppap_package", "PPAP Package"),
                             ("ppap-templates", "render_ppap_templates", "PPAP Element Templates")):
        assert f'("{route}", st.Page(ppap_page.{fn}, title="{title}"' in app
        assert f'"{route}":"PPAP"' in app and f'"{route}": "Engineering"' in app and f'"{title}": "{route}"' in app
        assert f'("{route}",' in app.split('"Engineering": (', 1)[1].split("),\n    \"", 1)[0]   # Engineering submenu
    assert '(PAGE_BY_PATH["ppap-home"], "PPAP", "Engineering"' in app and '"ppap"' not in app
    assert '"ppap-home", "ppap-package", "ppap-templates"' in (ROOT / "scripts/verify_phase1.py").read_text()
    page = (ROOT / "app_pages/ppap.py").read_text()
    assert "st.dataframe(" not in page and "db_message(exc)" in page and "module_permissions_with_source(profile, MODULE_KEY)" in page
    apqp = (ROOT / "app_pages/npd_apqp.py").read_text()
    assert "_render_apqp_ppap_link(repo, project, perms, parts)" in apqp and "pages[ROUTE_PACKAGE]" in apqp


def test_migration_is_additive_registered_and_complete():
    sql = MIGRATION.read_text()
    low = sql.lower()
    assert not re.search(r"\b(delete|update|drop|truncate)\b", low), "keep MCP-confirmation keywords out of the migration"
    tables = ("ppap_element_master", "ppap_customer_template_elements", "ppap_packages", "ppap_package_elements", "ppap_element_documents",
              "ppap_status_notes", "ppap_customer_submissions")
    for table in tables:
        assert f"create table if not exists public.{table}" in low
        assert f"'{table}'" in low  # RLS / policy loop + search map
    assert low.count("tenant_id uuid not null default public.current_tenant_id()") == 7
    assert "enable row level security" in low and "revoke all on public.%i from anon" in low and "pg_policies" in low
    assert "alter policy tenant_write on public.%i" in low and "public.qcms_ppap_can(''archive'')" in low
    assert "return public.qcms_effective_module_permission('ppap', v_perm)" in low and "qcms_effective_module_permission('npd_apqp', v_perm)" in low
    for column in ("default_department", "interim_submit_date", "full_submit_date", "responsible_department", "sort_order"):
        assert f"add column if not exists {column}" in low
    assert "on conflict (tenant_id, element_no) do nothing" in low and "(32, 'imds', true, 'quality')" in low
    assert "'ppap_document_due'" in low and "on conflict (tenant_id, schedule_key) do nothing" in low
    assert "select public.qcms_search_table_map()" in low
    assert "'ppap_packages', jsonb_build_array('ppap', 'ppap package', 'ppap-package')" in low
    assert "'ppap_customer_template_elements', jsonb_build_array('ppap', 'ppap customer template element', 'ppap-templates')" in low
    assert "grant execute on function public.qcms_ppap_due_notifications(uuid, date, integer) to service_role" in low
    assert "e.status in ('not_started', 'in_progress', 'rejected')" in low
    cron = low.split("-- 8. r38b", 1)[1]
    assert "exists (select 1 from pg_extension where extname = 'pg_cron')" in cron and "qcms_notification_scheduler_token" in cron
    assert "cron.schedule('qcms-ppap-due-notifier-daily', '20 2 * * *'" in cron and "/functions/v1/qcms-ppap-due-notifier" in cron
    files = (ROOT / "supabase/auto_apply.json").read_text()
    assert "supabase/migrations/20261028120000_qcms_r38_ppap.sql" in files and "20261028110000_qcms_r38_ppap" not in files
    assert not (ROOT / "supabase/migrations/20261028110000_qcms_r38_ppap.sql").exists()
