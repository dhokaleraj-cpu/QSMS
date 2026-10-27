"""R37b · price-history save fix, Purchase Order amendment, PO print (amendment + all price revisions)."""
from __future__ import annotations

import re
from copy import deepcopy
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase" / "migrations" / "20261027100000_qcms_r37b_po_amendment.sql"


# ----------------------------------------------------------------------------------------------
# Fake repository that enforces the live BEFORE INSERT/UPDATE guard of part_supplier_price_history
# (qcms_guard_part_supplier_price_history: auto-close on INSERT, non-overlap, unique start).
# ----------------------------------------------------------------------------------------------
def _d(v):
    return date.fromisoformat(str(v)[:10]) if v else None


class GuardRepo:
    def __init__(self, tables=None):
        self.tables = deepcopy(tables or {})
        self.writes = []
        self._n = 0

    def select(self, table, eq=None, **kwargs):
        rows = [deepcopy(r) for r in self.tables.get(table, []) if all(r.get(k) == v for k, v in (eq or {}).items())]
        if kwargs.get("order_by"):
            key = kwargs["order_by"]
            rows.sort(key=lambda r: (r.get(key) is None, str(r.get(key) or "")), reverse=bool(kwargs.get("desc")))
        return rows

    def find_one(self, table, eq):
        rows = self.select(table, eq=eq)
        return rows[0] if rows else None

    def get(self, table, key):
        return next((deepcopy(r) for r in self.tables.get(table, []) if str(r.get("id")) == str(key)), None)

    def _guard(self, row, *, inserting):
        if row.get("end_date") and _d(row["end_date"]) < _d(row["start_date"]):
            raise RuntimeError("Price End Date is before Start Date")
        rows = self.tables.setdefault("part_supplier_price_history", [])
        uom = str(row.get("uom") or "KGS").upper()
        same = [h for h in rows if h.get("id") != row.get("id") and h["part_id"] == row["part_id"] and h["supplier_id"] == row["supplier_id"]
                and str(h.get("uom") or "KGS").upper() == uom]
        if any(_d(h["start_date"]) == _d(row["start_date"]) for h in same):
            raise RuntimeError('duplicate key value violates unique constraint "uq_qcms_part_supplier_price_start"')
        same = [h for h in same if h.get("raw_material_detail_id") == row.get("raw_material_detail_id")]
        if inserting and str(row.get("status") or "ACTIVE") == "ACTIVE":
            for h in same:
                if h.get("status", "ACTIVE") == "ACTIVE" and _d(h["start_date"]) < _d(row["start_date"]) and (not h.get("end_date") or _d(h["end_date"]) >= _d(row["start_date"])):
                    h["end_date"] = (_d(row["start_date"]) - timedelta(days=1)).isoformat()
        if str(row.get("status") or "ACTIVE") == "ACTIVE":
            ns, ne = _d(row["start_date"]), _d(row.get("end_date")) or date.max
            for h in same:
                if h.get("status", "ACTIVE") != "ACTIVE":
                    continue
                hs, he = _d(h["start_date"]), _d(h.get("end_date")) or date.max
                if hs <= ne and ns <= he:
                    raise RuntimeError(f"Price period {ns:%d-%m-%Y} to open overlaps the existing period {hs:%d-%m-%Y} (price {h['price']}).")

    def insert(self, table, payload):
        self._n += 1
        row = {"id": f"{table}-{self._n}", **deepcopy(dict(payload))}
        if table == "part_supplier_price_history":
            row.setdefault("status", "ACTIVE")
            self._guard(row, inserting=True)
        self.tables.setdefault(table, []).append(row)
        self.writes.append(("insert", table, row["id"]))
        return deepcopy(row)

    def update(self, table, key, payload):
        for row in self.tables.get(table, []):
            if str(row.get("id")) == str(key):
                merged = {**row, **deepcopy(dict(payload))}
                if table == "part_supplier_price_history":
                    self._guard(merged, inserting=False)
                row.update(deepcopy(dict(payload)))
                self.writes.append(("update", table, key))
                return deepcopy(row)
        raise KeyError((table, key))

    def upsert_by(self, table, payload, *, natural_key):
        existing = self.find_one(table, natural_key)
        if existing:
            return self.update(table, existing["id"], payload), "updated"
        return self.insert(table, payload), "created"

    def rpc(self, *args, **kwargs):
        return {}


BASE = {"part_id": "part", "supplier_id": "sup", "raw_material_detail_id": "rm"}


def _history_repo():
    # Live shape of 26-10-2026 (part ffb785d5…): closed revisions + one open price from 01-09-2026.
    return GuardRepo({"part_supplier_price_history": [
        {"id": "h1", **BASE, "start_date": "2026-02-01", "end_date": "2026-04-30", "price": 925.48, "uom": "PCS", "currency": "INR", "status": "ACTIVE"},
        {"id": "h2", **BASE, "start_date": "2026-05-01", "end_date": "2026-07-31", "price": 1005.26, "uom": "PCS", "currency": "INR", "status": "ACTIVE"},
        {"id": "h3", **BASE, "start_date": "2026-09-01", "end_date": None, "price": 1034.61, "uom": "PCS", "currency": "INR", "status": "ACTIVE"},
    ]})


def _grid(repo, **changes_by_id):
    rows = []
    for r in repo.select("part_supplier_price_history", eq=BASE, order_by="start_date"):
        row = {k: r.get(k) for k in ("id", "start_date", "end_date", "price", "currency", "uom", "remarks", "status")}
        row.update(changes_by_id.get(r["id"], {}))
        rows.append(row)
    return rows


# ------------------------------------------------------------------ 1. price history save
def test_old_natural_key_save_reproduces_live_overlap_error():
    """Before R37b: Start Date 01-09-2026 -> 01-08-2026 became an INSERT next to the original row."""
    repo = _history_repo()
    grid = _grid(repo, h3={"start_date": "2026-08-01"})
    with pytest.raises(RuntimeError, match="overlaps the existing period 01-09-2026"):
        for row in grid:   # the pre-R37b _save_rows logic (no id in the grid)
            payload = {k: v for k, v in row.items() if k != "id"}
            natural = {**BASE, "uom": payload["uom"], "start_date": payload["start_date"]}
            repo.upsert_by("part_supplier_price_history", {**BASE, **payload}, natural_key=natural)


def test_edited_start_date_price_and_remark_now_save_on_the_same_row():
    from core.price_history_service import save_price_history_grid
    repo = _history_repo()
    grid = _grid(repo, h3={"start_date": "2026-08-01", "price": 1040.0, "remarks": "Steel increase", "currency": "INR"})
    plan = save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm", edited=grid)
    rows = {r["id"]: r for r in repo.tables["part_supplier_price_history"]}
    assert len(rows) == 3, "no duplicate row inserted"
    assert rows["h3"]["start_date"] == "2026-08-01" and rows["h3"]["price"] == 1040.0 and rows["h3"]["remarks"] == "Steel increase"
    assert not plan.inserts and [u[0] for u in plan.updates] == ["h3"]


def test_new_open_price_closes_previous_and_moving_start_earlier_auto_closes_untouched_row():
    from core.price_history_service import save_price_history_grid
    repo = _history_repo()
    grid = _grid(repo) + [{"start_date": "2026-11-01", "price": 1100, "uom": "PCS", "currency": "INR", "remarks": "Nov revision"}]
    plan = save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm", edited=grid)
    rows = sorted(repo.tables["part_supplier_price_history"], key=lambda r: r["start_date"])
    assert rows[2]["end_date"] == "2026-10-31" and rows[3].get("end_date") is None and rows[3]["price"] == 1100
    assert plan.auto_closed
    # Move 01-11 to 15-10: the 01-09 row (end not touched by the user) is re-closed first, no overlap error.
    new_id = rows[3]["id"]
    grid = _grid(repo, **{new_id: {"start_date": "2026-10-15"}})
    save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm", edited=grid)
    rows = {r["start_date"]: r for r in repo.tables["part_supplier_price_history"]}
    assert rows["2026-09-01"]["end_date"] == "2026-10-14" and rows["2026-10-15"].get("end_date") is None


def test_invalid_grid_is_rejected_before_any_write_with_clear_messages():
    from core.price_history_service import save_price_history_grid
    repo = _history_repo()
    with pytest.raises(ValueError, match="Two price rows start on 01-09-2026"):
        save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm",
                                edited=_grid(repo) + [{"start_date": "2026-09-01", "price": 1, "uom": "PCS"}])
    with pytest.raises(ValueError, match="End Date 01-01-2026 is before Start Date"):
        save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm",
                                edited=_grid(repo, h2={"end_date": "2026-01-01"}))
    with pytest.raises(ValueError, match="Basic Rate is required"):
        save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm",
                                edited=_grid(repo) + [{"start_date": "2026-12-01", "price": None, "uom": "PCS"}])
    with pytest.raises(ValueError, match="overlaps"):
        # user explicitly sets an overlapping End Date on the earlier row -> no silent change
        save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm",
                                edited=_grid(repo, h2={"end_date": "2026-09-10"}))
    assert repo.writes == [], "nothing is written when validation fails (no partial save)"


def test_blank_rows_and_unchanged_grid_write_nothing():
    from core.price_history_service import save_price_history_grid
    repo = _history_repo()
    plan = save_price_history_grid(repo, part_id="part", supplier_id="sup", raw_material_detail_id="rm",
                                   edited=_grid(repo) + [{"start_date": None, "price": None}])
    assert plan.changed == 0 and repo.writes == []


def test_part_master_grid_carries_hidden_row_id_and_uses_controlled_save():
    page = (ROOT / "app_pages" / "part_master.py").read_text(encoding="utf-8")
    block = page[page.index('price_section_perm = section_permissions'):page.index('JOMINY REQUIREMENT')]
    assert '"_id": str(r.get("id") or "")' in block and '"_id": None' in block
    assert "save_price_history_grid(" in block
    assert '_save_rows(repo, "part_supplier_price_history"' not in block
    assert "db_message(exc)" in block


def test_price_grid_apptest_renders_with_hidden_id():
    pytest.importorskip("pyarrow")
    from streamlit.testing.v1 import AppTest

    def script():
        import pandas as pd
        import streamlit as st
        frame = pd.DataFrame([{"Start Date": pd.Timestamp("2026-09-01"), "Basic Rate": 10.0, "_id": "h3"}])
        edited = st.data_editor(frame, column_config={"_id": None}, key="grid")
        st.write("ids:" + ",".join(edited["_id"].astype(str)))

    at = AppTest.from_function(script).run(timeout=30)
    assert not at.exception
    assert any("ids:h3" in str(m.value) for m in at.markdown)


# ------------------------------------------------------------------ 2. amendment service
def _po_repo():
    return GuardRepo({
        "supply_purchase_orders": [{"id": "po", "po_number": "PD9131000096", "po_type": "FORGING", "supplier_id": "sup", "order_date": "2026-09-10",
                                    "delivery_date": "2026-10-30", "company_branch_id": "br", "ship_to_source_type": "BRANCH", "ship_to_branch_id": "br",
                                    "currency": "INR", "status": "OPEN", "approval_status": "APPROVED", "approved_at": "2026-09-11T10:00:00+00:00", "amendment_no": 0,
                                    "payment_term": "NET 30"}],
        "supply_purchase_order_items": [{"id": "it1", "purchase_order_id": "po", "part_id": "part", "raw_material_detail_id": "rm", "item_no": "9346",
                                         "item_description": "Spider", "hsn_sac_code": "87085000", "quantity": 100, "unit_price": 1005.26, "uom": "NOS",
                                         "gst_percent": 18, "technical_data_snapshot": [], "price_history_snapshot": []}],
        "parties": [{"id": "sup", "party_name": "Forge Supplier", "state": "Maharashtra", "status": "ACTIVE"}],
        "company_branches": [{"id": "br", "branch_code": "D9", "branch_name": "FSI D9", "status": "ACTIVE"}],
        "parts": [{"id": "part", "part_number": "9346", "status": "ACTIVE"}],
        "part_raw_material_details": [{"id": "rm", "part_id": "part", "supplier_id": "sup", "status": "ACTIVE"}],
        "part_supplier_price_history": [
            {"id": "h2", **BASE, "start_date": "2026-05-01", "end_date": "2026-07-31", "price": 1005.26, "uom": "PCS", "status": "ACTIVE"},
            {"id": "h3", **BASE, "start_date": "2026-08-01", "end_date": None, "price": 1034.61, "uom": "PCS", "status": "ACTIVE"},
        ],
        "supply_po_amendments": [], "supply_po_amendment_lines": [],
    })


@pytest.fixture()
def svc(monkeypatch):
    import core.supply_chain_service as scs
    monkeypatch.setattr(scs, "current_employee_id", lambda refresh=False: "emp")
    service = scs.SupplyChainService(_po_repo())
    service.purchase_order_item_received_qty = lambda item_id: 40.0
    return service


def test_effective_date_price_pick_uses_price_history_including_pcs_rows_for_nos_po(svc):
    item = svc.purchase_order_items("po")[0]
    pick = svc.amendment_price_suggestion("po", item, "2026-10-01")
    assert pick["price"] == 1034.61 and pick["start_date"] == "2026-08-01" and pick["covering"]
    pick_old = svc.amendment_price_suggestion("po", item, "2026-06-15")
    assert pick_old["price"] == 1005.26 and pick_old["start_date"] == "2026-05-01"


def test_amendment_numbering_price_history_and_reapproval(svc):
    repo = svc.repo
    result = svc.amend_purchase_order("po", {"reason": "Steel increase Nov", "effective_date": "2026-11-01",
                                             "lines": [{"id": "it1", "unit_price": 1050.0, "quantity": 120}], "payment_term": "NET 45"})
    header = repo.get("supply_purchase_orders", "po")
    assert result["amendment_no"] == 1 and result["display_number"] == "PD9131000096 · Amd 1"
    assert header["amendment_no"] == 1 and header["amendment_effective_date"] == "2026-11-01" and header["amendment_date"] == date.today().isoformat()
    assert header["approval_status"] == "PENDING_APPROVAL", "amended approved PO goes back to approval"
    item = repo.get("supply_purchase_order_items", "it1")
    assert item["unit_price"] == 1050.0 and item["quantity"] == 120
    hist = sorted(repo.tables["part_supplier_price_history"], key=lambda r: r["start_date"])
    assert hist[-1]["start_date"] == "2026-11-01" and hist[-1]["price"] == 1050.0 and hist[-1]["end_date"] is None
    assert "PD9131000096 Amd 1" in hist[-1]["remarks"] and hist[-1]["source_purchase_order_item_id"] == "it1"
    assert hist[-2]["end_date"] == "2026-10-31", "previous revision closed the day before the Effective Date"
    assert len(item["price_history_snapshot"]) == 3
    amd = repo.tables["supply_po_amendments"][0]
    assert amd["reason"] == "Steel increase Nov" and amd["effective_date"] == "2026-11-01" and amd["header_changes"]["payment_term"]["new"] == "NET 45"
    line = repo.tables["supply_po_amendment_lines"][0]
    assert (line["old_unit_price"], line["new_unit_price"], line["old_quantity"], line["new_quantity"]) == (1005.26, 1050.0, 100.0, 120.0)
    assert line["price_history_id"] == hist[-1]["id"]
    # second amendment increments the number, PO number unchanged
    svc._invalidate_transactions()
    second = svc.amend_purchase_order("po", {"reason": "Qty cut", "effective_date": "2026-11-15", "lines": [{"id": "it1", "unit_price": 1050.0, "quantity": 90}]})
    assert second["amendment_no"] == 2 and repo.get("supply_purchase_orders", "po")["po_number"] == "PD9131000096"
    header = svc.purchase_order("po")
    assert [a["amendment_no"] for a in header["amendments"]] == [1, 2] and header["po_display_number"] == "PD9131000096 · Amd 2"


def test_amendment_quantity_never_below_received_and_mandatory_fields(svc):
    with pytest.raises(ValueError, match="below the quantity already received"):
        svc.amend_purchase_order("po", {"reason": "cut", "effective_date": "2026-11-01", "lines": [{"id": "it1", "quantity": 30}]})
    with pytest.raises(ValueError, match="Reason is required"):
        svc.amend_purchase_order("po", {"reason": " ", "effective_date": "2026-11-01", "lines": [{"id": "it1", "quantity": 120}]})
    with pytest.raises(ValueError, match="Effective Date is required"):
        svc.amend_purchase_order("po", {"reason": "x", "effective_date": "", "lines": [{"id": "it1", "quantity": 120}]})
    with pytest.raises(ValueError, match="Nothing was changed"):
        svc.amend_purchase_order("po", {"reason": "x", "effective_date": "2026-11-01", "lines": [{"id": "it1", "quantity": 100, "unit_price": 1005.26}]})
    assert svc.repo.tables["supply_po_amendments"] == [] and svc.repo.get("supply_purchase_orders", "po")["approval_status"] == "APPROVED"


def test_approval_marks_pending_amendments_approved(svc):
    svc.amend_purchase_order("po", {"reason": "price", "effective_date": "2026-11-01", "lines": [{"id": "it1", "unit_price": 1040}]})
    svc.repo.rpc = lambda name, params=None: {"id": "po", "approval_status": "APPROVED", "approved_at": "2026-11-02T09:00:00+00:00", "approver_employee_id": "boss"}
    svc.approve_purchase_order("po", "ok")
    amd = svc.repo.tables["supply_po_amendments"][0]
    assert amd["status"] == "APPROVED" and amd["approved_by_employee_id"] == "boss"


def test_amend_ui_wired_on_edit_page_and_po_list():
    page = (ROOT / "app_pages" / "supply_chain.py").read_text(encoding="utf-8")
    assert "def _render_po_amendment" in page and "_render_po_amendment(service, purchase_order_id=purchase_order_id" in page
    assert 'key="po_list_open_amend"' in page and 'st.session_state["supply_po_amend_open"] = selected' in page
    block = page[page.index("def _render_po_amendment"):page.index("def _render_purchase_order_edit(")]
    assert "verify_current_password(password)" in block and "amend_purchase_order(" in block and "db_message(exc)" in block
    assert "st.dataframe" not in block


# ------------------------------------------------------------------ 3. PO print
def _print_header(**extra):
    header = {"po_number": "PD9131000096", "order_date": "2026-09-10", "delivery_date": "2026-10-30", "po_type": "FORGING", "currency": "INR",
              "subtotal": 1000, "grand_total": 1180, "approval_status": "PENDING_APPROVAL",
              "vendor_snapshot": {"party_name": "Forge Supplier", "address": "Pune"},
              "amendment_no": 2, "amendment_date": "2026-10-27", "amendment_effective_date": "2026-11-01", "amendment_reason": "Steel increase",
              "amendments": [
                  {"amendment_no": 1, "amendment_date": "2026-10-20", "effective_date": "2026-10-25", "reason": "Quantity increase for schedule", "status": "APPROVED",
                   "header_changes": {}, "lines": [{"item_no": "9346", "uom": "NOS", "old_unit_price": 1005.26, "new_unit_price": 1005.26, "old_quantity": 100, "new_quantity": 120}]},
                  {"amendment_no": 2, "amendment_date": "2026-10-27", "effective_date": "2026-11-01", "reason": "Steel increase 3.55/kg", "status": "PENDING_APPROVAL",
                   "header_changes": {"payment_term": {"old": "NET 30", "new": "NET 45"}},
                   "lines": [{"item_no": "9346", "uom": "NOS", "old_unit_price": 1005.26, "new_unit_price": 1050.0, "old_quantity": 120, "new_quantity": 120}]},
              ]}
    header.update(extra)
    return header


def _items(n_rows=30, n_items=1):
    history = [{"start_date": f"{2020 + i // 12}-{i % 12 + 1:02d}-01", "end_date": None if i == n_rows - 1 else f"{2020 + (i + 1) // 12}-{(i + 1) % 12 + 1:02d}-01",
                "price": 100 + i, "remarks": f"REVROW{i:02d}"} for i in range(n_rows)]
    return [{"item_no": f"93{k}6", "item_description": "Spider", "quantity": 120, "unit_price": 1050, "uom": "NOS", "gst_percent": 18,
             "line_total": 126000, "hsn_sac_code": "87085000", "technical_data_snapshot": [{"heading": "Section", "value": "50 RD"}],
             "price_history_snapshot": history,
             "customer_source_rows": [{"customer": "SECRET CUSTOMER LTD", "customer_po_number": "CUST-PO-SECRET", "part_number": "CUSTPART-SECRET"}]}
            for k in range(n_items)]


def _pdf_text(header, items):
    from pypdf import PdfReader
    from core.purchase_order_reporting import purchase_order_pdf_bytes
    pdf = purchase_order_pdf_bytes(header, items, terms_path=ROOT / "templates" / "__missing__.pdf")
    reader = PdfReader(BytesIO(pdf))
    return [page.extract_text() or "" for page in reader.pages]


def test_pdf_prints_amendment_header_on_every_po_page_and_history_section():
    pages = _pdf_text(_print_header(), _items(30))
    caption = re.compile(r"Amendment No\. 2\s+·\s+Amendment Date 27-10-2026\s+·\s+Effective Date 01-11-2026")
    assert all(caption.search(p) for p in pages), "amendment no + dates on every page header"
    joined = "\n".join(pages)
    assert "AMENDMENT HISTORY" in joined
    assert "Steel increase 3.55/kg" in joined and "Quantity increase for schedule" in joined
    assert "1,005.26 -> 1,050.00" in joined and "100 -> 120" in joined and "Payment Term: NET 30 -> NET 45" in " ".join(joined.split())
    assert "SECRET CUSTOMER LTD" not in joined and "CUSTPART-SECRET" not in joined


def test_pdf_prints_all_30_price_revision_rows_for_every_line():
    for n_items in (1, 3):
        joined = "\n".join(_pdf_text(_print_header(), _items(30, n_items)))
        for i in range(30):
            assert joined.count(f"REVROW{i:02d}") == n_items, f"revision row {i} missing (items={n_items})"
        assert "30 REVISIONS · CONTINUED ON NEXT PAGE" in joined


def test_unamended_po_print_unchanged_and_excel_has_amendment_and_full_history():
    pages = _pdf_text(_print_header(amendment_no=0, amendments=[]), _items(4))
    assert not any("Amendment No." in p for p in pages) and "AMENDMENT HISTORY" not in "\n".join(pages)
    from openpyxl import load_workbook
    from core.purchase_order_reporting import purchase_order_excel_bytes
    ws = load_workbook(BytesIO(purchase_order_excel_bytes(_print_header(), _items(30))))["Purchase Order"]
    values = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    assert "Amd 2" in values and "01-11-2026" in values and "Amendment History" in values and "Steel increase 3.55/kg" in values
    assert all(f"REVROW{i:02d}" in values for i in range(30))
    assert "CUST-PO-SECRET" not in values and "SECRET CUSTOMER LTD" not in values


def test_print_item_history_treats_nos_and_pcs_as_one_unit():
    from core.supply_chain_service import SupplyChainService, same_price_uom
    assert same_price_uom("PCS", "NOS") and same_price_uom("nos", "PCS") and not same_price_uom("KGS", "NOS")
    repo = _po_repo()
    repo.tables["part_supplier_price_history"].append({"id": "hn", **BASE, "start_date": "2026-11-01", "end_date": None, "price": 1060, "uom": "NOS", "status": "ACTIVE"})
    rows = SupplyChainService(repo).price_history_for_po("part", "sup", uom="NOS", raw_material_detail_id="rm")
    assert [r["start_date"] for r in rows] == ["2026-05-01", "2026-08-01", "2026-11-01"]


def test_android_po_pdf_prints_amendment_caption():
    java = (ROOT / "mobile/android_qcms_app/app/src/main/java/com/fourstar/qcmsapp/SndPoPdf.java").read_text(encoding="utf-8")
    assert "static String amendmentCaption(JSONObject header)" in java
    assert java.count("amendmentCaption(header)") >= 2
    assert "Effective Date " in java and "Amendment Date " in java


# ------------------------------------------------------------------ 4. migration safety
def test_migration_is_additive_tenant_safe_and_has_no_destructive_words():
    sql = MIGRATION.read_text(encoding="utf-8")
    assert not re.search(r"\b(update|delete|drop|truncate)\b", sql, re.I)
    for token in ("add column if not exists amendment_no", "add column if not exists amendment_effective_date",
                  "create table if not exists public.supply_po_amendments", "create table if not exists public.supply_po_amendment_lines",
                  "tenant_id uuid not null default public.current_tenant_id()", "enable row level security", "revoke all on public.%I from anon",
                  "uq_supply_po_amendments_no", "log_row_change", "if not exists (select 1 from pg_policies"):
        assert token in sql, token
    auto = (ROOT / "supabase" / "auto_apply.json").read_text(encoding="utf-8")
    assert "20261027100000_qcms_r37b_po_amendment.sql" in auto
