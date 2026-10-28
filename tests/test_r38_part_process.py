"""R38 (agent D) · Part Processes, ECN on drawing replacement, ECN actions / notification, Sections applicable (text[])."""
from __future__ import annotations

from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase/migrations/20261028130000_qcms_r38_part_process_ecn.sql"
FULL = {k: True for k in ("can_view", "can_create", "can_edit", "can_archive", "can_validate", "can_approve")}


class Repo:
    tenant_id = "t1"

    def __init__(self, tables=None):
        self.tables = deepcopy(tables or {})

    def select(self, table, eq=None, order_by=None, desc=False, limit=None, **kwargs):
        rows = [r for r in self.tables.get(table, []) if all(str(r.get(k)) == str(v) for k, v in (eq or {}).items())]
        if order_by:
            rows = sorted(rows, key=lambda r: str(r.get(order_by) or ""), reverse=bool(desc))
        return deepcopy(rows[:limit] if limit else rows)

    def get(self, table, rid):
        return next((deepcopy(r) for r in self.tables.get(table, []) if str(r.get("id")) == str(rid)), None)

    def insert(self, table, payload):
        row = {"id": f"{table}-{len(self.tables.get(table, [])) + 1}", **deepcopy(payload)}
        self.tables.setdefault(table, []).append(row)
        return deepcopy(row)

    def update(self, table, rid, payload):
        for row in self.tables.get(table, []):
            if str(row.get("id")) == str(rid):
                row.update(deepcopy(payload))
        return self.get(table, rid)

    def rpc(self, name, params=None):
        raise RuntimeError("no rpc")


TABLES = {
    "parts": [{"id": "p1", "part_number": "7568", "part_name": "Shaft Diff. Pinion", "drawing_number": "40227568", "drawing_revision": "A"}],
    "processes": [{"id": "pr1", "process_code": "P10", "process_name": "Finish Machining", "process_type": "IN_HOUSE", "status": "ACTIVE"},
                  {"id": "pr2", "process_code": "P20", "process_name": "Plating - ENP", "process_type": "OUTSOURCED", "status": "ACTIVE"}],
    "employees": [{"id": "e1", "first_name": "Rajesh", "last_name": "Dhokale", "email": "q@fsi.test", "department": "Quality", "status": "ACTIVE"},
                  {"id": "e2", "first_name": "Nitin", "last_name": "Gawde", "email": "npd@fsi.test", "department": "NPD", "status": "ACTIVE"},
                  {"id": "e3", "first_name": "Store", "last_name": "Keeper", "email": "st@fsi.test", "department": "Stores", "status": "ACTIVE"}],
}


# --------------------------------------------------------------------------------------------- part processes
def test_process_section_save_load_and_mandatory_flag():
    from core.part_process_service import PartProcessService, mandatory_processes, part_process_payloads
    repo = Repo(TABLES)
    grid = pd.DataFrame([
        {"Op No.": 10, "Process": "P10 · Finish Machining", "Mandatory": True, "Source": "", "Remarks": "", "Status": "ACTIVE", "_id": None},
        {"Op No.": 20, "Process": "P20 · Plating - ENP", "Mandatory": "No", "Source": "", "Remarks": "ENP 10 micron", "Status": "ACTIVE", "_id": None},
    ])
    labels = {"P10 · Finish Machining": "pr1", "P20 · Plating - ENP": "pr2"}
    by_id = {p["id"]: p for p in TABLES["processes"]}
    payloads = part_process_payloads(grid, labels, by_id)
    assert [p["is_mandatory"] for p in payloads] == [True, False]
    assert payloads[1]["source_type"] == "OSP"  # outsourced process defaults to OSP
    service = PartProcessService(repo)
    assert service.save_routes("p1", payloads) == 2
    loaded = service.active_routes("p1")
    assert [r["operation_no"] for r in loaded] == [10, 20] and loaded[1]["remarks"] == "ENP 10 micron"
    mandatory = mandatory_processes("p1", repo)
    assert [m["process_code"] for m in mandatory] == ["P10"]
    # update keeps the row (no duplicate) and the flag can be switched
    payloads[1]["_id"] = loaded[1]["id"]; payloads[1]["is_mandatory"] = True
    payloads[0]["_id"] = loaded[0]["id"]
    service.save_routes("p1", payloads)
    assert len(repo.tables["part_process_routes"]) == 2 and len(service.mandatory_processes("p1")) == 2
    with pytest.raises(ValueError):
        part_process_payloads(pd.DataFrame([{"Op No.": 10, "Process": "P10 · Finish Machining"},
                                            {"Op No.": 10, "Process": "P20 · Plating - ENP"}]), labels, by_id)


# --------------------------------------------------------------------------------------------- ECN on replacement
def _ecn_form(**extra):
    from core.ecn_service import revision_defaults
    form = revision_defaults(TABLES["parts"][0], document_scope="FINISH_DRAWING", drawing_number="40227568",
                             current={"id": "att-old", "revision": "A", "revision_date": "2026-01-01"}, new_revision="B")
    form.update({"fsi_representative_ids": ["e1"], "plant_code": "D9", **extra})
    return form


def test_drawing_replacement_is_blocked_without_a_valid_ecn_and_keeps_old_revision():
    from core.ecn_service import ECNService, release_with_ecn
    repo = Repo(TABLES)
    released = []
    # ECN invalid (no change description) -> the new revision is never released
    with pytest.raises(ValueError):
        release_with_ecn(ECNService(repo), _ecn_form(revision_requested=""), lambda: released.append(1) or {"id": "att-new"})
    assert released == [] and not repo.tables.get("part_ecns")
    # valid ECN -> released and linked; old revision recorded for traceability (it stays as INACTIVE history)
    ecn, rev = release_with_ecn(ECNService(repo), _ecn_form(), lambda: {"id": "att-new", "superseded_id": "att-old"},
                                status="PENDING_APPROVAL")
    assert rev["id"] == "att-new" and ecn["document_attachment_id"] == "att-new" and ecn["previous_attachment_id"] == "att-old"
    assert ecn["previous_revision"] == "A" and ecn["new_revision"] == "B" and ecn["status"] == "PENDING_APPROVAL"
    assert ecn["ecn_number"].startswith("ECN-D9-") and ecn["ecn_number"].endswith("-0001")
    # release failure -> ECN cancelled, error raised
    with pytest.raises(RuntimeError):
        release_with_ecn(ECNService(repo), _ecn_form(), lambda: (_ for _ in ()).throw(RuntimeError("upload failed")))
    assert repo.tables["part_ecns"][-1]["status"] == "CANCELLED"


def test_part_master_replacement_path_queues_an_ecn_before_saving():
    text = (ROOT / "app_pages/part_master.py").read_text(encoding="utf-8")
    assert "_queue_pending_ecn(part_id" in text and '_render_pending_ecn(repo, existing, "PART_DRAWING", perms)' in text
    ecn_page = (ROOT / "app_pages/ecn.py").read_text(encoding="utf-8")
    assert "release_with_ecn(service, form, lambda: release_pending(repo, pending)" in ecn_page


def test_ecn_number_series():
    from core.ecn_service import format_ecn_number, next_from_existing
    on = date(2026, 10, 28)
    assert format_ecn_number("ECN", "D9", 7, on) == "ECN-D9-2026-0007"
    assert next_from_existing(["ECN-D9-2026-0001", "ECN-D9-2026-0004", "PCN-D9-2026-0009"], "ECN", "D9", on) == "ECN-D9-2026-0005"
    assert next_from_existing([], "PCN", "G27", on) == "PCN-G27-2026-0001"


# --------------------------------------------------------------------------------------------- ECN actions / print / mail
def test_ecn_actions_normalize_and_block_completion_while_open():
    from core.ecn_service import ECNService, default_actions, normalize_action_rows
    repo = Repo(TABLES)
    service = ECNService(repo)
    ecn = service.create(_ecn_form(supporting_documents={"pfd": True, "control_plan": True, "redline_drawing": True}), status="PENDING_APPROVAL")
    suggested = default_actions(ecn)
    assert [a["affected_document"] for a in suggested] == ["Drawing", "PFD", "Control Plan"]
    rows = normalize_action_rows([{"action_text": "Update PFD", "responsible_employee_id": "e2", "target_date": date(2026, 11, 5)},
                                  {"action_text": "", "remarks": ""}])
    assert len(rows) == 1 and rows[0]["status"] == "OPEN" and rows[0]["target_date"] == "2026-11-05"
    with pytest.raises(ValueError):
        normalize_action_rows([{"affected_document": "CP"}])  # action text missing
    with pytest.raises(ValueError):
        normalize_action_rows([{"action_text": "x", "status": "MAYBE"}])
    service.save_actions(ecn["id"], rows)
    service.transition(ecn["id"], "APPROVED", FULL)
    with pytest.raises(ValueError, match="implementation action"):
        service.transition(ecn["id"], "COMPLETED", FULL)
    action = service.actions(ecn["id"])[0]
    service.save_actions(ecn["id"], normalize_action_rows([{**action, "status": "DONE"}]))
    assert service.actions(ecn["id"])[0]["completed_date"]  # DONE stamps the date
    assert service.transition(ecn["id"], "COMPLETED", FULL)["status"] == "COMPLETED"
    with pytest.raises(ValueError):
        service.save_actions(ecn["id"], rows)  # locked after completion


def test_ecn_pdf_contains_format_labels_and_actions_page():
    pypdf = pytest.importorskip("pypdf")
    from core.ecn_pdf import build_ecn_pdf, build_ecn_xlsx
    ecn = {**_ecn_form(), "ecn_number": "ECN-D9-2026-0001", "ecn_date": "2026-10-28", "status": "PENDING_APPROVAL",
           "revision_requested": "Chamfer 1x45 added", "po_affected": True, "inventory_affected": True, "inventory_qty": 500}
    actions = [{"affected_document": "PFD", "action_text": "Update the Process Flow Diagram", "responsible_employee_id": "e2",
                "target_date": "2026-11-05", "status": "OPEN"}]
    pdf = build_ecn_pdf(ecn, employees={"e1": TABLES["employees"][0], "e2": TABLES["employees"][1]}, actions=actions)
    pages = pypdf.PdfReader(BytesIO(pdf)).pages
    assert len(pages) == 2
    text = " ".join(p.extract_text() for p in pages)
    for token in ("ENGINEERING CHANGE NOTE", "FSI/806/F05", "REV.DATE", "25.09.2024", "PAGE NO", "01 TO 02", "FSI REPRESENTATIVE",
                  "IDENTIFICATION", "DRAWING CRITERIA", "40227568", "7568", "PREVIOUS DESCRIPTION", "REVISION REQUESTED",
                  "SUPPORTING DOCUMENTS ATTACHED", "Purchase Orders Affected", "Current Inventory Affected", "Current Demands Affected",
                  "FSI USE ONLY", "Drawing change to occur", "Approval Signatures", "Completed By", "COMMENTS",
                  "IMPLEMENTATION ACTIONS", "Update the Process Flow Diagram", "Nitin Gawde", "05.11.2026"):
        assert token in text, token
    single = pypdf.PdfReader(BytesIO(build_ecn_pdf(ecn))).pages
    assert len(single) == 1 and "01 TO 01" in single[0].extract_text()
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(build_ecn_xlsx(ecn, actions=actions))).active
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert "IMPLEMENTATION ACTIONS" in values and "Update the Process Flow Diagram" in values


def test_new_ecn_notifies_quality_and_npd(monkeypatch):
    import core.ecn_service as svc
    repo = Repo(TABLES)
    queued = {}

    class Notifier:
        def enqueue(self, event, **kwargs):
            queued.update(event=event, **kwargs)
            return {"id": "o1"}

        def dispatch(self, rows):
            return {}

    monkeypatch.setattr(svc, "ecn_notifier", lambda repo=None: Notifier())
    ecn = svc.ECNService(repo).create(_ecn_form(), status="DRAFT")
    sent, message = svc.notify_ecn_created(repo, ecn)
    assert sent == 2 and message == ""
    assert queued["event"] == "ECN_CREATED" and queued["recipient_email"] == "q@fsi.test" and queued["cc_emails"] == ["npd@fsi.test"]
    assert queued["related_table"] == "part_ecns" and queued["include_supplier"] is False and queued["include_generated_pdf"] is True
    assert queued["subject"].startswith("Engineering · New ECN")
    page = (ROOT / "app_pages/ecn.py").read_text(encoding="utf-8")
    assert "notify_ecn_created(repo, ecn)" in page


# --------------------------------------------------------------------------------------------- sections applicable
def test_part_applicable_sections_contract():
    from core.part_process_service import (
        METALLURGICAL_SECTION_KEYS, PART_SECTION_KEYS, part_applicable_sections, section_is_applicable, sections_storage,
    )
    for key in ("CHEMISTRY", "MECHANICAL", "JOMINY", "GRAIN_SIZE", "INCLUSION", "MICROSTRUCTURE", "HARDNESS", "CASE_DEPTH",
                "HEAT_TREATMENT", "ULTRASONIC", "MACRO"):
        assert key in METALLURGICAL_SECTION_KEYS
    assert part_applicable_sections({}) == set(PART_SECTION_KEYS)  # nothing chosen -> all applicable
    assert part_applicable_sections({"applicable_sections": ["chemistry", "JOMINY", "BOGUS"]}) == {"CHEMISTRY", "JOMINY"}
    assert part_applicable_sections({"applicable_sections": "{HARDNESS,MACRO}"}) == {"HARDNESS", "MACRO"}
    assert part_applicable_sections({"applicable_sections": []}) == set()
    legacy = part_applicable_sections({"applicable_sections": None, "section_applicability": {"jominy": False, "osp_metlab": "no"}})
    assert "JOMINY" not in legacy and "PROCESS_SPECIFICATIONS" not in legacy and "MECHANICAL" in legacy
    assert part_applicable_sections("p1", Repo({"parts": [{"id": "p1", "applicable_sections": ["DRAWINGS"]}]})) == {"DRAWINGS"}
    assert section_is_applicable({"applicable_sections": ["DRAWINGS"]}, "UNKNOWN_KEY") is True
    assert section_is_applicable({"applicable_sections": ["DRAWINGS"]}, "jominy") is False
    stored = sections_storage({"DRAWINGS", "CHEMISTRY", "MECHANICAL"})
    assert stored["applicable_sections"] == ["DRAWINGS", "CHEMISTRY", "MECHANICAL"]
    assert stored["section_applicability"]["jominy"] is False and stored["section_applicability"]["controlled_drawings"] is True
    assert stored["section_applicability"]["chemistry"] is True
    # round trip through the legacy mapping keeps the RMTC side (core.section_applicability) consistent
    from core.section_applicability import default_rmtc_sections
    assert default_rmtc_sections([{"section_applicability": stored["section_applicability"]}])["jominy"] is False


def test_applicability_hides_sections_in_prints_and_requirement_rows():
    from core.part_process_service import filter_print_tables, requirement_section_key
    assert requirement_section_key("Effective Case Depth") == "CASE_DEPTH"
    assert requirement_section_key("Surface Hardness HRC") == "HARDNESS"
    assert requirement_section_key("Tensile Strength") == "MECHANICAL"
    assert requirement_section_key("Grain Size") == "GRAIN_SIZE" and requirement_section_key("Inclusion rating") == "INCLUSION"
    assert requirement_section_key("Remarks") is None
    applicable = {"STANDARDS", "FINAL_METALLURGICAL", "HARDNESS"}
    tables = {
        "Linked Customer Standards / Specifications": [{"x": 1}],
        "Jominy Requirements": [{"x": 2}],
        "Metallurgical Requirements": [{"parameter_name": "Surface Hardness"}, {"parameter_name": "Grain Size"}, {"parameter_name": "Note"}],
        "Something Else": [{"x": 3}],
    }
    printed = filter_print_tables(tables, applicable)
    assert "Jominy Requirements" not in printed and "Something Else" in printed
    assert [r["parameter_name"] for r in printed["Metallurgical Requirements"]] == ["Surface Hardness", "Note"]
    # generic record print / preview of a Part skips child tables of N/A sections
    from core import record_print
    repo = Repo({"parts": [{"id": "p1", "part_number": "7568", "applicable_sections": ["STANDARDS", "FINAL_METALLURGICAL", "HARDNESS"]}],
                 "part_jominy_requirements": [{"id": "j1", "part_id": "p1", "distance_label": "J1", "minimum_hrc": 40}],
                 "part_metallurgical_requirements": [{"id": "m1", "part_id": "p1", "parameter_name": "Grain Size", "minimum_spec": "5"},
                                                     {"id": "m2", "part_id": "p1", "parameter_name": "Surface Hardness", "minimum_spec": "58"}]})
    data = record_print.record_print_data("parts", "p1", repo)
    assert "JOMINY REQUIREMENTS" not in data["sections"]
    meta = data["sections"].get("METALLURGICAL REQUIREMENTS")
    assert meta is not None and len(meta) == 1


def test_part_master_uses_text_array_checklist_and_chem_grid():
    text = (ROOT / "app_pages/part_master.py").read_text(encoding="utf-8")
    assert "sections_storage(chosen)" in text and "PART_SECTION_REGISTRY" in text and "part_applicable_sections" in text
    assert "_ppsvc.filter_print_tables(" in text and "from core.chem_grid import jominy_distance_mm" in text


def test_migration_step3_additive_and_listed():
    sql = MIGRATION.read_text(encoding="utf-8")
    for token in ("add column if not exists applicable_sections text[]", "create table if not exists public.part_ecn_actions",
                  "tenant_id uuid not null default public.current_tenant_id()", "alter table public.part_ecn_actions enable row level security",
                  "revoke all on public.part_ecn_actions from anon", "qcms_guard_part_ecn_completion", "'ECN_CREATED'",
                  "on conflict (tenant_id, event_key) do nothing"):
        assert token in sql, token
    step3 = sql.split("R38 step 3 (agent D)", 1)[1].upper()
    for word in ("DELETE ", "DROP ", "UPDATE "):
        assert word not in step3.replace("UPDATED_", ""), word
    import json
    files = json.loads((ROOT / "supabase/auto_apply.json").read_text(encoding="utf-8"))["files"]
    assert "supabase/migrations/20261028130000_qcms_r38_part_process_ecn.sql" in files
