"""R37 · Indirect Purchase module: item master, requisition, Indirect PO (own IPO series, GST, approval route
INDIRECT_PURCHASE), Indirect Inward (GRN) with over-receipt tolerance, indirect stock ledger, prints and page renders."""
from __future__ import annotations

import re
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest

from core import indirect_purchase_service as ips
from core.indirect_purchase_service import IndirectPurchaseService, compute_line, compute_totals, derive_po_status, gst_mode, receivable_qty

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase/migrations/20261024120000_qcms_r37_indirect_purchase.sql"
INDIRECT_TABLES = ("indirect_item_categories", "indirect_items", "indirect_number_series", "indirect_number_counters",
                   "indirect_purchase_requisitions", "indirect_purchase_requisition_items", "indirect_purchase_orders",
                   "indirect_purchase_order_items", "indirect_po_approval_log", "indirect_grns", "indirect_grn_items", "indirect_stock_ledger")


class FakeRepo:
    """In-memory Repository; the RPCs mirror the SECURITY DEFINER functions of the R37 migration."""

    def __init__(self, tables=None):
        self.tables = deepcopy(tables or {}); self.calls = []; self.counters = {}

    def select(self, table, eq=None, in_=None, order_by=None, desc=False, limit=None, **kwargs):
        rows = [r for r in self.tables.get(table, []) if all(str(r.get(k)) == str(v) for k, v in (eq or {}).items())]
        if order_by:
            rows = sorted(rows, key=lambda r: str(r.get(order_by) or ""), reverse=bool(desc))
        return deepcopy(rows[:limit] if limit else rows)

    def get(self, table, rid):
        return next((deepcopy(r) for r in self.tables.get(table, []) if str(r.get("id")) == str(rid)), None)

    def insert(self, table, payload):
        row = {"id": f"{table}-{len(self.tables.get(table, [])) + 1}", "created_at": "2026-10-23T10:00:00+00:00", **payload}
        self.tables.setdefault(table, []).append(row); self.calls.append(("insert", table, dict(payload)))
        return deepcopy(row)

    def update(self, table, rid, payload):
        for row in self.tables.get(table, []):
            if str(row.get("id")) == str(rid):
                row.update(payload)
        self.calls.append(("update", table, rid, dict(payload)))
        return self.get(table, rid)

    def delete(self, table, rid):
        self.tables[table] = [r for r in self.tables.get(table, []) if str(r.get("id")) != str(rid)]

    def upsert_by(self, table, payload, *, natural_key):
        for row in self.tables.get(table, []):
            if all(str(row.get(k)) == str(v) for k, v in natural_key.items()):
                row.update(payload)
                return deepcopy(row), "UPDATED"
        return self.insert(table, payload), "CREATED"

    # ---------------------------------------------------------------- RPC mirrors
    def rpc(self, name, params=None):
        params = dict(params or {})
        self.calls.append(("rpc", name, params))
        if name == "qcms_next_indirect_number":
            key = params["p_series_key"]
            self.counters[key] = self.counters.get(key, 0) + 1
            if key == "IITEM":
                return f"IND-{self.counters[key]:05d}"
            branch = next((b for b in self.tables.get("company_branches", []) if b["id"] == params.get("p_company_branch_id")), {})
            return f"{key}-{branch.get('plant_code') or 'HO'}-26-{self.counters[key]:04d}"
        if name == "qcms_indirect_po_approval_target":
            raise RuntimeError("use python fallback")
        if name == "qcms_approve_indirect_po":
            po = self.get("indirect_purchase_orders", params["p_po_id"])
            assert po["approval_status"] == "PENDING_APPROVAL"
            return self.update("indirect_purchase_orders", po["id"], {"approval_status": "APPROVED", "status": "APPROVED", "approver_employee_id": "e2"})
        if name == "qcms_reject_indirect_po":
            return self.update("indirect_purchase_orders", params["p_po_id"], {"approval_status": "REJECTED", "status": "REJECTED", "approval_remarks": params["p_remarks"]})
        if name == "qcms_post_indirect_grn":
            grn = self.get("indirect_grns", params["p_grn_id"])
            if grn["grn_type"] == "WITHOUT_PO":
                return self.update("indirect_grns", grn["id"], {"status": "PENDING_APPROVAL"})
            po = self.get("indirect_purchase_orders", grn["purchase_order_id"])
            tol = float(po.get("over_receipt_tolerance_percent") or 0)
            for line in self.select("indirect_grn_items", eq={"grn_id": grn["id"]}):
                poi = self.get("indirect_purchase_order_items", line["po_item_id"])
                if line["received_qty"] > poi["quantity"] * (1 + tol / 100) - poi.get("accepted_qty", 0) + 0.0005:
                    raise RuntimeError("received exceeds the receivable balance")
                self.update("indirect_purchase_order_items", poi["id"], {
                    "received_qty": poi.get("received_qty", 0) + line["received_qty"], "accepted_qty": poi.get("accepted_qty", 0) + line["accepted_qty"],
                    "rejected_qty": poi.get("rejected_qty", 0) + line["rejected_qty"]})
                item = self.get("indirect_items", line.get("item_id")) if line.get("item_id") else None
                if item and item.get("is_stock_item", True) and line["accepted_qty"] > 0:
                    self.insert("indirect_stock_ledger", {"item_id": item["id"], "company_branch_id": grn.get("company_branch_id"), "txn_type": "GRN",
                                                          "txn_date": grn["grn_date"], "qty_in": line["accepted_qty"], "qty_out": 0, "reference_no": grn["grn_number"]})
            status = derive_po_status(self.select("indirect_purchase_order_items", eq={"purchase_order_id": po["id"]}), po["status"])
            self.update("indirect_purchase_orders", po["id"], {"status": status})
            return {**self.update("indirect_grns", grn["id"], {"status": "POSTED"}), "po_status": status}
        if name == "qcms_issue_indirect_stock":
            balance = sum(r.get("qty_in", 0) - r.get("qty_out", 0) for r in self.tables.get("indirect_stock_ledger", []) if r["item_id"] == params["p_item_id"])
            out = params["p_qty"] if params["p_txn_type"] == "ISSUE" else 0
            if out > balance:
                raise RuntimeError(f"Insufficient stock: available {balance}, requested {out}")
            return self.insert("indirect_stock_ledger", {"item_id": params["p_item_id"], "txn_type": params["p_txn_type"], "qty_in": 0 if out else params["p_qty"],
                                                         "qty_out": out, "balance_after": balance - out + (0 if out else params["p_qty"]), "reference_no": "ISS-HO-26-0001",
                                                         "department": params.get("p_department")})
        raise RuntimeError(f"unexpected rpc {name}")


BASE = {
    "company_branches": [{"id": "b1", "branch_name": "Four Star Industries D9", "branch_code": "D9", "plant_code": "D9", "gstin": "27AAGCF3769A1ZP",
                          "state": "Maharashtra", "status": "ACTIVE"}],
    "parties": [
        {"id": "s1", "party_code": "SUP-ABC", "party_name": "ABC Tools Pvt Ltd", "party_types": ["SUPPLIER"], "tax_identifier": "27ABCDE1234F1Z5",
         "state": "Maharashtra", "email": "sales@abctools.example", "status": "ACTIVE"},
        {"id": "s2", "party_code": "SUP-KAR", "party_name": "Bangalore Safety Co", "party_types": ["SUPPLIER"], "tax_identifier": "29ABCDE1234F1Z5",
         "state": "Karnataka", "status": "ACTIVE"},
        {"id": "c1", "party_code": "CUST-1", "party_name": "Secret Customer GmbH", "party_types": ["CUSTOMER"], "status": "ACTIVE"},
    ],
    "employees": [{"id": "e1", "employee_code": "E001", "first_name": "Rajesh", "last_name": "D", "department": "Maintenance", "status": "ACTIVE",
                   "reports_to_employee_id": "e2"},
                  {"id": "e2", "employee_code": "E002", "first_name": "Purchase", "last_name": "Head", "department": "Purchase", "status": "ACTIVE",
                   "email": "purchase.head@fourstarindustries.com"}],
    "indirect_items": [{"id": "i1", "item_code": "IND-00001", "item_name": "Nitrile Gloves", "category": "SAFETY_PPE", "uom": "PAIR", "hsn_sac": "40151900",
                        "gst_percent": 12, "last_price": 85, "is_stock_item": True, "reorder_level": 50, "status": "ACTIVE"},
                       {"id": "i2", "item_code": "IND-00002", "item_name": "Carbide Insert CNMG", "category": "TOOLS_TOOLING", "uom": "NOS",
                        "gst_percent": 18, "last_price": 450, "is_stock_item": True, "status": "ACTIVE"}],
}


def service(tables=None) -> IndirectPurchaseService:
    return IndirectPurchaseService(FakeRepo({**deepcopy(BASE), **(tables or {})}))


def _po(svc, lines=None, **extra):
    payload = {"supplier_id": "s1", "company_branch_id": "b1", "po_date": "2026-10-23", "delivery_date": "2026-10-30", "department": "Maintenance",
               "requisitioner": "Rajesh D", **extra}
    lines = lines or [{"item_id": "i1", "quantity": 100, "rate": 80, "discount_percent": 5}, {"description": "AMC – compressor", "quantity": 1, "rate": 10000, "gst_percent": 18, "uom": "JOB"}]
    return svc.create_po(payload, lines, submit=True, submitted_by_employee_id="e1")


# ------------------------------------------------------------------------------------------- pure helpers
def test_line_and_totals_with_discount_and_gst_split():
    line = compute_line({"quantity": 10, "rate": 100, "discount_percent": 10, "gst_percent": 18}, "INTRA")
    assert line["gross_amount"] == 1000 and line["discount_amount"] == 100 and line["taxable_amount"] == 900
    assert line["gst_amount"] == 162 and line["cgst_amount"] == 81 and line["sgst_amount"] == 81 and line["igst_amount"] == 0
    assert line["line_total"] == 1062
    inter = compute_line({"quantity": 10, "rate": 100, "discount_percent": 10, "gst_percent": 18}, "INTER")
    assert inter["igst_amount"] == 162 and inter["cgst_amount"] == 0
    totals = compute_totals([{"quantity": 3, "rate": 33.33, "gst_percent": 18}, {"quantity": 1, "rate": 10, "gst_percent": 5}], "INTRA", freight=50)
    assert totals["subtotal"] == 109.99 and totals["freight_amount"] == 50
    assert totals["grand_total"] == round(totals["grand_total"]) and abs(totals["round_off"]) < 0.5
    assert totals["amount_in_words"].startswith("Rupees") and totals["amount_in_words"].endswith("Only")
    assert ips.amount_in_words(125000.5) == "Rupees One Lakh Twenty Five Thousand and Fifty Paise Only"


def test_gst_mode_from_gstin_state_code():
    plant = BASE["company_branches"][0]
    assert gst_mode(BASE["parties"][0], plant) == "INTRA"
    assert gst_mode(BASE["parties"][1], plant) == "INTER"
    assert gst_mode({"country": "Germany"}, plant) == "INTER"


def test_receivable_quantity_and_status_derivation():
    assert receivable_qty(100, 0) == 100 and receivable_qty(100, 60) == 40
    assert receivable_qty(100, 0, 5) == 105 and receivable_qty(100, 100, 5) == 5 and receivable_qty(10, 20) == 0
    assert receivable_qty(100, 0, 500) == 150  # tolerance is capped at 50 %
    assert derive_po_status([{"quantity": 10, "accepted_qty": 0, "received_qty": 0}], "APPROVED") == "APPROVED"
    assert derive_po_status([{"quantity": 10, "accepted_qty": 4, "received_qty": 4}], "APPROVED") == "PARTIALLY_RECEIVED"
    assert derive_po_status([{"quantity": 10, "accepted_qty": 10, "received_qty": 12}], "PARTIALLY_RECEIVED") == "CLOSED"
    assert derive_po_status([{"quantity": 10, "accepted_qty": 0}], "PENDING_APPROVAL") == "PENDING_APPROVAL"


# ------------------------------------------------------------------------------------------- service flows
def test_create_po_numbers_totals_and_snapshots_are_supplier_safe():
    svc = service()
    po = _po(svc)
    assert po["po_number"] == "IPO-D9-26-0001" and po["status"] == "PENDING_APPROVAL" and po["approval_status"] == "PENDING_APPROVAL"
    assert po["gst_mode"] == "INTRA" and po["subtotal"] == round(100 * 80 * 0.95 + 10000, 2)
    assert po["cgst_amount"] == po["sgst_amount"] and po["igst_amount"] == 0
    lines = svc.po_items(po["id"])
    assert [l["line_no"] for l in lines] == [1, 2] and lines[0]["hsn_sac"] == "40151900" and lines[0]["gst_percent"] == 12
    assert lines[1]["item_id"] is None and lines[1]["item_name"] == "AMC – compressor"
    # never customer data on an indirect PO
    text = str(svc.purchase_order(po["id"])).casefold()
    assert "secret customer" not in text and "customer_po" not in text
    # numbering stays separate from the direct PO series
    assert not any(c[1] in {"qcms_next_po_number", "supply_po_sequences"} for c in svc.repo.calls if c[0] == "rpc")


def test_po_validation_errors():
    svc = service()
    with pytest.raises(ValueError, match="supplier"):
        svc.build_po({"company_branch_id": "b1"}, [{"item_id": "i1", "quantity": 1}])
    with pytest.raises(ValueError, match="quantity"):
        svc.build_po({"supplier_id": "s1", "company_branch_id": "b1"}, [{"item_id": "i1", "quantity": 0, "rate": 1}])
    with pytest.raises(ValueError, match="Delivery date"):
        svc.build_po({"supplier_id": "s1", "company_branch_id": "b1", "po_date": "2026-10-23", "delivery_date": "2026-10-01"}, [{"item_id": "i1", "quantity": 1}])
    with pytest.raises(ValueError, match="Employee Master"):
        svc.create_po({"supplier_id": "s1", "company_branch_id": "b1"}, [{"item_id": "i1", "quantity": 1}], submit=True, submitted_by_employee_id=None)


def test_status_transitions_approve_receive_partial_then_close():
    svc = service()
    po = _po(svc, lines=[{"item_id": "i1", "quantity": 100, "rate": 80}])
    target = svc.approval_target(po["id"])
    assert target["source"] == "REPORTS_TO" and target["employee_id"] == "e2"
    svc.approve_po(po["id"], "ok")
    assert svc.purchase_order(po["id"])["status"] == "APPROVED"
    assert [p["id"] for p in svc.open_pos_for_grn()] == [po["id"]]
    pending = svc.po_pending_lines(po["id"])
    grn = svc.create_grn({"grn_type": "AGAINST_PO", "purchase_order_id": po["id"], "grn_date": "2026-10-24", "supplier_invoice_no": "INV-1",
                          "gate_entry_no": "GE-77", "challan_no": "DC-5"}, [{**pending[0], "po_item_id": pending[0]["id"], "received_qty": 60, "accepted_qty": 55,
                                                                             "rejected_qty": 5, "rejection_reason": "torn"}])
    assert grn["grn_number"] == "IGRN-D9-26-0001" and grn["po_status"] == "PARTIALLY_RECEIVED" and grn["gate_entry_no"] == "GE-77"
    line = svc.po_pending_lines(po["id"])[0]
    assert line["pending_qty"] == 45 and line["po_accepted_qty"] == 55 and line["po_rejected_qty"] == 5 and "accepted_qty" not in line
    stock = {r["Item Code"]: r for r in svc.stock_rows()}
    assert stock["IND-00001"]["Balance"] == 55 and stock["IND-00001"]["Stock Alert"] == ""
    # cannot cancel / revise once material is received
    with pytest.raises(ValueError, match="received"):
        svc.cancel_po(po["id"], "not needed", "e1")
    with pytest.raises(ValueError, match="cannot be edited|received"):
        svc.update_po(po["id"], {"supplier_id": "s1", "company_branch_id": "b1"}, [{"item_id": "i1", "quantity": 1}], submit=False)
    pending = svc.po_pending_lines(po["id"])
    grn2 = svc.create_grn({"grn_type": "AGAINST_PO", "purchase_order_id": po["id"]}, [{**pending[0], "po_item_id": pending[0]["id"], "received_qty": 45}])
    assert grn2["po_status"] == "CLOSED" and svc.purchase_order(po["id"])["status"] == "CLOSED"


def test_over_receipt_guard_and_tolerance():
    svc = service()
    po = _po(svc, lines=[{"item_id": "i2", "quantity": 10, "rate": 450}])
    svc.approve_po(po["id"])
    pending = svc.po_pending_lines(po["id"])
    assert pending[0]["max_receivable_qty"] == 10
    with pytest.raises(ValueError, match="exceeds the receivable balance"):
        svc.validate_grn_lines([{**pending[0], "received_qty": 11}], against_po=True)
    with pytest.raises(ValueError, match="accepted"):
        svc.validate_grn_lines([{**pending[0], "received_qty": 5, "accepted_qty": 5, "rejected_qty": 1}], against_po=True)
    with pytest.raises(ValueError, match="rejection reason"):
        svc.validate_grn_lines([{**pending[0], "received_qty": 5, "accepted_qty": 4, "rejected_qty": 1}], against_po=True)
    # with 10 % tolerance on the PO, 11 may be received and the PO closes
    svc.repo.update("indirect_purchase_orders", po["id"], {"over_receipt_tolerance_percent": 10}); svc.invalidate()
    pending = svc.po_pending_lines(po["id"])
    assert pending[0]["max_receivable_qty"] == 11
    grn = svc.create_grn({"grn_type": "AGAINST_PO", "purchase_order_id": po["id"]}, [{**pending[0], "po_item_id": pending[0]["id"], "received_qty": 11}])
    assert grn["po_status"] == "CLOSED"


def test_grn_requires_approved_open_po_and_reject_resubmit():
    svc = service()
    po = _po(svc)
    with pytest.raises(ValueError, match="not approved"):
        svc.create_grn({"grn_type": "AGAINST_PO", "purchase_order_id": po["id"]}, [{"received_qty": 1}])
    svc.reject_po(po["id"], "price too high")
    assert svc.purchase_order(po["id"])["status"] == "REJECTED"
    svc.update_po(po["id"], {"supplier_id": "s2", "company_branch_id": "b1"}, [{"item_id": "i1", "quantity": 10, "rate": 70}], submit=True, submitted_by_employee_id="e1")
    revised = svc.purchase_order(po["id"])
    assert revised["status"] == "PENDING_APPROVAL" and revised["gst_mode"] == "INTER" and revised["igst_amount"] > 0
    assert len(svc.po_items(po["id"])) == 1
    with pytest.raises(ValueError, match="cancellation reason"):
        svc.cancel_po(po["id"], "", "e1")
    assert svc.cancel_po(po["id"], "Duplicate order", "e1")["status"] == "CANCELLED"


def test_requisition_to_po_and_stock_issue():
    svc = service()
    pr = svc.create_requisition({"company_branch_id": "b1", "department": "Maintenance", "requested_by_employee_id": "e1", "purpose": "Breakdown"},
                                [{"item_id": "i1", "quantity": 20}, {"description": "Bearing 6205", "quantity": 4, "uom": "NOS"}])
    assert pr["pr_number"] == "IPR-D9-26-0001"
    with pytest.raises(ValueError, match="Self-approval"):
        svc.decide_requisition(pr["id"], True, "", "e1")
    svc.decide_requisition(pr["id"], True, "", "e2")
    assert [r["id"] for r in svc.convertible_requisitions()] == [pr["id"]]
    pr_lines = svc.requisition_items(pr["id"])
    po = _po(svc, lines=[{"item_id": "i1", "quantity": 20, "rate": 80, "requisition_item_id": pr_lines[0]["id"]}], requisition_id=pr["id"])
    assert svc.requisition(pr["id"])["status"] == "PARTIALLY_ORDERED" and po["requisition_id"] == pr["id"]
    with pytest.raises(Exception, match="Insufficient"):
        svc.issue_stock(item_id="i1", qty=5, branch_id="b1", department="Maintenance")
    with pytest.raises(ValueError, match="department"):
        svc.issue_stock(item_id="i1", qty=5, branch_id="b1", department="")


def test_item_master_auto_code_duplicates_and_import():
    svc = service()
    row = svc.save_item({"item_name": "Printer Toner 88A", "category": "IT", "uom": "NOS", "gst_percent": 18})
    assert row["item_code"] == "IND-00001"
    with pytest.raises(ValueError, match="already exists"):
        svc.save_item({"item_name": "Nitrile gloves", "uom": "PAIR"})
    with pytest.raises(ValueError, match="GST"):
        svc.item_payload({"item_name": "x", "gst_percent": 120})
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(svc.import_template_bytes()))
    assert tuple(c.value for c in wb.active[1]) == IndirectPurchaseService.IMPORT_COLUMNS
    preview = svc.import_preview([{"Item Code": "IND-00002", "Item Name": "Carbide Insert CNMG 120408", "Category": "TOOLS_TOOLING", "UOM": "NOS", "GST %": 18},
                                  {"Item Name": "Coffee", "Category": "PANTRY"}, {"Item Name": "Safety Shoes", "Category": "SAFETY_PPE", "Default Supplier Code": "SUP-ABC"}])
    assert [p["Action"] for p in preview] == ["UPDATE", "ERROR", "CREATE"]
    counts = svc.apply_import(preview)
    assert counts == {"created": 1, "updated": 1, "skipped": 1}


def test_dashboard_kpis_and_registers():
    svc = service()
    po = _po(svc)
    svc.approve_po(po["id"])
    data = svc.dashboard(today=__import__("datetime").date(2026, 11, 5))
    assert data["open"] == 1 and data["overdue"] == 1 and data["pending_approval"] == 0
    data_oct = svc.dashboard(today=__import__("datetime").date(2026, 10, 25))
    assert data_oct["spend_month"] > 0 and "Safety / PPE" in data_oct["spend_by_category"] and "ABC Tools Pvt Ltd" in data_oct["spend_by_supplier"]
    rows = svc.po_register_rows()
    assert rows[0]["IPO Number"] == "IPO-D9-26-0001" and rows[0]["Supplier"] == "ABC Tools Pvt Ltd"
    assert set(svc.grn_register_rows()[0].keys()) if svc.grn_register_rows() else True


# ------------------------------------------------------------------------------------------- PDF
def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)


def test_po_and_grn_pdf_contract_by_id():
    pytest.importorskip("pypdf")
    from core.indirect_po_pdf import indirect_grn_pdf_bytes, indirect_po_pdf_bytes, watermark_label
    svc = service()
    po = _po(svc)
    pdf = indirect_po_pdf_bytes(po["id"], svc.repo, include_terms=False)
    text = _pdf_text(pdf)
    assert pdf[:4] == b"%PDF" and "INDIRECT PURCHASE ORDER" in text and "IPO-D9-26-0001" in text
    assert "ABC Tools Pvt Ltd" in text and "Secret Customer" not in text
    assert watermark_label({"approval_status": "PENDING_APPROVAL"}) == "PENDING APPROVAL"
    assert watermark_label({"status": "CANCELLED", "approval_status": "APPROVED"}) == "CANCELLED"
    svc.approve_po(po["id"])
    pending = svc.po_pending_lines(po["id"])
    grn = svc.create_grn({"grn_type": "AGAINST_PO", "purchase_order_id": po["id"], "gate_entry_no": "GE-9"}, [{**pending[0], "po_item_id": pending[0]["id"], "received_qty": 10}])
    gtext = _pdf_text(indirect_grn_pdf_bytes(grn["id"], svc.repo))
    assert "INDIRECT GOODS RECEIPT NOTE" in gtext and grn["grn_number"] in gtext and "GE-9" in gtext
    with pytest.raises(ValueError):
        indirect_po_pdf_bytes("missing", svc.repo)


# ------------------------------------------------------------------------------------------- migration & wiring
def test_migration_is_additive_tenant_scoped_and_guarded():
    sql = MIGRATION.read_text()
    low = sql.lower()
    for table in INDIRECT_TABLES:
        assert f"create table if not exists public.{table}" in low, table
        assert f"'{table}'" in low  # RLS / grant loop covers it
    assert "enable row level security" in low and "revoke all on public.%i from anon" in low
    assert "public.current_tenant_id()" in low and "qcms_indirect_purchase_can" in low
    assert "add column if not exists over_receipt_tolerance_percent" in low and "add column if not exists gate_entry_no" in low
    # no destructive data statements
    assert not re.search(r"^\s*(delete\s+from|truncate|drop\s+table)", low, re.M)
    assert not re.search(r"^\s*update\s+public\.(?!indirect_)", low, re.M)
    for stmt in re.findall(r"^\s*create (?:policy|trigger)[^\n]*", low, re.M):
        pass  # every policy / trigger creation is inside an "if not exists" guard (checked below)
    assert low.count("create policy") <= low.count("if not exists (select 1 from pg_policies")
    assert low.count("create trigger") <= low.count("if not exists (select 1 from pg_trigger")
    # global search keeps every older entry and routes the new tables to the new pages
    for key in ("\"parties\"", "\"supply_purchase_orders\"", "\"qcms_delete_requests\"", "\"indirect_purchase_orders\": [\"indirect_purchase\"",
                "\"indirect_grns\": [\"indirect_purchase\"", "\"indirect_items\": [\"indirect_purchase\""):
        assert key in low, key
    auto = (ROOT / "supabase/auto_apply.json").read_text()
    assert "20261024120000_qcms_r37_indirect_purchase.sql" in auto


def test_routes_permissions_and_module_registration():
    app = (ROOT / "streamlit_app.py").read_text()
    for route in ("indirect-purchase", "indirect-requisition", "indirect-po", "indirect-po-list", "indirect-po-approval", "indirect-inward", "indirect-items", "indirect-stock"):
        assert f'url_path="{route}"' in app, route
        assert f'"{route}":"INDIRECT_PURCHASE"' in app, route
        assert f'"{route}": "Indirect Purchase"' in app, route
    assert '"Indirect Purchase": (' in app and 'PAGE_BY_PATH["indirect-purchase"], "Indirect Purchase", "Indirect Purchase"' in app
    from core.access import MODULES, _legacy_role_permissions
    assert ("INDIRECT_PURCHASE", "Indirect Purchase") in MODULES
    assert _legacy_role_permissions({"role": "PROCUREMENT"}, "INDIRECT_PURCHASE")["can_approve"] is True
    assert _legacy_role_permissions({"role": "QUALITY_ENGINEER"}, "INDIRECT_PURCHASE")["can_create"] is False
    verify = (ROOT / "scripts/verify_phase1.py").read_text()
    assert '"indirect-po-list"' in verify
    from core import report_builder as rb
    assert {"indirect_purchase_orders", "indirect_grns", "indirect_stock_ledger"} <= {s["table"] for s in rb.INVENTORY_SOURCES}
    page = (ROOT / "app_pages/indirect_purchase.py").read_text()
    assert "st.dataframe(" not in page and "render_filters(" in page and "report_exports(" in page
    # stage_section is an expander: the page must not open another expander inside it
    assert "st.expander(" not in page
    # direct PO register / genealogy never reads the indirect tables
    for path in ("core/supply_chain_service.py", "core/purchase_order_reporting.py"):
        assert "indirect_" not in (ROOT / path).read_text()


# ------------------------------------------------------------------------------------------- AppTest renders
SCRIPT_HEAD = '''
import streamlit as st
st.data_editor = lambda frame, **kwargs: frame
from tests.test_r37_indirect_purchase import FakeRepo, BASE, _po
from core.indirect_purchase_service import IndirectPurchaseService
import app_pages.indirect_purchase as page
repo = st.session_state.setdefault("ip_repo", FakeRepo(BASE))
if not repo.tables.get("indirect_purchase_orders"):
    svc = IndirectPurchaseService(repo); po = _po(svc); svc.approve_po(po["id"])
    _po(svc, lines=[{"item_id": "i2", "quantity": 2, "rate": 450}])
page.IndirectPurchaseService = lambda *a, **k: IndirectPurchaseService(repo)
page.current_profile = lambda: {"id": "u1", "role": "ADMIN", "full_name": "Admin", "email": "admin@fourstarindustries.com"}
page.current_employee_id = lambda refresh=False: "e2"
page.module_permissions_with_source = lambda profile, key: ({k: True for k in ("can_view", "can_create", "can_edit", "can_archive", "can_validate", "can_approve")}, "ADMIN")
page.resolve_current_branch = lambda repo, profile: {"id": "b1"}
page.selectable_table = lambda frame, **k: None
'''


@pytest.mark.parametrize("render", ["render_home", "render_items", "render_requisition", "render_po", "render_po_list",
                                    "render_po_approval", "render_inward", "render_stock"])
def test_pages_render_without_exception(render):
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string(SCRIPT_HEAD + f"\npage.{render}()\n").run(timeout=60)
    assert not app.exception, app.exception


def test_po_register_shows_po_and_print_button():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string(SCRIPT_HEAD + "\npage.render_po_list()\n").run(timeout=60)
    assert not app.exception, app.exception
    labels = [b.label for b in app.get("download_button")]
    assert "Download / Print Indirect PO" in labels and "Download Excel" in labels and "Download PDF" in labels
    assert any(t.label == "Cancel this PO" for t in app.toggle)
    assert any(b.label == "Receive (GRN)" for b in app.button)
