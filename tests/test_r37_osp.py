"""R37 · OSP Material Out returnable delivery challan, vendor returns, OSP MetLAB pending queue and
OSP layout process fallback (pre-inward + post-receipt, Dimensional + MetLAB)."""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

import pytest

from core.osp_challan import (OSPChallanService, balance_summary, build_print_payload, challan_notifier, challan_pages, challan_totals,
                              email_body, merge_returns, next_number_from, osp_challan_pdf_bytes)
from core.osp_service import OSPService

ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "supabase/migrations/20261024110000_qcms_r37_osp.sql"


def read(rel: str) -> str:
    return (ROOT / rel).read_text()


class FakeRepo:
    def __init__(self, tables=None):
        self.tables = deepcopy(tables or {}); self.calls = []

    def select(self, table, eq=None, in_=None, **kwargs):
        rows = [r for r in self.tables.get(table, []) if all(str(r.get(k)) == str(v) for k, v in (eq or {}).items())
                and all(str(r.get(k)) in {str(x) for x in v} for k, v in (in_ or {}).items())]
        return deepcopy(rows)

    def get(self, table, rid):
        return next((deepcopy(r) for r in self.tables.get(table, []) if str(r.get("id")) == str(rid)), None)

    def insert(self, table, payload):
        row = {"id": f"{table}-{len(self.tables.get(table, [])) + 1}", **payload}
        self.tables.setdefault(table, []).append(row); self.calls.append(("insert", table, dict(payload)))
        return deepcopy(row)

    def update(self, table, rid, payload):
        for row in self.tables.get(table, []):
            if str(row.get("id")) == str(rid):
                row.update(payload)
        self.calls.append(("update", table, rid, dict(payload)))
        return self.get(table, rid)

    def delete(self, table, rid):
        self.tables[table] = [r for r in self.tables.get(table, []) if str(r.get("id")) != str(rid)]

    def rpc(self, name, params=None):
        raise RuntimeError("RPC functions are unavailable in controlled preview mode.")


JOB = {"id": "j1", "material_out_id": "g1", "material_out_line_no": 1, "part_id": "p1", "vendor_id": "v1", "process_id": "pr1",
       "quantity_dispatched": 1500, "heat_number": "A42539", "osp_batch_code": "OSP-D9-2026-00081-BATCH", "process_name": "Case Carburizing",
       "part_number": "40257237", "part_name": "SIDE GEAR", "dispatch_challan": "JWCD901026"}
TABLES = {
    "osp_challans": [{"id": "c1", "material_out_id": "g1", "challan_number": "JWCD901026", "challan_date": "2026-10-19", "branch_id": "b1", "vendor_id": "v1",
                      "user_name": "Suresh Jadhav", "ship_via": "Road", "incoterm": "EXWORKS", "delivery_date": "2026-10-19", "payment_term": "NET 30 DAYS AFTER GRN",
                      "vehicle_number": "MH04KF0030", "shipping_amount": 0, "other_amount": 0}],
    "osp_challan_lines": [{"id": "l1", "challan_id": "c1", "osp_job_id": "j1", "unit_price": 100, "gst_percent": 18, "uom": "NOS", "batch_code": "LC", "line_no": 1}],
    "osp_challan_returns": [],
    "osp_jobs": [JOB],
    "v_qsms_osp_register": [JOB],
    "parts": [{"id": "p1", "part_number": "40257237", "fsi_part_number": "7237-10", "part_name": "SIDE GEAR", "customer_id": "cust", "hsn_sac_code": "8708"}],
    "company_branches": [{"id": "b1", "branch_name": "Four Star Industries Private Limited D9", "plant_code": "D9", "gstin": "27AAGCF3769A1ZP", "email": "orders@fourstarindustries.com", "status": "ACTIVE", "is_default": True}],
    "parties": [{"id": "v1", "party_name": "PUNE METAL TREATER PVT LTD", "email": "pmt@example.com", "tax_identifier": "27AANCP7622B1ZE"}],
    "osp_receipts": [{"id": "r1", "osp_job_id": "j1", "receipt_number": "OSPR-1", "receipt_challan": "PMT-501", "receipt_date": "2026-10-25", "quantity_received": 700}],
}


# ------------------------------------------------------------------------------------------- migration
def test_migration_is_additive_and_has_contracts():
    sql = MIGRATION.read_text()
    low = sql.lower()
    for token in ("create table if not exists public.osp_challans", "create table if not exists public.osp_challan_lines",
                  "create table if not exists public.osp_challan_returns", "qcms_next_osp_challan_number", "qcms_osp_layout_acceptable",
                  "create or replace function public.enforce_osp_dimensional_source", "create or replace function public.enforce_osp_metlab_source",
                  "enable row level security", "revoke all on public.%i from anon", "tenant_id uuid not null default public.current_tenant_id()",
                  "osp_challan_returns_receipt_uq"):
        assert token.lower() in low, token
    assert "drop " not in low and "truncate" not in low and "delete from" not in low
    assert not re.search(r"^\s*update\s+public\.", low, re.M)
    # the layout check is the only change in the copied trigger definitions
    assert low.count("qcms_osp_layout_acceptable(plan_row.id, job_row.id") == 2
    assert "plan_row.process_id<>job_row.process_id" not in low
    files = read("supabase/auto_apply.json")
    assert "supabase/migrations/20261024110000_qcms_r37_osp.sql" in files


# ------------------------------------------------------------------------------------------- numbering / maths
def test_challan_number_series_per_branch():
    assert next_number_from(["JWCD901026", "JWCD901033", "jwcd900999", "CH-1", None], "D9") == "JWCD901034"
    assert next_number_from([], "THANE") == "JWCTHANE00001"
    assert next_number_from(["JWCD901033"], "D-9") == "JWCD901034"


def test_returns_merge_never_double_counts_osp_inward():
    receipts = [{"id": "r1", "osp_job_id": "j1", "receipt_number": "R1", "receipt_challan": "VC-1", "receipt_date": "2026-10-25", "quantity_received": 700}]
    entries = [{"id": "e1", "osp_receipt_id": "r1", "vehicle_number": "MH12", "quantity_returned": 999, "quantity_rejected": 5},
               {"id": "e2", "vendor_challan_number": "VC-2", "vendor_challan_date": "2026-10-27", "quantity_returned": 300},
               {"id": "e3", "osp_receipt_id": "gone", "quantity_returned": 50}]
    rows = merge_returns(receipts, entries)
    assert [r["qty_returned"] for r in rows] == [700, 300]
    assert rows[0]["vehicle_no"] == "MH12" and rows[0]["challan_no"] == "VC-1" and rows[0]["source"] == "OSP Inward"
    assert balance_summary(1500, rows) == {"dispatched": 1500, "returned": 1000, "rejected": 5, "balance": 495}
    assert balance_summary(10, [{"qty_returned": 12}])["balance"] == 0


def test_totals_and_pages():
    t = challan_totals([{"quantity": 1500, "unit_price": 100, "gst_percent": 18}], shipping=50, other=0)
    assert t == {"subtotal": 150000.0, "tax": 27000.0, "shipping": 50.0, "other": 0.0, "total": 177050.0}
    assert [len(p) for p in challan_pages(range(6))] == [6]
    assert [len(p) for p in challan_pages(range(7))] == [6, 1]
    assert [len(p) for p in challan_pages(range(40))] == [13, 19, 8]


# ------------------------------------------------------------------------------------------- payload / PDF / e-mail
def _service():
    return OSPChallanService(FakeRepo(TABLES))


def test_print_payload_is_vendor_safe_and_uses_fsi_item():
    payload = _service().print_payload_for("c1")
    item = payload["lines"][0]
    assert item["item_code"] == "7237-10" and item["batch_code"] == "LC" and item["heat_number"] == "A42539"
    flat = repr(payload)
    assert "40257237" not in flat and "cust" not in flat
    assert payload["totals"]["total"] == 177000.0
    assert payload["summary"]["returned"] == 700 and payload["summary"]["balance"] == 800
    assert payload["process_text"] == "CASE CARBURIZING" and payload["plant"]["gstin"] == "27AAGCF3769A1ZP"


def _pdf_text(pdf: bytes) -> str:
    from io import BytesIO
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)


def test_pdf_has_reference_layout_sections_and_no_customer_part():
    pdf = osp_challan_pdf_bytes(_service().print_payload_for("c1"))
    assert pdf.startswith(b"%PDF")
    text = _pdf_text(pdf)
    for token in ("RETURNABLE DELIVERY CHALLAN", "PLANT ADDRESS", "VENDOR", "SHIP TO", "REFERENCE DETAILS", "OLD PO DETAILS", "USER", "PAYMENT TERM",
                  "PROCESS TO BE DONE", "Supplier Details for Returnable Challan", "Authorised Signatory", "JWCD901026", "PMT-501", "Batch Code :- LC",
                  "7237-10", "177,000.00"):
        assert token in text, token
    assert "40257237" not in text


def test_pdf_prints_several_returns_as_table():
    payload = build_print_payload(TABLES["osp_challans"][0], [dict(JOB, fsi_part_number="7237-10")], TABLES["osp_challan_lines"],
                                  merge_returns(TABLES["osp_receipts"], [{"id": "e2", "vendor_challan_number": "VC-2", "vendor_challan_date": "2026-10-27", "quantity_returned": 300}]))
    text = _pdf_text(osp_challan_pdf_bytes(payload))
    assert "Vendor Challan" in text and "VC-2" in text and "Balance 500" in text


def test_email_notifier_attaches_challan_pdf():
    notifier = challan_notifier(FakeRepo(TABLES))
    name, content = notifier._generated_report("osp_challans", "c1")
    assert name == "JWCD901026.pdf" and content.startswith(b"%PDF")
    subject, body = email_body(_service().print_payload_for("c1"))
    assert "JWCD901026" in subject and "7237-10" in body and "40257237" not in body
    page = read("app_pages/osp_transactions.py")
    assert 'notifier.notify("OSP_MATERIAL_OUT_CHALLAN"' in page and 'related_table="osp_challans"' in page and "include_generated_pdf=True" in page


# ------------------------------------------------------------------------------------------- service writes
def test_save_creates_header_and_lines_and_return_rules():
    repo = FakeRepo({**TABLES, "osp_challans": [], "osp_challan_lines": []})
    svc = OSPChallanService(repo)
    saved = svc.save({"is_group": True, "material_out_id": "g9"}, {"challan_number": " jwcd901040 ", "shipping_amount": 10},
                     [{"osp_job_id": "j1", "unit_price": 12.5, "gst_percent": 18, "batch_code": "LC"}])
    assert repo.tables["osp_challans"][0]["material_out_id"] == "g9" and repo.tables["osp_challans"][0]["challan_number"] == "JWCD901040"
    assert repo.tables["osp_challan_lines"][0]["challan_id"] == saved["id"]
    with pytest.raises(ValueError):
        svc.save({"id": "j1"}, {"challan_number": ""}, [])
    with pytest.raises(ValueError):
        svc.clean_lines([{"osp_job_id": "j1", "gst_percent": 120}])
    linked = svc.clean_return({"osp_receipt_id": "r1", "quantity_returned": 700, "vehicle_number": "MH12"})
    assert linked["quantity_returned"] == 0.0  # receipt quantity is the returned quantity -> never stored twice
    with pytest.raises(ValueError):
        svc.clean_return({"quantity_returned": 0})
    with pytest.raises(ValueError):
        svc.clean_return({"quantity_returned": 5})  # needs vendor challan or invoice number


def test_osp_inward_receipt_links_to_challan():
    repo = FakeRepo(TABLES)
    svc = OSPChallanService(repo)
    assert svc.link_receipt(JOB, "r1") is None  # nothing extra to store
    svc.link_receipt(JOB, "r1", vehicle_number="MH12AB1234", quantity_rejected=3)
    svc.link_receipt(JOB, "r1", vehicle_number="MH12AB9999")
    entries = repo.tables["osp_challan_returns"]
    assert len(entries) == 1 and entries[0]["osp_receipt_id"] == "r1" and entries[0]["vehicle_number"] == "MH12AB9999"
    page = read("app_pages/osp_transactions.py")
    assert "link_receipt(job" in page and "Return Vehicle No" in page


def test_register_rows_balance():
    from app_pages.osp_transactions import material_out_register
    docs = [{"id": "g1", "material_out_number": "OUT-1", "dispatch_challan": "JWCD901026", "quantity_dispatched": 1500, "quantity_received": 700, "vendor_name": "PMT"}]
    rows = material_out_register(docs, TABLES["osp_challans"], [{"challan_id": "c1", "osp_receipt_id": None, "quantity_returned": 300, "quantity_rejected": 0},
                                                              {"challan_id": "c1", "osp_receipt_id": "r1", "quantity_returned": 0, "quantity_rejected": 5}])
    assert rows[0]["Returned pcs"] == 1000 and rows[0]["Rejected pcs"] == 5 and rows[0]["Balance pcs"] == 495 and rows[0]["Challan"] == "Saved"


# ------------------------------------------------------------------------------------------- bug 2: MetLAB pending queue
def _queue_service(rows):
    obj = OSPService.__new__(OSPService)
    obj.repo = FakeRepo({"lab_tests": [], "inspection_plans": []})
    obj.register = lambda: deepcopy(rows)
    return obj


def test_completed_batches_leave_pending_receipt_metlab_queue():
    rows = [
        # live case OSP-D9-2026-00020: released (COMPLETED / receipt gate ACCEPTED) but its draft receipt MetLAB was deleted later
        {"id": "done", "status": "COMPLETED", "receipt_quality_disposition": "ACCEPTED", "metlab_required": True, "quantity_dispatched": 1500,
         "quantity_received": 1500, "receipt_metlab_disposition": "PENDING", "sample_gate_status": "ACCEPTED"},
        {"id": "open", "status": "PART_RECEIVED", "receipt_quality_disposition": "PENDING", "metlab_required": True, "quantity_dispatched": 100,
         "quantity_received": 100, "receipt_metlab_disposition": "PENDING", "sample_gate_status": "ACCEPTED"},
    ]
    svc = _queue_service(rows)
    assert [r["id"] for r in svc.jobs_for_inspection("OSP_RECEIPT", "METLAB")] == ["open"]
    assert [r["id"] for r in svc.released_without_report("OSP_RECEIPT", "METLAB")] == ["done"]
    assert OSPService.inspection_gate_closed({"sample_gate_status": "ACCEPTED"}, "OSP_SAMPLE")
    assert not OSPService.inspection_gate_closed({"sample_gate_status": "PENDING", "status": "AT_VENDOR"}, "OSP_SAMPLE")


def test_decision_column_shows_actual_report_decision():
    from app_pages.osp_inspections import decision_text
    assert decision_text({}, {"disposition": "ACCEPTED", "status": "DRAFT", "report_number": "MLAB-1"}, "OSP_RECEIPT") == "ACCEPTED · DRAFT · MLAB-1"
    assert decision_text({"status": "COMPLETED", "receipt_quality_disposition": "ACCEPTED"}, None, "OSP_RECEIPT") == "No report · gate ACCEPTED"
    assert decision_text({"status": "PART_RECEIVED"}, None, "OSP_RECEIPT") == "No report yet"
    page = read("app_pages/osp_inspections.py")
    assert "osp.latest_reports(report_type, scope" in page and "released_without_report" in page


# ------------------------------------------------------------------------------------------- bug 3: layout fallback
def test_layout_fallback_matches_db_rule():
    job = {"process_name": "Case Carburizing"}
    named = {"id": "a", "layout_name": "CASE CARBURIZING 9265", "process_id": "qt"}
    assert OSPService.fallback_layouts([named], job, {"cc"}) == [named]
    other = {"id": "b", "layout_name": "Layout 9265", "process_id": "qt", "inspection_stage_id": "osp"}
    assert OSPService.fallback_layouts([other], job, {"cc"}, {"osp"}) == [other]
    assert OSPService.fallback_layouts([other], job, {"cc", "qt"}, {"osp"}) == []  # QT is a real OSP process of the part
    assert OSPService.fallback_layouts([other, dict(other, id="c")], job, {"cc"}, {"osp"}) == []


def test_layout_page_offers_admin_correction_with_password():
    page = read("app_pages/osp_inspections.py")
    assert "Correct layout process to" in page and "verify_current_password(password)" in page
    assert 'inspection.repo.update("inspection_plans", str(plan["id"]), {"process_id": job.get("process_id")})' in page
    assert "OSPService.fallback_layouts(part_layouts, job" in page and "OSPService.name_matched_layouts" in page


# ------------------------------------------------------------------------------------------- UI
def test_returns_section_renders(monkeypatch):
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
from tests.test_r37_osp import FakeRepo, TABLES, JOB
from core.osp_challan import OSPChallanService
from app_pages.osp_transactions import _render_returns_section
svc = OSPChallanService(FakeRepo(TABLES))
doc = {"id": "g1", "is_group": True, "lines": [JOB], **JOB}
_render_returns_section(svc, doc, TABLES["osp_challans"][0], {"can_edit": True, "can_create": True, "can_archive": True})
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception
    assert any(b.label == "Save Return" for b in app.button)


def test_material_out_page_renders_with_challan_module():
    pytest.importorskip("pyarrow")
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
from tests.test_r37_osp import FakeRepo, TABLES, JOB
from core.osp_challan import OSPChallanService
from core.osp_service import OSPService
import app_pages.osp_transactions as page
page.OSPChallanService = lambda *a, **k: OSPChallanService(FakeRepo(TABLES))
class Svc(OSPService):
    def __init__(self): self.repo = FakeRepo(TABLES)
    def register(self): return [dict(JOB, vendor_name="PMT", status="AT_VENDOR", material_out_number="OUT-1")]
st.session_state["osp_mo_open"] = "g1"
page._render_challan_module(Svc(), {"can_edit": True, "can_create": True, "can_archive": True})
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception


def test_material_out_challan_module_without_pyarrow():
    """Same page with st.data_editor stubbed (pyarrow is not installed in the sandbox)."""
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
st.data_editor = lambda frame, **kwargs: frame
from tests.test_r37_osp import FakeRepo, TABLES, JOB
from core.osp_challan import OSPChallanService
from core.osp_service import OSPService
import app_pages.osp_transactions as page
repo = FakeRepo(TABLES)
page.OSPChallanService = lambda *a, **k: OSPChallanService(repo)
page.selectable_table = lambda frame, **k: 0
class Svc(OSPService):
    def __init__(self): self.repo = repo
    def register(self): return [dict(JOB, vendor_name="PMT", status="AT_VENDOR", material_out_number="OUT-1")]
page._render_challan_module(Svc(), {"can_edit": True, "can_create": True, "can_archive": True})
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception, app.exception
    # R38: the challan is entered with the Material Out; saved details are corrected behind one toggle (no second entry form).
    assert "Update Challan" not in [b.label for b in app.button]
    next(t for t in app.toggle if t.label == "Correct challan details").set_value(True).run(timeout=30)
    assert not app.exception, app.exception
    labels = [b.label for b in app.button]
    assert "Update Challan" in labels and "Send Challan E-mail" in labels and "Save Return" in labels
    assert any(b.label == "Download Challan PDF" for b in app.get("download_button"))


def test_material_out_entry_saves_challan_with_lines_without_pyarrow():
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
st.data_editor = lambda frame, **kwargs: frame
from tests.test_r37_osp import FakeRepo, TABLES
from core.osp_challan import OSPChallanService
from core.osp_service import OSPService
import app_pages.osp_transactions as page
repo = st.session_state.setdefault("challan_repo", FakeRepo({**TABLES, "osp_challans": [], "osp_challan_lines": []}))
page.OSPChallanService = lambda *a, **k: OSPChallanService(repo)
class Rpc:
    def rpc(self, name, args):
        st.session_state["submitted_rpc"] = {"name": name, "args": args}
        return {"id": "doc", "material_out_number": "OUT-TEST", "lines": [{"id": f"job{i}", **r} for i, r in enumerate(args["p_lines"])]}
class Service(OSPService):
    def __init__(self): self.repo = Rpc()
    def dispatch_candidates(self): return [{"candidate_key": "INWARD:a", "inward_lot_id": "a", "part_id": "part", "part_number": "PN", "heat_number": "HA", "inward_number": "IN-A", "osp_available_quantity_pcs": 10}]
    def specifications(self, part): return [{"id": "spec", "process_id": "process", "sample_quantity": 1}]
    def processes(self): return {"process": {"id": "process", "process_name": "Case Carburizing"}}
    def vendors(self): return [{"id": "vendor", "party_name": "OSP Vendor"}]
page._render_material_out_entry(Service(), {"can_create": True})
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception, app.exception
    app.multiselect[0].select("INWARD:a").run()
    challan_box = next(x for x in app.text_input if x.label == "Material Out Challan Number")
    assert challan_box.value.startswith("JWCD9")  # suggested from the branch series
    challan_box.set_value("JWCD901050")
    next(x for x in app.text_input if x.label == "Vehicle No").set_value("MH04KF0030")
    next(x for x in app.button if x.label == "Create OSP Material Out").click().run(timeout=30)
    assert not app.exception, app.exception
    assert app.session_state["submitted_rpc"]["args"]["p_header"]["dispatch_challan"] == "JWCD901050"
    repo = app.session_state["challan_repo"]
    header = repo.tables["osp_challans"][0]
    assert header["challan_number"] == "JWCD901050" and header["material_out_id"] == "doc" and header["vehicle_number"] == "MH04KF0030"
    assert header["process_text"] == "CASE CARBURIZING"
    assert repo.tables["osp_challan_lines"][0]["osp_job_id"] == "job0"


# ------------------------------------------------------------------------------------------- R37 core/osp_challan_pdf contract
def _cust_tables():
    job = dict(JOB, customer_name="SECRET CUSTOMER LTD", customer_part_number="CUST-PN-999", customer_po_number="CPO-777")
    return {**TABLES, "osp_jobs": [job], "v_qsms_osp_register": [job],
            "osp_challan_returns": [{"id": "e9", "challan_id": "c1", "osp_job_id": "j1", "vendor_challan_number": "VC-9", "vendor_challan_date": "2026-10-28",
                                     "vendor_invoice_number": "INV-9", "quantity_returned": 100, "vehicle_number": "MH14AB1234", "remarks": "unprocessed"}]}


def test_challan_pdf_by_job_id_contains_vendor_item_heat_process_returns_not_customer():
    from core.osp_challan_pdf import osp_challan_pdf_bytes as by_job
    repo = FakeRepo(_cust_tables())
    text = _pdf_text(by_job("j1", repo))
    for token in ("PUNE METAL TREATER PVT LTD", "7237-10", "A42539", "CASE CARBURIZING", "Return details", "PMT-501", "VC-9", "MH14AB1234"):
        assert token in text, token
    for secret in ("SECRET CUSTOMER", "CUST-PN-999", "CPO-777", "40257237"):
        assert secret not in text, secret
    # the challan id and the Material Out id resolve to the same challan
    assert "JWCD901026" in _pdf_text(by_job("c1", repo))


def test_draft_challan_for_job_without_saved_challan_and_status_watermark():
    from core.osp_challan_pdf import challan_status_label, osp_challan_payload
    tables = {**_cust_tables(), "osp_challans": [], "osp_challan_lines": [], "osp_challan_returns": [], "osp_receipts": []}
    payload = osp_challan_payload("j1", FakeRepo(tables))
    assert payload["status"] == "DRAFT" and payload["saved"] is False and payload["challan_number"] == "JWCD901026"
    assert payload["summary"]["balance"] == 1500 and "SECRET" not in repr(payload)
    assert challan_status_label(payload) == "DRAFT"
    assert challan_status_label({"saved": True, "status": "ISSUED"}) == "ISSUED"
    assert challan_status_label({"saved": True, "status": "CANCELLED"}) == "CANCELLED"
    assert challan_status_label({"saved": True, "status": "ISSUED", "returns": [{}], "summary": {"dispatched": 10, "balance": 0}}) == "CLOSED"
    with pytest.raises(ValueError):
        osp_challan_payload("missing", FakeRepo(tables))


def test_challan_excel_has_sections_and_no_customer():
    from io import BytesIO
    from openpyxl import load_workbook
    from core.osp_challan_pdf import osp_challan_excel_bytes
    wb = load_workbook(BytesIO(osp_challan_excel_bytes("j1", FakeRepo(_cust_tables()))))
    text = " ".join(str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None)
    for token in ("RETURNABLE DELIVERY CHALLAN", "PUNE METAL TREATER PVT LTD", "7237-10", "A42539", "PROCESS TO BE DONE", "CASE CARBURIZING",
                  "SUPPLIER DETAILS FOR RETURNABLE CHALLAN", "VC-9", "Balance Quantity", "Authorised Signatory"):
        assert token in text, token
    assert "SECRET" not in text and "CUST-PN-999" not in text and "40257237" not in text


def test_pages_offer_challan_pdf_excel_and_records_view():
    page = read("app_pages/osp_transactions.py")
    assert "Download Challan Excel" in page and "Download DRAFT Challan PDF" in page
    # R38: OSP Records opens the SAME Material Out / challan record view (Details · Print · Vendor returns) - no separate challan section.
    assert "def render_material_out_record" in page and "render_material_out_record(service, challans, doc, challan, perms, allow_delete=False)" in page
    src = read("core/osp_challan_pdf.py")
    assert "def osp_challan_pdf_bytes(job_id" in src and "_watermark_pdf_bytes" in src and "saved_layout(\"PURCHASE_ORDER\")" in src


def test_material_out_and_records_challan_buttons_apptest():
    """R38: the shared record view (used by OSP Records) offers the challan PDF / Excel on its Print tab."""
    from streamlit.testing.v1 import AppTest
    script = '''
import streamlit as st
st.data_editor = lambda frame, **kwargs: frame
from tests.test_r37_osp import FakeRepo, TABLES, JOB
from core.osp_challan import OSPChallanService
from core.osp_service import OSPService
import app_pages.osp_transactions as page
repo = FakeRepo(TABLES)
svc = OSPChallanService(repo)
class Svc(OSPService):
    def __init__(self): self.repo = repo
doc = OSPService.material_out_documents([dict(JOB, osp_job_number="OSP-1", material_out_number="OUT-1")])[0]
page.render_material_out_record(Svc(), svc, doc, svc.challan_for(doc), {"can_edit": True, "can_create": True, "can_archive": True}, allow_delete=False)
'''
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception, app.exception
    labels = [b.label for b in app.get("download_button")]
    assert "Download Challan PDF" in labels and "Download Challan Excel" in labels
