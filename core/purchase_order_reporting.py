from __future__ import annotations

from io import BytesIO
import re
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile
from pathlib import Path
from typing import Any, Mapping, Sequence

from reportlab.lib.colors import Color, HexColor, black, white
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

try:
    from pypdf import PdfReader, PdfWriter
except Exception:  # pragma: no cover - surfaced at runtime with a clear message
    PdfReader = PdfWriter = None  # type: ignore[assignment]

try:
    import fitz  # PyMuPDF - used only to compact the authoritative terms into portrait pages.
except Exception:  # pragma: no cover - updater installs requirements; fallback preserves portrait source pages.
    fitz = None  # type: ignore[assignment]


NAVY = HexColor("#0B2E63")
LIGHT_BLUE = HexColor("#D7E1F1")
LIGHT_GREY = HexColor("#E4E5E7")
YELLOW = HexColor("#FFF700")
TEXT = HexColor("#202124")

PLANT = {
    "name": "Four Star Industries Private Limited D9",
    "address1": "Plot No.D9, Chakan MIDC PH II",
    "address2": "Bhamboli, Khed",
    "address3": "Pune 410501",
    "tax_identifier": "27AAGCF3769A1ZP",
    "phone": "022 40104412",
    "email": "orders@fourstarindustries.com",
}

DEFAULT_SPECIAL_INSTRUCTIONS = "\n".join(
    [
        "1. PO number should be mentioned on your Invoice.",
        "2. Material must be packed properly to avoid water damage and free from dent, damages & rust.",
        "3. All QC related reports & TCs to be attached with the Invoices.",
        "4. All reports to be sent to quality@fourstarindustries.com",
        "5. Refer STANDARD PURCHASE ORDER TERMS AND CONDITIONS no FSI/703/F04.",
    ]
)


def _s(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _n(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _date(value: Any) -> str:
    text = _s(value)
    if not text:
        return ""
    try:
        y, m, d = text[:10].split("-")
        return f"{d}-{m}-{y}"
    except Exception:
        return text[:10]


def _money(value: Any) -> str:
    return f"{_n(value):,.2f}"


def _draw_text(c: canvas.Canvas, x: float, y: float, text: Any, *, size: float = 8.0, bold: bool = False, color=TEXT, max_width: float | None = None) -> None:
    value = _s(text)
    if max_width and value:
        font = "Helvetica-Bold" if bold else "Helvetica"
        if stringWidth(value, font, size) > max_width:
            while len(value) > 3 and stringWidth(value + "...", font, size) > max_width:
                value = value[:-1]
            value = value.rstrip() + "..."
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    c.drawString(x, y, value)


def _wrap(c: canvas.Canvas, x: float, y: float, text: Any, width: float, *, size: float = 7.4, leading: float = 9.2, bold: bool = False, max_lines: int = 8) -> float:
    font = "Helvetica-Bold" if bold else "Helvetica"
    words = _s(text).split()
    line = ""
    lines: list[str] = []
    for word in words:
        test = f"{line} {word}".strip()
        if stringWidth(test, font, size) <= width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
            if len(lines) >= max_lines:
                break
    if line and len(lines) < max_lines:
        lines.append(line)
    for index, value in enumerate(lines):
        c.setFillColor(TEXT)
        c.setFont(font, size)
        c.drawString(x, y - index * leading, value)
    return y - len(lines) * leading


def _bar(c: canvas.Canvas, x: float, y: float, width: float, title: str, *, height: float = 13) -> None:
    c.setFillColor(NAVY)
    c.rect(x, y - height, width, height, stroke=0, fill=1)
    _draw_text(c, x + 4, y - height + 3.1, title, size=7.2, bold=True, color=white)


def _party_lines(snapshot: Mapping[str, Any]) -> list[str]:
    # Supports both Party Master snapshots (address/city/state/country) and
    # legacy plant snapshots (address1/address2/address3) so historical POs
    # keep printing correctly while v4.14.10 Ship-To can come from any master.
    party_name = _s(snapshot.get("party_name") or snapshot.get("name"))
    address1 = _s(snapshot.get("address") or snapshot.get("address1"))
    locality = ", ".join(v for v in (_s(snapshot.get("city")), _s(snapshot.get("state")), _s(snapshot.get("country"))) if v)
    legacy_locality = ", ".join(v for v in (_s(snapshot.get("address2")), _s(snapshot.get("address3"))) if v)
    return [
        party_name,
        address1,
        locality or legacy_locality,
        _s(snapshot.get("tax_identifier")),
        _s(snapshot.get("contact_person")),
        _s(snapshot.get("phone")),
        _s(snapshot.get("email")),
    ]


def _block(c: canvas.Canvas, x: float, y_top: float, width: float, title: str, lines: list[str], *, max_lines: int = 7) -> float:
    _bar(c, x, y_top, width, title)
    y = y_top - 24
    for line in [v for v in lines if _s(v)][:max_lines]:
        y = _wrap(c, x + 4, y, line, width - 8, size=7.0, leading=8.2, max_lines=2)
    return y


FIRST_HISTORY_ROWS = 3
CONT_HISTORY_ROWS = 28
HISTORY_ONLY_ROWS = 42


def _technical_pairs(item: Mapping[str, Any], *, po_type: str = "FORGING", limit: int = 12) -> list[tuple[str, str]]:
    """Return supplier-facing controlled technical data in display order.

    For Raw Material POs, FSI forging-only standard fields remain suppressed, but
    every supplier-controlled custom row marked Include on PO is printed. This
    preserves the RM-only print contract while allowing rows such as PACKING, RM
    RATE, CONVERSION COST, CUTTING COST and SHOT BLASTING to appear on the PDF.
    """
    raw = item.get("technical_data_snapshot") or []
    if not isinstance(raw, list):
        return []
    po_kind = _s(po_type).upper()
    rm_allowed_standard = {
        "raw material type", "raw material section", "material grade", "section size",
        "supplier rm item code", "supplier lead time",
    }
    rm_forging_only = {
        "supplier forging part no.", "supplier forging part no", "forge wt", "gross wt",
        "input wt", "forging route",
    }
    priority_names = (
        "Raw Material Type", "Raw Material Section", "Material Grade", "Section Size",
        "Supplier RM Item Code", "Supplier Lead Time", "Forge wt", "Gross wt",
        "Input wt", "Forging Route",
    )
    priority = {name.casefold(): idx for idx, name in enumerate(priority_names)}
    prepared: list[tuple[int, str, str]] = []
    for idx, row in enumerate(raw):
        if not isinstance(row, Mapping):
            continue
        heading = _s(row.get("heading")); value = _s(row.get("value"))
        if not heading or not value:
            continue
        heading_key = heading.casefold()
        source = _s(row.get("source")).upper()
        if po_kind == "RAW_MATERIAL":
            # Explicit custom rows always print when they were snapshotted from an
            # Include-on-PO technical-data row. Legacy snapshots without source are
            # treated as custom unless they are known forging-only standard fields.
            if source == "STANDARD" and heading_key not in rm_allowed_standard:
                continue
            if source != "CUSTOM" and heading_key in rm_forging_only:
                continue
        if heading_key == "raw material section":
            heading = "Raw Material Type"; heading_key = "raw material type"
        prepared.append((priority.get(heading_key, 100 + idx), heading, value))
    prepared.sort(key=lambda v: v[0])
    return [(h, v) for _, h, v in prepared[:limit]]


def _technical_height(item: Mapping[str, Any], *, po_type: str) -> float:
    pairs = _technical_pairs(item, po_type=po_type, limit=12)
    rows = max(1, (len(pairs) + 2) // 3)
    return max(42.0, 22.0 + rows * 13.0)


def _price_history_rows(item: Mapping[str, Any]) -> list[dict[str, Any]]:
    raw = item.get("price_history_snapshot") or []
    if not isinstance(raw, list):
        return []
    rows: list[dict[str, Any]] = []
    for row in raw:
        if not isinstance(row, Mapping):
            continue
        rows.append({
            "start_date": row.get("start_date"),
            "end_date": row.get("end_date"),
            "price": row.get("price"),
            "freight": row.get("freight"),
            "tool_cost": row.get("tool_cost"),
            "packing_forwarding": row.get("packing_forwarding"),
            "profit": row.get("profit"),
            "icc_rejection": row.get("icc_rejection"),
            "currency": row.get("currency") or "INR",
            "uom": row.get("uom") or item.get("uom") or "",
            "remarks": row.get("remarks") or row.get("remark") or "",
            "status": row.get("status") or "ACTIVE",
        })
    rows.sort(key=lambda r: (str(r.get("start_date") or ""), str(r.get("end_date") or "9999-12-31")))
    return rows



def _customer_source_rows(item: Mapping[str, Any]) -> list[dict[str, Any]]:
    raw = item.get("customer_source_rows") or []
    if not isinstance(raw, list):
        return []
    return [dict(row) for row in raw if isinstance(row, Mapping)]


def _customer_reference_height(item: Mapping[str, Any]) -> float:
    rows = _customer_source_rows(item)
    shown = min(max(len(rows), 1), 4)
    return 30.0 + shown * 14.0


def _draw_customer_reference(c: canvas.Canvas, item: Mapping[str, Any], *, top: float, left: float, right: float, page_width: float) -> float:
    """Print supplier-safe source traceability without customer part/delivery disclosure.

    Customer identity, customer Part Number and customer Delivery Date stay inside QCMS
    genealogy and are intentionally excluded from the supplier-facing Purchase Order.
    """
    total_w = page_width - left - right
    rows = _customer_source_rows(item)
    shown = rows[:4]
    height = _customer_reference_height(item)
    bottom = top - height
    _bar(c, left, top, total_w, "PO SOURCE REFERENCE", height=15)
    y = top - 15
    widths = [130, 55, 245, 72, total_w - 502]
    titles = ["CUSTOMER PO NO.", "POS", "PART DESCRIPTION", "QTY", "UOM"]
    x = left
    for sw, title in zip(widths, titles):
        c.setFillColor(HexColor("#E5E7EB")); c.rect(x, y-14, sw, 14, stroke=1, fill=1)
        _draw_text(c, x+3, y-9.8, title, size=5.0, bold=True, color=HexColor("#111827"), max_width=sw-6)
        x += sw
    y -= 14
    if not shown:
        c.setFillColor(white); c.rect(left, y-14, total_w, 14, stroke=1, fill=1)
        _draw_text(c, left+4, y-9.8, "No linked order / schedule source stored for this historical PO line.", size=5.3, color=HexColor("#6B7280"), max_width=total_w-8)
        y -= 14
    else:
        for row in shown:
            x = left
            values = [
                row.get("customer_po_number"), row.get("po_position"), row.get("part_description"),
                f"{_n(row.get('quantity')):,.3f}".rstrip("0").rstrip("."), row.get("uom"),
            ]
            for sw, value in zip(widths, values):
                c.setFillColor(white); c.rect(x, y-14, sw, 14, stroke=1, fill=1)
                _draw_text(c, x+3, y-9.8, value, size=5.0, max_width=sw-6)
                x += sw
            y -= 14
        if len(rows) > len(shown):
            _draw_text(c, left+4, bottom+3, f"+ {len(rows)-len(shown)} additional linked source(s) retained in QCMS genealogy.", size=5.0, bold=True, color=NAVY, max_width=total_w-8)
    return bottom


# No vertical grid lines in the PO item body; item identity and values use open-column spacing.
def _draw_item_row(c: canvas.Canvas, item: Mapping[str, Any], *, y_top: float, left: float, right: float, page_width: float) -> float:
    widths = [218, 44, 62, 42, 40, 70, 63]
    titles = ["ITEM #", "QTY", "UNIT PRICE", "UNIT", "GST%", "GST AMOUNT", "TOTAL"]
    x = left
    for sw, title in zip(widths, titles):
        _bar(c, x, y_top, sw, title, height=14); x += sw
    row_top = y_top - 14
    row_bottom = row_top - 42
    c.setFillColor(white); c.rect(left, row_bottom, page_width-left-right, 42, stroke=1, fill=1)
    item_display = " ".join(v for v in (_s(item.get("item_no") or item.get("fsi_part_number_snapshot")), _s(item.get("item_description"))) if v)
    vals = [item_display, f"{_n(item.get('quantity')):,.2f}".rstrip("0").rstrip("."), _money(item.get("unit_price")), item.get("uom"), f"{_n(item.get('gst_percent')):g}%", _money(item.get("gst_amount")), _money(item.get("line_total"))]
    x = left
    for idx, (sw, value) in enumerate(zip(widths, vals)):
        if idx == 0:
            _draw_text(c, x+5, row_top-10, value, size=6.7, bold=True, max_width=sw-10)
            part_description = _s(item.get("part_description_master"))
            if part_description:
                _draw_text(c, x+5, row_top-21, f"Part Description: {part_description}", size=5.7, color=HexColor("#374151"), max_width=sw-10)
            hsn = _s(item.get("hsn_sac_code"))
            if hsn:
                _draw_text(c, x+5, row_top-34, f"HSN / SAC: {hsn}", size=5.6, bold=True, color=HexColor("#4B5563"), max_width=sw-10)
        else:
            _draw_text(c, x+4, row_top-20, value, size=6.7, max_width=sw-8)
        x += sw
    return row_bottom


def _draw_technical(c: canvas.Canvas, item: Mapping[str, Any], *, po_type: str, top: float, bottom: float, left: float, right: float, page_width: float) -> None:
    total_w = page_width-left-right
    is_rm = _s(po_type).upper() == "RAW_MATERIAL"
    title = "RAW MATERIAL DETAILS & SUPPLIER TECHNICAL DATA" if is_rm else "RAW MATERIAL / FORGING PARAMETERS & SUPPLIER TECHNICAL DATA"
    c.setFillColor(HexColor("#F8FAFC")); c.rect(left, bottom, total_w, max(top-bottom, 0), stroke=1, fill=1)
    _bar(c, left, top, total_w, title, height=15)
    pairs = _technical_pairs(item, po_type=po_type, limit=12)
    if not pairs:
        _draw_text(c, left+5, top-29, "No controlled supplier technical data configured", size=5.8, color=HexColor("#6B7280"), max_width=310)
        return

    # Compact 3-pair grid: HEADING | VALUE repeated three times. This keeps the
    # Part Master technical-data table visible on the same supplier item page.
    row_h = 13.0
    inner_left = left + 3
    inner_w = total_w - 6
    pair_w = inner_w / 3.0
    label_w = pair_w * 0.43
    first_row_top = top - 17
    for row_index in range((len(pairs) + 2) // 3):
        y_top = first_row_top - row_index * row_h
        y_bottom = y_top - row_h
        if y_bottom < bottom + 2:
            break
        for col_index in range(3):
            pair_index = row_index * 3 + col_index
            px = inner_left + col_index * pair_w
            if pair_index >= len(pairs):
                c.setFillColor(white); c.rect(px, y_bottom, pair_w, row_h, stroke=1, fill=1)
                continue
            heading, value = pairs[pair_index]
            c.setFillColor(HexColor("#E9EDF1")); c.rect(px, y_bottom, label_w, row_h, stroke=1, fill=1)
            c.setFillColor(white); c.rect(px+label_w, y_bottom, pair_w-label_w, row_h, stroke=1, fill=1)
            _draw_text(c, px+3, y_bottom+4.1, heading, size=5.2, bold=True, color=HexColor("#30363B"), max_width=label_w-6)
            _draw_text(c, px+label_w+3, y_bottom+4.1, value, size=5.2, color=HexColor("#202124"), max_width=pair_w-label_w-6)


def _draw_price_history(c: canvas.Canvas, item: Mapping[str, Any], history: Sequence[Mapping[str, Any]], *, top: float, bottom: float, left: float, right: float, page_width: float, max_rows: int) -> int:
    """Draw all item-wise price revisions in the compact Part Master format.

    The approved print contract follows the supplied table: Start Date / End Date / Price /
    Remark. Closed revisions remain visible and the open-ended row prints as CURRENT.
    """
    _bar(c, left, top, page_width-left-right, f"PRICE REVISION HISTORY · FSI PART {_s(item.get('item_no') or item.get('fsi_part_number_snapshot'))}", height=15)
    y = top - 15
    total_w = page_width-left-right
    widths = [92, 92, 100, total_w-284]
    headers = ["START DATE", "END DATE", "PRICE", "REMARK"]
    x = left
    for sw, title in zip(widths, headers):
        c.setFillColor(HexColor("#E5E7EB"))
        c.rect(x, y-15, sw, 15, stroke=1, fill=1)
        _draw_text(c, x+3, y-10.4, title, size=5.4, bold=True, color=HexColor("#111827"), max_width=sw-6)
        x += sw
    y -= 15
    available_rows = max(int((y-bottom) // 15), 0)
    count = min(len(history), max_rows, available_rows)
    if count == 0 and not history and y-15 >= bottom:
        c.setFillColor(white); c.rect(left, y-15, total_w, 15, stroke=1, fill=1)
        _draw_text(c, left+4, y-10.5, "No controlled price revision history is recorded for this Supplier / FSI Part.", size=5.8, color=HexColor("#6B7280"), max_width=total_w-8)
        return 0
    for row in list(history)[:count]:
        x = left
        vals = [
            _date(row.get("start_date")),
            _date(row.get("end_date")) if row.get("end_date") else "CURRENT",
            _money(row.get("price")) if row.get("price") not in (None, "") else "-",
            _s(row.get("remarks") or row.get("remark")),
        ]
        for sw, value in zip(widths, vals):
            c.setFillColor(white); c.rect(x, y-15, sw, 15, stroke=1, fill=1)
            _draw_text(c, x+3, y-10.5, value, size=5.35, max_width=sw-6)
            x += sw
        y -= 15
    remaining = len(history) - count
    if remaining > 0 and y-12 >= bottom:
        _draw_text(c, left+4, y-9, f"Continued on next controlled item page · {remaining} more revision(s)", size=5.6, bold=True, color=NAVY, max_width=total_w-8)
    return count


def _draw_continuation_header(c: canvas.Canvas, header: Mapping[str, Any], *, title: str) -> tuple[float, float, float, float]:
    w, h = A4; left, right = 28, 28
    logo_path = Path(__file__).resolve().parent.parent / "assets" / "fsi_logo.png"
    if logo_path.exists():
        try: c.drawImage(str(logo_path), left, h-58, width=105, height=39, preserveAspectRatio=True, mask="auto")
        except Exception: pass
    c.setFillColor(black); c.setFont("Helvetica-Bold", 16.5); c.drawRightString(w-right, h-33, title)
    _draw_text(c, w-176, h-51, f"PO # {_s(header.get('po_number'))}", size=7.2, bold=True)
    _draw_text(c, w-176, h-64, f"DATE {_date(header.get('order_date'))}", size=7.2)
    return w, h, left, right


def _history_continuation_pages(c: canvas.Canvas, header: Mapping[str, Any], item: Mapping[str, Any], history: Sequence[Mapping[str, Any]]) -> None:
    remaining = list(history)
    while remaining:
        w, h, left, right = _draw_continuation_header(c, header, title="PURCHASE ORDER · PRICE HISTORY CONTINUED")
        _draw_text(c, left, h-88, f"FSI PART: {_s(item.get('item_no') or item.get('fsi_part_number_snapshot'))} · {_s(item.get('item_description'))}", size=7.2, bold=True, color=NAVY, max_width=w-left-right)
        rendered = _draw_price_history(c, item, remaining, top=h-102, bottom=50, left=left, right=right, page_width=w, max_rows=HISTORY_ONLY_ROWS)
        if rendered <= 0:
            break
        remaining = remaining[rendered:]
        c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica",5.8); c.drawCentredString(w/2,27,"QCMS controlled Purchase Order · Price revision history is item-specific and supplier-specific.")
        c.showPage()


def _first_page_bytes(header: Mapping[str, Any], items: Sequence[Mapping[str, Any]]) -> bytes:
    out = BytesIO(); c = canvas.Canvas(out, pagesize=A4); w, h = A4; left, right = 28, 28
    logo_path = Path(__file__).resolve().parent.parent / "assets" / "fsi_logo.png"
    if logo_path.exists():
        try: c.drawImage(str(logo_path), left, h-58, width=105, height=39, preserveAspectRatio=True, mask="auto")
        except Exception: pass

    c.setFillColor(black); c.setFont("Helvetica-Bold", 24); c.drawRightString(w-right, h-34, "PURCHASE ORDER")
    _draw_text(c, w-137, h-51, "DATE", size=7.4); c.rect(w-91, h-58, 62, 12, stroke=1, fill=0); _draw_text(c, w-86, h-55, _date(header.get("order_date")), size=7.2)
    _draw_text(c, w-137, h-65, "PO #", size=7.4); c.rect(w-91, h-72, 62, 12, stroke=1, fill=0); _draw_text(c, w-86, h-69, header.get("po_number"), size=7.2)

    col_w = 230; left_x = left; right_x = w-right-184; y0 = h-62
    plant = dict(header.get("plant_snapshot") or PLANT)
    _block(c, left_x, y0, col_w, "PLANT / COMPANY BRANCH", [plant.get("name") or plant.get("branch_name"), plant.get("address1"), plant.get("address2"), plant.get("address3"), plant.get("tax_identifier") or plant.get("gstin"), plant.get("phone"), plant.get("email")])
    vendor = dict(header.get("vendor_snapshot") or {}); _block(c, left_x, h-154, col_w, "VENDOR", _party_lines(vendor), max_lines=7)
    _block(c, right_x, h-86, 184, "REFERENCE DETAILS", [f"QUOTATION DATE: {_date(header.get('quotation_date'))}", _s(header.get("quotation_reference"))], max_lines=3)
    _block(c, right_x, h-136, 184, "OLD PO DETAILS", [_s(header.get("old_po_reference"))], max_lines=2)
    ship = dict(header.get("ship_to_snapshot") or PLANT)
    _block(c, right_x, h-174, 184, "SHIP TO", _party_lines(ship), max_lines=7)

    y_strip = h-284; strip_widths = [78,106,82,93,w-left-right-359]
    strip_titles = ["REQUISITIONER","SHIP VIA","INCOTERM","DELIVERY DATE","PAYMENT TERM"]
    strip_values = [header.get("requisitioner"),header.get("ship_via"),header.get("incoterm"),_date(header.get("delivery_date")),header.get("payment_term")]
    x = left
    for sw,title,value in zip(strip_widths,strip_titles,strip_values):
        _bar(c,x,y_strip,sw,title,height=14); c.rect(x,y_strip-34,sw,20,stroke=1,fill=0); _draw_text(c,x+4,y_strip-27,value,size=6.6,max_width=sw-8); x += sw

    # One complete item pocket on the first page. Additional items continue on controlled
    # item pages so each supplier line keeps its own technical data and full price history.
    item = dict(items[0])
    y_item_top = y_strip-45
    row_bottom = _draw_item_row(c, item, y_top=y_item_top, left=left, right=right, page_width=w)
    source_bottom = _draw_customer_reference(c, item, top=row_bottom, left=left, right=right, page_width=w)
    po_type = _s(header.get("po_type") or "FORGING").upper()
    tech_bottom = source_bottom-_technical_height(item, po_type=po_type)
    _draw_technical(c, item, po_type=po_type, top=source_bottom, bottom=tech_bottom, left=left, right=right, page_width=w)
    history = _price_history_rows(item)
    _draw_price_history(c, item, history, top=tech_bottom-4, bottom=245, left=left, right=right, page_width=w, max_rows=FIRST_HISTORY_ROWS)

    _draw_text(c,left+2,236,"Remarks:",size=7.6,bold=True); c.setFillColor(YELLOW); c.rect(left,211,386,21,stroke=0,fill=1); _draw_text(c,left+4,220,header.get("remarks") or "PART WILL BE SUPPLIED AS PER DRAWING.",size=7.0,bold=True,max_width=378)
    c.setFillColor(LIGHT_GREY); c.rect(left,192,386,14,stroke=0,fill=1); _draw_text(c,left+3,196,"Comments or Special Instructions",size=6.8,bold=True); _wrap(c,left+2,181,header.get("special_instructions") or DEFAULT_SPECIAL_INSTRUCTIONS,374,size=6.2,leading=11.3,max_lines=7)
    total_x=417; ytot=205
    for label,val in [("SUBTOTAL",header.get("subtotal")),("CGST 9%",header.get("cgst_amount")),("SGST 9%",header.get("sgst_amount")),("IGST",header.get("igst_amount")),("OTHER",header.get("other_amount"))]:
        _draw_text(c,total_x,ytot,label,size=6.9); c.rect(total_x+55,ytot-5,95,13,stroke=1,fill=0); _draw_text(c,total_x+61,ytot-1,_money(val),size=6.8); ytot -= 16
    c.setFont("Helvetica-Bold",8.0); c.drawString(total_x,ytot,"TOTAL"); c.drawString(total_x+61,ytot,f"INR {_money(header.get('grand_total'))}")
    _draw_text(c,w-143,81,"Authorised Signatory",size=7.0,bold=True); _draw_text(c,w-174,67,plant.get("name") or plant.get("branch_name") or PLANT["name"],size=6.2)
    c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica",6.0)
    c.drawCentredString(w/2,47,"If you have any questions about this purchase order, please contact")
    c.drawCentredString(w/2,36,"FSI · connect@fourstarindustries.com")
    c.showPage(); c.save(); return out.getvalue()


def _continuation_items_bytes(header: Mapping[str, Any], items: Sequence[Mapping[str, Any]]) -> bytes:
    """One clean supplier item per continuation page, including full price history."""
    out = BytesIO(); c = canvas.Canvas(out, pagesize=A4)
    for item_src in items:
        item = dict(item_src)
        w,h,left,right = _draw_continuation_header(c, header, title="PURCHASE ORDER · ITEM CONTINUED")
        row_bottom = _draw_item_row(c, item, y_top=h-92, left=left, right=right, page_width=w)
        source_bottom = _draw_customer_reference(c, item, top=row_bottom, left=left, right=right, page_width=w)
        po_type = _s(header.get("po_type") or "FORGING").upper()
        tech_bottom = source_bottom-_technical_height(item, po_type=po_type)
        _draw_technical(c, item, po_type=po_type, top=source_bottom, bottom=tech_bottom, left=left, right=right, page_width=w)
        history = _price_history_rows(item)
        rendered = _draw_price_history(c, item, history, top=tech_bottom-6, bottom=50, left=left, right=right, page_width=w, max_rows=CONT_HISTORY_ROWS)
        c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica",5.8); c.drawCentredString(w/2,27,"QCMS controlled continuation · Technical data and Price Revision History are item-specific and supplier-specific.")
        c.showPage()
        if rendered < len(history):
            _history_continuation_pages(c, header, item, history[rendered:])
    c.save(); return out.getvalue()


def _terms_with_dynamic_header(terms_path: Path, *, po_number: str, order_date: str) -> list[Any]:
    if PdfReader is None:
        raise RuntimeError("pypdf is required for the controlled Purchase Order terms pages.")
    reader = PdfReader(str(terms_path))
    result = []
    for source_page in reader.pages:
        overlay = BytesIO()
        w = float(source_page.mediabox.width)
        h = float(source_page.mediabox.height)
        c = canvas.Canvas(overlay, pagesize=(w, h))
        # Controlled terms pages are US Letter. Cover the original DATE / PO area
        # from the reference document, then stamp this saved QCMS PO identity.
        c.setFillColor(white)
        # Original FSI/703/F04 Letter template places DATE/PO at x≈492..580,
        # 65..84 points from the top. Replace only that small cell area.
        c.rect(487, h - 89, 111, 31, stroke=0, fill=1)
        _draw_text(c, 492, h - 73, "DATE", size=7.4)
        c.rect(532, h - 75, 52, 12, stroke=1, fill=0)
        _draw_text(c, 536, h - 72, _date(order_date), size=6.8)
        _draw_text(c, 495, h - 84, "PO #", size=7.4)
        c.rect(532, h - 86, 52, 12, stroke=1, fill=0)
        _draw_text(c, 536, h - 83, po_number, size=6.8)
        c.save(); overlay.seek(0)
        overlay_page = PdfReader(overlay).pages[0]
        page = source_page
        page.merge_page(overlay_page)
        result.append(page)
    return result


def _stamped_terms_pdf_bytes(terms_path: Path, *, po_number: str, order_date: str) -> bytes:
    """Return the authoritative terms with the live PO/date stamp applied."""
    if PdfWriter is None:
        raise RuntimeError("pypdf is required for the controlled Purchase Order terms pages.")
    writer = PdfWriter()
    for page in _terms_with_dynamic_header(terms_path, po_number=po_number, order_date=order_date):
        writer.add_page(page)
    out = BytesIO(); writer.write(out); return out.getvalue()


def _compact_terms_portrait(terms_path: Path, *, po_number: str, order_date: str) -> list[Any]:
    """Compact the controlled terms into portrait A4 pages without changing clause wording.

    v4.14.33 removes blank vertical space by clipping the visible text blocks from the
    authoritative FSI/703/F04 source and packing them sequentially on portrait A4. The
    original clause order and visible wording are retained. A known duplicate carry-over
    source page is skipped only when the duplicate signature is present in the template.
    If PyMuPDF is unavailable, the stamped source pages are returned unchanged (still portrait).
    """
    source_pages = _terms_with_dynamic_header(terms_path, po_number=po_number, order_date=order_date)
    if fitz is None:
        return source_pages

    # Open a stamped in-memory copy so every compact page carries the current PO/date.
    stamped = _stamped_terms_pdf_bytes(terms_path, po_number=po_number, order_date=order_date)
    src = fitz.open(stream=stamped, filetype="pdf")
    out = fitz.open()
    a4_w, a4_h = A4
    margin_x = 18.0
    footer_h = 24.0
    gap = 3.0
    usable_w = float(a4_w) - 2 * margin_x

    # The controlled 2023 source contains one accidental carry-over page where clause 6
    # and clauses 1-5 are repeated. Skip it only when both the duplicate and the next-page
    # clause-6 continuation signatures are present. No other page is suppressed.
    skip_pages: set[int] = set()
    if len(src) >= 3:
        duplicate_text = " ".join(src[1].get_text("text").split()).casefold()
        next_text = " ".join(src[2].get_text("text").split()).casefold()
        if (
            "6. supplier quality and development" in duplicate_text
            and "1. acceptance:" in duplicate_text
            and "6. supplier quality and development" in next_text
        ):
            skip_pages.add(1)

    segments: list[tuple[int, Any]] = []
    for pno, page in enumerate(src):
        if pno in skip_pages:
            continue
        for block in page.get_text("blocks", sort=True):
            x0, y0, x1, y1, text, *_ = block
            if not " ".join(str(text or "").split()):
                continue
            # Remove the repeated per-page PURCHASE ORDER header and controlled footer.
            # The full title area on the first source page is re-used once as the compact
            # terms title. All other visible terms content remains in sequence.
            if y0 >= 738:
                continue
            if pno == 0 and y0 < 150:
                continue
            clip_y0 = max(float(y0), 86.0)
            clip_y1 = min(float(y1), 735.0)
            if clip_y1 - clip_y0 < 4.0:
                continue
            rect = fitz.Rect(
                max(20.0, float(x0) - 2.0),
                max(84.0, clip_y0 - 1.5),
                min(595.0, float(x1) + 2.0),
                min(735.0, clip_y1 + 1.5),
            )
            segments.append((pno, rect))

    page = None
    y = 0.0
    page_no = 0

    def _new_compact_page(first: bool) -> tuple[Any, float]:
        nonlocal page_no
        target_page = out.new_page(width=float(a4_w), height=float(a4_h))
        page_no += 1
        if first:
            # Full title + plant + dynamic PO/date from terms source page 1.
            header_clip = fitz.Rect(20.0, 35.0, 595.0, 148.0)
            target = fitz.Rect(
                margin_x, 18.0, float(a4_w) - margin_x,
                18.0 + header_clip.height * (usable_w / header_clip.width),
            )
            target_page.show_pdf_page(target, src, 0, clip=header_clip)
            return target_page, target.y1 + 7.0
        # Use a source page whose header is cleanly separated from the first clause.
        normal_header_source = 9 if len(src) > 9 else 0
        header_clip = fitz.Rect(20.0, 35.0, 595.0, 84.5)
        target = fitz.Rect(
            margin_x, 15.0, float(a4_w) - margin_x,
            15.0 + header_clip.height * (usable_w / header_clip.width),
        )
        target_page.show_pdf_page(target, src, normal_header_source, clip=header_clip)
        return target_page, target.y1 + 6.0

    for pno, clip in segments:
        # 0.86 retains comfortable printed readability while reducing the current terms
        # from 12 physical source pages to about 7 portrait A4 pages on the 2023 template.
        scale = min(0.86, usable_w / float(clip.width))
        h = float(clip.height) * scale
        if page is None:
            page, y = _new_compact_page(True)
        bottom_limit = float(a4_h) - footer_h - 10.0
        if y + h > bottom_limit:
            page, y = _new_compact_page(False)
        target = fitz.Rect(margin_x, y, margin_x + float(clip.width) * scale, y + h)
        page.show_pdf_page(target, src, pno, clip=clip)
        y = target.y1 + gap

    total_pages = len(out)
    for idx, target_page in enumerate(out):
        fy = float(a4_h) - 12.0
        target_page.insert_text((18.0, fy), "STANDARD PURCHASE ORDER TERMS AND CONDITIONS", fontsize=6.3, fontname="helv")
        target_page.insert_text((237.0, fy), "FSI/703/F04", fontsize=6.3, fontname="helv")
        target_page.insert_text((315.0, fy), "1st April 2023", fontsize=6.3, fontname="helv")
        target_page.insert_text((500.0, fy), f"Terms {idx + 1} of {total_pages}", fontsize=6.3, fontname="helv")

    packed = out.tobytes(garbage=3, deflate=True)
    out.close(); src.close()
    return list(PdfReader(BytesIO(packed)).pages)


def _po_pdf_filename(header: Mapping[str, Any], position: int, used_names: set[str]) -> str:
    """Make a cross-platform, collision-safe basename; never a path inside the ZIP."""
    number = str(header.get("po_number") or header.get("id") or f"PO_{position:03d}").strip()
    number = unicodedata.normalize("NFKC", number)
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", number).strip(" ._")[:100] or f"PO_{position:03d}"
    if stem.lower().endswith(".pdf"):
        stem = stem[:-4] or f"PO_{position:03d}"
    # Reserved Windows device names are invalid even with an extension.
    if stem.split(".", 1)[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
        stem = "PO_" + stem
    candidate = f"{stem}.pdf"
    suffix = 2
    while candidate.casefold() in used_names:
        candidate = f"{stem}__{suffix}.pdf"
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def purchase_order_pdf_files(
    purchase_orders: Sequence[tuple[Mapping[str, Any], Sequence[Mapping[str, Any]]]],
    *, copies_per_order: int = 1, terms_path: str | Path | None = None,
) -> list[tuple[str, bytes]]:
    """Generate one complete PDF per distinct PO, with no inter-order merging.

    Multiple print copies repeat pages *inside that PO's file only*. Duplicate
    selections of the same saved PO id are ignored, but different PO ids that
    share a display number still receive separate, collision-safe filenames.
    A missing/invalid document stops generation; never return an incomplete ZIP
    presented as a successful full batch. This function does not mutate records.
    """
    if isinstance(copies_per_order, bool) or not isinstance(copies_per_order, int) or not 1 <= copies_per_order <= 5:
        raise ValueError("Copies per PO must be a whole number between 1 and 5.")
    if not purchase_orders:
        raise ValueError("Select at least one Purchase Order.")
    files: list[tuple[str, bytes]] = []
    used_names: set[str] = set()
    seen_ids: set[str] = set()
    for position, (header, items) in enumerate(purchase_orders, 1):
        if not isinstance(header, Mapping) or not header:
            raise ValueError(f"Selected PO {position}: the Purchase Order no longer exists or is not accessible.")
        po_id = str(header.get("id") or "").strip()
        if po_id and po_id in seen_ids:
            continue
        po_number = str(header.get("po_number") or po_id or position)
        if not items:
            raise ValueError(f"{po_number}: no printable Purchase Order items were found.")
        pdf = purchase_order_pdf_bytes(header, list(items), terms_path=terms_path)
        if copies_per_order > 1:
            # Never pass more than this one PO to a combined document writer.
            if PdfReader is None or PdfWriter is None:
                raise RuntimeError("pypdf is required to generate print copies.")
            reader = PdfReader(BytesIO(pdf))
            writer = PdfWriter()
            for _ in range(copies_per_order):
                for page in reader.pages:
                    writer.add_page(page)
            out = BytesIO()
            writer.write(out)
            pdf = out.getvalue()
        files.append((_po_pdf_filename(header, position, used_names), pdf))
        if po_id:
            seen_ids.add(po_id)
    return files


def purchase_order_files_zip_bytes(files: Sequence[tuple[str, bytes]]) -> bytes:
    """Package already generated per-PO PDFs without changing their contents."""
    if not files:
        raise ValueError("No Purchase Order PDFs were generated.")
    buffer = BytesIO()
    names: set[str] = set()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED, compresslevel=6) as archive:
        for name, pdf in files:
            if (not name or "/" in name or "\\" in name or name.startswith(".")
                    or not name.lower().endswith(".pdf") or name.casefold() in names):
                raise ValueError("Purchase Order PDF filenames must be safe, unique basenames.")
            if not isinstance(pdf, bytes) or not pdf.startswith(b"%PDF-"):
                raise ValueError(f"{name}: generated content is not a PDF.")
            names.add(name.casefold())
            archive.writestr(name, pdf)
    return buffer.getvalue()


def batch_purchase_order_pdf_bytes(
    purchase_orders: Sequence[tuple[Mapping[str, Any], Sequence[Mapping[str, Any]]]],
    *, copies_per_order: int = 1, terms_path: str | Path | None = None,
) -> bytes:
    """Combine multiple controlled Purchase Orders into one print-ready PDF.

    Each selected PO remains a complete controlled document with its own portrait terms.
    copies_per_order supports physical duplicate copies without requiring repeated downloads.
    """
    if PdfReader is None or PdfWriter is None:
        raise RuntimeError("pypdf is not installed. Add pypdf to requirements.txt.")
    copies = max(1, min(int(copies_per_order or 1), 10))
    writer = PdfWriter()
    for header, items in purchase_orders:
        pdf = purchase_order_pdf_bytes(header, list(items), terms_path=terms_path)
        pages = list(PdfReader(BytesIO(pdf)).pages)
        for _copy in range(copies):
            for page in pages:
                writer.add_page(page)
    out = BytesIO(); writer.write(out); return out.getvalue()


def _po_approval_watermark(header: Mapping[str, Any]) -> str:
    return "APPROVED" if _s(header.get("approval_status")).upper() == "APPROVED" else "PENDING APPROVAL"


def _watermark_pdf_bytes(pdf_data: bytes, label: str) -> bytes:
    """Apply a light diagonal approval watermark to every PDF page."""
    if PdfReader is None or PdfWriter is None:
        raise RuntimeError("pypdf is required for Purchase Order watermarks.")
    reader = PdfReader(BytesIO(pdf_data))
    writer = PdfWriter()
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_buf = BytesIO()
        wc = canvas.Canvas(overlay_buf, pagesize=(width, height))
        wc.saveState()
        try:
            wc.setFillAlpha(0.055)
        except Exception:
            pass
        wc.setFillColor(HexColor("#98A2B3"))
        wc.setFont("Helvetica-Bold", 46 if label == "APPROVED" else 34)
        wc.translate(width / 2.0, height / 2.0)
        wc.rotate(38)
        text_width = stringWidth(label, "Helvetica-Bold", 46 if label == "APPROVED" else 34)
        wc.drawString(-text_width / 2.0, -10, label)
        wc.restoreState()
        wc.save(); overlay_buf.seek(0)
        overlay_page = PdfReader(overlay_buf).pages[0]
        page.merge_page(overlay_page)
        writer.add_page(page)
    out = BytesIO(); writer.write(out); return out.getvalue()


def purchase_order_excel_bytes(header: Mapping[str, Any], items: Mapping[str, Any] | Sequence[Mapping[str, Any]]) -> bytes:
    """Return a supplier-facing Purchase Order workbook with print-page watermark headers.

    The Excel print intentionally omits customer identity, customer Part Number and
    customer Delivery Date while preserving Customer PO / Position traceability,
    Part Master description, purchased item identity and controlled commercial data.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    normalized = [dict(items)] if isinstance(items, Mapping) else [dict(v) for v in items]
    if not normalized:
        raise ValueError("At least one Purchase Order line is required.")
    wb = Workbook(); ws = wb.active; ws.title = "Purchase Order"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    watermark = _po_approval_watermark(header)
    for header_obj in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
        header_obj.center.text = watermark
        header_obj.center.font = "Arial,Bold"
        header_obj.center.size = 28
        header_obj.center.color = "D9DDE3"
    ws.oddFooter.center.text = "Four Star Industries Pvt. Ltd. | QCMS Controlled Purchase Order"
    ws.evenFooter.center.text = ws.oddFooter.center.text
    ws.firstFooter.center.text = ws.oddFooter.center.text

    navy = "0B2E63"; light = "E8EEF7"; border = Side(style="thin", color="B8C0CC")
    ws.merge_cells("A1:G1"); ws["A1"] = "PURCHASE ORDER"; ws["A1"].font = Font(size=16, bold=True, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor=navy); ws["A1"].alignment = Alignment(horizontal="center")
    header_rows = [
        ("PO Number", header.get("po_number"), "Order Date", _date(header.get("order_date"))),
        ("Supplier", (header.get("supplier_snapshot") or {}).get("party_name") or header.get("supplier_name"), "PO Type", _s(header.get("po_type")).replace("_", " ")),
        ("Approval Status", watermark, "Currency", header.get("currency") or "INR"),
    ]
    r=3
    for l1,v1,l2,v2 in header_rows:
        ws[f"A{r}"]=l1; ws[f"B{r}"]=v1 or ""; ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
        ws[f"D{r}"]=l2; ws[f"E{r}"]=v2 or ""; ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=7)
        for cell in (ws[f"A{r}"],ws[f"D{r}"]): cell.font=Font(bold=True,color=navy)
        r+=1
    r+=1
    titles=["ITEM / FSI PART", "PART DESCRIPTION", "QTY", "UNIT", "UNIT PRICE", "GST %", "TOTAL"]
    for cidx,title in enumerate(titles,1):
        cell=ws.cell(r,cidx,title); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(bottom=border)
    r+=1
    for item in normalized:
        vals=[item.get("item_no") or item.get("fsi_part_number_snapshot"), item.get("part_description_master") or item.get("item_description"), item.get("quantity"), item.get("uom"), item.get("unit_price"), item.get("gst_percent"), item.get("line_total")]
        for cidx,val in enumerate(vals,1):
            cell=ws.cell(r,cidx,val or 0 if cidx in (3,5,6,7) else val or ""); cell.alignment=Alignment(vertical="top",wrap_text=True); cell.border=Border(bottom=border)
        # Supplier-safe source traceability: no customer part number or customer delivery date.
        src=_customer_source_rows(item)
        if src:
            r+=1; ws.cell(r,1,"Source Ref").font=Font(bold=True,color=navy)
            refs=[]
            for row in src:
                refs.append("PO {po} | Pos {pos} | {desc} | Qty {qty} {uom}".format(po=_s(row.get("customer_po_number")) or "-", pos=_s(row.get("po_position")) or "-", desc=_s(row.get("part_description")) or "-", qty=f"{_n(row.get('quantity')):g}", uom=_s(row.get("uom"))))
            ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7); ws.cell(r,2,"\n".join(refs)); ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
        r+=1
    r+=1; ws.cell(r,1,"Special Instructions").font=Font(bold=True,color=navy); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7); ws.cell(r,2,_s(header.get("special_instructions") or DEFAULT_SPECIAL_INSTRUCTIONS)); ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
    widths={"A":22,"B":34,"C":12,"D":13,"E":15,"F":10,"G":16}
    for col,width in widths.items(): ws.column_dimensions[col].width=width
    ws.freeze_panes="A8"; ws.print_area=f"A1:G{r}"
    out=BytesIO(); wb.save(out); return out.getvalue()


def purchase_order_pdf_bytes(header: Mapping[str, Any], items: Mapping[str, Any] | Sequence[Mapping[str, Any]], *, terms_path: str | Path | None = None) -> bytes:
    """Return the controlled FSI Purchase Order PDF.

    Page 1 prints one supplier-facing item line plus a supplier-safe source reference
    table containing Customer PO Number, PO Position, Part Description and allocated quantity.
    Customer identity, customer Part Number and customer Delivery Date are intentionally
    omitted from the supplier print.
    Each supplier item retains its technical data and Price Revision History. Standard terms
    are compacted only into portrait A4 pages by moving visible controlled content into unused
    white space; clause wording and order remain unchanged.
    """
    if PdfReader is None or PdfWriter is None:
        raise RuntimeError("pypdf is not installed. Add pypdf to requirements.txt.")
    normalized = [dict(items)] if isinstance(items, Mapping) else [dict(v) for v in items]
    if not normalized:
        raise ValueError("At least one Purchase Order line is required.")
    first = _first_page_bytes(header, normalized)
    writer = PdfWriter(); writer.add_page(PdfReader(BytesIO(first)).pages[0])
    first_history = _price_history_rows(normalized[0])
    if len(first_history) > FIRST_HISTORY_ROWS:
        overflow = BytesIO(); overflow_canvas = canvas.Canvas(overflow, pagesize=A4)
        _history_continuation_pages(overflow_canvas, header, normalized[0], first_history[FIRST_HISTORY_ROWS:])
        overflow_canvas.save(); overflow.seek(0)
        for page in PdfReader(overflow).pages: writer.add_page(page)
    if len(normalized) > 1:
        continuation = PdfReader(BytesIO(_continuation_items_bytes(header, normalized[1:])))
        for page in continuation.pages: writer.add_page(page)
    path = Path(terms_path) if terms_path else Path(__file__).resolve().parent.parent / "templates" / "FSI_STANDARD_PO_TERMS_2023.pdf"
    if path.exists():
        for page in _compact_terms_portrait(path, po_number=_s(header.get("po_number")), order_date=_s(header.get("order_date"))):
            writer.add_page(page)
    out = BytesIO(); writer.write(out)
    return _watermark_pdf_bytes(out.getvalue(), _po_approval_watermark(header))

