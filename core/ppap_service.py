"""R38 · PPAP Document Manager service (linked to APQP projects).

Data model (migration 20261028120000_qcms_r38_ppap.sql, module key ``PPAP``):

* ``ppap_element_master``      - the tenant's default PPAP element template (32 elements of the customer PPAP tracker,
                                 editable: add / rename / reorder / default responsible department / default lead days).
* ``ppap_customer_template_elements`` - optional customer-specific element template (used instead of the default one).
* ``ppap_packages``            - one PPAP package per part / customer PPAP no., optionally linked to an APQP project
                                 (``apqp_project_id`` -> ``ppap_projects``, the APQP register).
* ``ppap_package_elements``    - element status / due date / responsible employee / completed date per package.
* ``ppap_element_documents``   - uploaded files per element with revision history (``is_latest`` = current file).
* ``ppap_status_notes``        - free-text status log per package (the tracker's "status" sheet).
* ``ppap_customer_submissions`` - every package e-mailed / linked / downloaded for the customer.

Pure helpers (status mapping, progress %, due / overdue logic, tracker Excel import / template, ZIP layout,
submission method) are module functions so they can be unit tested without a database.
"""
from __future__ import annotations

import hashlib
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Iterable, Mapping, Sequence

# ----------------------------------------------------------------------------- constants
DEFAULT_ELEMENTS: tuple[tuple[int, str, bool], ...] = (
    (1, "Supplier Timing", True), (2, "Fixtures/Tooling/Gaging", True), (3, "Supply Chain Map", True),
    (4, "Tiered Risk Assessment", True), (5, "DFMEA", False), (6, "DVP", False), (7, "Design Record", True),
    (8, "Process Flow", True), (9, "Floor Plan Layout", True), (10, "PFMEA", True), (11, "Process Control Plan", True),
    (12, "PM Plan", True), (13, "Error proofing", True), (14, "Vision Systems", False), (15, "Operator Instructions", True),
    (16, "Visual Aids", True), (17, "MSA", True), (18, "Sub Supplier PPAP Approvals", True), (19, "Dimensional Results", True),
    (20, "Material Certifications", True), (21, "Shipping / Packaging", True), (22, "Capability Studies", True),
    (23, "SP12 (Containment)", True), (24, "DVP&R", True), (25, "Run at Rate", True), (26, "Launch Readiness Review", True),
    (27, "Interim approval action register", True), (28, "Interim Submit", True), (29, "Interim approved", True),
    (30, "Full Submit", True), (31, "Approval", True), (32, "IMDS", True),
)

ELEMENT_STATUSES = ["NOT_STARTED", "IN_PROGRESS", "SUBMITTED", "COMPLETED", "NA", "REJECTED"]
UI_ELEMENT_STATUSES = ["NOT_STARTED", "IN_PROGRESS", "COMPLETED", "NA", "REJECTED", "SUBMITTED"]   # Pending first
ELEMENT_STATUS_LABELS = {"NOT_STARTED": "Pending", "IN_PROGRESS": "In progress", "SUBMITTED": "Submitted", "COMPLETED": "Completed",
                         "NA": "NA", "REJECTED": "Rejected"}
CLOSED_STATUSES = {"COMPLETED", "NA"}
DUE_STATUSES = {"NOT_STARTED", "IN_PROGRESS", "REJECTED"}   # open work of the responsible employee (due e-mail / KPIs)
MODULE_KEY = "PPAP"
ROUTE_HOME, ROUTE_PACKAGE, ROUTE_TEMPLATES = "ppap-home", "ppap-package", "ppap-templates"
DEPARTMENTS = ["Quality", "Engineering", "NPD", "Production", "Purchase", "Maintenance", "Logistics", "Sales"]
DEFAULT_DEPARTMENTS: dict[int, str] = {
    1: "NPD", 2: "Engineering", 3: "Purchase", 4: "Quality", 5: "Engineering", 6: "Engineering", 7: "Engineering", 8: "Engineering",
    9: "Production", 10: "Quality", 11: "Quality", 12: "Maintenance", 13: "Engineering", 14: "Engineering", 15: "Production",
    16: "Quality", 17: "Quality", 18: "Purchase", 19: "Quality", 20: "Quality", 21: "Logistics", 22: "Quality", 23: "Quality",
    24: "Engineering", 25: "Production", 26: "NPD", 27: "Quality", 28: "Quality", 29: "Quality", 30: "Quality", 31: "Quality", 32: "Quality",
}
PACKAGE_STATUSES = ["IN_PROGRESS", "SUBMITTED", "APPROVED", "ON_HOLD", "REJECTED", "CANCELLED", "CLOSED"]
INACTIVE_PACKAGE_STATUSES = {"APPROVED", "CANCELLED", "CLOSED"}
PSW_STATUSES = ["NOT_SUBMITTED", "INTERIM_SUBMITTED", "INTERIM_APPROVED", "FULL_SUBMITTED", "APPROVED", "REJECTED"]
SUBMISSION_LEVELS = [1, 2, 3, 4, 5]

DUE_SOON_DAYS = 3                       # reminder window (qcms-ppap-due-notifier default, schedule PPAP_DOCUMENT_DUE days_ahead)
MAX_EMAIL_BYTES = 10 * 1024 * 1024      # above this the customer gets a secure download link instead of attachments
MAX_EMAIL_FILES = 18                    # qcms-send-email attaches at most 20 files (index PDF + documents)
SIGNED_LINK_SECONDS = 7 * 24 * 3600     # secure link validity (7 days)
BUCKET = "quality-documents"
OPEN_REQUEST_KEY = "ppap_package_open_request_id"      # Records Centre / Global Search / APQP page -> open a package
CREATE_FROM_APQP_KEY = "ppap_create_from_apqp_id"     # APQP page -> prefill a new package from an APQP project

STATUS_COLORS: dict[str, tuple[str, str]] = {   # state -> (text, background)
    "completed": ("#166534", "#DCFCE7"), "na": ("#475569", "#E2E8F0"), "overdue": ("#991B1B", "#FEE2E2"),
    "due_soon": ("#9A3412", "#FFEDD5"), "submitted": ("#1D4ED8", "#DBEAFE"), "rejected": ("#9F1239", "#FFE4E6"),
    "open": ("#92400E", "#FEF3C7"), "not_planned": ("#334155", "#F1F5F9"),
}

_STATUS_ALIASES = {
    "COMPLETED": "COMPLETED", "COMPLETE": "COMPLETED", "DONE": "COMPLETED", "OK": "COMPLETED", "CLOSED": "COMPLETED",
    "APPROVED": "COMPLETED", "YES": "COMPLETED", "Y": "COMPLETED",
    "NA": "NA", "N/A": "NA", "NOT APPLICABLE": "NA", "NOTAPPLICABLE": "NA", "NOT_APPLICABLE": "NA", "-NA-": "NA",
    "PENDING": "NOT_STARTED", "IN PROGRESS": "IN_PROGRESS", "IN_PROGRESS": "IN_PROGRESS", "INPROGRESS": "IN_PROGRESS",
    "WIP": "IN_PROGRESS", "OPEN": "IN_PROGRESS", "ONGOING": "IN_PROGRESS", "UNDER PROGRESS": "IN_PROGRESS",
    "SUBMITTED": "SUBMITTED", "UPLOADED": "SUBMITTED", "SUBMIT": "SUBMITTED",
    "REJECTED": "REJECTED", "REJECT": "REJECTED", "NOT OK": "REJECTED", "NG": "REJECTED",
    "NOT STARTED": "NOT_STARTED", "NOT_STARTED": "NOT_STARTED", "": "NOT_STARTED",
}


# ----------------------------------------------------------------------------- small helpers
def normalize_status(value: Any) -> str:
    """Tracker / free-text status -> element status code (Completed, NA, Pending ... -> COMPLETED, NA, NOT_STARTED = Pending)."""
    text = re.sub(r"\s+", " ", str(value or "").strip()).upper()
    if text in _STATUS_ALIASES:
        return _STATUS_ALIASES[text]
    if text.replace(" ", "_") in ELEMENT_STATUSES:
        return text.replace(" ", "_")
    if text.startswith("COMPLET"):
        return "COMPLETED"
    if text.startswith("PEND"):
        return "NOT_STARTED"
    if "PROGRESS" in text:
        return "IN_PROGRESS"
    if text.startswith("N/A") or text == "N.A." or text.startswith("NOT APPL"):
        return "NA"
    return "IN_PROGRESS" if text else "NOT_STARTED"


_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d/%m/%y", "%Y/%m/%d")


def parse_date(value: Any) -> date | None:
    """Excel / ISO / Indian date -> date; '-', 'NA', blanks and bad text -> None (never today)."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial date
        if 20000 < float(value) < 80000:
            return date(1899, 12, 30) + timedelta(days=int(value))
        return None
    text = str(value).strip()
    if not text or text.upper() in {"-", "NA", "N/A", "TBD", "NONE", "NAN", "NAT"}:
        return None
    text = text[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", text) else text
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def iso(value: Any) -> str | None:
    parsed = parse_date(value)
    return parsed.isoformat() if parsed else None


def dmy(value: Any) -> str:
    parsed = parse_date(value)
    return parsed.strftime("%d-%m-%Y") if parsed else ""


def status_label(code: Any) -> str:
    """Element status code -> label (NOT_STARTED = Pending); other codes (package / PSW status) -> Title Case."""
    text = str(code or "").upper()
    if text in ELEMENT_STATUS_LABELS:
        return ELEMENT_STATUS_LABELS[text]
    return text.replace("_", " ").title()


def status_from_label(value: Any) -> str:
    """Grid label ('Pending', 'In progress' ...) or code -> element status code."""
    text = str(value or "").strip()
    for code, label in ELEMENT_STATUS_LABELS.items():
        if text == label or text.upper() == code:
            return code
    return normalize_status(text)


def default_department(element_no: Any) -> str:
    try:
        return DEFAULT_DEPARTMENTS.get(int(float(element_no)), "Quality")
    except (TypeError, ValueError):
        return "Quality"


def element_state(element: Mapping[str, Any], today: date | None = None, *, due_soon_days: int = DUE_SOON_DAYS) -> str:
    """Colour state of one element: completed / na / overdue / due_soon / submitted / rejected / open / not_planned."""
    today = today or date.today()
    status = str(element.get("status") or "NOT_STARTED").upper()
    if status == "COMPLETED":
        return "completed"
    if status == "NA":
        return "na"
    due = parse_date(element.get("due_date"))
    if due and due < today:
        return "overdue"
    if status == "REJECTED":
        return "rejected"
    if due and due <= today + timedelta(days=due_soon_days):
        return "due_soon"
    if status == "SUBMITTED":
        return "submitted"
    if due is None and status == "NOT_STARTED":
        return "not_planned"
    return "open"


def is_due_item(element: Mapping[str, Any], today: date | None = None, *, days_ahead: int = DUE_SOON_DAYS) -> bool:
    """Reminder rule (mirrors qcms-ppap-due-notifier): Pending / In progress / Rejected element due within N days or overdue."""
    today = today or date.today()
    if str(element.get("status") or "NOT_STARTED").upper() not in DUE_STATUSES:
        return False
    due = parse_date(element.get("due_date"))
    return bool(due and due <= today + timedelta(days=days_ahead))


def due_items(elements: Iterable[Mapping[str, Any]], today: date | None = None, *, days_ahead: int = DUE_SOON_DAYS) -> list[dict]:
    today = today or date.today()
    rows = [dict(e) for e in elements if is_due_item(e, today, days_ahead=days_ahead)]
    rows.sort(key=lambda e: (parse_date(e.get("due_date")) or date.max, int(e.get("element_no") or 0)))
    return rows


def is_overdue(element: Mapping[str, Any], today: date | None = None) -> bool:
    return element_state(element, today) == "overdue"


def package_progress(elements: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Progress % = completed / applicable elements (NA elements are excluded from the base)."""
    total = len(elements)
    na = sum(str(e.get("status") or "").upper() == "NA" for e in elements)
    completed = sum(str(e.get("status") or "").upper() == "COMPLETED" for e in elements)
    applicable = total - na
    percent = 0.0 if total == 0 else (100.0 if applicable == 0 else round(completed * 100.0 / applicable, 1))
    return {"total": total, "na": na, "applicable": applicable, "completed": completed,
            "open": applicable - completed, "percent": percent}


def portfolio_kpis(packages: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]], today: date | None = None) -> dict[str, Any]:
    today = today or date.today()
    active_ids = {str(p.get("id")) for p in packages if str(p.get("status") or "IN_PROGRESS").upper() not in INACTIVE_PACKAGE_STATUSES}
    active_elements = [e for e in elements if str(e.get("package_id")) in active_ids]
    week_end = today + timedelta(days=7)
    overdue = sum(1 for e in active_elements if element_state(e, today) == "overdue")
    due_week = sum(1 for e in active_elements if str(e.get("status") or "").upper() not in CLOSED_STATUSES
                   and (parse_date(e.get("due_date")) or date.max) >= today and (parse_date(e.get("due_date")) or date.max) <= week_end)
    pending = sum(1 for e in active_elements if str(e.get("status") or "NOT_STARTED").upper() in DUE_STATUSES)
    progress = package_progress(elements)
    return {"packages": len(packages), "active": len(active_ids), "overdue": overdue, "due_week": due_week, "pending": pending,
            "percent": progress["percent"], "completed": progress["completed"], "applicable": progress["applicable"]}


def _active_package_map(packages: Sequence[Mapping[str, Any]]) -> dict[str, Mapping[str, Any]]:
    return {str(p.get("id")): p for p in packages if str(p.get("status") or "IN_PROGRESS").upper() not in INACTIVE_PACKAGE_STATUSES}


def employee_name(employee: Mapping[str, Any] | None) -> str:
    employee = employee or {}
    return " ".join(x for x in (str(employee.get("first_name") or "").strip(), str(employee.get("last_name") or "").strip()) if x) \
        or str(employee.get("employee_code") or employee.get("email") or "")


def by_responsible(packages: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]],
                   employees: Sequence[Mapping[str, Any]], today: date | None = None) -> list[dict]:
    """Open PPAP work per responsible employee (active packages): open / overdue / due <= 7 days / next due date."""
    today = today or date.today()
    active = _active_package_map(packages)
    emp = {str(e.get("id")): e for e in employees}
    stats: dict[str, dict] = {}
    for element in elements:
        if str(element.get("package_id")) not in active or str(element.get("status") or "NOT_STARTED").upper() not in DUE_STATUSES:
            continue
        key = str(element.get("responsible_employee_id") or "")
        row = stats.setdefault(key, {"employee_id": key, "Responsible": employee_name(emp.get(key)) if key else "Not assigned",
                                     "Department": str((emp.get(key) or {}).get("department") or element.get("responsible_department") or ""),
                                     "Open": 0, "Overdue": 0, "Due <= 7 days": 0, "_next": None})
        row["Open"] += 1
        due = parse_date(element.get("due_date"))
        if due and due < today:
            row["Overdue"] += 1
        elif due and due <= today + timedelta(days=7):
            row["Due <= 7 days"] += 1
        if due and (row["_next"] is None or due < row["_next"]):
            row["_next"] = due
    out = []
    for row in sorted(stats.values(), key=lambda r: (-r["Overdue"], -r["Due <= 7 days"], -r["Open"], r["Responsible"])):
        row["Next Due"] = row.pop("_next").strftime("%d-%m-%Y") if row.get("_next") else ""
        out.append(row)
    return out


def due_digest_subject(department: Any, report_date: date, count: int) -> str:
    """'Quality · PPAP elements due · 3 item(s) · 28-10-2026' (department first, like every QCMS reminder)."""
    try:
        from core.notification_service import department_label
        label = department_label(department) or ""
    except Exception:
        label = ""
    label = label or str(department or "").strip() or "Quality"
    return f"{label} · PPAP elements due · {count} item(s) · {report_date.strftime('%d-%m-%Y')}"


def due_digests(packages: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]], employees: Sequence[Mapping[str, Any]],
                today: date | None = None, *, days_ahead: int = DUE_SOON_DAYS) -> list[dict]:
    """One e-mail digest per responsible employee (mirrors supabase/functions/qcms-ppap-due-notifier):
    Pending / In progress / Rejected elements of active packages due within ``days_ahead`` days or overdue; the reporting
    manager (employees.reports_to_employee_id) is CC'd when at least one row is overdue. Employees without e-mail are skipped."""
    today = today or date.today()
    active = _active_package_map(packages)
    emp = {str(e.get("id")): e for e in employees}
    grouped: dict[str, list[dict]] = {}
    for element in elements:
        package = active.get(str(element.get("package_id")))
        owner = str(element.get("responsible_employee_id") or "")
        if not package or not owner or not is_due_item(element, today, days_ahead=days_ahead):
            continue
        due = parse_date(element.get("due_date"))
        grouped.setdefault(owner, []).append({
            "package_id": str(package.get("id")), "element_id": str(element.get("id") or ""),
            "part": cell_text(package.get("customer_part_no")) or "-", "ppap_no": cell_text(package.get("customer_ppap_no")) or "-",
            "package_no": str(package.get("package_no") or ""),
            "element": f"{int(element.get('element_no') or 0):02d} {element.get('element_name') or ''}".strip(),
            "due_date": due.isoformat() if due else "", "days_overdue": max(0, (today - due).days) if due else 0,
            "status": status_label(element.get("status") or "NOT_STARTED"),
        })
    digests = []
    for employee_id, rows in grouped.items():
        employee = emp.get(employee_id) or {}
        email = str(employee.get("email") or "").strip()
        if "@" not in email:
            continue
        rows.sort(key=lambda r: (-r["days_overdue"], r["due_date"], r["part"], r["element"]))
        cc: list[str] = []
        if any(r["days_overdue"] > 0 for r in rows):
            manager = emp.get(str(employee.get("reports_to_employee_id") or "")) or {}
            manager_email = str(manager.get("email") or "").strip()
            if "@" in manager_email and manager_email.casefold() != email.casefold():
                cc.append(manager_email)
        department = str(employee.get("department") or "").strip()
        digests.append({"employee_id": employee_id, "email": email, "name": employee_name(employee), "department": department,
                        "cc": cc, "rows": rows, "overdue": sum(r["days_overdue"] > 0 for r in rows),
                        "subject": due_digest_subject(department, today, len(rows)),
                        "dedupe_key": f"PPAP_ELEMENT_DUE:{today.isoformat()}:{employee_id}"})
    digests.sort(key=lambda d: d["email"].casefold())
    return digests


def next_revision(current: Any) -> str:
    """A -> B, Z -> AA, 01 -> 02, 3 -> 4, '' -> A."""
    text = str(current or "").strip().upper()
    if not text:
        return "A"
    match = re.match(r"^(.*?)(\d+)$", text)
    if match:
        prefix, number = match.groups()
        return f"{prefix}{str(int(number) + 1).zfill(len(number))}"
    if re.fullmatch(r"[A-Z]+", text):
        chars = list(text)
        i = len(chars) - 1
        while i >= 0:
            if chars[i] != "Z":
                chars[i] = chr(ord(chars[i]) + 1)
                return "".join(chars)
            chars[i] = "A"
            i -= 1
        return "A" + "".join(chars)
    return text + "-1"


def safe_name(text: Any, *, keep_spaces: bool = True, limit: int = 90) -> str:
    raw = str(text or "").strip()
    pattern = r"[^A-Za-z0-9 ._&()+-]+" if keep_spaces else r"[^A-Za-z0-9._-]+"
    cleaned = re.sub(pattern, "_" if not keep_spaces else " ", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._") if keep_spaces else re.sub(r"_+", "_", cleaned).strip("_.")
    return cleaned[:limit] or "file"


def element_folder(element_no: Any, element_name: Any) -> str:
    """ZIP folder per element: '01 Supplier Timing', '21 Shipping - Packaging'."""
    try:
        number = int(float(element_no))
    except (TypeError, ValueError):
        number = 0
    name = str(element_name or "Element").replace("/", " - ")
    return f"{number:02d} {safe_name(name)}"


def build_zip(entries: Sequence[Mapping[str, Any]], *, extra_files: Sequence[tuple[str, bytes]] = ()) -> bytes:
    """ZIP with one folder per element ("NN Element name/file"); duplicate names get " (2)"; extra files at the root."""
    buffer = BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in extra_files:
            path = safe_name(name)
            used.add(path.casefold())
            archive.writestr(path, content)
        for entry in entries:
            folder = element_folder(entry.get("element_no"), entry.get("element_name"))
            file_name = safe_name(entry.get("file_name") or "document")
            path = f"{folder}/{file_name}"
            stem, dot, ext = file_name.rpartition(".")
            counter = 2
            while path.casefold() in used:
                alt = f"{stem} ({counter}).{ext}" if dot else f"{file_name} ({counter})"
                path = f"{folder}/{alt}"
                counter += 1
            used.add(path.casefold())
            archive.writestr(path, bytes(entry.get("content") or b""))
    return buffer.getvalue()


def template_sort_key(row: Mapping[str, Any]) -> tuple[int, int]:
    """Template / element order: sort_order (reorder) first, element number second."""
    number = int(row.get("element_no") or 0)
    try:
        order = int(row.get("sort_order")) if row.get("sort_order") not in (None, "") else number
    except (TypeError, ValueError):
        order = number
    return (order, number)


def storage_path(tenant_id: Any, package_id: Any, element: Mapping[str, Any], file_name: str) -> str:
    """quality-documents/<tenant>/ppap/<package>/<NN_Element_name>/<file> (storage keys without spaces)."""
    folder = safe_name(element_folder(element.get("element_no"), element.get("element_name")), keep_spaces=False)
    return f"{tenant_id}/ppap/{package_id}/{folder}/{file_name}"


def own_documents(documents: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]]) -> list[Mapping[str, Any]]:
    """Only documents that belong to the given package elements (a package ZIP never carries another package's files)."""
    element_ids = {str(e.get("id")) for e in elements}
    package_ids = {str(e.get("package_id")) for e in elements if e.get("package_id")}
    out = []
    for doc in documents:
        if str(doc.get("package_element_id")) not in element_ids:
            continue
        if package_ids and doc.get("package_id") and str(doc.get("package_id")) not in package_ids:
            continue
        out.append(doc)
    return out


def package_cover_xlsx(package: Mapping[str, Any], elements: Sequence[Mapping[str, Any]], documents: Sequence[Mapping[str, Any]] = (),
                       *, employee_labels: Mapping[str, str] | None = None, notes: Sequence[str] = ()) -> bytes:
    """Cover sheet Excel: the package in the customer tracker layout (Sheet1) + an element / document register sheet."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    master = [{"element_no": e.get("element_no"), "element_name": e.get("element_name")} for e in sorted(elements, key=template_sort_key)]
    data = tracker_workbook_bytes(master, [package], {str(package.get("id")): list(elements)}, notes)
    wb = load_workbook(BytesIO(data))
    ws = wb.create_sheet("Package register", 0)
    labels = employee_labels or {}
    head = [("PPAP Package No.", package.get("package_no")), ("Customer PPAP No.", package.get("customer_ppap_no")),
            ("Customer Part No.", package.get("customer_part_no")), ("Submission Level", package.get("submission_level")),
            ("PSW Status", status_label(package.get("psw_status"))), ("Customer Due Date", dmy(package.get("customer_due_date"))),
            ("Interim Submit", dmy(package.get("interim_submit_date"))), ("Full Submit", dmy(package.get("full_submit_date"))),
            ("Progress", f"{package_progress(elements)['percent']:.0f}%")]
    for r, (label, value) in enumerate(head, start=1):
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, "" if value is None else value)
    start = len(head) + 2
    columns = ["No.", "Element", "Status", "Due Date", "Completed", "Responsible", "Department", "Documents (current)", "Remarks"]
    for c, label in enumerate(columns, start=1):
        cell = ws.cell(start, c, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    docs_by_element: dict[str, list[str]] = {}
    for doc in documents:
        if doc.get("is_latest", True):
            docs_by_element.setdefault(str(doc.get("package_element_id")), []).append(f"{doc.get('file_name')} (Rev. {doc.get('revision') or '-'})")
    for r, e in enumerate(sorted(elements, key=template_sort_key), start=start + 1):
        values = [int(e.get("element_no") or 0), e.get("element_name"), status_label(e.get("status")), dmy(e.get("due_date")),
                  dmy(e.get("completed_date")), labels.get(str(e.get("responsible_employee_id") or ""), ""),
                  e.get("responsible_department") or "", "; ".join(docs_by_element.get(str(e.get("id")), [])), e.get("remarks") or ""]
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value)
    for col, width in zip("ABCDEFGHI", (8, 34, 13, 12, 12, 26, 14, 60, 40)):
        ws.column_dimensions[col].width = width
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def submission_method(documents: Sequence[Mapping[str, Any]], *, max_bytes: int = MAX_EMAIL_BYTES, max_files: int = MAX_EMAIL_FILES) -> str:
    """EMAIL_ATTACHMENTS when the selected files fit in one e-mail, else SECURE_LINK (7-day signed ZIP link)."""
    total = sum(int(d.get("size_bytes") or 0) for d in documents)
    return "SECURE_LINK" if total > max_bytes or len(documents) > max_files else "EMAIL_ATTACHMENTS"


def split_emails(text: Any) -> list[str]:
    out: list[str] = []
    for token in re.split(r"[;,\n\s]+", str(text or "")):
        token = token.strip().strip("<>")
        if token and re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", token) and token.casefold() not in {x.casefold() for x in out}:
            out.append(token)
    return out


def party_emails(party: Mapping[str, Any] | None) -> list[str]:
    if not party:
        return []
    return split_emails(f"{party.get('email') or ''};{party.get('notification_emails') or ''}")


def foreign_party_recipients(addresses: Sequence[str], parties: Sequence[Mapping[str, Any]], customer_id: Any) -> list[str]:
    """Addresses that belong to another party (supplier / other customer): the package is confidential to its customer."""
    own = str(customer_id or "")
    by_email: dict[str, set[str]] = {}
    for party in parties:
        for email in party_emails(party):
            by_email.setdefault(email.casefold(), set()).add(str(party.get("id")))
    bad = []
    for address in addresses:
        owners = by_email.get(str(address).casefold())
        if owners and own not in owners:
            bad.append(address)
    return bad


def digits(text: Any) -> str:
    return re.sub(r"\D", "", str(text or ""))


def leading_part_token(text: Any) -> str:
    """'40256626 TRMF' -> '40256626'; 40257237 -> '40257237'."""
    raw = str(text or "").strip()
    if isinstance(text, float) and text.is_integer():
        raw = str(int(text))
    match = re.match(r"^\s*([A-Za-z0-9.\-/]+)", raw)
    return match.group(1) if match else raw


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def match_part(token: Any, parts: Sequence[Mapping[str, Any]]) -> dict | None:
    """Tracker PART NO. -> Part Master row (part number, FSI part number, then digit match)."""
    text = str(token or "").strip().casefold()
    if not text:
        return None
    for key in ("part_number", "fsi_part_number"):
        for part in parts:
            if str(part.get(key) or "").strip().casefold() == text:
                return dict(part)
    want = digits(text)
    if len(want) >= 4:
        for part in parts:
            if digits(part.get("part_number")) == want:
                return dict(part)
    return None


def package_number(ppap_no: Any, part_text: Any, existing: Iterable[str] = ()) -> str:
    slug = re.sub(r"[^A-Z0-9]+", "-", cell_text(part_text).upper()).strip("-")[:40] or "PART"
    base = f"PPAP-{cell_text(ppap_no)}-{slug}" if cell_text(ppap_no) else f"PPAP-{slug}"
    taken = {str(x).upper() for x in existing}
    candidate, n = base, 2
    while candidate.upper() in taken:
        candidate = f"{base}-{n}"
        n += 1
    return candidate


# ----------------------------------------------------------------------------- tracker import / template
@dataclass
class TrackerColumn:
    part_text: str
    ppap_no: str
    column: int
    elements: list[dict] = field(default_factory=list)

    @property
    def part_token(self) -> str:
        return leading_part_token(self.part_text)


@dataclass
class TrackerImport:
    sheet: str
    columns: list[TrackerColumn]
    element_names: dict[int, str]
    notes: list[str]
    warnings: list[str] = field(default_factory=list)


def _norm_label(value: Any) -> str:
    return re.sub(r"[^A-Z]", "", str(value or "").upper())


def parse_tracker(data: bytes | BytesIO | Any) -> TrackerImport:
    """Read the customer PPAP tracker (or the QCMS template): parts as columns (PART NO. / PPAP NO.), one
    Status + Due Date pair per part, rows = PPAP elements; other sheets = free-text status notes."""
    from openpyxl import load_workbook

    stream = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    workbook = load_workbook(stream, data_only=True)
    sheet = None
    anchor: tuple[int, int] | None = None
    for ws in workbook.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
            for cell in row:
                if _norm_label(cell.value) == "PARTNO":
                    sheet, anchor = ws, (cell.row, cell.column)
                    break
            if anchor:
                break
        if anchor:
            break
    if sheet is None or anchor is None:
        raise ValueError("The workbook has no 'PART NO.' row. Use the QCMS PPAP tracker template.")
    part_row, label_col = anchor
    ppap_row = None
    for r in range(part_row + 1, part_row + 4):
        if _norm_label(sheet.cell(r, label_col).value) in {"PPAPNO", "PPAPNUMBER"}:
            ppap_row = r
            break
    header_row = None
    for r in range(part_row + 1, part_row + 6):
        if any(_norm_label(sheet.cell(r, c).value) == "STATUS" for c in range(label_col + 1, sheet.max_column + 1)):
            header_row = r
            break
    if header_row is None:
        raise ValueError("The tracker has no 'Status' / 'Due Date' header row under PART NO. / PPAP NO.")
    no_col, name_col = label_col - 1, label_col
    for c in range(1, label_col + 1):
        label = _norm_label(sheet.cell(header_row, c).value)
        if label in {"ELEMENTNUMBER", "ELEMENTNO", "NO", "SRNO"}:
            no_col = c
        elif label in {"ELEMENTDESCRIPTION", "ELEMENT", "ELEMENTNAME", "DESCRIPTION"}:
            name_col = c

    element_rows: list[tuple[int, int, str]] = []
    blank_run = 0
    for r in range(header_row + 1, sheet.max_row + 1):
        raw_no, raw_name = sheet.cell(r, no_col).value if no_col >= 1 else None, sheet.cell(r, name_col).value
        if raw_no in (None, "") and raw_name in (None, ""):
            blank_run += 1
            if blank_run >= 3:
                break
            continue
        blank_run = 0
        try:
            number = int(float(str(raw_no).strip()))
        except (TypeError, ValueError):
            continue
        element_rows.append((r, number, cell_text(raw_name) or f"Element {number}"))

    warnings: list[str] = []
    columns: list[TrackerColumn] = []
    for c in range(label_col + 1, sheet.max_column + 1):
        if _norm_label(sheet.cell(header_row, c).value) != "STATUS":
            continue
        part_text = cell_text(sheet.cell(part_row, c).value)
        ppap_no = cell_text(sheet.cell(ppap_row, c).value) if ppap_row else ""
        due_col = c + 1 if _norm_label(sheet.cell(header_row, c + 1).value) in {"DUEDATE", "DUE", "TARGETDATE", "DATE"} else None
        if not part_text:
            if any(sheet.cell(r, c).value not in (None, "") for r, _, _ in element_rows):
                warnings.append(f"Column {sheet.cell(header_row, c).column_letter} has statuses but no PART NO. - skipped.")
            continue
        column = TrackerColumn(part_text=part_text, ppap_no=ppap_no, column=c)
        for r, number, name in element_rows:
            raw_status = sheet.cell(r, c).value
            raw_due = sheet.cell(r, due_col).value if due_col else None
            status = normalize_status(raw_status)
            column.elements.append({
                "element_no": number, "element_name": name, "status": status,
                "raw_status": cell_text(raw_status), "due_date": iso(raw_due),
            })
        columns.append(column)

    notes: list[str] = []
    for ws in workbook.worksheets:
        if ws is sheet:
            continue
        for row in ws.iter_rows():
            for cell in row:
                text = cell_text(cell.value)
                if len(text) >= 3 and text != NOTES_SHEET_HEADER and not isinstance(cell.value, (int, float, datetime, date)):
                    notes.append(re.sub(r"\s+", " ", text))
    return TrackerImport(sheet=sheet.title, columns=columns, element_names={n: name for _, n, name in element_rows},
                         notes=notes, warnings=warnings)


NOTES_SHEET_HEADER = "Status notes (one per row; mention the part number, e.g. 'Part No. 6128 ...')"


def note_due_date(note: str) -> date | None:
    """'... Customer due date is 15-Jul-2026.' -> 2026-07-15 (customer due date written in a status note)."""
    match = re.search(r"due\s*date\s*(?:is|:|-)?\s*([0-9]{1,2}[-/. ][A-Za-z]{3}[-/. ][0-9]{2,4}|[0-9]{1,2}[-/.][0-9]{1,2}[-/.][0-9]{4}|[0-9]{4}-[0-9]{2}-[0-9]{2})",
                      str(note or ""), re.I)
    if not match:
        return None
    return parse_date(re.sub(r"[ /.]", "-", match.group(1)) if re.search(r"[A-Za-z]", match.group(1)) else match.group(1))


def note_targets(notes: Sequence[str], packages: Sequence[Mapping[str, Any]]) -> list[tuple[str, list[int]]]:
    """Assign each status note to package indexes: a note naming a part number / FSI number / PPAP no. goes to that
    package; a following note without a number continues the previous one; notes before any match go to all."""
    result: list[tuple[str, list[int]]] = []
    previous: list[int] | None = None
    for note in notes:
        tokens = [t for t in re.findall(r"\d{4,}", note)]
        matched: list[int] = []
        for index, package in enumerate(packages):
            part_digits = digits(leading_part_token(package.get("customer_part_no")))
            fsi = [digits(x) for x in re.split(r"[/,; ]+", str(package.get("fsi_part_number") or "")) if digits(x)]
            ppap = digits(package.get("customer_ppap_no"))
            for token in tokens:
                if token == ppap or token == part_digits or (part_digits and part_digits.endswith(token)) or token in fsi:
                    matched.append(index)
                    break
        if not matched:
            matched = list(previous) if previous is not None else list(range(len(packages)))
        else:
            previous = matched
        result.append((note, matched))
    return result


def tracker_workbook_bytes(element_master: Sequence[Mapping[str, Any]], packages: Sequence[Mapping[str, Any]] = (),
                           elements_by_package: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
                           notes: Sequence[str] = ()) -> bytes:
    """Excel in the customer tracker layout (template when ``packages`` is empty, export of the matrix otherwise)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    thin = Side(style="thin", color="94A3B8")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    sub_fill = PatternFill("solid", fgColor="DBEAFE")
    fills = {"COMPLETED": "DCFCE7", "NA": "E2E8F0", "IN_PROGRESS": "FEF3C7", "SUBMITTED": "DBEAFE", "REJECTED": "FFE4E6"}
    columns = list(packages) or [{"customer_part_no": "", "customer_ppap_no": ""} for _ in range(3)]
    last_col = 3 + 2 * len(columns)
    ws.merge_cells(start_row=2, start_column=2, end_row=4, end_column=last_col)
    ws.cell(2, 2, "PPAP Tracker").font = Font(bold=True, size=16, color="FFFFFF")
    ws.cell(2, 2).fill = head_fill
    ws.cell(2, 2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(5, 3, "PART NO.").font = Font(bold=True)
    ws.cell(6, 3, "PPAP NO.").font = Font(bold=True)
    ws.cell(7, 2, "Element Number").font = Font(bold=True)
    ws.cell(7, 3, "Element Description").font = Font(bold=True)
    for index, package in enumerate(columns):
        c = 4 + 2 * index
        ws.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c + 1)
        ws.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c + 1)
        ws.cell(5, c, package.get("customer_part_no") or None).alignment = Alignment(horizontal="center")
        ws.cell(6, c, package.get("customer_ppap_no") or None).alignment = Alignment(horizontal="center")
        ws.cell(7, c, "Status").font = Font(bold=True)
        ws.cell(7, c + 1, "Due Date").font = Font(bold=True)
        for col in (c, c + 1):
            ws.cell(5, col).fill = sub_fill
            ws.cell(6, col).fill = sub_fill
            ws.column_dimensions[get_column_letter(col)].width = 13
    master = sorted(element_master, key=lambda e: int(e.get("element_no") or 0)) or [
        {"element_no": n, "element_name": name} for n, name, _ in DEFAULT_ELEMENTS]
    for offset, element in enumerate(master):
        r = 8 + offset
        ws.cell(r, 2, int(element.get("element_no") or 0))
        ws.cell(r, 3, element.get("element_name"))
        for index, package in enumerate(packages):
            c = 4 + 2 * index
            rows = (elements_by_package or {}).get(str(package.get("id")), [])
            row = next((e for e in rows if int(e.get("element_no") or 0) == int(element.get("element_no") or 0)), None)
            if not row:
                continue
            status = str(row.get("status") or "NOT_STARTED").upper()
            ws.cell(r, c, status_label(status))
            due = parse_date(row.get("due_date"))
            ws.cell(r, c + 1, due if due else ("NA" if status == "NA" else None))
            if due:
                ws.cell(r, c + 1).number_format = "DD-MM-YYYY"
            if status in fills:
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fills[status])
    for row in ws.iter_rows(min_row=5, max_row=7 + len(master), min_col=2, max_col=last_col):
        for cell in row:
            cell.border = box
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 34
    ws.freeze_panes = "D8"
    notes_ws = wb.create_sheet("status")
    notes_ws.cell(1, 1, NOTES_SHEET_HEADER).font = Font(bold=True)
    for index, note in enumerate(notes):
        notes_ws.cell(3 + index, 3, note)
    notes_ws.column_dimensions["C"].width = 120
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- PDF index / PSW
PSW_DECLARATION = (
    "We hereby affirm that the samples represented by this warrant are representative of our parts, were made by a "
    "process that meets all Production Part Approval Process requirements, and that the documents listed are the "
    "current revisions submitted for approval.")


def index_frames(elements: Sequence[Mapping[str, Any]], documents: Sequence[Mapping[str, Any]],
                 employee_labels: Mapping[str, str] | None = None):
    import pandas as pd

    labels = employee_labels or {}
    docs_by_element: dict[str, list[Mapping[str, Any]]] = {}
    for doc in documents:
        docs_by_element.setdefault(str(doc.get("package_element_id")), []).append(doc)
    element_rows = []
    for element in sorted(elements, key=lambda e: int(e.get("element_no") or 0)):
        files = docs_by_element.get(str(element.get("id")), [])
        element_rows.append({
            "No.": int(element.get("element_no") or 0), "Element": element.get("element_name"),
            "Status": status_label(element.get("status")), "Due Date": dmy(element.get("due_date")),
            "Completed": dmy(element.get("completed_date")),
            "Responsible": labels.get(str(element.get("responsible_employee_id") or ""), ""),
            "Documents": len(files),
        })
    element_no_by_id = {str(e.get("id")): int(e.get("element_no") or 0) for e in elements}
    element_name_by_id = {str(e.get("id")): e.get("element_name") for e in elements}
    doc_rows = [{
        "No.": element_no_by_id.get(str(d.get("package_element_id")), 0),
        "Element": element_name_by_id.get(str(d.get("package_element_id")), ""),
        "File": d.get("file_name"), "Rev.": d.get("revision"),
        "Uploaded by": d.get("uploaded_by_name") or "", "Uploaded on": dmy(d.get("uploaded_at")),
        "Size (KB)": round(int(d.get("size_bytes") or 0) / 1024, 1),
    } for d in sorted(documents, key=lambda d: (element_no_by_id.get(str(d.get("package_element_id")), 0), str(d.get("file_name"))))]
    return pd.DataFrame(element_rows), pd.DataFrame(doc_rows)


def package_index_pdf(package: Mapping[str, Any], elements: Sequence[Mapping[str, Any]], documents: Sequence[Mapping[str, Any]], *,
                      part: Mapping[str, Any] | None = None, customer: Mapping[str, Any] | None = None,
                      apqp: Mapping[str, Any] | None = None, owner: str = "", employee_labels: Mapping[str, str] | None = None,
                      company: Mapping[str, Any] | None = None) -> bytes:
    """PSW / package index in the company PO print look (core.reporting.transaction_record_pdf_bytes, Design B)."""
    from core.reporting import transaction_record_pdf_bytes

    part = part or {}
    customer = customer or {}
    progress = package_progress(elements)
    fields = [
        ("PPAP Package No.", package.get("package_no")), ("Customer PPAP No.", package.get("customer_ppap_no")),
        ("Customer Part No.", package.get("customer_part_no") or part.get("part_number")), ("FSI Part No.", part.get("fsi_part_number")),
        ("Part Description", part.get("part_name")), ("Drawing / Revision", " / ".join(x for x in (str(part.get("drawing_number") or ""), str(part.get("drawing_revision") or "")) if x)),
        ("Customer", customer.get("party_name")), ("Submission Level", f"Level {package.get('submission_level') or 3}"),
        ("PSW Status", status_label(package.get("psw_status"))), ("Customer Due Date", dmy(package.get("customer_due_date"))),
        ("Interim Submit / Approved", " / ".join(dmy(package.get(k)) or "-" for k in ("interim_submit_date", "interim_approved_date"))),
        ("Full Submit / Approval", " / ".join(dmy(package.get(k)) or "-" for k in ("full_submit_date", "approval_date"))),
        ("APQP Project", (apqp or {}).get("project_code")), ("PPAP Owner", owner),
        ("Progress", f"{progress['percent']:.0f}% · {progress['completed']}/{progress['applicable']} applicable elements complete"),
        ("Documents in package", len(documents)),
        ("Declaration", PSW_DECLARATION),
    ]
    element_frame, doc_frame = index_frames(elements, documents, employee_labels)
    return transaction_record_pdf_bytes(
        "PPAP SUBMISSION PACKAGE · PSW INDEX", fields,
        {"PPAP ELEMENTS": element_frame, "DOCUMENTS IN THIS PACKAGE": doc_frame},
        document_number=str(package.get("package_no") or ""), document_date=date.today().isoformat(),
        company=dict(company or {}), status=str(package.get("psw_status") or "NOT_SUBMITTED"),
        approval={"Prepared by": owner} if owner else {},
    )


# ----------------------------------------------------------------------------- service
class PpapService:
    """Repository-backed PPAP operations (RLS: module PPAP create / edit rights; falls back to NPD & APQP rights while the
    PPAP module has no explicit user / role / department permission row)."""

    def __init__(self, repo: Any = None) -> None:
        if repo is None:
            from core.repository import Repository
            repo = Repository()
        self.repo = repo

    # ------------------------------------------------------------- reads
    def element_master(self, *, active_only: bool = False) -> list[dict]:
        rows = self.repo.select("ppap_element_master", order_by="element_no", limit=500) or []
        if active_only:
            rows = [r for r in rows if str(r.get("status") or "ACTIVE").upper() == "ACTIVE"]
        return sorted(rows, key=lambda r: int(r.get("element_no") or 0))

    def ensure_master(self) -> list[dict]:
        """Seed the 32 default elements when the tenant has none (the migration seeds existing tenants)."""
        rows = self.element_master()
        if rows:
            return rows
        for number, name, required in DEFAULT_ELEMENTS:
            self.repo.insert("ppap_element_master", {"element_no": number, "element_name": name, "default_required": required,
                                                     "default_lead_days": 0, "default_department": default_department(number),
                                                     "sort_order": number, "status": "ACTIVE"})
        return self.element_master()

    # ------------------------------------------------------------- element templates (default + per customer)
    def customer_template(self, customer_id: str | None) -> list[dict]:
        if not customer_id:
            return []
        try:
            rows = self.repo.select("ppap_customer_template_elements", eq={"customer_id": str(customer_id)}, limit=500) or []
        except Exception:
            return []
        return sorted(rows, key=template_sort_key)

    def template_customers(self) -> list[str]:
        try:
            rows = self.repo.select("ppap_customer_template_elements", limit=20000) or []
        except Exception:
            return []
        return sorted({str(r.get("customer_id")) for r in rows if r.get("customer_id")})

    def template_for(self, customer_id: str | None = None, *, active_only: bool = True) -> list[dict]:
        """Element template for a new package: the customer's own template when it exists, else the default master."""
        rows = self.customer_template(customer_id) or sorted(self.ensure_master(), key=template_sort_key)
        if active_only:
            rows = [r for r in rows if str(r.get("status") or "ACTIVE").upper() == "ACTIVE"]
        return rows

    def save_template(self, rows: Sequence[Mapping[str, Any]], *, customer_id: str | None = None) -> dict[str, int]:
        """Save the default (customer_id None -> ppap_element_master) or a customer template: add / rename / reorder;
        rows missing from ``rows`` are removed. Returns counts."""
        table = "ppap_customer_template_elements" if customer_id else "ppap_element_master"
        current = {str(r["id"]): r for r in (self.customer_template(customer_id) if customer_id else self.element_master())}
        seen: set[int] = set()
        kept: set[str] = set()
        counts = {"saved": 0, "added": 0, "removed": 0}
        for order, row in enumerate(rows, start=1):
            name = str(row.get("element_name") or "").strip()
            if not name or row.get("element_no") in (None, ""):
                continue
            number = int(float(row["element_no"]))
            if number in seen:
                raise ValueError(f"Element No. {number} is duplicated.")
            seen.add(number)
            lead = int(float(row.get("default_lead_days") or 0))
            if lead < 0:
                raise ValueError(f"Element {number:02d} - lead days cannot be negative.")
            payload = {"element_no": number, "element_name": name, "description": str(row.get("description") or "").strip() or None,
                       "default_required": bool(row.get("default_required", True)), "default_lead_days": lead,
                       "default_department": str(row.get("default_department") or "").strip() or None,
                       "sort_order": int(float(row["sort_order"])) if row.get("sort_order") not in (None, "") else order,
                       "status": "INACTIVE" if str(row.get("status") or "ACTIVE").upper() == "INACTIVE" else "ACTIVE"}
            if customer_id:
                payload["customer_id"] = str(customer_id)
            rid = str(row.get("id") or "")
            if rid in current:
                self.repo.update(table, rid, payload)
                kept.add(rid)
                counts["saved"] += 1
            else:
                kept.add(str(self.repo.insert(table, payload)["id"]))
                counts["added"] += 1
        for rid in set(current) - kept:
            self.repo.delete(table, rid)
            counts["removed"] += 1
        return counts

    def copy_template_to_customer(self, customer_id: str) -> int:
        """Start a customer-specific template from the default element master."""
        if not customer_id:
            raise ValueError("Customer - select the customer for the template.")
        if self.customer_template(customer_id):
            raise ValueError("This customer already has its own PPAP template.")
        rows = [{k: m.get(k) for k in ("element_no", "element_name", "description", "default_required", "default_lead_days",
                                       "default_department", "sort_order", "status")} for m in sorted(self.ensure_master(), key=template_sort_key)]
        self.save_template(rows, customer_id=customer_id)
        return len(rows)

    def packages(self) -> list[dict]:
        return self.repo.select("ppap_packages", order_by="package_no", limit=3000) or []

    def package(self, package_id: str) -> dict | None:
        return self.repo.get("ppap_packages", package_id)

    def elements(self, package_id: str) -> list[dict]:
        rows = self.repo.select("ppap_package_elements", eq={"package_id": str(package_id)}, order_by="element_no", limit=500) or []
        return sorted(rows, key=template_sort_key)

    def all_elements(self, package_ids: Sequence[str] | None = None) -> list[dict]:
        if package_ids is not None:
            ids = [str(x) for x in package_ids if x]
            if not ids:
                return []
            rows: list[dict] = []
            for start in range(0, len(ids), 150):
                rows.extend(self.repo.select("ppap_package_elements", in_={"package_id": ids[start:start + 150]}, limit=20000) or [])
            return rows
        return self.repo.select("ppap_package_elements", limit=20000) or []

    def documents(self, package_id: str, *, latest_only: bool = False) -> list[dict]:
        rows = self.repo.select("ppap_element_documents", eq={"package_id": str(package_id), "status": "ACTIVE"},
                                order_by="uploaded_at", desc=True, limit=2000) or []
        return [r for r in rows if r.get("is_latest")] if latest_only else rows

    def notes(self, package_id: str) -> list[dict]:
        return self.repo.select("ppap_status_notes", eq={"package_id": str(package_id)}, order_by="created_at", desc=True, limit=500) or []

    def submissions(self, package_id: str) -> list[dict]:
        return self.repo.select("ppap_customer_submissions", eq={"package_id": str(package_id)}, order_by="submitted_at", desc=True, limit=500) or []

    def packages_for_apqp(self, project_id: str) -> list[dict]:
        if not project_id:
            return []
        return self.repo.select("ppap_packages", eq={"apqp_project_id": str(project_id)}, order_by="package_no", limit=200) or []

    # ------------------------------------------------------------- writes
    def refresh_progress(self, package_id: str) -> float:
        percent = package_progress(self.elements(package_id))["percent"]
        try:
            self.repo.update("ppap_packages", str(package_id), {"completion_percent": percent})
        except Exception:
            pass
        return percent

    def create_package(self, header: Mapping[str, Any], *, default_due: date | None = None, responsible_employee_id: str | None = None) -> dict:
        """Insert the package and one element row per ACTIVE master element (non-required elements start as NA)."""
        payload = {k: v for k, v in dict(header).items() if v not in ("",)}
        if not str(payload.get("package_no") or "").strip():
            payload["package_no"] = package_number(payload.get("customer_ppap_no"), payload.get("customer_part_no") or "PART",
                                                   [p.get("package_no") for p in self.packages()])
        payload.setdefault("status", "IN_PROGRESS")
        payload.setdefault("psw_status", "NOT_SUBMITTED")
        payload.setdefault("submission_level", 3)
        saved = self.repo.insert("ppap_packages", payload)
        package_id = str(saved["id"])
        template = self.template_for(payload.get("customer_id"))
        from_master = not self.customer_template(payload.get("customer_id"))
        for order, master in enumerate(template, start=1):
            lead = int(master.get("default_lead_days") or 0)
            due = (date.today() + timedelta(days=lead)) if lead else default_due
            required = bool(master.get("default_required", True))
            self.repo.insert("ppap_package_elements", {
                "package_id": package_id, "element_master_id": master.get("id") if from_master else None,
                "element_no": int(master.get("element_no") or 0), "element_name": master.get("element_name"),
                "status": "NOT_STARTED" if required else "NA",
                "due_date": due.isoformat() if isinstance(due, date) and required else None,
                "responsible_employee_id": (responsible_employee_id or None) if required else None,
                "responsible_department": master.get("default_department") or default_department(master.get("element_no")),
                "sort_order": template_sort_key(master)[0] if master.get("sort_order") not in (None, "") else order,
            })
        self.refresh_progress(package_id)
        return saved

    def header_from_apqp(self, project: Mapping[str, Any], parts: Sequence[Mapping[str, Any]] = ()) -> dict:
        part = next((p for p in parts if str(p.get("id")) == str(project.get("part_id"))), {}) if project.get("part_id") else {}
        level = digits(project.get("submission_level")) or "3"
        return {
            "apqp_project_id": project.get("id"), "part_id": project.get("part_id"), "customer_id": project.get("customer_id"),
            "customer_part_no": part.get("part_number") or None, "submission_level": max(1, min(5, int(level[:1]))),
            "customer_due_date": iso(project.get("target_submission_date")),
            "owner_employee_id": project.get("coordinator_employee_id") or None,
            "remarks": f"Created from APQP project {project.get('project_code') or ''}".strip(),
        }

    def save_elements(self, package_id: str, rows: Sequence[Mapping[str, Any]]) -> list[dict]:
        """Save edited element rows; returns [{element, employee_id}] for newly assigned responsible employees."""
        current = {str(e["id"]): e for e in self.elements(package_id)}
        assigned: list[dict] = []
        for row in rows:
            element_id = str(row.get("id") or "")
            status = status_from_label(row.get("status") or "NOT_STARTED")
            completed = parse_date(row.get("completed_date"))
            if status == "COMPLETED" and completed is None:
                completed = date.today()
            if status != "COMPLETED":
                completed = None
            payload = {
                "status": status, "due_date": iso(row.get("due_date")),
                "responsible_employee_id": str(row.get("responsible_employee_id") or "") or None,
                "completed_date": completed.isoformat() if completed else None,
                "remarks": str(row.get("remarks") or "").strip() or None,
            }
            if row.get("element_name"):
                payload["element_name"] = str(row.get("element_name")).strip()
            if "responsible_department" in row:
                payload["responsible_department"] = str(row.get("responsible_department") or "").strip() or None
            if element_id in current:
                before = current[element_id]
                changes = {k: v for k, v in payload.items() if (before.get(k) or None) != (v or None)}
                if changes:
                    saved = self.repo.update("ppap_package_elements", element_id, changes)
                    if changes.get("responsible_employee_id"):
                        assigned.append({"element": {**before, **saved}, "employee_id": changes["responsible_employee_id"]})
            elif row.get("element_no") not in (None, ""):
                saved = self.repo.insert("ppap_package_elements", {"package_id": str(package_id), "element_no": int(float(row["element_no"])),
                                                                   "element_name": str(row.get("element_name") or f"Element {row['element_no']}"), **payload})
                if payload.get("responsible_employee_id"):
                    assigned.append({"element": saved, "employee_id": payload["responsible_employee_id"]})
        self.refresh_progress(package_id)
        return assigned

    def bulk_assign(self, package_id: str, element_ids: Sequence[str], *, responsible_employee_id: str | None = None,
                    due_date: date | None = None, status: str | None = None) -> list[dict]:
        rows = []
        for element in self.elements(package_id):
            if str(element["id"]) not in {str(x) for x in element_ids}:
                continue
            row = dict(element)
            if responsible_employee_id:
                row["responsible_employee_id"] = responsible_employee_id
            if due_date:
                row["due_date"] = due_date
            if status:
                row["status"] = status
            rows.append(row)
        return self.save_elements(package_id, rows)

    def notify_assignments(self, package: Mapping[str, Any], assigned: Sequence[Mapping[str, Any]]) -> int:
        """In-app notification (bell / mobile push) to each newly assigned responsible employee."""
        sent = 0
        for item in assigned:
            element = item.get("element") or {}
            try:
                self.repo.rpc("qcms_notify", {
                    "p_recipient_profile_id": None, "p_recipient_employee_id": item.get("employee_id"),
                    "p_title": f"PPAP task: {int(element.get('element_no') or 0):02d} {element.get('element_name') or ''}",
                    "p_body": f"Package {package.get('package_no') or ''} · due {dmy(element.get('due_date')) or 'not set'} · upload the document in PPAP → Package.",
                    "p_category": "TASK", "p_module_key": MODULE_KEY, "p_related_table": "ppap_package_elements",
                    "p_related_id": element.get("id"), "p_web_route": ROUTE_PACKAGE,
                })
                sent += 1
            except Exception:
                continue
        return sent

    def add_note(self, package_id: str, note: str, *, note_date: date | None = None, author: str = "", source: str = "MANUAL") -> dict:
        text = str(note or "").strip()
        if not text:
            raise ValueError("Status note - enter the note text.")
        return self.repo.insert("ppap_status_notes", {"package_id": str(package_id), "note": text, "author_name": author or None,
                                                      "note_date": (note_date or date.today()).isoformat(), "source": source})

    # ------------------------------------------------------------- files
    def _storage(self):
        from core.database import get_session_client
        client = get_session_client()
        if client is None:
            raise RuntimeError("A live Supabase session is required for PPAP document storage.")
        return client.storage.from_(BUCKET)

    def upload_document(self, package: Mapping[str, Any], element: Mapping[str, Any], *, file_name: str, content: bytes,
                        mime_type: str | None = None, revision: str = "", remarks: str = "", uploader_name: str = "",
                        uploader_employee_id: str | None = None, mark_submitted: bool = True) -> dict:
        if not content:
            raise ValueError("The selected file is empty.")
        previous = [d for d in self.repo.select("ppap_element_documents", eq={"package_element_id": str(element["id"]), "status": "ACTIVE"}, limit=500) or []]
        latest = next((d for d in previous if d.get("is_latest")), None)
        revision = str(revision or "").strip() or (next_revision(latest.get("revision")) if latest else "A")
        token = uuid.uuid4().hex[:12]
        path = storage_path(self.repo.tenant_id, package["id"], element, f"{token}_{safe_name(file_name, keep_spaces=False)}")
        self._storage().upload(path, content, {"content-type": mime_type or "application/octet-stream", "upsert": "false"})
        try:
            saved = self.repo.insert("ppap_element_documents", {
                "package_id": str(package["id"]), "package_element_id": str(element["id"]), "revision": revision,
                "file_name": str(file_name), "object_path": path, "mime_type": mime_type, "size_bytes": len(content),
                "checksum": hashlib.sha256(content).hexdigest(), "is_latest": True, "status": "ACTIVE", "uploaded_by_name": uploader_name or None,
                "uploaded_by_employee_id": uploader_employee_id or None, "remarks": str(remarks or "").strip() or None,
            })
        except Exception:
            try:
                self._storage().remove([path])
            except Exception:
                pass
            raise
        for doc in previous:
            if doc.get("is_latest"):
                self.repo.update("ppap_element_documents", str(doc["id"]), {"is_latest": False})
        if mark_submitted and str(element.get("status") or "").upper() in {"NOT_STARTED", "REJECTED"}:
            # a document is in: Pending / Rejected -> In progress (the responsible sets Completed after review)
            self.repo.update("ppap_package_elements", str(element["id"]), {"status": "IN_PROGRESS"})
            self.refresh_progress(str(package["id"]))
        return saved

    def set_latest(self, document: Mapping[str, Any]) -> None:
        for doc in self.repo.select("ppap_element_documents", eq={"package_element_id": str(document.get("package_element_id"))}, limit=500) or []:
            want = str(doc["id"]) == str(document.get("id"))
            if bool(doc.get("is_latest")) != want:
                self.repo.update("ppap_element_documents", str(doc["id"]), {"is_latest": want})

    def archive_document(self, document: Mapping[str, Any]) -> None:
        self.repo.update("ppap_element_documents", str(document["id"]), {"status": "ARCHIVED", "is_latest": False})
        remaining = [d for d in self.repo.select("ppap_element_documents", eq={"package_element_id": str(document.get("package_element_id")), "status": "ACTIVE"}, limit=500) or []]
        if remaining and not any(d.get("is_latest") for d in remaining):
            newest = sorted(remaining, key=lambda d: str(d.get("uploaded_at") or d.get("created_at") or ""))[-1]
            self.repo.update("ppap_element_documents", str(newest["id"]), {"is_latest": True})

    def document_bytes(self, document: Mapping[str, Any]) -> bytes:
        path = str(document.get("object_path") or "").strip()
        if not path:
            raise ValueError("Document storage path is missing.")
        return bytes(self._storage().download(path))

    def zip_entries(self, documents: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]]) -> list[dict]:
        by_id = {str(e.get("id")): e for e in elements}
        entries = []
        for doc in own_documents(documents, elements):
            element = by_id.get(str(doc.get("package_element_id")), {})
            entries.append({"element_no": element.get("element_no") or 0, "element_name": element.get("element_name") or "Element",
                            "file_name": doc.get("file_name"), "content": self.document_bytes(doc)})
        return entries

    def package_zip(self, package: Mapping[str, Any], documents: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]],
                    *, index_pdf: bytes | None = None, cover_xlsx: bytes | None = None) -> bytes:
        """ZIP 'NN Element name/file' + cover sheets at the root (00 PSW Index PDF, 00 PPAP Tracker Excel)."""
        extra = []
        if index_pdf:
            extra.append((f"00 PSW Index {safe_name(package.get('package_no'))}.pdf", index_pdf))
        if cover_xlsx:
            extra.append((f"00 PPAP Tracker {safe_name(package.get('package_no'))}.xlsx", cover_xlsx))
        return build_zip(self.zip_entries(documents, elements), extra_files=extra)

    # ------------------------------------------------------------- customer submission
    def log_submission(self, package: Mapping[str, Any], documents: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]], *,
                       sent_to: str, cc: Sequence[str] = (), method: str = "DOWNLOAD", status: str = "DOWNLOADED",
                       zip_path: str | None = None, link_expires: datetime | None = None, outbox_id: str | None = None,
                       message: str = "", sender: str = "") -> dict:
        by_id = {str(e.get("id")): e for e in elements}
        docs = [{"element_no": (by_id.get(str(d.get("package_element_id"))) or {}).get("element_no"),
                 "element_name": (by_id.get(str(d.get("package_element_id"))) or {}).get("element_name"),
                 "file_name": d.get("file_name"), "revision": d.get("revision"), "document_id": d.get("id"),
                 "size_bytes": d.get("size_bytes")} for d in documents]
        return self.repo.insert("ppap_customer_submissions", {
            "package_id": str(package["id"]), "sent_to": sent_to or "-", "cc_emails": "; ".join(cc) or None, "method": method,
            "documents": docs, "document_count": len(docs), "total_bytes": sum(int(d.get("size_bytes") or 0) for d in documents),
            "zip_object_path": zip_path, "link_expires_at": link_expires.isoformat() if link_expires else None,
            "outbox_id": outbox_id, "status": status, "message": message or None, "sent_by_name": sender or None,
        })

    def send_to_customer(self, package: Mapping[str, Any], documents: Sequence[Mapping[str, Any]], elements: Sequence[Mapping[str, Any]], *,
                         to: Sequence[str], cc: Sequence[str] = (), subject: str = "", message: str = "", sender: str = "",
                         index_pdf: bytes | None = None, customer_name: str = "") -> dict:
        """E-mail the package: attachments when small, else a 7-day signed link to the uploaded ZIP. Logged always."""
        from core.notification_service import NotificationService

        to = [x for x in to if x]
        if not to:
            raise ValueError("Send to - enter at least one customer e-mail address.")
        if not documents:
            raise ValueError("Documents - select at least one document for the package.")
        foreign = [d for d in documents if str(d.get("package_id") or package.get("id")) != str(package.get("id"))]
        if foreign:
            raise ValueError("Documents - only documents of this PPAP package can be sent to its customer.")
        method = submission_method(documents)
        manifest: list[dict] = []
        zip_path = None
        expires = None
        link = ""
        if method == "SECURE_LINK":
            data = self.package_zip(package, documents, elements, index_pdf=index_pdf)
            zip_path = f"{self.repo.tenant_id}/ppap/{package['id']}/submissions/{datetime.now().strftime('%Y%m%d%H%M%S')}_{safe_name(package.get('package_no'), keep_spaces=False)}.zip"
            storage = self._storage()
            storage.upload(zip_path, data, {"content-type": "application/zip", "upsert": "true"})
            signed = storage.create_signed_url(zip_path, SIGNED_LINK_SECONDS)
            link = str((signed or {}).get("signedURL") or (signed or {}).get("signedUrl") or (signed or {}).get("signed_url") or "") if isinstance(signed, Mapping) else str(signed or "")
            if not link:
                raise RuntimeError("The secure download link could not be created.")
            expires = datetime.now() + timedelta(seconds=SIGNED_LINK_SECONDS)
        else:
            manifest = [{"bucket": BUCKET, "object_path": d.get("object_path"), "file_name": d.get("file_name"),
                         "mime_type": d.get("mime_type") or "application/octet-stream"} for d in documents]

        class _PackageMail(NotificationService):
            def attachment_manifest(self, *args: Any, **kwargs: Any) -> list[dict]:  # noqa: D401 - explicit package files
                return list(manifest)

        service = _PackageMail(self.repo)
        if index_pdf:
            stored = service._persist_generated_pdf(f"PSW_Index_{safe_name(package.get('package_no'), keep_spaces=False)}.pdf", index_pdf)
            if stored:
                manifest.insert(0, stored)
        lines = [message.strip() or f"Dear {customer_name or 'Customer'},\n\nPlease find the PPAP submission package below."]
        lines += ["", f"PPAP package: {package.get('package_no') or '-'}", f"PPAP no.: {package.get('customer_ppap_no') or '-'}",
                  f"Part no.: {package.get('customer_part_no') or '-'}", f"Submission level: {package.get('submission_level') or 3}",
                  f"Documents: {len(documents)}"]
        if link:
            lines += ["", "The package is larger than an e-mail allows. Download it (ZIP, valid for 7 days):", link]
        else:
            lines += ["", "Documents attached:"] + [f"- {d.get('file_name')} (Rev. {d.get('revision') or '-'})" for d in documents]
        lines += ["", "Regards,", sender or "Quality Department"]
        row = service.enqueue(
            "PPAP_PACKAGE_SUBMISSION", subject=subject or f"PPAP Submission · {package.get('customer_ppap_no') or package.get('package_no')} · Part {package.get('customer_part_no') or ''}".strip(),
            body_text="\n".join(lines), recipient_email=to[0], recipient_name=customer_name or None, cc_emails=list(cc),
            include_generated_pdf=False, include_record_attachments=False,
            context={"to_emails": list(to), "package_no": package.get("package_no"), "ppap_no": package.get("customer_ppap_no"),
                     "document_count": len(documents), "method": method},
        )
        status = "QUEUED"
        if row:
            result = service.dispatch([row])
            status = "SENT" if int((result or {}).get("sent") or 0) > 0 else ("FAILED" if (result or {}).get("failed") or (result or {}).get("error") else "QUEUED")
        submission = self.log_submission(package, documents, elements, sent_to="; ".join(to), cc=cc, method=method, status=status,
                                         zip_path=zip_path, link_expires=expires, outbox_id=(row or {}).get("id"),
                                         message=link or message, sender=sender)
        if status in {"SENT", "QUEUED"} and str(package.get("status") or "").upper() == "IN_PROGRESS":
            try:
                self.repo.update("ppap_packages", str(package["id"]), {"status": "SUBMITTED"})
            except Exception:
                pass
        return {"submission": submission, "method": method, "status": status, "link": link}

    # ------------------------------------------------------------- tracker import
    def import_tracker(self, parsed: TrackerImport, *, parts: Sequence[Mapping[str, Any]] = (), selected: Sequence[int] | None = None,
                       owner_employee_id: str | None = None, author: str = "") -> dict[str, Any]:
        packages = self.packages()
        taken = [p.get("package_no") for p in packages]
        master = {int(m.get("element_no") or 0): m for m in self.ensure_master()}
        summary = {"created": 0, "updated": 0, "elements": 0, "notes": 0, "packages": []}
        chosen = [c for i, c in enumerate(parsed.columns) if selected is None or i in set(selected)]
        imported: list[dict] = []
        for column in chosen:
            existing = next((p for p in packages if cell_text(p.get("customer_part_no")) == column.part_text
                             and cell_text(p.get("customer_ppap_no")) == column.ppap_no), None)
            part = match_part(column.part_token, parts)
            if existing is None:
                header = {"package_no": package_number(column.ppap_no, column.part_text, taken), "customer_part_no": column.part_text,
                          "customer_ppap_no": column.ppap_no or None, "part_id": (part or {}).get("id"),
                          "customer_id": (part or {}).get("customer_id"), "owner_employee_id": owner_employee_id or None,
                          "status": "IN_PROGRESS", "psw_status": "NOT_SUBMITTED", "submission_level": 3,
                          "remarks": f"Imported from PPAP tracker ({parsed.sheet})"}
                existing = self.repo.insert("ppap_packages", header)
                taken.append(header["package_no"])
                packages.append(existing)
                summary["created"] += 1
            else:
                summary["updated"] += 1
            current = {int(e.get("element_no") or 0): e for e in self.elements(str(existing["id"]))}
            for item in column.elements:
                number = int(item["element_no"])
                payload = {"status": item["status"], "due_date": item["due_date"],
                           "completed_date": item["due_date"] if item["status"] == "COMPLETED" else None}
                if number in current:
                    row = current[number]
                    if item["status"] == "COMPLETED" and row.get("completed_date"):
                        payload["completed_date"] = row.get("completed_date")
                    self.repo.update("ppap_package_elements", str(row["id"]), payload)
                else:
                    self.repo.insert("ppap_package_elements", {"package_id": str(existing["id"]), "element_no": number,
                                                               "element_master_id": (master.get(number) or {}).get("id"),
                                                               "element_name": item["element_name"], "sort_order": number,
                                                               "responsible_department": (master.get(number) or {}).get("default_department") or default_department(number),
                                                               **payload})
                summary["elements"] += 1
            self.refresh_progress(str(existing["id"]))
            imported.append({**existing, "fsi_part_number": (part or {}).get("fsi_part_number")})
            summary["packages"].append(existing.get("package_no"))
        for note, targets in note_targets(parsed.notes, imported):
            for index in targets:
                package = imported[index]
                known = {str(n.get("note") or "").strip() for n in self.notes(str(package["id"]))}
                if note.strip() in known:
                    continue
                self.add_note(str(package["id"]), note, author=author, source="IMPORT")
                summary["notes"] += 1
                due = note_due_date(note)
                if due and not package.get("customer_due_date"):
                    self.repo.update("ppap_packages", str(package["id"]), {"customer_due_date": due.isoformat()})
                    package["customer_due_date"] = due.isoformat()
        return summary
