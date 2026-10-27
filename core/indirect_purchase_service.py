"""QCMS R37 · Indirect Purchase service (separate from the Direct / Supply Chain Purchase Order).

Item master + categories, Purchase Requisition → Indirect PO (own IPO-<branch>-<yy>-<seq> series, value-based multi-level
approval via qcms_module_approval_routes module key INDIRECT_PURCHASE), Indirect Inward (GRN) and an indirect-items stock
ledger with Issue-to-department. Approval, GRN posting and stock movements run in SECURITY DEFINER RPCs (migration
20261024120000_qcms_r37_indirect_purchase.sql) so balances and approvals cannot be bypassed from the browser.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from typing import Any, Iterable, Mapping, Sequence

from core.repository import Repository

MODULE_KEY = "INDIRECT_PURCHASE"

DEFAULT_CATEGORIES: tuple[tuple[str, str], ...] = (
    ("CONSUMABLES", "Consumables"),
    ("TOOLS_TOOLING", "Tools & Tooling"),
    ("SPARES_MAINTENANCE", "Spares / Maintenance"),
    ("STATIONERY", "Stationery"),
    ("IT", "IT / Computers"),
    ("SERVICES", "Services"),
    ("CAPITAL_ASSET", "Capital / Asset"),
    ("PACKING", "Packing Material"),
    ("SAFETY_PPE", "Safety / PPE"),
    ("OTHERS", "Others"),
)
UOMS = ("NOS", "PCS", "SET", "PAIR", "BOX", "PKT", "KG", "GM", "LTR", "ML", "MTR", "MM", "SQM", "ROLL", "REAM", "LOT", "JOB", "HRS", "DAYS", "MONTH", "YEAR", "LS")
GST_RATES = (0.0, 5.0, 12.0, 18.0, 28.0)
SUPPLIER_TYPES = {"SUPPLIER", "OSP_VENDOR", "VENDOR", "SERVICE_PROVIDER", "CONTRACTOR", "STEEL_MILL"}
OPEN_PO_STATUSES = {"APPROVED", "PARTIALLY_RECEIVED"}
EDITABLE_PO_STATUSES = {"DRAFT", "PENDING_APPROVAL", "REJECTED", "APPROVED"}
STATUS_LABELS = {
    "DRAFT": "Draft", "PENDING_APPROVAL": "Pending Approval", "APPROVED": "Approved (Open)", "PARTIALLY_RECEIVED": "Partially Received",
    "CLOSED": "Closed", "CANCELLED": "Cancelled", "REJECTED": "Rejected",
}
DEFAULT_TERMS = "Prices are firm for the order quantity. Goods / services are subject to inspection and acceptance at our premises."

# ------------------------------------------------------------------------------------------------ pure helpers

def number(value: Any) -> float:
    try:
        if value is None or value == "":
            return 0.0
        result = float(value)
        return 0.0 if result != result else result  # NaN → 0
    except (TypeError, ValueError):
        return 0.0


def clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


_ONES = ("", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve", "Thirteen",
         "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen")
_TENS = ("", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety")


def _below_thousand(n: int) -> str:
    words: list[str] = []
    if n >= 100:
        words.append(f"{_ONES[n // 100]} Hundred")
        n %= 100
    if n >= 20:
        words.append(_TENS[n // 10] + (f" {_ONES[n % 10]}" if n % 10 else ""))
    elif n:
        words.append(_ONES[n])
    return " ".join(words)


def indian_words(n: int) -> str:
    """Integer in words with the Indian system (Thousand, Lakh, Crore)."""
    n = int(n)
    if n == 0:
        return "Zero"
    parts: list[str] = []
    crore, n = divmod(n, 10_000_000)
    lakh, n = divmod(n, 100_000)
    thousand, n = divmod(n, 1000)
    if crore:
        parts.append(f"{indian_words(crore)} Crore")
    if lakh:
        parts.append(f"{_below_thousand(lakh)} Lakh")
    if thousand:
        parts.append(f"{_below_thousand(thousand)} Thousand")
    if n:
        parts.append(_below_thousand(n))
    return " ".join(parts)


def amount_in_words(amount: Any, currency: str = "INR") -> str:
    value = round(abs(number(amount)) + 1e-9, 2)
    whole = int(value)
    fraction = int(round((value - whole) * 100))
    unit, sub = {"INR": ("Rupees", "Paise"), "USD": ("US Dollars", "Cents"), "EUR": ("Euros", "Cents")}.get(str(currency or "INR").upper(), (str(currency or "INR").upper(), "Cents"))
    text = f"{unit} {indian_words(whole)}"
    if fraction:
        text += f" and {indian_words(fraction)} {sub}"
    return text + " Only"


def _state_code(gstin: Any) -> str:
    text = clean(gstin).upper()
    return text[:2] if len(text) >= 2 and text[:2].isdigit() else ""


def gst_mode(supplier: Mapping[str, Any] | None, plant: Mapping[str, Any] | None) -> str:
    """INTRA (CGST + SGST) when the supplier is in the plant's state, else INTER (IGST) – like the direct PO."""
    s_code, p_code = _state_code((supplier or {}).get("tax_identifier") or (supplier or {}).get("gstin")), _state_code((plant or {}).get("gstin") or (plant or {}).get("tax_identifier"))
    if s_code and p_code:
        return "INTRA" if s_code == p_code else "INTER"
    s_state = clean((supplier or {}).get("state")).casefold()
    p_state = clean((plant or {}).get("state")).casefold() or "maharashtra"
    aliases = {"mh": "maharashtra", "ka": "karnataka", "gj": "gujarat", "tn": "tamil nadu", "dl": "delhi"}
    s_state = aliases.get(s_state, s_state); p_state = aliases.get(p_state, p_state)
    country = clean((supplier or {}).get("country")).casefold()
    if country and country not in {"india", "in", "bharat"}:
        return "INTER"
    if not s_state:
        return "INTRA"
    return "INTRA" if s_state == p_state else "INTER"


def compute_line(line: Mapping[str, Any], mode: str = "INTRA") -> dict[str, Any]:
    qty = round(number(line.get("quantity")), 4)
    rate = round(number(line.get("rate")), 4)
    disc_pct = min(max(number(line.get("discount_percent")), 0.0), 100.0)
    gross = round(qty * rate, 2)
    discount = round(gross * disc_pct / 100.0, 2)
    taxable = round(gross - discount, 2)
    gst_pct = max(number(line.get("gst_percent")), 0.0)
    gst = round(taxable * gst_pct / 100.0, 2)
    if str(mode).upper() == "INTER":
        cgst = sgst = 0.0; igst = gst
    else:
        cgst = round(gst / 2.0, 2); sgst = round(gst - cgst, 2); igst = 0.0
    return {**dict(line), "quantity": qty, "rate": rate, "discount_percent": disc_pct, "gross_amount": gross, "discount_amount": discount,
            "taxable_amount": taxable, "gst_percent": gst_pct, "gst_amount": gst, "cgst_amount": cgst, "sgst_amount": sgst,
            "igst_amount": igst, "line_total": round(taxable + gst, 2)}


def compute_totals(lines: Sequence[Mapping[str, Any]], mode: str = "INTRA", *, freight: Any = 0, other: Any = 0, round_to_rupee: bool = True, currency: str = "INR") -> dict[str, Any]:
    computed = [compute_line(line, mode) for line in lines]
    gross = round(sum(l["gross_amount"] for l in computed), 2)
    discount = round(sum(l["discount_amount"] for l in computed), 2)
    taxable = round(sum(l["taxable_amount"] for l in computed), 2)
    cgst = round(sum(l["cgst_amount"] for l in computed), 2)
    sgst = round(sum(l["sgst_amount"] for l in computed), 2)
    igst = round(sum(l["igst_amount"] for l in computed), 2)
    freight_v = round(max(number(freight), 0.0), 2); other_v = round(number(other), 2)
    before = round(taxable + cgst + sgst + igst + freight_v + other_v, 2)
    grand = float(round(before)) if round_to_rupee else before
    return {"lines": computed, "gst_mode": str(mode).upper(), "gross_amount": gross, "discount_amount": discount, "subtotal": taxable,
            "cgst_amount": cgst, "sgst_amount": sgst, "igst_amount": igst, "freight_amount": freight_v, "other_amount": other_v,
            "round_off": round(grand - before, 2), "grand_total": round(grand, 2), "amount_in_words": amount_in_words(grand, currency)}


def financial_year(on: date | None = None) -> str:
    d = on or date.today()
    start = d.year if d.month >= 4 else d.year - 1
    return f"{start % 100:02d}"


def preview_number(series: str, branch_code: str = "", *, on: date | None = None) -> str:
    prefix = {"IITEM": "IND"}.get(series.upper(), series.upper())
    if series.upper() == "IITEM":
        return f"{prefix}-NEW"
    return "-".join(v for v in (prefix, clean(branch_code).upper() or "HO", financial_year(on), "NEW") if v)


def receivable_qty(ordered: Any, accepted: Any, tolerance_percent: Any = 0) -> float:
    """Quantity that may still be received on a PO line: ordered × (1 + tolerance %) − accepted (never negative)."""
    tol = min(max(number(tolerance_percent), 0.0), 50.0)
    return round(max(number(ordered) * (1.0 + tol / 100.0) - number(accepted), 0.0), 4)


def derive_po_status(items: Sequence[Mapping[str, Any]], current: str) -> str:
    """Receipt status of an approved PO from its line quantities (mirrors qcms_post_indirect_grn)."""
    if str(current).upper() not in OPEN_PO_STATUSES | {"CLOSED"}:
        return str(current).upper()
    if items and all(number(i.get("accepted_qty")) + 0.0005 >= number(i.get("quantity")) for i in items):
        return "CLOSED"
    if any(number(i.get("received_qty")) > 0 for i in items):
        return "PARTIALLY_RECEIVED"
    return "APPROVED"


def employee_name(row: Mapping[str, Any] | None) -> str:
    row = row or {}
    return " ".join(v for v in (clean(row.get("first_name")), clean(row.get("last_name"))) if v) or clean(row.get("employee_code"))


def _iso(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return value or None


def _branch_snapshot(row: Mapping[str, Any] | None) -> dict[str, Any]:
    from core.branch_context import branch_snapshot
    return branch_snapshot(row)


def _party_snapshot(row: Mapping[str, Any] | None) -> dict[str, Any]:
    row = row or {}
    return {key: row.get(key) for key in ("party_code", "party_name", "address", "city", "state", "country", "tax_identifier", "phone", "email", "contact_person")}


# ------------------------------------------------------------------------------------------------ service
class IndirectPurchaseService:
    def __init__(self, repo: Repository | Any | None = None):
        self.repo = repo or Repository()
        self._cache: dict[tuple, Any] = {}

    def _memo(self, key: tuple, loader):
        if key not in self._cache:
            self._cache[key] = loader()
        return self._cache[key]

    def invalidate(self) -> None:
        self._cache.clear()

    def _select(self, table: str, **kwargs: Any) -> list[dict]:
        try:
            return list(self.repo.select(table, **kwargs) or [])
        except Exception:
            return []

    # -------------------------------------------------------------- masters / lookups
    def parties(self) -> list[dict]:
        return self._memo(("parties",), lambda: self._select("parties", eq={"status": "ACTIVE"}, order_by="party_name", limit=5000))

    def suppliers(self) -> list[dict]:
        return [r for r in self.parties() if {str(v).upper() for v in (r.get("party_types") or [])} & SUPPLIER_TYPES]

    def party(self, party_id: Any) -> dict:
        pid = str(party_id or "")
        return next((r for r in self.parties() if str(r.get("id")) == pid), None) or (self.repo.get("parties", pid) or {} if pid else {})

    def branches(self) -> list[dict]:
        return self._memo(("branches",), lambda: self._select("company_branches", eq={"status": "ACTIVE"}, order_by="branch_code", limit=500))

    def branch(self, branch_id: Any) -> dict:
        bid = str(branch_id or "")
        return next((r for r in self.branches() if str(r.get("id")) == bid), {}) if bid else {}

    def employees(self) -> list[dict]:
        return self._memo(("employees",), lambda: self._select("employees", eq={"status": "ACTIVE"}, order_by="first_name", limit=5000))

    def employee(self, employee_id: Any) -> dict:
        eid = str(employee_id or "")
        if not eid:
            return {}
        return next((r for r in self.employees() if str(r.get("id")) == eid), None) or (self.repo.get("employees", eid) or {})

    def departments(self) -> list[str]:
        values = {clean(r.get("department")) for r in self.employees() if clean(r.get("department"))}
        return sorted(values, key=str.casefold)

    def categories(self) -> list[dict]:
        rows = self._memo(("categories",), lambda: self._select("indirect_item_categories", order_by="category_name", limit=500))
        by_code = {str(code): {"category_code": code, "category_name": name, "status": "ACTIVE", "_default": True} for code, name in DEFAULT_CATEGORIES}
        for row in rows:
            by_code[str(row.get("category_code") or "").upper()] = dict(row)
        return [r for r in by_code.values() if str(r.get("status") or "ACTIVE").upper() == "ACTIVE"]

    def category_labels(self) -> dict[str, str]:
        return {str(r.get("category_code")): str(r.get("category_name") or r.get("category_code")) for r in self.categories()}

    def save_category(self, code: str, name: str, description: str = "", status: str = "ACTIVE") -> dict:
        key = "".join(ch if ch.isalnum() else "_" for ch in clean(code).upper()).strip("_")
        if not key or not clean(name):
            raise ValueError("Category code and name are required.")
        payload = {"category_code": key, "category_name": clean(name), "description": clean(description) or None, "status": status}
        row, _ = self.repo.upsert_by("indirect_item_categories", payload, natural_key={"category_code": key})
        self.invalidate()
        return row

    def items(self, *, active_only: bool = False) -> list[dict]:
        rows = self._memo(("items",), lambda: self._select("indirect_items", order_by="item_code", limit=10000))
        return [r for r in rows if str(r.get("status") or "ACTIVE").upper() == "ACTIVE"] if active_only else rows

    def item(self, item_id: Any) -> dict:
        iid = str(item_id or "")
        return next((r for r in self.items() if str(r.get("id")) == iid), None) or (self.repo.get("indirect_items", iid) or {} if iid else {})

    def next_number(self, series: str, branch_id: Any = None) -> str:
        value = clean(self.repo.rpc("qcms_next_indirect_number", {"p_series_key": series.upper(), "p_company_branch_id": str(branch_id) if branch_id else None}))
        if not value:
            raise RuntimeError("The Indirect Purchase number series did not return a number. Run the R37 database step.")
        return value

    def item_payload(self, payload: Mapping[str, Any]) -> dict[str, Any]:
        p = dict(payload)
        name = clean(p.get("item_name"))
        if not name:
            raise ValueError("Item name is required.")
        gst = number(p.get("gst_percent"))
        if gst < 0 or gst > 100:
            raise ValueError("GST % must be between 0 and 100.")
        for key in ("reorder_level", "min_stock", "max_stock"):
            if p.get(key) not in (None, "") and number(p.get(key)) < 0:
                raise ValueError(f"{key.replace('_', ' ').title()} cannot be negative.")
        if p.get("min_stock") not in (None, "") and p.get("max_stock") not in (None, "") and number(p.get("max_stock")) and number(p["min_stock"]) > number(p["max_stock"]):
            raise ValueError("Minimum stock cannot be more than maximum stock.")
        return {
            "item_code": clean(p.get("item_code")).upper() or None, "item_name": name,
            "category": clean(p.get("category")).upper() or "OTHERS", "description": clean(p.get("description")) or None,
            "specification": str(p.get("specification") or "").strip() or None, "uom": clean(p.get("uom")).upper() or "NOS",
            "hsn_sac": clean(p.get("hsn_sac")) or None, "gst_percent": gst,
            "default_supplier_id": str(p.get("default_supplier_id") or "") or None,
            "last_price": number(p.get("last_price")) if p.get("last_price") not in (None, "") else None,
            "is_stock_item": bool(p.get("is_stock_item", True)),
            "reorder_level": number(p.get("reorder_level")) if p.get("reorder_level") not in (None, "") else None,
            "min_stock": number(p.get("min_stock")) if p.get("min_stock") not in (None, "") else None,
            "max_stock": number(p.get("max_stock")) if p.get("max_stock") not in (None, "") else None,
            "storage_location": clean(p.get("storage_location")) or None, "status": clean(p.get("status")).upper() or "ACTIVE",
            "remarks": clean(p.get("remarks")) or None,
        }

    def save_item(self, payload: Mapping[str, Any], record_id: str | None = None) -> dict:
        data = self.item_payload(payload)
        for row in self.items():
            if record_id and str(row.get("id")) == str(record_id):
                continue
            if data["item_code"] and clean(row.get("item_code")).upper() == data["item_code"]:
                raise ValueError(f"Item code {data['item_code']} already exists.")
            if clean(row.get("item_name")).casefold() == data["item_name"].casefold() and clean(row.get("uom")).upper() == data["uom"] and str(row.get("status") or "ACTIVE") == "ACTIVE":
                raise ValueError(f"An active item '{data['item_name']}' ({data['uom']}) already exists as {row.get('item_code')}.")
        if record_id:
            if not data["item_code"]:
                data.pop("item_code")
            row = self.repo.update("indirect_items", str(record_id), data)
        else:
            if not data["item_code"]:
                data["item_code"] = self.next_number("IITEM")
            row = self.repo.insert("indirect_items", data)
        self.invalidate()
        return row

    IMPORT_COLUMNS = ("Item Code", "Item Name", "Category", "Description", "Specification", "UOM", "HSN/SAC", "GST %",
                      "Default Supplier Code", "Last Price", "Stock Item (Y/N)", "Reorder Level", "Min Stock", "Max Stock", "Storage Location", "Status")

    def import_template_bytes(self) -> bytes:
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Indirect Items"
        ws.append(list(self.IMPORT_COLUMNS))
        ws.append(["", "Safety Gloves Nitrile", "SAFETY_PPE", "Nitrile coated", "EN388", "PAIR", "39261019", 12, "", 85, "Y", 50, 20, 200, "Store Rack A1", "ACTIVE"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B6E70")
            ws.column_dimensions[cell.column_letter].width = 18
        guide = wb.create_sheet("Guide")
        guide.append(["Column", "Rule"])
        guide.append(["Item Code", "Blank = next number from the IND series; existing code = update that item"])
        guide.append(["Category", "One of: " + ", ".join(c for c, _ in DEFAULT_CATEGORIES) + " (or a category added in the master)"])
        guide.append(["UOM", ", ".join(UOMS)])
        guide.append(["Default Supplier Code", "Party Code of a Supplier in Party Master (optional)"])
        out = BytesIO(); wb.save(out); return out.getvalue()

    def import_preview(self, rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
        codes = {clean(r.get("item_code")).upper(): r for r in self.items()}
        suppliers = {clean(r.get("party_code")).upper(): r for r in self.suppliers() if r.get("party_code")}
        categories = set(self.category_labels())
        preview: list[dict[str, Any]] = []
        seen: set[str] = set()
        for index, raw in enumerate(rows, start=2):
            get = lambda k: raw.get(k) if not isinstance(raw.get(k), float) or raw.get(k) == raw.get(k) else None  # NaN guard
            name = clean(get("Item Name"))
            if not name and not clean(get("Item Code")):
                continue
            code = clean(get("Item Code")).upper()
            category = clean(get("Category")).upper().replace(" ", "_").replace("&", "").replace("/", "_") or "OTHERS"
            supplier_code = clean(get("Default Supplier Code")).upper()
            errors: list[str] = []
            if not name:
                errors.append("Item Name is required")
            if category not in categories:
                errors.append(f"unknown category {category}")
            if supplier_code and supplier_code not in suppliers:
                errors.append(f"supplier {supplier_code} not found")
            key = code or f"{name.casefold()}|{clean(get('UOM')).upper()}"
            if key in seen:
                errors.append("duplicate row in file")
            seen.add(key)
            payload = {
                "item_code": code, "item_name": name, "category": category, "description": get("Description"),
                "specification": get("Specification"), "uom": clean(get("UOM")).upper() or "NOS", "hsn_sac": clean(get("HSN/SAC")),
                "gst_percent": number(get("GST %")) if get("GST %") not in (None, "") else 18,
                "default_supplier_id": (suppliers.get(supplier_code) or {}).get("id"), "last_price": get("Last Price"),
                "is_stock_item": clean(get("Stock Item (Y/N)")).upper() not in {"N", "NO", "FALSE", "0"},
                "reorder_level": get("Reorder Level"), "min_stock": get("Min Stock"), "max_stock": get("Max Stock"),
                "storage_location": get("Storage Location"), "status": clean(get("Status")).upper() or "ACTIVE",
            }
            try:
                self.item_payload(payload)
            except ValueError as exc:
                errors.append(str(exc))
            existing = codes.get(code) if code else None
            preview.append({"Row": index, "Action": "ERROR" if errors else ("UPDATE" if existing else "CREATE"), "Item Code": code or "(auto)",
                            "Item Name": name, "Category": category, "UOM": payload["uom"], "Issues": "; ".join(errors),
                            "_payload": payload, "_id": (existing or {}).get("id")})
        return preview

    def apply_import(self, preview: Sequence[Mapping[str, Any]]) -> dict[str, int]:
        counts = {"created": 0, "updated": 0, "skipped": 0}
        for row in preview:
            if row.get("Action") == "ERROR":
                counts["skipped"] += 1
                continue
            self.save_item(row["_payload"], record_id=row.get("_id") or None)
            counts["updated" if row.get("_id") else "created"] += 1
        return counts

    # -------------------------------------------------------------- number series settings
    SERIES = (("IPO", "Indirect Purchase Order"), ("IGRN", "Indirect GRN"), ("IPR", "Purchase Requisition"), ("IITEM", "Indirect Item Code"), ("ISS", "Stock Issue"))

    def series_settings(self) -> dict[str, dict]:
        rows = {str(r.get("series_key")).upper(): r for r in self._select("indirect_number_series", limit=50)}
        result = {}
        for key, label in self.SERIES:
            default = {"series_key": key, "prefix": "IND" if key == "IITEM" else key, "include_branch": key != "IITEM", "include_year": key != "IITEM",
                       "year_mode": "FY", "digits": 5 if key == "IITEM" else 4, "separator": "-"}
            result[key] = {**default, **rows.get(key, {}), "label": label}
        return result

    def save_series(self, key: str, payload: Mapping[str, Any]) -> dict:
        data = {"series_key": key.upper(), "prefix": clean(payload.get("prefix")).upper() or key.upper(),
                "include_branch": bool(payload.get("include_branch")), "include_year": bool(payload.get("include_year")),
                "year_mode": "CALENDAR" if str(payload.get("year_mode")).upper() == "CALENDAR" else "FY",
                "digits": int(min(max(number(payload.get("digits")) or 4, 3), 8)), "separator": str(payload.get("separator") or "-")[:2] or "-"}
        row, _ = self.repo.upsert_by("indirect_number_series", data, natural_key={"series_key": key.upper()})
        return row

    # -------------------------------------------------------------- requisitions
    def requisitions(self) -> list[dict]:
        return self._memo(("requisitions",), lambda: self._select("indirect_purchase_requisitions", order_by="created_at", desc=True, limit=3000))

    def requisition(self, pr_id: Any) -> dict:
        rid = str(pr_id or "")
        return next((r for r in self.requisitions() if str(r.get("id")) == rid), None) or (self.repo.get("indirect_purchase_requisitions", rid) or {} if rid else {})

    def requisition_items(self, pr_id: Any) -> list[dict]:
        return self._select("indirect_purchase_requisition_items", eq={"requisition_id": str(pr_id)}, order_by="line_no", limit=500) if pr_id else []

    def create_requisition(self, header: Mapping[str, Any], lines: Sequence[Mapping[str, Any]]) -> dict:
        valid = [l for l in lines if number(l.get("quantity")) > 0 and (l.get("item_id") or clean(l.get("description")) or clean(l.get("item_name")))]
        if not valid:
            raise ValueError("Add at least one requisition line with a quantity.")
        if not clean(header.get("department")):
            raise ValueError("Department is required.")
        number_value = self.next_number("IPR", header.get("company_branch_id"))
        pr = self.repo.insert("indirect_purchase_requisitions", {
            "pr_number": number_value, "pr_date": _iso(header.get("pr_date")) or date.today().isoformat(),
            "company_branch_id": header.get("company_branch_id") or None, "department": clean(header.get("department")),
            "cost_centre": clean(header.get("cost_centre")) or None, "requested_by_employee_id": header.get("requested_by_employee_id") or None,
            "requested_by_name": clean(header.get("requested_by_name")) or None, "required_date": _iso(header.get("required_date")),
            "purpose": clean(header.get("purpose")) or None, "priority": clean(header.get("priority")).upper() or "NORMAL",
            "status": "SUBMITTED", "remarks": clean(header.get("remarks")) or None,
        })
        for index, line in enumerate(valid, start=1):
            item = self.item(line.get("item_id")) if line.get("item_id") else {}
            self.repo.insert("indirect_purchase_requisition_items", {
                "requisition_id": pr["id"], "line_no": index, "item_id": line.get("item_id") or None,
                "item_code": item.get("item_code") or clean(line.get("item_code")) or None,
                "item_name": item.get("item_name") or clean(line.get("item_name")) or clean(line.get("description")),
                "description": clean(line.get("description")) or item.get("description"), "uom": clean(line.get("uom")).upper() or item.get("uom") or "NOS",
                "quantity": number(line.get("quantity")), "estimated_rate": number(line.get("estimated_rate")) or item.get("last_price"),
                "required_date": _iso(line.get("required_date")) or _iso(header.get("required_date")), "remarks": clean(line.get("remarks")) or None,
            })
        self.invalidate()
        return pr

    def decide_requisition(self, pr_id: str, approve: bool, remarks: str, employee_id: str | None) -> dict:
        pr = self.requisition(pr_id)
        if str(pr.get("status") or "").upper() != "SUBMITTED":
            raise ValueError("Only a submitted requisition can be approved or rejected.")
        if not approve and len(clean(remarks)) < 3:
            raise ValueError("Enter the rejection remarks.")
        if employee_id and str(pr.get("requested_by_employee_id") or "") == str(employee_id):
            raise ValueError("Self-approval is not permitted for a requisition.")
        row = self.repo.update("indirect_purchase_requisitions", str(pr_id), {
            "status": "APPROVED" if approve else "REJECTED", "approver_employee_id": employee_id or None,
            "approved_at": datetime.now(timezone.utc).isoformat(), "approval_remarks": clean(remarks) or None})
        self.invalidate()
        return row

    def convertible_requisitions(self) -> list[dict]:
        return [r for r in self.requisitions() if str(r.get("status") or "").upper() in {"APPROVED", "PARTIALLY_ORDERED"}]

    # -------------------------------------------------------------- purchase orders
    def purchase_orders(self) -> list[dict]:
        return self._memo(("pos",), lambda: self._select("indirect_purchase_orders", order_by="created_at", desc=True, limit=5000))

    def all_po_items(self) -> list[dict]:
        return self._memo(("po_items",), lambda: self._select("indirect_purchase_order_items", order_by="line_no", limit=50000))

    def po_items(self, po_id: Any) -> list[dict]:
        pid = str(po_id or "")
        if not pid:
            return []
        rows = [r for r in self.all_po_items() if str(r.get("purchase_order_id")) == pid]
        if not rows:
            rows = self._select("indirect_purchase_order_items", eq={"purchase_order_id": pid}, order_by="line_no", limit=500)
        return sorted(rows, key=lambda r: int(number(r.get("line_no"))))

    def purchase_order(self, po_id: Any) -> dict:
        pid = str(po_id or "")
        if not pid:
            return {}
        row = next((r for r in self.purchase_orders() if str(r.get("id")) == pid), None) or self.repo.get("indirect_purchase_orders", pid) or {}
        return self._enrich_approver(row)

    def _enrich_approver(self, row: Mapping[str, Any]) -> dict:
        item = dict(row or {})
        emp = self.employee(item.get("approver_employee_id")) if item.get("approver_employee_id") else {}
        if emp:
            item["approver_employee_name"] = employee_name(emp)
            item["approver_employee_code"] = clean(emp.get("employee_code"))
        return item

    def build_po(self, payload: Mapping[str, Any], lines: Sequence[Mapping[str, Any]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        """Validate + compute header and lines (no database write). Used for preview and save."""
        p = dict(payload)
        supplier = self.party(p.get("supplier_id"))
        if not supplier:
            raise ValueError("Select the supplier.")
        branch = self.branch(p.get("company_branch_id"))
        if not branch and self.branches():
            raise ValueError("Select the Company Branch / Plant.")
        plant = _branch_snapshot(branch) if branch else {}
        valid: list[dict[str, Any]] = []
        for raw in lines:
            line = dict(raw)
            if number(line.get("quantity")) <= 0 and not line.get("item_id") and not clean(line.get("description")):
                continue
            item = self.item(line.get("item_id")) if line.get("item_id") else {}
            name = clean(line.get("item_name")) or clean(item.get("item_name")) or clean(line.get("description"))
            if not name:
                raise ValueError(f"Line {len(valid) + 1}: select an item or enter a description.")
            if number(line.get("quantity")) <= 0:
                raise ValueError(f"Line {len(valid) + 1} ({name}): quantity must be more than zero.")
            if number(line.get("rate")) < 0:
                raise ValueError(f"Line {len(valid) + 1} ({name}): rate cannot be negative.")
            if not 0 <= number(line.get("discount_percent")) <= 100:
                raise ValueError(f"Line {len(valid) + 1} ({name}): discount % must be 0–100.")
            valid.append({
                "item_id": str(line.get("item_id") or "") or None, "item_code": clean(line.get("item_code")) or clean(item.get("item_code")) or None,
                "item_name": name, "description": clean(line.get("description")) or clean(item.get("description")) or None,
                "category": clean(line.get("category")) or clean(item.get("category")) or None,
                "hsn_sac": clean(line.get("hsn_sac")) or clean(item.get("hsn_sac")) or None,
                "uom": clean(line.get("uom")).upper() or clean(item.get("uom")).upper() or "NOS",
                "quantity": number(line.get("quantity")), "rate": number(line.get("rate")),
                "discount_percent": number(line.get("discount_percent")),
                "gst_percent": number(line.get("gst_percent")) if line.get("gst_percent") not in (None, "") else number(item.get("gst_percent")),
                "required_date": _iso(line.get("required_date")), "requisition_item_id": line.get("requisition_item_id") or None,
                "remarks": clean(line.get("remarks")) or None,
            })
        if not valid:
            raise ValueError("Add at least one PO line.")
        mode = clean(p.get("gst_mode")).upper() or gst_mode(supplier, plant)
        currency = clean(p.get("currency")).upper() or "INR"
        totals = compute_totals(valid, mode, freight=p.get("freight_amount"), other=p.get("other_amount"), round_to_rupee=bool(p.get("round_to_rupee", True)), currency=currency)
        ship_source = clean(p.get("ship_to_source_type")).upper() or "BRANCH"
        if ship_source == "BRANCH":
            ship_branch = self.branch(p.get("ship_to_branch_id") or p.get("company_branch_id"))
            ship_snapshot = _branch_snapshot(ship_branch) if ship_branch else plant
        else:
            ship_snapshot = _party_snapshot(self.party(p.get("ship_to_party_id")))
        if p.get("delivery_date") and p.get("po_date") and str(_iso(p.get("delivery_date"))) < str(_iso(p.get("po_date"))):
            raise ValueError("Delivery date cannot be before the PO date.")
        header = {
            "po_date": _iso(p.get("po_date")) or date.today().isoformat(), "company_branch_id": str(p.get("company_branch_id") or "") or None,
            "plant_snapshot": plant, "supplier_id": str(supplier.get("id")), "vendor_snapshot": _party_snapshot(supplier),
            "ship_to_source_type": ship_source, "ship_to_branch_id": (str(p.get("ship_to_branch_id") or p.get("company_branch_id") or "") or None) if ship_source == "BRANCH" else None,
            "ship_to_party_id": (str(p.get("ship_to_party_id") or "") or None) if ship_source != "BRANCH" else None, "ship_to_snapshot": ship_snapshot,
            "department": clean(p.get("department")) or None, "cost_centre": clean(p.get("cost_centre")) or None,
            "requisitioner": clean(p.get("requisitioner")) or None, "requisitioner_employee_id": p.get("requisitioner_employee_id") or None,
            "requisition_id": p.get("requisition_id") or None, "delivery_date": _iso(p.get("delivery_date")),
            "payment_term": clean(p.get("payment_term")) or None, "incoterm": clean(p.get("incoterm")) or None, "ship_via": clean(p.get("ship_via")) or None,
            "quotation_reference": clean(p.get("quotation_reference")) or None, "quotation_date": _iso(p.get("quotation_date")),
            "currency": currency, "gst_mode": totals["gst_mode"],
            **{k: totals[k] for k in ("gross_amount", "discount_amount", "subtotal", "cgst_amount", "sgst_amount", "igst_amount", "freight_amount", "other_amount", "round_off", "grand_total", "amount_in_words")},
            "terms": str(p.get("terms") or "").strip() or None, "special_instructions": str(p.get("special_instructions") or "").strip() or None,
            "remarks": clean(p.get("remarks")) or None, "print_options": dict(p.get("print_options") or {}),
            "over_receipt_tolerance_percent": min(max(number(p.get("over_receipt_tolerance_percent")), 0.0), 50.0),
        }
        items = []
        for index, line in enumerate(totals["lines"], start=1):
            items.append({**{k: line.get(k) for k in ("item_id", "item_code", "item_name", "description", "category", "hsn_sac", "uom", "quantity", "rate",
                                                     "discount_percent", "discount_amount", "taxable_amount", "gst_percent", "gst_amount", "cgst_amount",
                                                     "sgst_amount", "igst_amount", "line_total", "required_date", "requisition_item_id", "remarks")},
                          "line_no": index})
        return header, items

    def preview_po(self, payload: Mapping[str, Any], lines: Sequence[Mapping[str, Any]]) -> tuple[dict, list[dict]]:
        header, items = self.build_po(payload, lines)
        branch = self.branch(header.get("company_branch_id"))
        header.update({"po_number": preview_number("IPO", branch.get("plant_code") or branch.get("branch_code") or ""), "status": "DRAFT", "approval_status": "DRAFT"})
        return header, items

    def _insert_lines(self, po_id: str, items: Sequence[Mapping[str, Any]]) -> None:
        for item in items:
            self.repo.insert("indirect_purchase_order_items", {**dict(item), "purchase_order_id": po_id})

    def create_po(self, payload: Mapping[str, Any], lines: Sequence[Mapping[str, Any]], *, submit: bool = True, submitted_by_employee_id: str | None = None) -> dict:
        header, items = self.build_po(payload, lines)
        if submit and not submitted_by_employee_id:
            raise ValueError("Your login must be linked to an ACTIVE Employee Master record to submit an Indirect PO for approval.")
        header["po_number"] = self.next_number("IPO", header.get("company_branch_id"))
        status = "PENDING_APPROVAL" if submit else "DRAFT"
        header.update({"status": status, "approval_status": status, "approval_level": None,
                       "submitted_by_employee_id": submitted_by_employee_id or None,
                       "submitted_at": datetime.now(timezone.utc).isoformat() if submit else None})
        po = self.repo.insert("indirect_purchase_orders", header)
        try:
            self._insert_lines(str(po["id"]), items)
        except Exception:
            try:
                self.repo.update("indirect_purchase_orders", str(po["id"]), {"status": "CANCELLED", "approval_status": "CANCELLED",
                                                                             "cancellation_reason": "Auto-cancelled: PO lines could not be saved"})
            except Exception:
                pass
            raise
        self._mark_requisition(header.get("requisition_id"), str(po["id"]), items)
        self.invalidate()
        return {**po, "_items": items}

    def _mark_requisition(self, pr_id: Any, po_id: str, items: Sequence[Mapping[str, Any]]) -> None:
        if not pr_id:
            return
        try:
            ordered: dict[str, float] = {}
            for item in items:
                if item.get("requisition_item_id"):
                    ordered[str(item["requisition_item_id"])] = ordered.get(str(item["requisition_item_id"]), 0.0) + number(item.get("quantity"))
            pr_lines = self.requisition_items(pr_id)
            for line in pr_lines:
                if str(line.get("id")) in ordered:
                    self.repo.update("indirect_purchase_requisition_items", str(line["id"]), {"ordered_qty": number(line.get("ordered_qty")) + ordered[str(line["id"])]})
            fully = all(number(l.get("ordered_qty")) + ordered.get(str(l.get("id")), 0.0) + 0.0005 >= number(l.get("quantity")) for l in pr_lines) if pr_lines else True
            self.repo.update("indirect_purchase_requisitions", str(pr_id), {"status": "CONVERTED" if fully else "PARTIALLY_ORDERED", "converted_po_id": po_id})
        except Exception:
            pass

    def has_receipts(self, po_id: str) -> bool:
        if any(number(i.get("received_qty")) > 0 for i in self.po_items(po_id)):
            return True
        return any(str(g.get("purchase_order_id")) == str(po_id) and str(g.get("status") or "").upper() in {"POSTED", "DRAFT"} for g in self.grns())

    def update_po(self, po_id: str, payload: Mapping[str, Any], lines: Sequence[Mapping[str, Any]], *, submit: bool = True, submitted_by_employee_id: str | None = None) -> dict:
        current = self.purchase_order(po_id)
        status = str(current.get("status") or "").upper()
        if status not in EDITABLE_PO_STATUSES:
            raise ValueError(f"An Indirect PO in status {STATUS_LABELS.get(status, status)} cannot be edited.")
        if self.has_receipts(po_id):
            raise ValueError("Material has already been received against this PO; it cannot be revised. Cancel the balance and raise a new PO.")
        if submit and not submitted_by_employee_id:
            raise ValueError("Your login must be linked to an ACTIVE Employee Master record to submit an Indirect PO for approval.")
        header, items = self.build_po(payload, lines)
        new_status = "PENDING_APPROVAL" if submit else "DRAFT"
        header.update({"status": new_status, "approval_status": new_status, "approval_level": None,
                       "submitted_by_employee_id": submitted_by_employee_id or current.get("submitted_by_employee_id"),
                       "submitted_at": datetime.now(timezone.utc).isoformat() if submit else current.get("submitted_at"),
                       "revision_no": int(number(current.get("revision_no"))) + (1 if status == "APPROVED" else 0)})
        header.pop("requisition_id", None)
        row = self.repo.update("indirect_purchase_orders", str(po_id), header)
        for old in self.po_items(po_id):
            self.repo.delete("indirect_purchase_order_items", str(old["id"]))
        self._insert_lines(str(po_id), items)
        self.invalidate()
        return row

    def submit_po(self, po_id: str, employee_id: str | None) -> dict:
        current = self.purchase_order(po_id)
        if str(current.get("status") or "").upper() not in {"DRAFT", "REJECTED"}:
            raise ValueError("Only a Draft or Rejected Indirect PO can be submitted.")
        if not employee_id:
            raise ValueError("Your login must be linked to an ACTIVE Employee Master record.")
        row = self.repo.update("indirect_purchase_orders", str(po_id), {"status": "PENDING_APPROVAL", "approval_status": "PENDING_APPROVAL", "approval_level": None,
                                                                          "submitted_by_employee_id": employee_id, "submitted_at": datetime.now(timezone.utc).isoformat()})
        self.invalidate()
        return row

    def cancel_po(self, po_id: str, reason: str, employee_id: str | None) -> dict:
        current = self.purchase_order(po_id)
        if str(current.get("status") or "").upper() in {"CLOSED", "CANCELLED"}:
            raise ValueError("This Indirect PO is already closed / cancelled.")
        if len(clean(reason)) < 3:
            raise ValueError("Enter the cancellation reason.")
        if any(number(i.get("received_qty")) > 0 for i in self.po_items(po_id)):
            raise ValueError("Material has been received against this PO; it cannot be cancelled. Short-close is recorded by a final GRN.")
        row = self.repo.update("indirect_purchase_orders", str(po_id), {"status": "CANCELLED", "approval_status": "CANCELLED", "cancellation_reason": clean(reason),
                                                                          "cancelled_by_employee_id": employee_id or None, "cancelled_at": datetime.now(timezone.utc).isoformat()})
        self.invalidate()
        return row

    def set_print_options(self, po_id: str, options: Mapping[str, Any]) -> dict:
        return self.repo.update("indirect_purchase_orders", str(po_id), {"print_options": {k: bool(v) for k, v in dict(options or {}).items() if str(k).startswith("show_")}})

    # -------------------------------------------------------------- approval
    def approval_target(self, po_id: str) -> dict:
        try:
            result = self.repo.rpc("qcms_indirect_po_approval_target", {"p_po_id": str(po_id)})
            if isinstance(result, Mapping):
                return dict(result)
        except Exception:
            pass
        po = self.purchase_order(po_id)
        submitter = self.employee(po.get("submitted_by_employee_id"))
        dept = clean(po.get("department") or submitter.get("department")).casefold()
        routes = [r for r in self.approval_routes() if r.get("required", True) and r.get("employee_id")
                  and number(r.get("min_amount")) <= number(po.get("grand_total"))
                  and (not clean(r.get("department")) or clean(r.get("department")).casefold() == dept)]
        levels = sorted({int(number(r.get("level_no")) or 1) for r in routes})
        if levels:
            level = next((l for l in levels if l >= int(number(po.get("approval_level")))), levels[-1])
            route = sorted([r for r in routes if int(number(r.get("level_no")) or 1) == level], key=lambda r: 0 if clean(r.get("department")) else 1)[0]
            emp = self.employee(route.get("employee_id"))
            return {"source": "CONFIGURED_ROUTE", "employee_id": emp.get("id"), "employee_name": employee_name(emp), "email": emp.get("email"),
                    "level_no": level, "level_name": clean(route.get("level_name")) or f"Level {level} Approval", "levels": levels}
        manager = self.employee(submitter.get("reports_to_employee_id"))
        if manager:
            return {"source": "REPORTS_TO", "employee_id": manager.get("id"), "employee_name": employee_name(manager), "email": manager.get("email"),
                    "level_no": 1, "level_name": "Reports-To Manager Approval", "levels": [1]}
        return {"source": "PERMISSION_FALLBACK", "employee_id": None, "employee_name": None, "email": None, "level_no": 1,
                "level_name": "Any different employee with Indirect Purchase Approve permission", "levels": [1]}

    def approve_po(self, po_id: str, remarks: str = "") -> dict:
        result = dict(self.repo.rpc("qcms_approve_indirect_po", {"p_po_id": str(po_id), "p_remarks": clean(remarks) or None}) or {})
        self.invalidate()
        return self._enrich_approver(result)

    def reject_po(self, po_id: str, remarks: str) -> dict:
        if len(clean(remarks)) < 3:
            raise ValueError("Enter the rejection remarks.")
        result = dict(self.repo.rpc("qcms_reject_indirect_po", {"p_po_id": str(po_id), "p_remarks": clean(remarks)}) or {})
        self.invalidate()
        return result

    def approval_log(self, po_id: str) -> list[dict]:
        return self._select("indirect_po_approval_log", eq={"purchase_order_id": str(po_id)}, order_by="created_at", limit=200)

    def approval_routes(self) -> list[dict]:
        return self._memo(("routes",), lambda: [r for r in self._select("qcms_module_approval_routes", eq={"module_key": MODULE_KEY}, order_by="level_no", limit=200)
                                                if str(r.get("status") or "ACTIVE").upper() == "ACTIVE"])

    def save_route(self, *, level_no: int, employee_id: str, level_name: str, department: str = "", min_amount: float = 0.0, overdue_after_hours: int = 24) -> dict:
        if not 1 <= int(level_no) <= 3:
            raise ValueError("Approval level must be 1, 2 or 3.")
        if not employee_id:
            raise ValueError("Select the approver employee.")
        dept = clean(department) or None
        payload = {"module_key": MODULE_KEY, "department": dept, "level_no": int(level_no), "level_name": clean(level_name) or f"Level {level_no} Approval",
                   "employee_id": employee_id, "required": True, "status": "ACTIVE", "min_amount": max(number(min_amount), 0.0),
                   "overdue_after_hours": int(min(max(number(overdue_after_hours) or 24, 1), 720))}
        row, _ = self.repo.upsert_by("qcms_module_approval_routes", payload, natural_key={"module_key": MODULE_KEY, "department": dept, "level_no": int(level_no)})
        self.invalidate()
        return row

    def deactivate_route(self, route_id: str) -> dict:
        row = self.repo.update("qcms_module_approval_routes", str(route_id), {"status": "INACTIVE"})
        self.invalidate()
        return row

    # -------------------------------------------------------------- registers
    def po_register_rows(self) -> list[dict[str, Any]]:
        items_by_po: dict[str, list[dict]] = {}
        for item in self.all_po_items():
            items_by_po.setdefault(str(item.get("purchase_order_id")), []).append(item)
        supplier_names = {str(p.get("id")): p.get("party_name") for p in self.parties()}
        branch_codes = {str(b.get("id")): b.get("plant_code") or b.get("branch_code") for b in self.branches()}
        today = date.today().isoformat()
        rows = []
        for po in self.purchase_orders():
            items = items_by_po.get(str(po.get("id")), [])
            ordered = sum(number(i.get("quantity")) for i in items)
            accepted = sum(number(i.get("accepted_qty")) for i in items)
            status = str(po.get("status") or "").upper()
            overdue = status in OPEN_PO_STATUSES and str(po.get("delivery_date") or "9999") < today
            rows.append({
                "IPO Number": po.get("po_number"), "PO Date": po.get("po_date"), "Branch": branch_codes.get(str(po.get("company_branch_id")), "-"),
                "Supplier": (po.get("vendor_snapshot") or {}).get("party_name") or supplier_names.get(str(po.get("supplier_id")), "-"),
                "Department": po.get("department") or "-", "Cost Centre": po.get("cost_centre") or "-",
                "Items": ", ".join(clean(i.get("item_name")) for i in items[:3]) + (f" +{len(items) - 3}" if len(items) > 3 else ""),
                "Lines": len(items), "Grand Total": number(po.get("grand_total")), "Currency": po.get("currency") or "INR",
                "Delivery Date": po.get("delivery_date"), "Received %": round(100.0 * accepted / ordered, 1) if ordered else 0.0,
                "Status": status, "Approval Status": str(po.get("approval_status") or "").upper(), "Delivery": "OVERDUE" if overdue else "",
                "Requisitioner": po.get("requisitioner") or "-", "_id": po.get("id"),
            })
        return rows

    def pending_approval_rows(self) -> list[dict[str, Any]]:
        rows = []
        now = datetime.now(timezone.utc)
        for po in self.purchase_orders():
            if str(po.get("approval_status") or "").upper() != "PENDING_APPROVAL":
                continue
            target = self.approval_target(str(po["id"]))
            try:
                submitted = datetime.fromisoformat(str(po.get("submitted_at") or po.get("created_at")).replace("Z", "+00:00"))
                days = max(0, (now - submitted.astimezone(timezone.utc)).days)
            except Exception:
                days = 0
            rows.append({"IPO Number": po.get("po_number"), "PO Date": po.get("po_date"), "Supplier": (po.get("vendor_snapshot") or {}).get("party_name") or "-",
                         "Department": po.get("department") or "-", "Grand Total": number(po.get("grand_total")),
                         "Requisitioner": po.get("requisitioner") or "-", "Approver": target.get("employee_name") or "ANY APPROVER",
                         "Approver Email": target.get("email") or "-", "Approval Level": target.get("level_name") or "-",
                         "Pending Days": days, "Status": "PENDING_APPROVAL", "_id": po.get("id")})
        return rows

    def open_pos_for_grn(self) -> list[dict]:
        return [po for po in self.purchase_orders() if str(po.get("approval_status") or "").upper() == "APPROVED" and str(po.get("status") or "").upper() in OPEN_PO_STATUSES]

    def po_pending_lines(self, po_id: str) -> list[dict[str, Any]]:
        tolerance = number(self.purchase_order(po_id).get("over_receipt_tolerance_percent"))
        rows = []
        for item in self.po_items(po_id):
            pending = round(number(item.get("quantity")) - number(item.get("accepted_qty")), 4)
            # Cumulative PO-line receipts are renamed (po_*) so they are never mistaken for this GRN's received / accepted quantities.
            line = {k: v for k, v in item.items() if k not in {"received_qty", "accepted_qty", "rejected_qty"}}
            rows.append({**line, "po_received_qty": number(item.get("received_qty")), "po_accepted_qty": number(item.get("accepted_qty")),
                         "po_rejected_qty": number(item.get("rejected_qty")), "pending_qty": max(pending, 0.0), "tolerance_percent": tolerance,
                         "max_receivable_qty": receivable_qty(item.get("quantity"), item.get("accepted_qty"), tolerance)})
        return rows

    # -------------------------------------------------------------- GRN
    def grns(self) -> list[dict]:
        return self._memo(("grns",), lambda: self._select("indirect_grns", order_by="created_at", desc=True, limit=5000))

    def grn(self, grn_id: Any) -> dict:
        gid = str(grn_id or "")
        return next((r for r in self.grns() if str(r.get("id")) == gid), None) or (self.repo.get("indirect_grns", gid) or {} if gid else {})

    def grn_items(self, grn_id: Any) -> list[dict]:
        return self._select("indirect_grn_items", eq={"grn_id": str(grn_id)}, order_by="line_no", limit=500) if grn_id else []

    @staticmethod
    def validate_grn_lines(lines: Sequence[Mapping[str, Any]], *, against_po: bool) -> list[dict[str, Any]]:
        valid: list[dict[str, Any]] = []
        for raw in lines:
            line = dict(raw)
            received = number(line.get("received_qty"))
            if received <= 0:
                continue
            accepted = number(line.get("accepted_qty")) if line.get("accepted_qty") not in (None, "") else received
            rejected = number(line.get("rejected_qty")) if line.get("rejected_qty") not in (None, "") else max(received - accepted, 0.0)
            label = clean(line.get("item_code")) or clean(line.get("item_name")) or f"line {len(valid) + 1}"
            if accepted < 0 or rejected < 0:
                raise ValueError(f"{label}: accepted / rejected cannot be negative.")
            if abs(accepted + rejected - received) > 0.0005:
                raise ValueError(f"{label}: accepted ({accepted:g}) + rejected ({rejected:g}) must equal received ({received:g}).")
            if rejected > 0 and not clean(line.get("rejection_reason")):
                raise ValueError(f"{label}: enter the rejection reason.")
            if against_po:
                pending = number(line.get("pending_qty"))
                if line.get("max_receivable_qty") not in (None, ""):
                    limit = number(line.get("max_receivable_qty"))
                    if received > limit + 0.0005:
                        tol = number(line.get("tolerance_percent"))
                        raise ValueError(f"{label}: received {received:g} exceeds the receivable balance {limit:g} "
                                         f"(pending {pending:g} + over-receipt tolerance {tol:g}%).")
                elif line.get("pending_qty") not in (None, "") and received > pending + 0.0005:
                    raise ValueError(f"{label}: received {received:g} exceeds the pending PO quantity {pending:g}.")
            elif not line.get("item_id") and not clean(line.get("item_name")):
                raise ValueError("Select an item (or enter a description) for every received line.")
            valid.append({**line, "received_qty": received, "accepted_qty": accepted, "rejected_qty": rejected})
        if not valid:
            raise ValueError("Enter a received quantity on at least one line.")
        return valid

    def create_grn(self, header: Mapping[str, Any], lines: Sequence[Mapping[str, Any]]) -> dict:
        h = dict(header)
        against_po = str(h.get("grn_type") or "AGAINST_PO").upper() == "AGAINST_PO"
        po = self.purchase_order(h.get("purchase_order_id")) if against_po else {}
        if against_po and not po:
            raise ValueError("Select the approved Indirect PO.")
        if against_po and (str(po.get("approval_status")).upper() != "APPROVED" or str(po.get("status")).upper() not in OPEN_PO_STATUSES):
            raise ValueError(f"Indirect PO {po.get('po_number')} is not approved / open for receipt.")
        if not against_po and not (h.get("supplier_id") or clean(h.get("supplier_name"))):
            raise ValueError("Enter the supplier / shop name for a GRN without PO.")
        valid = self.validate_grn_lines(lines, against_po=against_po)
        branch_id = h.get("company_branch_id") or po.get("company_branch_id")
        supplier = self.party(h.get("supplier_id") or po.get("supplier_id"))
        subtotal = round(sum(number(l.get("accepted_qty")) * number(l.get("rate")) for l in valid), 2)
        gst = round(sum(number(l.get("accepted_qty")) * number(l.get("rate")) * number(l.get("gst_percent")) / 100.0 for l in valid), 2)
        grn = self.repo.insert("indirect_grns", {
            "grn_number": self.next_number("IGRN", branch_id), "grn_date": _iso(h.get("grn_date")) or date.today().isoformat(),
            "grn_type": "AGAINST_PO" if against_po else "WITHOUT_PO", "purchase_order_id": po.get("id") if against_po else None,
            "company_branch_id": branch_id or None, "supplier_id": supplier.get("id") or None,
            "supplier_name": supplier.get("party_name") or clean(h.get("supplier_name")) or None,
            "supplier_invoice_no": clean(h.get("supplier_invoice_no")) or None, "supplier_invoice_date": _iso(h.get("supplier_invoice_date")),
            "challan_no": clean(h.get("challan_no")) or None, "challan_date": _iso(h.get("challan_date")), "vehicle_no": clean(h.get("vehicle_no")).upper() or None,
            "gate_entry_no": clean(h.get("gate_entry_no")) or None, "gate_entry_date": _iso(h.get("gate_entry_date")) if clean(h.get("gate_entry_no")) else None,
            "payment_mode": clean(h.get("payment_mode")).upper() or ("CREDIT" if against_po else "PETTY_CASH"),
            "department": clean(h.get("department") or po.get("department")) or None, "cost_centre": clean(h.get("cost_centre") or po.get("cost_centre")) or None,
            "storage_location": clean(h.get("storage_location")) or None, "received_by_employee_id": h.get("received_by_employee_id") or None,
            "received_by_name": clean(h.get("received_by_name")) or None, "subtotal": subtotal, "gst_amount": gst, "grand_total": round(subtotal + gst, 2),
            "status": "DRAFT", "approval_status": "NOT_REQUIRED" if against_po else "PENDING", "remarks": clean(h.get("remarks")) or None,
        })
        try:
            for index, line in enumerate(valid, start=1):
                item = self.item(line.get("item_id")) if line.get("item_id") else {}
                self.repo.insert("indirect_grn_items", {
                    "grn_id": grn["id"], "line_no": index, "po_item_id": line.get("po_item_id") or line.get("id") if against_po else None,
                    "item_id": line.get("item_id") or None, "item_code": clean(line.get("item_code")) or item.get("item_code"),
                    "item_name": clean(line.get("item_name")) or item.get("item_name"), "description": clean(line.get("description")) or None,
                    "uom": clean(line.get("uom")) or item.get("uom"), "ordered_qty": number(line.get("quantity")) if against_po else None,
                    "received_qty": line["received_qty"], "accepted_qty": line["accepted_qty"], "rejected_qty": line["rejected_qty"],
                    "rejection_reason": clean(line.get("rejection_reason")) or None, "rate": number(line.get("rate")),
                    "gst_percent": number(line.get("gst_percent")), "amount": round(line["accepted_qty"] * number(line.get("rate")), 2),
                    "storage_location": clean(line.get("storage_location")) or clean(h.get("storage_location")) or item.get("storage_location"),
                    "remarks": clean(line.get("remarks")) or None,
                })
            result = dict(self.repo.rpc("qcms_post_indirect_grn", {"p_grn_id": str(grn["id"])}) or {})
        except Exception:
            try:
                self.repo.update("indirect_grns", str(grn["id"]), {"status": "CANCELLED", "remarks": "Auto-cancelled: GRN could not be posted"})
            except Exception:
                pass
            self.invalidate()
            raise
        self.invalidate()
        return {**grn, **result}

    def decide_grn(self, grn_id: str, approve: bool, remarks: str = "") -> dict:
        if not approve and len(clean(remarks)) < 3:
            raise ValueError("Enter the rejection remarks.")
        result = dict(self.repo.rpc("qcms_decide_indirect_grn", {"p_grn_id": str(grn_id), "p_approve": bool(approve), "p_remarks": clean(remarks) or None}) or {})
        self.invalidate()
        return result

    def grn_register_rows(self) -> list[dict[str, Any]]:
        po_numbers = {str(p.get("id")): p.get("po_number") for p in self.purchase_orders()}
        return [{"GRN Number": g.get("grn_number"), "GRN Date": g.get("grn_date"), "Type": str(g.get("grn_type") or "").replace("_", " ").title(),
                 "Indirect PO": po_numbers.get(str(g.get("purchase_order_id")), "-"), "Supplier": g.get("supplier_name") or "-",
                 "Invoice / Challan": " / ".join(v for v in (clean(g.get("supplier_invoice_no")), clean(g.get("challan_no"))) if v) or "-",
                 "Gate Entry": clean(g.get("gate_entry_no")) or "-",
                 "Department": g.get("department") or "-", "Value": number(g.get("grand_total")), "Received By": g.get("received_by_name") or "-",
                 "Status": str(g.get("status") or "").upper(), "_id": g.get("id")} for g in self.grns()]

    # -------------------------------------------------------------- stock
    def ledger(self, *, item_id: str | None = None, branch_id: str | None = None) -> list[dict]:
        eq: dict[str, Any] = {}
        if item_id:
            eq["item_id"] = item_id
        if branch_id:
            eq["company_branch_id"] = branch_id
        return self._select("indirect_stock_ledger", eq=eq or None, order_by="created_at", limit=20000)

    def stock_rows(self, *, branch_id: str | None = None) -> list[dict[str, Any]]:
        balances: dict[str, float] = {}
        last_move: dict[str, str] = {}
        for row in self.ledger(branch_id=branch_id):
            key = str(row.get("item_id"))
            balances[key] = balances.get(key, 0.0) + number(row.get("qty_in")) - number(row.get("qty_out"))
            last_move[key] = str(row.get("txn_date") or "")
        labels = self.category_labels()
        rows = []
        for item in self.items():
            if not item.get("is_stock_item", True):
                continue
            key = str(item.get("id"))
            balance = round(balances.get(key, 0.0), 4)
            reorder = item.get("reorder_level")
            alert = "REORDER" if reorder not in (None, "") and balance <= number(reorder) else ("BELOW MIN" if item.get("min_stock") not in (None, "") and balance < number(item.get("min_stock")) else "")
            if item.get("max_stock") not in (None, "") and number(item.get("max_stock")) and balance > number(item.get("max_stock")):
                alert = "ABOVE MAX"
            rows.append({"Item Code": item.get("item_code"), "Item Name": item.get("item_name"), "Category": labels.get(str(item.get("category")), item.get("category")),
                         "UOM": item.get("uom"), "Balance": balance, "Last Price": number(item.get("last_price")),
                         "Stock Value": round(balance * number(item.get("last_price")), 2), "Reorder Level": item.get("reorder_level"),
                         "Min": item.get("min_stock"), "Max": item.get("max_stock"), "Location": item.get("storage_location") or "-",
                         "Last Movement": last_move.get(key, "-"), "Stock Alert": alert, "Status": str(item.get("status") or "ACTIVE"), "_id": key})
        return rows

    def issue_stock(self, *, item_id: str, qty: float, branch_id: str | None, department: str, cost_centre: str = "", employee_id: str | None = None,
                    issue_date: Any = None, remarks: str = "", txn_type: str = "ISSUE") -> dict:
        if not item_id:
            raise ValueError("Select the item.")
        if txn_type.upper() == "ISSUE" and not clean(department):
            raise ValueError("Select the department the stock is issued to.")
        if number(qty) == 0 or (txn_type.upper() != "ADJUSTMENT" and number(qty) < 0):
            raise ValueError("Enter a quantity greater than zero.")
        result = dict(self.repo.rpc("qcms_issue_indirect_stock", {
            "p_item_id": item_id, "p_qty": number(qty), "p_company_branch_id": branch_id or None, "p_department": clean(department) or None,
            "p_cost_centre": clean(cost_centre) or None, "p_issued_to_employee_id": employee_id or None, "p_issue_date": _iso(issue_date),
            "p_remarks": clean(remarks) or None, "p_txn_type": txn_type.upper()}) or {})
        self.invalidate()
        return result

    # -------------------------------------------------------------- dashboard
    def dashboard(self, today: date | None = None) -> dict[str, Any]:
        today = today or date.today()
        month = today.strftime("%Y-%m")
        pos = self.purchase_orders()
        items_by_po: dict[str, list[dict]] = {}
        for item in self.all_po_items():
            items_by_po.setdefault(str(item.get("purchase_order_id")), []).append(item)
        labels = self.category_labels()
        pending = [p for p in pos if str(p.get("approval_status") or "").upper() == "PENDING_APPROVAL"]
        open_pos = [p for p in pos if str(p.get("status") or "").upper() in OPEN_PO_STATUSES]
        overdue = [p for p in open_pos if p.get("delivery_date") and str(p.get("delivery_date")) < today.isoformat()]
        grns_month = [g for g in self.grns() if str(g.get("status") or "").upper() == "POSTED" and str(g.get("grn_date") or "").startswith(month)]
        spend: dict[str, float] = {}
        spend_supplier: dict[str, float] = {}
        for po in pos:
            if str(po.get("approval_status") or "").upper() != "APPROVED" or not str(po.get("po_date") or "").startswith(month):
                continue
            supplier = clean((po.get("vendor_snapshot") or {}).get("party_name")) or "-"
            spend_supplier[supplier] = round(spend_supplier.get(supplier, 0.0) + number(po.get("subtotal")), 2)
            for item in items_by_po.get(str(po.get("id")), []):
                key = labels.get(str(item.get("category") or "OTHERS"), str(item.get("category") or "Others"))
                spend[key] = round(spend.get(key, 0.0) + number(item.get("taxable_amount")), 2)
        return {
            "pending_approval": len(pending), "pending_value": round(sum(number(p.get("grand_total")) for p in pending), 2),
            "open": len(open_pos), "open_value": round(sum(number(p.get("grand_total")) for p in open_pos), 2),
            "overdue": len(overdue), "received_month": len(grns_month), "received_month_value": round(sum(number(g.get("grand_total")) for g in grns_month), 2),
            "spend_month": round(sum(spend.values()), 2), "spend_by_category": dict(sorted(spend.items(), key=lambda kv: -kv[1])),
            "spend_by_supplier": dict(sorted(spend_supplier.items(), key=lambda kv: -kv[1])[:10]),
            "requisitions_open": len([r for r in self.requisitions() if str(r.get("status") or "").upper() in {"SUBMITTED", "APPROVED", "PARTIALLY_ORDERED"}]),
            "grn_pending_approval": len([g for g in self.grns() if str(g.get("status") or "").upper() == "PENDING_APPROVAL"]),
            "overdue_rows": overdue,
        }


# ------------------------------------------------------------------------------------------------ e-mail
def indirect_notifier(repo: Any | None = None):
    """NotificationService that attaches the Indirect PO / GRN print as the generated PDF."""
    from core.notification_service import NotificationService

    class IndirectNotificationService(NotificationService):
        def _generated_report(self, related_table: str, related_id: str):  # type: ignore[override]
            try:
                if related_table == "indirect_purchase_orders":
                    from core.indirect_po_pdf import build_indirect_po_pdf
                    service = IndirectPurchaseService(self.repo)
                    header = service.purchase_order(related_id)
                    if not header:
                        return None
                    if getattr(self, "po_print_options", None):
                        header = {**header, "print_options": dict(self.po_print_options)}
                    return f"{header.get('po_number') or 'Indirect_PO'}.pdf", build_indirect_po_pdf(header, service.po_items(related_id))
                if related_table == "indirect_grns":
                    from core.indirect_po_pdf import indirect_grn_pdf_bytes
                    grn = IndirectPurchaseService(self.repo).grn(related_id)
                    if not grn:
                        return None
                    return f"{grn.get('grn_number') or 'Indirect_GRN'}.pdf", indirect_grn_pdf_bytes(related_id, self.repo)
            except Exception:
                return None
            return super()._generated_report(related_table, related_id)

    return IndirectNotificationService(repo)


def po_email_text(header: Mapping[str, Any], *, purpose: str = "SUPPLIER") -> tuple[str, str]:
    number_value = clean(header.get("po_number"))
    supplier = clean((header.get("vendor_snapshot") or {}).get("party_name")) or "Supplier"
    plant = dict(header.get("plant_snapshot") or {})
    company = clean(plant.get("name") or plant.get("branch_name")) or "Four Star Industries"
    total = f"{clean(header.get('currency')) or 'INR'} {number(header.get('grand_total')):,.2f}"
    if purpose == "APPROVAL":
        subject = f"Approval required · Indirect Purchase Order {number_value} · {supplier} · {total}"
        body = (f"Dear Approver,\n\nIndirect Purchase Order {number_value} for {supplier} ({total}) is waiting for your approval.\n"
                f"Department: {clean(header.get('department')) or '-'} · Cost centre: {clean(header.get('cost_centre')) or '-'}\n"
                f"Requisitioner: {clean(header.get('requisitioner')) or '-'}\n\nThe draft PO (PENDING APPROVAL watermark) is attached. "
                "Approve or reject it in STAWN QCMS → Indirect Purchase → PO Approval.\n\nRegards,\nSTAWN QCMS")
        return subject, body
    subject = f"Indirect Purchase Order {number_value} · {company}"
    body = (f"Dear {supplier},\n\nPlease find attached our Indirect Purchase Order {number_value} dated {clean(header.get('po_date'))} "
            f"for {total}.\nDelivery date: {clean(header.get('delivery_date')) or 'as agreed'} · Payment term: {clean(header.get('payment_term')) or 'as agreed'}.\n\n"
            "Kindly acknowledge the order and quote the PO number on your invoice and delivery challan.\n\n"
            f"Regards,\n{clean(header.get('requisitioner')) or 'Purchase Department'}\n{company}")
    return subject, body
