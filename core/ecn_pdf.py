"""QCMS R38 · Engineering Change Note / Process Change Note print (format FSI/806/F05 rev 00).

Layout follows the customer workbook (sheets "ECN Form" / "PCN"): format box (FMT NO. / REV. DATE / REV. NO. / PAGE NO),
Date, FSI Representative(s), To / CC / From identification (contact name, address, phone, fax), requested priority &
response, drawing criteria (drawing no, part no, part name), previous description, revision requested, supporting
documents attached, purchase orders / inventory / demands affected, FSI use only (approved / disapproved, drawing
change to occur, customer approval), approval signatures + dates, completed by / date, comments and benefit of change.
The look is the company Purchase Order print (core.purchase_order_reporting Design B: logo, accent title, navy section
bars, open ruled boxes, status watermark behind the content).

Public contract:
* ``build_ecn_pdf(ecn, *, employees=None, part=None, layout=None, actions=None) -> bytes``  (unsaved data / preview;
  ``actions`` = part_ecn_actions rows -> page 2 "IMPLEMENTATION ACTIONS")
* ``ecn_pdf_bytes(ecn_id, repo=None) -> bytes``                                (saved register row)
* ``watermark_label(ecn) -> str``
* ``build_ecn_xlsx(ecn, *, employees=None, part=None, actions=None) -> bytes``     Excel copy in the FSI/806/F05 sheet layout
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Mapping

from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from core import purchase_order_reporting as por
from core.ecn_service import (
    CONTACT_FIELDS, DRAWING_CHANGE_MODES, FORMAT_NO, FORMAT_REV_DATE, FORMAT_REV_NO, SUPPORTING_DOCUMENTS,
    change_title, clean, clean_contact, clean_supporting, employee_name, normalize_status, status_label,
)

LEFT = 28.0
RIGHT = 28.0
MUTED = HexColor("#4B5563")


def watermark_label(ecn: Mapping[str, Any]) -> str:
    status = normalize_status(ecn.get("status"))
    return {"PENDING_APPROVAL": "PENDING APPROVAL", "APPROVED": "APPROVED", "DISAPPROVED": "DISAPPROVED", "COMPLETED": "COMPLETED",
            "CANCELLED": "CANCELLED"}.get(status, "DRAFT")


def _date(value: Any) -> str:
    text = clean(value)[:10]
    if len(text) == 10 and text[4] == "-" and text[7] == "-":
        return f"{text[8:10]}.{text[5:7]}.{text[0:4]}"
    return text


def _qty(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    try:
        return f"{float(text):,.3f}".rstrip("0").rstrip(".")
    except ValueError:
        return text


def _lines(text: Any, width: float, *, size: float, bold: bool = False, max_lines: int = 12) -> list[str]:
    font = "Helvetica-Bold" if bold else "Helvetica"
    out: list[str] = []
    for paragraph in str(text or "").splitlines() or [""]:
        line = ""
        for word in paragraph.split():
            test = f"{line} {word}".strip()
            if stringWidth(test, font, size) <= width:
                line = test
                continue
            if line:
                out.append(line)
            while stringWidth(word, font, size) > width and len(word) > 1:
                cut = len(word)
                while cut > 1 and stringWidth(word[:cut], font, size) > width:
                    cut -= 1
                out.append(word[:cut]); word = word[cut:]
            line = word
        if line:
            out.append(line)
    if len(out) > max_lines:
        out = out[:max_lines]
        out[-1] = out[-1][: max(len(out[-1]) - 3, 1)] + "..."
    return out


def _text_box(c: canvas.Canvas, x: float, y_top: float, width: float, height: float, text: Any, *, size: float = 7.0) -> float:
    por._rules(c)
    c.setFillColor(white); c.rect(x, y_top - height, width, height, stroke=1, fill=1)
    leading = size + 1.8
    max_lines = max(1, int((height - 6) // leading))
    y = y_top - size - 3
    for line in _lines(text, width - 8, size=size, max_lines=max_lines) or ["-"]:
        por._draw_text(c, x + 4, y, line, size=size)
        y -= leading
    return y_top - height


def _checkbox(c: canvas.Canvas, x: float, y: float, checked: bool, label: str, *, size: float = 6.8) -> float:
    por._rules(c)
    c.setFillColor(white); c.rect(x, y - 1.5, 7.5, 7.5, stroke=1, fill=1)
    if checked:
        c.setStrokeColor(por._accent()); c.setLineWidth(1.2)
        c.line(x + 1.4, y + 2.6, x + 3.2, y + 0.4); c.line(x + 3.2, y + 0.4, x + 6.4, y + 5.4)
        por._rules(c)
    por._draw_text(c, x + 11, y, label, size=size, bold=checked)
    return x + 11 + stringWidth(label, "Helvetica-Bold" if checked else "Helvetica", size) + 12


def _cell(c: canvas.Canvas, x: float, y_top: float, width: float, height: float, label: str, value: Any, *,
          label_width: float = 70.0, size: float = 6.9) -> None:
    por._rules(c)
    c.setFillColor(por._fill_or_white("label_fill")); c.rect(x, y_top - height, label_width, height, stroke=1, fill=1)
    c.setFillColor(white); c.rect(x + label_width, y_top - height, width - label_width, height, stroke=1, fill=1)
    por._draw_text(c, x + 3, y_top - height + 4.2, label, size=6.4, bold=True, max_width=label_width - 5)
    por._draw_text(c, x + label_width + 3, y_top - height + 4.2, value, size=size, max_width=width - label_width - 6)


def _logo(c: canvas.Canvas, h: float) -> None:
    logo_path = Path(__file__).resolve().parent.parent / "assets" / "fsi_logo.png"
    if logo_path.exists():
        try:
            c.drawImage(str(logo_path), LEFT, h - 58, width=105, height=39, preserveAspectRatio=True, mask="auto")
        except Exception:
            pass


def _header(c: canvas.Canvas, ecn: Mapping[str, Any], w: float, h: float, page_text: str = "01 TO 01") -> float:
    _logo(c, h)
    title = change_title(ecn.get("change_type"))
    c.setFillColor(por._accent()); c.setFont("Helvetica-Bold", 15.5)
    c.drawCentredString(w / 2 - 20, h - 40, title)
    # format box (top right)
    bx, bw, rh = w - RIGHT - 150, 150.0, 11.0
    rows = (("FMT NO.-", FORMAT_NO), ("REV.DATE", FORMAT_REV_DATE), ("REV. NO.-", FORMAT_REV_NO), ("PAGE NO", page_text))
    y = h - 18
    for label, value in rows:
        _cell(c, bx, y, bw, rh, label, value, label_width=58, size=6.6)
        y -= rh
    # identity strip
    y_strip = h - 74
    labels = ("ECN / PCN NO.", "DATE", "TYPE", "PLANT", "STATUS")
    values = (ecn.get("ecn_number"), _date(ecn.get("ecn_date")), clean(ecn.get("change_type")) or "ECN",
              ecn.get("plant_code"), status_label(ecn.get("status")).upper())
    widths = (150.0, 85.0, 70.0, 80.0, w - LEFT - RIGHT - 385.0)
    x = LEFT
    for sw, label, value in zip(widths, labels, values):
        por._bar(c, x, y_strip, sw, label, height=13)
        por._rules(c); c.setFillColor(white); c.rect(x, y_strip - 30, sw, 17, stroke=1, fill=1)
        por._draw_text(c, x + 4, y_strip - 25, value, size=7.4, bold=label == "ECN / PCN NO.", max_width=sw - 8)
        x += sw
    return y_strip - 36


def _representatives(ecn: Mapping[str, Any], employees: Mapping[str, Any]) -> str:
    names = []
    for emp_id in ecn.get("fsi_representative_ids") or []:
        emp = employees.get(clean(emp_id))
        name = employee_name(emp) if isinstance(emp, Mapping) else clean(emp)
        if name:
            names.append(name)
    extra = clean(ecn.get("fsi_representatives_text"))
    if extra:
        names.append(extra)
    return ", ".join(dict.fromkeys(names)) or "-"


def _employee(employees: Mapping[str, Any], emp_id: Any) -> str:
    emp = employees.get(clean(emp_id))
    return employee_name(emp) if isinstance(emp, Mapping) else clean(emp)


def _draw(c: canvas.Canvas, ecn: Mapping[str, Any], employees: Mapping[str, Any], part: Mapping[str, Any],
          page_text: str = "01 TO 01") -> None:
    w, h = A4
    full = w - LEFT - RIGHT
    y = _header(c, ecn, w, h, page_text)

    # FSI representatives
    por._bar(c, LEFT, y, full, "FSI REPRESENTATIVE(S)")
    y = _text_box(c, LEFT, y - 13, full, 22, _representatives(ecn, employees), size=7.0) - 6

    # Identification: To / CC / From
    label_w = 78.0
    col_w = (full - label_w) / 3
    por._bar(c, LEFT, y, label_w, "IDENTIFICATION")
    for idx, title in enumerate(("TO", "CC", "FROM")):
        por._bar(c, LEFT + label_w + idx * col_w, y, col_w, title)
    y -= 13
    contacts = [clean_contact(ecn.get(k)) for k in ("to_contact", "cc_contact", "from_contact")]
    for key, label in CONTACT_FIELDS:
        row_h = 22.0 if key == "address" else 13.0
        por._rules(c)
        c.setFillColor(por._fill_or_white("label_fill")); c.rect(LEFT, y - row_h, label_w, row_h, stroke=1, fill=1)
        por._draw_text(c, LEFT + 3, y - 9, label, size=6.6, bold=True)
        for idx, contact in enumerate(contacts):
            x = LEFT + label_w + idx * col_w
            c.setFillColor(white); c.rect(x, y - row_h, col_w, row_h, stroke=1, fill=1)
            ly = y - 9
            for line in _lines(contact[key], col_w - 6, size=6.6, max_lines=2 if key == "address" else 1):
                por._draw_text(c, x + 3, ly, line, size=6.6); ly -= 8.2
        y -= row_h
    y -= 6

    # Requested priority & response
    third = full / 3
    _cell(c, LEFT, y, third, 14, "Requested Priority", ecn.get("requested_priority"), label_width=80)
    _cell(c, LEFT + third, y, third, 14, "Response By", _date(ecn.get("response_by")), label_width=60)
    _cell(c, LEFT + 2 * third, y, third, 14, "Response From", ecn.get("response_by_party") or "FSI", label_width=66)
    y -= 20

    # Drawing criteria
    por._bar(c, LEFT, y, full, "DRAWING CRITERIA · Fill in the information completely to expedite the process.")
    y -= 13
    half = full / 2
    _cell(c, LEFT, y, half, 14, "Drawing No.", ecn.get("drawing_number"), label_width=64)
    _cell(c, LEFT + half, y, half, 14, "Revision", f"{clean(ecn.get('previous_revision')) or '-'}  ->  {clean(ecn.get('new_revision')) or '-'}", label_width=64)
    y -= 14
    _cell(c, LEFT, y, half, 14, "Part No.", ecn.get("part_number") or part.get("part_number"), label_width=64)
    _cell(c, LEFT + half, y, half, 14, "Part Name", ecn.get("part_name") or part.get("part_name"), label_width=64)
    y -= 20

    por._bar(c, LEFT, y, full, "PREVIOUS DESCRIPTION")
    y = _text_box(c, LEFT, y - 13, full, 96, ecn.get("previous_description"), size=6.6) - 5
    por._bar(c, LEFT, y, full, "REVISION REQUESTED")
    y = _text_box(c, LEFT, y - 13, full, 112, ecn.get("revision_requested"), size=6.6) - 5

    # Supporting documents
    is_pcn = clean(ecn.get("change_type")).upper() == "PCN"
    por._bar(c, LEFT, y, full, "SUPPORTING DOCUMENTS REQUIRED (Purchase Orders, PPAP, PFD, CP, PFMEA, inventory)" if is_pcn
             else "SUPPORTING DOCUMENTS ATTACHED (redline drawings, datasheets, etc.)")
    y -= 13
    por._rules(c); c.setFillColor(white); c.rect(LEFT, y - 30, full, 30, stroke=1, fill=1)
    docs = clean_supporting(ecn.get("supporting_documents"))
    step = full / 4
    for index, (key, label) in enumerate(SUPPORTING_DOCUMENTS):
        row, col = divmod(index, 4)
        _checkbox(c, LEFT + 6 + col * step, y - 11 - row * 13, docs[key], label)
    y -= 36

    # Impact
    por._rules(c); c.setFillColor(white); c.rect(LEFT, y - 42, full, 42, stroke=1, fill=1)
    impacts = (
        ("Purchase Orders Affected", bool(ecn.get("po_affected")), None),
        ("Current Inventory Affected", bool(ecn.get("inventory_affected")), ecn.get("inventory_qty")),
        ("Current Demands Affected", bool(ecn.get("demands_affected")), ecn.get("demands_qty")),
    )
    iy = y - 11
    for label, flag, qty in impacts:
        por._draw_text(c, LEFT + 6, iy, label, size=6.9, bold=True)
        _checkbox(c, LEFT + 150, iy, flag, "Yes")
        _checkbox(c, LEFT + 200, iy, not flag, "No")
        if label != "Purchase Orders Affected":
            por._draw_text(c, LEFT + 250, iy, f"Qty.: {_qty(qty) if flag else '-'}", size=6.9)
        iy -= 13
    y -= 48

    # FSI use only
    por._bar(c, LEFT, y, full, "FSI USE ONLY (Check the appropriate box(s))")
    y -= 13
    por._rules(c); c.setFillColor(white); c.rect(LEFT, y - 44, full, 44, stroke=1, fill=1)
    decision = clean(ecn.get("fsi_decision")).upper()
    por._draw_text(c, LEFT + 6, y - 11, "FSI", size=6.9, bold=True)
    _checkbox(c, LEFT + 110, y - 11, decision == "APPROVED", "Approved")
    _checkbox(c, LEFT + 190, y - 11, decision == "DISAPPROVED", "Disapproved")
    customer = clean(ecn.get("customer_approval")).upper()
    cust_label = "Customer" + (f" ({clean(ecn.get('customer_name_text'))})" if clean(ecn.get("customer_name_text")) else "")
    por._draw_text(c, LEFT + 290, y - 11, cust_label, size=6.9, bold=True, max_width=95)
    cx = _checkbox(c, LEFT + 390, y - 11, customer == "APPROVED", "Approved")
    _checkbox(c, cx - 4, y - 11, customer == "DISAPPROVED", "Disapproved")
    mode = clean(ecn.get("drawing_change_mode")).upper()
    por._draw_text(c, LEFT + 6, y - 25, "Drawing change to occur:", size=6.9, bold=True)
    x = LEFT + 110
    for key, label in DRAWING_CHANGE_MODES.items():
        x = _checkbox(c, x, y - 25, mode == key, label)
    if customer in {"PENDING", "NOT_REQUIRED"}:
        por._draw_text(c, LEFT + 6, y - 38, f"Customer approval: {customer.replace('_', ' ').title()}", size=6.6, color=MUTED)
    y -= 50

    # Signatures
    sig_rows = (
        ("Approval Signatures: FSI", _employee(employees, ecn.get("fsi_approver_employee_id")), _date(ecn.get("fsi_approval_date"))),
        ("Approval Signatures: Customer", ecn.get("customer_approver_name"), _date(ecn.get("customer_approval_date"))),
        ("Completed By", _employee(employees, ecn.get("completed_by_employee_id")), _date(ecn.get("completed_date"))),
    )
    for label, name, when in sig_rows:
        _cell(c, LEFT, y, full * 0.66, 15, label, name, label_width=128)
        _cell(c, LEFT + full * 0.66, y, full * 0.34, 15, "Date" if label != "Completed By" else "Completed Date", when, label_width=72)
        y -= 15
    y -= 5

    por._bar(c, LEFT, y, full, "COMMENTS / BENEFIT OF CHANGE")
    text = clean(ecn.get("comments"))
    benefit = clean(ecn.get("benefit_of_change"))
    if benefit:
        text = (text + "\n" if text else "") + f"Benefit to this Change: {benefit}"
    box_h = max(30.0, y - 13 - 62)
    _text_box(c, LEFT, y - 13, full, box_h, text or "-", size=6.9)

    c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica", 6.0)
    c.drawCentredString(w / 2, 47, f"QCMS controlled {change_title(ecn.get('change_type')).title()} · {FORMAT_NO} Rev {FORMAT_REV_NO}")
    c.drawCentredString(w / 2, 36, por._ACTIVE.get("footer_text") or "FSI · connect@fourstarindustries.com")


ACTION_COLUMNS: tuple[tuple[str, str, float], ...] = (
    ("No.", "sequence_no", 0.05), ("Affected Document", "affected_document", 0.14), ("Implementation Action", "action_text", 0.33),
    ("Responsible", "responsible", 0.15), ("Target Date", "target_date", 0.09), ("Status", "status", 0.10),
    ("Done On", "completed_date", 0.09), ("Remarks", "remarks", 0.05),
)


def _action_value(row: Mapping[str, Any], key: str, employees: Mapping[str, Any], index: int) -> str:
    if key == "sequence_no":
        return str(index)
    if key == "responsible":
        return _employee(employees, row.get("responsible_employee_id")) or clean(row.get("responsible_text")) or "-"
    if key in {"target_date", "completed_date"}:
        return _date(row.get(key))
    if key == "status":
        return clean(row.get("status")).replace("_", " ").title()
    return clean(row.get(key))


def _draw_actions(c: canvas.Canvas, ecn: Mapping[str, Any], employees: Mapping[str, Any], actions: list, page_text: str) -> None:
    """Continuation page: IMPLEMENTATION ACTIONS (part_ecn_actions) in the same PO look."""
    w, h = A4
    full = w - LEFT - RIGHT
    y = _header(c, ecn, w, h, page_text)
    por._bar(c, LEFT, y, full, "IMPLEMENTATION ACTIONS (AFFECTED DOCUMENTS · RESPONSIBLE · TARGET DATE)")
    y -= 13
    x = LEFT
    for title, _key, share in ACTION_COLUMNS:
        cw = full * share
        por._rules(c); c.setFillColor(por._fill_or_white("label_fill")); c.rect(x, y - 14, cw, 14, stroke=1, fill=1)
        por._draw_text(c, x + 2, y - 10, title, size=6.2, bold=True, max_width=cw - 4)
        x += cw
    y -= 14
    for index, row in enumerate(actions, start=1):
        cells = []
        height = 14.0
        for _title, key, share in ACTION_COLUMNS:
            lines = _lines(_action_value(row, key, employees, index), full * share - 4, size=6.4, max_lines=6)
            cells.append(lines)
            height = max(height, 5 + 8 * max(1, len(lines)))
        if y - height < 70:
            break
        x = LEFT
        for (_title, _key, share), lines in zip(ACTION_COLUMNS, cells):
            cw = full * share
            por._rules(c); c.setFillColor(white); c.rect(x, y - height, cw, height, stroke=1, fill=1)
            for li, line in enumerate(lines):
                por._draw_text(c, x + 2, y - 9 - li * 8, line, size=6.4, max_width=cw - 4)
            x += cw
        y -= height
    c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica", 6.0)
    c.drawCentredString(w / 2, 47, f"QCMS controlled {change_title(ecn.get('change_type')).title()} · {FORMAT_NO} Rev {FORMAT_REV_NO}")


def build_ecn_pdf(ecn: Mapping[str, Any], *, employees: Mapping[str, Any] | None = None, part: Mapping[str, Any] | None = None,
                  layout: Mapping[str, Any] | None = None, actions: list | None = None) -> bytes:
    """Controlled ECN / PCN print in the company PO look (saved Print Layout Designer theme)."""
    if layout is None:
        try:
            layout = por._saved_layout("PURCHASE_ORDER")
        except Exception:
            layout = None
    por.set_po_design(None, layout)
    try:
        out = BytesIO()
        c = canvas.Canvas(out, pagesize=A4)
        title = change_title(ecn.get("change_type")).title()
        c.setTitle(f"{title} {clean(ecn.get('ecn_number'))}".strip())
        pages = 2 if actions else 1
        _draw(c, dict(ecn or {}), dict(employees or {}), dict(part or {}), f"01 TO {pages:02d}")
        c.showPage()
        if actions:
            _draw_actions(c, dict(ecn or {}), dict(employees or {}), [dict(a) for a in actions], f"02 TO {pages:02d}")
            c.showPage()
        c.save()
        data = out.getvalue()
    finally:
        por.set_po_design(None)
    if por.PdfReader is not None and por.PdfWriter is not None:
        try:
            return por._watermark_pdf_bytes(data, watermark_label(ecn))
        except Exception:
            return data
    return data


def ecn_pdf_bytes(ecn_id: Any, repo: Any | None = None) -> bytes:
    from core.ecn_service import ECNService
    service = ECNService(repo)
    ecn = service.get(ecn_id)
    if not ecn:
        raise ValueError("ECN was not found.")
    ids = {clean(v) for v in (ecn.get("fsi_representative_ids") or [])}
    ids |= {clean(ecn.get("fsi_approver_employee_id")), clean(ecn.get("completed_by_employee_id"))}
    actions = service.actions(ecn_id)
    ids |= {clean(a.get("responsible_employee_id")) for a in actions}
    employees = {clean(r.get("id")): r for r in service.employees() if clean(r.get("id")) in ids}
    part = service.repo.get("parts", clean(ecn.get("part_id"))) or {} if clean(ecn.get("part_id")) else {}
    return build_ecn_pdf(ecn, employees=employees, part=part, actions=actions)




# ---------------------------------------------------------------------------------------------------------------- Excel
def _box(flag: bool) -> str:
    return "\u2611" if flag else "\u2610"


def build_ecn_xlsx(ecn: Mapping[str, Any], *, employees: Mapping[str, Any] | None = None, part: Mapping[str, Any] | None = None,
                   actions: list | None = None) -> bytes:
    """Excel copy laid out like the customer workbook (sheet "ECN Form" / "PCN", columns A:AD, same anchors / merges)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ecn = dict(ecn or {})
    employees = dict(employees or {})
    part = dict(part or {})
    is_pcn = clean(ecn.get("change_type")).upper() == "PCN"
    wb = Workbook()
    ws = wb.active
    ws.title = "PCN" if is_pcn else "ECN Form"
    for col in range(1, 31):
        ws.column_dimensions[get_column_letter(col)].width = 4.4
    thin = Side(style="thin", color="1F2937")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bar = PatternFill("solid", fgColor="1F3A5F")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def put(ref: str, value: Any, *, merge: str | None = None, bold: bool = False, size: float = 9, fill=None, align=None, white_text=False):
        cell = ws[ref]
        cell.value = value if value not in (None, "") else None
        cell.font = Font(bold=bold, size=size, color="FFFFFF" if white_text else "111827")
        cell.alignment = align or Alignment(vertical="center", wrap_text=True)
        if fill is not None:
            cell.fill = fill
        if merge:
            ws.merge_cells(f"{ref}:{merge}")
            for row in ws[f"{ref}:{merge}"]:
                for c in row:
                    c.border = border
        else:
            cell.border = border

    def section(row: int, title: str) -> None:
        put(f"A{row}", title, merge=f"AD{row}", bold=True, fill=bar, white_text=True)

    def emp(emp_id: Any) -> str:
        row = employees.get(clean(emp_id))
        return employee_name(row) if isinstance(row, Mapping) else clean(row)

    reps = [emp(v) for v in (ecn.get("fsi_representative_ids") or [])]
    if clean(ecn.get("fsi_representatives_text")):
        reps.append(clean(ecn.get("fsi_representatives_text")))
    docs = clean_supporting(ecn.get("supporting_documents"))
    decision = clean(ecn.get("fsi_decision")).upper()
    customer = clean(ecn.get("customer_approval")).upper()
    mode = clean(ecn.get("drawing_change_mode")).upper()

    # format box + title (rows 1-4)
    for row, (label, value) in enumerate((("FMT NO.-", FORMAT_NO), ("REV.DATE", FORMAT_REV_DATE), ("REV. NO.-", FORMAT_REV_NO),
                                          ("PAGE NO", "01 TO 01")), start=1):
        put(f"Y{row}", label, merge=f"Z{row}", bold=True, size=8)
        put(f"AA{row}", value, merge=f"AD{row}", size=8)
    put("A1", "STAWN QCMS", merge="X2", bold=True, size=9)
    put("A3", change_title(ecn.get("change_type")), merge="X4", bold=True, size=16, align=center)
    put("A5", "Date", merge="B5", bold=True)
    put("C5", _date(ecn.get("ecn_date")), merge="H5")
    put("I5", "No.", merge="K5", bold=True)
    put("L5", ecn.get("ecn_number"), merge="S5", bold=True)
    put("T5", "Status", merge="V5", bold=True)
    put("W5", status_label(ecn.get("status")), merge="AD5")
    put("A7", "FSI Representative(s)", merge="H7", bold=True)
    put("I7", ", ".join(dict.fromkeys(r for r in reps if r)) or None, merge="AD7")
    # identification (rows 9-14)
    put("A9", " Identification", merge="H9", bold=True)
    for offset, (key, title) in enumerate((("to_contact", "To:"), ("cc_contact", "CC:"), ("from_contact", "From:"))):
        col = ("I", "Q", "Y")[offset]
        end = ("P", "X", "AD")[offset]
        put(f"{col}9", title, merge=f"{end}9", bold=True)
    contacts = [clean_contact(ecn.get(k)) for k in ("to_contact", "cc_contact", "from_contact")]
    for r, (key, label) in enumerate(CONTACT_FIELDS, start=10):
        put(f"A{r}", label, merge=f"H{r}", bold=True)
        for offset, contact in enumerate(contacts):
            put(f"{('I', 'Q', 'Y')[offset]}{r}", contact[key], merge=f"{('P', 'X', 'AD')[offset]}{r}")
    put("A15", "Requested Priority & Response:", merge="H15", bold=True)
    put("I15", ecn.get("requested_priority"), merge="P15")
    put("Q15", "Request Response By:", merge="W15", bold=True)
    put("X15", " ".join(v for v in (_date(ecn.get("response_by")), clean(ecn.get("response_by_party")) or "FSI") if v), merge="AD15")
    section(17, "Drawing Criteria:          Fill in the information completely to expedite the process.")
    put("A18", "Drawing No.", merge="F18", bold=True)
    put("G18", ecn.get("drawing_number"), merge="O18")
    put("P18", "Revision", merge="T18", bold=True)
    put("U18", f"{clean(ecn.get('previous_revision')) or '-'} -> {clean(ecn.get('new_revision')) or '-'}", merge="AD18")
    put("A19", "Part No.", merge="F19", bold=True)
    put("G19", ecn.get("part_number") or part.get("part_number"), merge="AD19")
    put("A20", "Part Name", merge="F20", bold=True)
    put("G20", ecn.get("part_name") or part.get("part_name"), merge="AD20")
    section(22, "Previous Description")
    put("A23", ecn.get("previous_description"), merge="AD28", align=wrap)
    section(30, "Revision requested")
    put("A31", ecn.get("revision_requested"), merge="AD38", align=wrap)
    if is_pcn:
        section(40, "  Supporting Documents Required")
        pcn_docs = (("Purchase Orders", docs["purchase_order"]), ("PPAP", docs["ppap"]), ("PFD", docs["pfd"]), ("CP", docs["control_plan"]),
                    ("PFMEA", docs["pfmea"]), ("MAINTAIN INVENTORY", docs["maintain_inventory"]))
        for index, (label, flag) in enumerate(pcn_docs):
            row = 41 + (index // 2) * 2
            col, end, ans, ans_end = (("C", "J", "K", "N") if index % 2 == 0 else ("P", "W", "X", "AA"))
            put(f"{col}{row}", label, merge=f"{end}{row}", bold=True)
            put(f"{ans}{row}", f"{_box(flag)} Yes   {_box(not flag)} No", merge=f"{ans_end}{row}")
    else:
        section(40, "  Supporting Documents Attached (redline drawings, datasheets, etc.)")
        attached = ", ".join(label for key, label in SUPPORTING_DOCUMENTS if docs[key]) or "-"
        put("A41", attached, merge="AD41")
        impacts = (("Purchase Orders Affected", bool(ecn.get("po_affected")), None),
                   ("Current Inventory Affected", bool(ecn.get("inventory_affected")), ecn.get("inventory_qty")),
                   ("Current Demands Affected", bool(ecn.get("demands_affected")), ecn.get("demands_qty")))
        for index, (label, flag, qty) in enumerate(impacts):
            row = 43 + index * 2
            put(f"C{row}", label, merge=f"K{row}", bold=True)
            put(f"L{row}", f"{_box(flag)} Yes", merge=f"M{row}")
            if label != "Purchase Orders Affected":
                put(f"N{row}", "Qty.", merge=f"O{row}", bold=True)
                put(f"P{row}", _qty(qty) if flag else None, merge=f"T{row}")
            put(f"U{row}", f"{_box(not flag)} No", merge=f"W{row}")
    section(50, "FSI Use ONLY (Check the appropriate box(s))")
    put("C51", f"{_box(decision == 'APPROVED')} Approved", merge="P51")
    put("R51", f"{_box(decision == 'DISAPPROVED')} Disapproved", merge="AD51")
    customer_label = "Customer" + (f" ({clean(ecn.get('customer_name_text'))})" if clean(ecn.get("customer_name_text")) else "")
    put("A52", customer_label, merge="B52", bold=True)
    put("C52", f"{_box(customer == 'APPROVED')} Approved", merge="P52")
    put("R52", f"{_box(customer == 'DISAPPROVED')} Disapproved", merge="AD52")
    put("C53", "Drawing change to occur:", merge="AD53", bold=True)
    put("D55", f"{_box(mode == 'AS_SUGGESTED')} as suggested", merge="H55")
    put("I55", f"{_box(mode == 'NOT_AS_SUGGESTED')} not as suggested", merge="Q55")
    put("R55", f"{_box(mode == 'OTHER')} Other (See Comments)", merge="AD55")
    put("B57", "Approval Signatures: FSI", merge="H57", bold=True)
    put("I57", emp(ecn.get("fsi_approver_employee_id")), merge="R57")
    put("T57", "Approval Date:", merge="X57", bold=True)
    put("Y57", _date(ecn.get("fsi_approval_date")), merge="AB57")
    put("B58", "Approval Signatures: Customer", merge="H58", bold=True)
    put("I58", ecn.get("customer_approver_name"), merge="R58")
    put("T58", "Approval Date:", merge="X58", bold=True)
    put("Y58", _date(ecn.get("customer_approval_date")), merge="AB58")
    put("B59", "Completed By:", merge="G59", bold=True)
    put("H59", emp(ecn.get("completed_by_employee_id")), merge="R59")
    put("T59", "Completed Date:", merge="X59", bold=True)
    put("Y59", _date(ecn.get("completed_date")), merge="AB59")
    comments = clean(ecn.get("comments"))
    if clean(ecn.get("benefit_of_change")):
        comments = (comments + "\n" if comments else "") + f"Benefit to this Change : {clean(ecn.get('benefit_of_change'))}"
    put("B62", "Comments:", merge="E62", bold=True)
    put("F62", comments, merge="AD64", align=wrap)
    for row in (23, 31):
        ws.row_dimensions[row].height = 16
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    last_row = 64
    if actions:
        section(66, "IMPLEMENTATION ACTIONS")
        heads = (("A", "B", "No."), ("C", "G", "Affected Document"), ("H", "Q", "Implementation Action"), ("R", "V", "Responsible"),
                 ("W", "Y", "Target Date"), ("Z", "AB", "Status"), ("AC", "AD", "Done On"))
        for first, last, title in heads:
            put(f"{first}67", title, merge=f"{last}67" if first != last else None, bold=True, align=center)
        for index, action in enumerate(actions, start=1):
            r = 67 + index
            values = (str(index), clean(action.get("affected_document")), clean(action.get("action_text")),
                      emp(action.get("responsible_employee_id")) or clean(action.get("responsible_text")),
                      _date(action.get("target_date")), clean(action.get("status")).replace("_", " ").title(), _date(action.get("completed_date")))
            for (first, last, _title), value in zip(heads, values):
                put(f"{first}{r}", value, merge=f"{last}{r}" if first != last else None, align=wrap)
            ws.row_dimensions[r].height = 26
            last_row = r
    ws.print_area = f"A1:AD{last_row}"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


__all__ = ["build_ecn_pdf", "ecn_pdf_bytes", "watermark_label", "build_ecn_xlsx"]
