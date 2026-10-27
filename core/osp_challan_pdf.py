"""R37 · OSP Material Out "RETURNABLE DELIVERY CHALLAN" print (PDF + Excel).

Same Design-B family as the Purchase Order print (core/purchase_order_reporting.py drawing helpers, the saved
PURCHASE_ORDER layout from core/print_designer.py, company logo, low-visibility status watermark).

Public contract (used by the OSP pages and by the universal PDF service for table ``osp_jobs``):

    osp_challan_pdf_bytes(job_id)      -> bytes   # job id, Material Out id or osp_challans id; or a print payload
    osp_challan_excel_bytes(job_id)    -> bytes
    osp_challan_payload(job_id)        -> dict    # vendor-safe print payload

Customer confidentiality (R36b): vendor-facing - FSI part number / description only, never customer name, code,
PO, order number, customer part number or customer delivery date (core.osp_challan.vendor_safe).
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Mapping, Sequence

from core.osp_challan import (CHALLAN_TITLE, DEFAULT_UOM, OSPChallanService, _date_text, _f, _qty, _s, build_print_payload,
                              line_amounts, vendor_safe)

# Item rows per page: (rows when this is the last page - room for process/totals/return box, rows otherwise).
FIRST_PAGE_ROWS = (6, 13)
NEXT_PAGE_ROWS = (12, 19)


def challan_pages(lines: Sequence[Any]) -> list[list[Any]]:
    """Split item lines so the last page always has room for the process bar, totals and supplier box."""
    remaining = list(lines); pages: list[list[Any]] = []; caps = FIRST_PAGE_ROWS
    while True:
        last_cap, full_cap = caps
        if len(remaining) <= last_cap:
            pages.append(remaining)
            return pages
        take = min(full_cap, len(remaining) - 1)
        pages.append(remaining[:take]); remaining = remaining[take:]; caps = NEXT_PAGE_ROWS


def _render_plain(payload: Mapping[str, Any]) -> bytes:
    """RETURNABLE DELIVERY CHALLAN (A4 portrait, same Design-B family as the Purchase Order print), no watermark."""
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from core import purchase_order_reporting as po

    data = vendor_safe(payload)
    try:
        from core.print_designer import saved_layout
        po.set_po_design(None, saved_layout("PURCHASE_ORDER"))
    except Exception:
        po.set_po_design(None)
    try:
        out = BytesIO(); c = canvas.Canvas(out, pagesize=A4); w, h = A4; left, right = 28, 28
        c.setTitle(f"{CHALLAN_TITLE} {data.get('challan_number') or ''}".strip())
        lines = list(data.get("lines") or [])
        row_h = 34
        pages = challan_pages(lines)
        for page_index, page_lines in enumerate(pages):
            last = page_index == len(pages) - 1
            po._rules(c)
            logo = Path(__file__).resolve().parent.parent / "assets" / "fsi_logo.png"
            if logo.exists():
                try:
                    c.drawImage(str(logo), left, h - 58, width=105, height=39, preserveAspectRatio=True, mask="auto")
                except Exception:
                    pass
            c.setFillColor(HexColor("#111827")); c.setFont("Helvetica-Bold", 19); c.drawRightString(w - right, h - 36, CHALLAN_TITLE); po._rules(c)
            po._draw_text(c, w - 160, h - 53, "DATE", size=7.4); c.rect(w - 112, h - 58, 84, 12, stroke=1, fill=0)
            po._draw_text(c, w - 107, h - 55, _date_text(data.get("challan_date")), size=7.2)
            po._draw_text(c, w - 160, h - 67, "CHALLAN #", size=7.4); c.rect(w - 112, h - 72, 84, 12, stroke=1, fill=0)
            po._draw_text(c, w - 107, h - 69, data.get("challan_number"), size=7.2, bold=True, max_width=78)
            if page_index == 0:
                col_w = 260; rx = w - right - 220
                plant = dict(data.get("plant") or {})
                if not plant:
                    plant = {**{k: po.PLANT.get(k) for k in ("name", "address1", "address2", "address3", "phone", "email")}, "gstin": po.PLANT.get("tax_identifier")}
                po._block(c, left, h - 64, col_w, "PLANT ADDRESS", [plant.get("name"), plant.get("address1"), plant.get("address2"), plant.get("address3"), plant.get("gstin"), plant.get("phone"), plant.get("email")])
                po._block(c, left, h - 160, col_w, "VENDOR", po._party_lines(dict(data.get("vendor") or {})), max_lines=7)
                po._block(c, rx, h - 84, 220, "REFERENCE DETAILS", [_s(data.get("reference_details"))], max_lines=2)
                po._block(c, rx, h - 120, 220, "OLD PO DETAILS", [_s(data.get("old_po_details"))], max_lines=2)
                po._block(c, rx, h - 160, 220, "SHIP TO", po._party_lines(dict(data.get("ship_to") or {})), max_lines=7)
                y_strip = h - 262
                widths = [92, 104, 84, 90, w - left - right - 370]
                titles = ["USER", "SHIP VIA", "INCOTERM", "DELIVERY DATE", "PAYMENT TERM"]
                values = [data.get("user_name"), data.get("ship_via"), data.get("incoterm"), _date_text(data.get("delivery_date")), data.get("payment_term")]
                x = left
                for sw, title, value in zip(widths, titles, values):
                    po._bar(c, x, y_strip, sw, title, height=14); c.rect(x, y_strip - 34, sw, 20, stroke=1, fill=0)
                    po._draw_text(c, x + 4, y_strip - 27, value, size=6.8, max_width=sw - 8); x += sw
                y = y_strip - 44
            else:
                po._draw_text(c, left, h - 80, f"Continued · Challan {data.get('challan_number') or ''}", size=8, bold=True, color=po._accent())
                y = h - 92
            # item grid
            cols = [("ITEM #", 165), ("QTY", 52), ("UNIT PRICE", 62), ("UNIT", 42), ("GST%", 42), ("GST AMOUNT", 88), ("TAXABLE AMOUNT", w - left - right - 451)]
            x = left
            for title, cw in cols:
                po._bar(c, x, y, cw, title, height=14); x += cw
            y -= 14
            grid_top = y
            for item in page_lines:
                amounts = line_amounts(item)
                values = [None, _qty(item.get("quantity")), f"{_f(item.get('unit_price')):,.2f}".rstrip("0").rstrip("."), _s(item.get("uom")) or DEFAULT_UOM,
                          f"{_f(item.get('gst_percent')):g}%", f"{amounts['gst']:,.2f}", f"{amounts['taxable']:,.2f}"]
                po._draw_text(c, left + 4, y - 10, _s(item.get("item_code")), size=7.6, bold=True, max_width=160)
                po._draw_text(c, left + 4, y - 19, _s(item.get("item_description")) + (f" · HSN {item.get('hsn_code')}" if _s(item.get("hsn_code")) else ""), size=6.2, max_width=160)
                x = left + cols[0][1]
                for (title, cw), value in zip(cols[1:], values[1:]):
                    if title in {"GST AMOUNT", "TAXABLE AMOUNT"}:
                        c.setFillColor(po.TEXT); c.setFont("Helvetica", 7.2); c.drawRightString(x + cw - 4, y - 12, str(value))
                    else:
                        po._draw_text(c, x + 4, y - 12, value, size=7.2, max_width=cw - 8)
                    x += cw
                label = f"Heat No.{_s(item.get('heat_number')) or '-'}   Batch Code :- {_s(item.get('batch_code')) or '-'}"
                c.setFillColor(HexColor("#FFF700")); c.rect(left + 1, y - row_h + 2, min(250, 12 + len(label) * 4.1), 9.5, stroke=0, fill=1)
                po._draw_text(c, left + 4, y - row_h + 4.5, label, size=6.8, bold=True, max_width=244)
                y -= row_h
            if last:
                po._bar(c, left, y - 2, 365, "PROCESS TO BE DONE", height=14)
                y = po._wrap(c, left + 4, y - 26, data.get("process_text") or "-", 357, size=8.0, leading=10, max_lines=3)
                if _s(data.get("remarks")):
                    y = po._wrap(c, left + 4, y - 2, "Remarks: " + _s(data.get("remarks")), 357, size=6.8, leading=8.4, max_lines=2)
            grid_bottom = 250 if last else 60
            po._rules(c); c.rect(left, grid_bottom, w - left - right, grid_top - grid_bottom, stroke=1, fill=0)
            if not last:
                c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica", 6); c.drawCentredString(w / 2, 40, "Continued on next page")
                c.showPage(); continue
            # totals
            totals = dict(data.get("totals") or {})
            tx = w - right - 170; ty = 238
            for label, value in (("SUBTOTAL", totals.get("subtotal")), ("TAX", totals.get("tax")), ("SHIPPING", totals.get("shipping")), ("OTHER", totals.get("other"))):
                po._draw_text(c, tx, ty, label, size=7.0); c.rect(tx + 70, ty - 4, 100, 12, stroke=1, fill=0)
                amount = _f(value); po._draw_text(c, tx + 74, ty - 1, f"{amount:,.2f}" if amount else "-", size=7.0); ty -= 15
            c.setStrokeColor(HexColor("#111827")); c.setLineWidth(1.4); c.line(tx, ty + 9, tx + 170, ty + 9); po._rules(c)
            c.setFillColor(HexColor("#111827")); c.setFont("Helvetica-Bold", 8.2); c.drawString(tx, ty - 2, "TOTAL"); c.drawRightString(tx + 168, ty - 2, f"INR {_f(totals.get('total')):,.2f}")
            _draw_supplier_box(c, po, data, x=left, top=244, width=300)
            company = (data.get("plant") or {}).get("name") or po.PLANT["name"]
            po._draw_text(c, w - right - 105, 92, "Authorised Signatory", size=7.4, bold=True)
            po._draw_text(c, w - right - 150, 80, f"For {company}", size=6.2, max_width=150)
            c.setFillColor(HexColor("#6B7280")); c.setFont("Helvetica", 6.0)
            c.drawCentredString(w / 2, 47, "If you have any questions about this document, please contact")
            contact = (data.get("plant") or {}).get("email") or po.PLANT["email"]
            c.drawCentredString(w / 2, 36, f"FSI, {contact}")
            c.showPage()
        c.save()
        return out.getvalue()
    finally:
        po.set_po_design(None)


def _draw_supplier_box(c: Any, po: Any, data: Mapping[str, Any], *, x: float, top: float, width: float) -> None:
    """Bottom-left 'Supplier Details for Returnable Challan' box: key/value for 0-1 return, small table for several."""
    from reportlab.lib.colors import HexColor
    returns = list(data.get("returns") or [])
    summary = dict(data.get("summary") or {})
    c.setFillColor(HexColor("#BFBFBF")); c.rect(x, top - 13, width, 13, stroke=1, fill=1)
    po._draw_text(c, x + 3, top - 9.5, "Supplier Details for Returnable Challan / Return details", size=7.2, bold=True)
    y = top - 13
    if len(returns) <= 1:
        r = returns[0] if returns else {}
        pairs = [
            ("Challan No", r.get("challan_no")), ("Challan Date", _date_text(r.get("challan_date"))),
            ("Invoice No / Date", " / ".join(v for v in (_s(r.get("invoice_no")), _date_text(r.get("invoice_date"))) if v)),
            ("Quantity dispatch", _qty(summary.get("dispatched"))),
            ("Quantity returned", _qty(summary.get("returned")) if returns else ""),
            ("Quantity rejected", _qty(summary.get("rejected")) if summary.get("rejected") else ""),
            ("Balance Quantity", _qty(summary.get("balance")) if returns else ""),
            ("Vehicle No", r.get("vehicle_no") or data.get("vehicle_number")),
            ("Remarks", r.get("remarks")),
        ]
        po._rules(c)
        for label, value in pairs:
            c.rect(x, y - 13, width, 13, stroke=1, fill=0)
            po._draw_text(c, x + 3, y - 9.5, label, size=6.8); po._draw_text(c, x + 98, y - 9.5, f":  {_s(value)}", size=6.8, max_width=width - 102)
            y -= 13
        return
    cols = [("Vendor Challan", 70), ("Date", 46), ("Invoice", 62), ("Qty Ret.", 36), ("Rej.", 28), ("Vehicle", width - 242)]
    shown = returns[-7:]
    po._rules(c)
    cx = x
    for title, cw in cols:
        c.setFillColor(HexColor("#E4E5E7")); c.rect(cx, y - 11, cw, 11, stroke=1, fill=1); po._draw_text(c, cx + 2, y - 8, title, size=6.0, bold=True); cx += cw
    y -= 11
    for r in shown:
        values = [r.get("challan_no"), _date_text(r.get("challan_date")), r.get("invoice_no"), _qty(r.get("qty_returned")), _qty(r.get("qty_rejected")) if _f(r.get("qty_rejected")) else "-", r.get("vehicle_no")]
        cx = x
        for (title, cw), value in zip(cols, values):
            c.rect(cx, y - 11, cw, 11, stroke=1, fill=0); po._draw_text(c, cx + 2, y - 8, value, size=6.0, max_width=cw - 4); cx += cw
        y -= 11
    note = f"  ({len(returns) - len(shown)} earlier returns not shown)" if len(returns) > len(shown) else ""
    c.rect(x, y - 12, width, 12, stroke=1, fill=0)
    po._draw_text(c, x + 3, y - 8.5, f"Dispatched {_qty(summary.get('dispatched'))} · Returned {_qty(summary.get('returned'))} · Rejected {_qty(summary.get('rejected'))} · Balance {_qty(summary.get('balance'))}{note}", size=6.4, bold=True, max_width=width - 6)
    y -= 12
    c.rect(x, y - 12, width, 12, stroke=1, fill=0)
    po._draw_text(c, x + 3, y - 8.5, f"Dispatch Vehicle No : {_s(data.get('vehicle_number')) or '-'}", size=6.4, max_width=width - 6)


def challan_status_label(payload: Mapping[str, Any]) -> str:
    """Watermark text: DRAFT (not saved yet), CANCELLED, CLOSED (everything back from the vendor) or ISSUED."""
    status = _s(payload.get("status")).upper()
    if not payload.get("saved", True) or status == "DRAFT":
        return "DRAFT"
    if status in {"CANCELLED", "CLOSED"}:
        return status
    summary = dict(payload.get("summary") or {})
    if payload.get("returns") and _f(summary.get("dispatched")) > 0 and _f(summary.get("balance")) <= 0:
        return "CLOSED"
    return "ISSUED"


def render_challan_pdf(payload: Mapping[str, Any]) -> bytes:
    """Challan PDF from a print payload, with the PO-style status watermark merged behind the content."""
    pdf = _render_plain(payload)
    try:
        from core.purchase_order_reporting import _watermark_pdf_bytes
        return _watermark_pdf_bytes(pdf, challan_status_label(payload))
    except Exception:
        return pdf


def osp_challan_payload(job_id: Any, repo: Any = None) -> dict:
    """Vendor-safe print payload for an OSP job (or its Material Out / challan id).

    A saved challan (osp_challans) prints as saved. A Material Out without saved challan details prints a
    DRAFT challan from the Material Out itself (challan no = dispatch challan, delivery = expected return)."""
    service = OSPChallanService(repo)
    key = _s(job_id)
    if not key:
        raise ValueError("Select an OSP job to print its Returnable Delivery Challan.")
    job = service.repo.get("osp_jobs", key)
    group = None
    if not job:
        challan = service.repo.get("osp_challans", key)
        if challan:
            return service.print_payload_for(str(challan["id"]))
        group = service.repo.get("osp_material_outs", key)
        if not group:
            raise ValueError("OSP job / Material Out not found.")
    material_out_id = (job or {}).get("material_out_id") or (group or {}).get("id")
    document = {"is_group": bool(material_out_id), "material_out_id": material_out_id, "id": material_out_id or (job or {}).get("id")}
    challan = service.challan_for(document)
    if challan:
        return service.print_payload_for(str(challan["id"]))
    source = job or {}
    if not source and material_out_id:
        first = service._select("osp_jobs", eq={"material_out_id": str(material_out_id)}, order_by="material_out_line_no", limit=1)
        source = first[0] if first else {}
    draft = {"challan_number": source.get("dispatch_challan") or (group or {}).get("material_out_number"), "challan_date": source.get("dispatch_date"),
             "delivery_date": source.get("expected_return_date"), "vendor_id": source.get("vendor_id"), "status": "DRAFT",
             **({"material_out_id": material_out_id} if material_out_id else {"osp_job_id": source.get("id")})}
    rows = service.document_rows(draft)
    branches = service.branches()
    branch = next((b for b in branches if b.get("is_default")), branches[0] if branches else None)
    vendor = service.repo.get("parties", str(source.get("vendor_id") or "")) if source.get("vendor_id") else None
    returns = service.returns_for(None, [str(r.get("id")) for r in rows])
    return build_print_payload(draft, rows, [], returns, branch=branch, vendor=vendor)


def _payload(source: Any, repo: Any = None) -> dict:
    return dict(source) if isinstance(source, Mapping) else osp_challan_payload(source, repo)


def osp_challan_pdf_bytes(job_id: Any, repo: Any = None) -> bytes:
    """RETURNABLE DELIVERY CHALLAN PDF for an OSP job id (also accepts a Material Out / challan id or a payload)."""
    return render_challan_pdf(_payload(job_id, repo))


def osp_challan_excel_bytes(job_id: Any, repo: Any = None) -> bytes:
    """Excel copy of the challan (same sections as the PDF, watermark in the print header)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    data = vendor_safe(_payload(job_id, repo))
    label = challan_status_label(data)
    wb = Workbook(); ws = wb.active; ws.title = "Returnable Challan"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    for header_obj in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
        header_obj.center.text = label; header_obj.center.font = "Arial,Bold"; header_obj.center.size = 28; header_obj.center.color = "D9DDE3"
    navy = "0B2E63"; light = "D9E1F2"; thin = Side(style="thin", color="B8C0CC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bar_font = Font(bold=True, color="FFFFFF"); bar_fill = PatternFill("solid", fgColor=navy)

    def bar(row: int, col: int, text: str, span: int = 1) -> None:
        if span > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
        cell = ws.cell(row, col, text); cell.font = bar_font; cell.fill = bar_fill; cell.alignment = Alignment(horizontal="center", vertical="center")

    def put(row: int, col: int, value: Any, *, bold: bool = False, span: int = 1, border: bool = False, number: bool = False) -> None:
        if span > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
        cell = ws.cell(row, col, value if value not in (None, "") else "")
        cell.font = Font(bold=bold); cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right" if number else "left")
        if number:
            cell.number_format = "#,##0.00"
        if border:
            cell.border = box

    ws.merge_cells("A1:G1"); ws["A1"] = CHALLAN_TITLE; ws["A1"].font = Font(size=16, bold=True, color="FFFFFF"); ws["A1"].fill = bar_fill; ws["A1"].alignment = Alignment(horizontal="center")
    put(2, 5, "DATE", bold=True); put(2, 6, _date_text(data.get("challan_date")), span=2, border=True)
    put(3, 5, "CHALLAN #", bold=True); put(3, 6, data.get("challan_number"), span=2, border=True, bold=True)
    put(4, 5, "STATUS", bold=True); put(4, 6, label, span=2, border=True)

    def party(p: Mapping[str, Any]) -> list[str]:
        p = dict(p or {})
        return [v for v in (_s(p.get("party_name") or p.get("name")), _s(p.get("address") or p.get("address1")), _s(p.get("address2")),
                            " ".join(x for x in (_s(p.get("city")), _s(p.get("state")), _s(p.get("address3"))) if x),
                            _s(p.get("tax_identifier") or p.get("gstin")), _s(p.get("phone")), _s(p.get("email"))) if v]

    r = 6
    bar(r, 1, "PLANT ADDRESS", 3); bar(r, 5, "REFERENCE DETAILS", 3)
    plant = party(data.get("plant") or {}); put(r + 1, 1, "\n".join(plant), span=3); put(r + 1, 5, data.get("reference_details"), span=3)
    bar(r + 2, 5, "OLD PO DETAILS", 3); put(r + 3, 5, data.get("old_po_details"), span=3)
    ws.row_dimensions[r + 1].height = 15 * max(len(plant), 1)
    r += 5
    bar(r, 1, "VENDOR", 3); bar(r, 5, "SHIP TO", 3)
    vendor, ship = party(data.get("vendor") or {}), party(data.get("ship_to") or {})
    put(r + 1, 1, "\n".join(vendor), span=3); put(r + 1, 5, "\n".join(ship), span=3)
    ws.row_dimensions[r + 1].height = 15 * max(len(vendor), len(ship), 1)
    r += 3
    strip = [("USER", data.get("user_name")), ("SHIP VIA", data.get("ship_via")), ("INCOTERM", data.get("incoterm")),
             ("DELIVERY DATE", _date_text(data.get("delivery_date"))), ("PAYMENT TERM", data.get("payment_term")),
             ("VEHICLE NO", data.get("vehicle_number")), ("E-WAY BILL", data.get("eway_bill_number"))]
    for col, (title, value) in enumerate(strip, start=1):
        bar(r, col, title); put(r + 1, col, value, border=True)
    r += 3
    for col, title in enumerate(("ITEM #", "QTY", "UNIT PRICE", "UNIT", "GST%", "GST AMOUNT", "TAXABLE AMOUNT"), start=1):
        bar(r, col, title)
    r += 1
    for item in data.get("lines") or []:
        amounts = line_amounts(item)
        put(r, 1, f"{_s(item.get('item_code'))}\n{_s(item.get('item_description'))}".strip(), bold=True, border=True)
        put(r, 2, _f(item.get("quantity")), border=True, number=True); put(r, 3, _f(item.get("unit_price")), border=True, number=True)
        put(r, 4, _s(item.get("uom")) or DEFAULT_UOM, border=True); put(r, 5, f"{_f(item.get('gst_percent')):g}%", border=True)
        put(r, 6, amounts["gst"], border=True, number=True); put(r, 7, amounts["taxable"], border=True, number=True)
        r += 1
        put(r, 1, f"Heat No.{_s(item.get('heat_number')) or '-'}   Batch Code :- {_s(item.get('batch_code')) or '-'}", bold=True, span=3)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFF700")
        r += 1
    r += 1
    bar(r, 1, "PROCESS TO BE DONE", 4); r += 1
    put(r, 1, data.get("process_text") or "-", span=4)
    totals = dict(data.get("totals") or {})
    tr = r - 1
    for title, key in (("SUBTOTAL", "subtotal"), ("TAX", "tax"), ("SHIPPING", "shipping"), ("OTHER", "other"), ("TOTAL", "total")):
        put(tr, 6, title, bold=title == "TOTAL"); put(tr, 7, _f(totals.get(key)), border=True, number=True, bold=title == "TOTAL"); tr += 1
    if _s(data.get("remarks")):
        r += 1; put(r, 1, "Remarks: " + _s(data.get("remarks")), span=4)
    r = max(r, tr) + 2
    bar(r, 1, "SUPPLIER DETAILS FOR RETURNABLE CHALLAN / RETURN DETAILS", 7); r += 1
    heads = ("Return Challan No", "Return Challan Date", "Vendor Invoice No / Date", "Qty Returned", "Qty Rejected", "Vehicle No", "Remarks")
    for col, title in enumerate(heads, start=1):
        put(r, col, title, bold=True, border=True); ws.cell(r, col).fill = PatternFill("solid", fgColor=light)
    r += 1
    returns = list(data.get("returns") or [])
    for ret in returns or [{}] * 3:
        values = [ret.get("challan_no"), _date_text(ret.get("challan_date")),
                  " / ".join(v for v in (_s(ret.get("invoice_no")), _date_text(ret.get("invoice_date"))) if v),
                  _qty(ret.get("qty_returned")) if ret else "", _qty(ret.get("qty_rejected")) if _f(ret.get("qty_rejected")) else "",
                  ret.get("vehicle_no"), ret.get("remarks")]
        for col, value in enumerate(values, start=1):
            put(r, col, value, border=True)
        r += 1
    summary = dict(data.get("summary") or {})
    put(r, 1, f"Quantity dispatch {_qty(summary.get('dispatched'))} · Returned {_qty(summary.get('returned'))} · Rejected {_qty(summary.get('rejected'))} · Balance Quantity {_qty(summary.get('balance'))}", bold=True, span=7)
    r += 3
    put(r, 5, "Authorised Signatory", bold=True, span=3)
    for col, width in zip("ABCDEFG", (30, 16, 18, 12, 14, 16, 22)):
        ws.column_dimensions[col].width = width
    ws.print_area = f"A1:G{r}"
    out = BytesIO(); wb.save(out); return out.getvalue()
