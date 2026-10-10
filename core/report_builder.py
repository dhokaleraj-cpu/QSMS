"""QCMS R16 · Report Builder + Master List Excel exports.

Pure pandas logic (testable without Streamlit):
* ``build_report`` – choose columns, filters, group-by with aggregates, sort and row limit.
* ``excel_bytes`` – clean, printable Excel (title rows, bold header, filter, frozen header,
  column widths, no colour fills except the header band).
* Saved report definitions live in the tenant ``master_value_catalog`` table
  (field_key ``qcms.report_builder.saved``) — no SQL migration needed.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any, Callable, Mapping, Sequence

import pandas as pd

SAVED_KEY = "qcms.report_builder.saved"
HIDDEN_COLUMNS = {"tenant_id", "created_by", "updated_by", "search_vector", "password_hash"}
OPERATORS = {
    "contains": "contains",
    "equals": "equals",
    "not_equals": "is not",
    "starts_with": "starts with",
    "gt": "greater than",
    "gte": "greater or equal",
    "lt": "less than",
    "lte": "less or equal",
    "between": "between (dates / numbers)",
    "is_empty": "is empty",
    "not_empty": "is not empty",
}
AGGREGATES = {"count": "Count", "sum": "Sum", "mean": "Average", "min": "Minimum", "max": "Maximum", "nunique": "Distinct count"}

# Which QCMS module permission protects each master list (View permission required).
MASTER_MODULE = {
    "parties": "REFERENCE_MASTERS",
    "parts": "PART_MASTER",
    "material_grades": "MATERIAL_GRADE",
    "material_grade_elements": "MATERIAL_GRADE",
    "employees": "EMPLOYEE_MASTER",
}

LABEL_FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "part_id": ("parts", ("part_number", "part_name")),
    "customer_id": ("parties", ("party_code", "party_name")),
    "supplier_id": ("parties", ("party_code", "party_name")),
    "party_id": ("parties", ("party_code", "party_name")),
    "vendor_id": ("parties", ("party_code", "party_name")),
    "osp_vendor_id": ("parties", ("party_code", "party_name")),
    "steel_mill_id": ("parties", ("party_code", "party_name")),
    "process_id": ("processes", ("process_code", "process_name")),
    "material_grade_id": ("material_grades", ("grade_code",)),
    "grade_id": ("material_grades", ("grade_code",)),
    "inspection_stage_id": ("inspection_stages", ("stage_code", "stage_name")),
    "inspection_plan_id": ("inspection_plans", ("plan_number",)),
    "layout_plan_id": ("inspection_plans", ("plan_number",)),
    "prepared_by_employee_id": ("employees", ("employee_code", "first_name", "last_name")),
    "validated_by_employee_id": ("employees", ("employee_code", "first_name", "last_name")),
    "approved_by_employee_id": ("employees", ("employee_code", "first_name", "last_name")),
}


def pretty(column: str) -> str:
    text = str(column).replace("_id", "").replace("_", " ").strip()
    return " ".join(w.upper() if w in {"po", "osp", "rmtc", "fsi", "grn", "hsn", "id", "npd", "ppap", "uom"} else w.capitalize() for w in text.split())


def flatten(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value if v not in (None, ""))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, default=str)[:500]
    return value


def label_maps(load: Callable[[str], list[dict]], columns: Sequence[str]) -> dict[str, dict[str, str]]:
    """Return {fk_column: {id: label}} for the foreign-key columns present."""
    cache: dict[str, list[dict]] = {}
    out: dict[str, dict[str, str]] = {}
    for column in columns:
        spec = LABEL_FIELDS.get(column)
        if not spec:
            continue
        table, fields = spec
        if table not in cache:
            try:
                cache[table] = load(table)
            except Exception:
                cache[table] = []
        out[column] = {str(r.get("id")): " · ".join(str(r.get(f) or "").strip() for f in fields if str(r.get(f) or "").strip()) for r in cache[table] if r.get("id")}
    return out


def readable_frame(rows: Sequence[Mapping[str, Any]], load: Callable[[str], list[dict]] | None = None, *, keep_ids: bool = False) -> pd.DataFrame:
    """Rows → DataFrame with lists/dicts flattened, id columns replaced by labels and system columns hidden."""
    frame = pd.DataFrame([dict(r) for r in rows])
    if frame.empty:
        return frame
    frame = frame[[c for c in frame.columns if c not in HIDDEN_COLUMNS]]
    for column in frame.columns:
        if frame[column].map(lambda v: isinstance(v, (list, tuple, set, dict))).any():
            frame[column] = frame[column].map(flatten)
    if load is not None:
        for column, mapping in label_maps(load, list(frame.columns)).items():
            frame[column] = frame[column].map(lambda v, m=mapping: m.get(str(v), v) if v not in (None, "") else v)
    if not keep_ids and "id" in frame.columns:
        frame = frame.drop(columns=["id"])
    return frame


def _as_number(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def _as_date(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, errors="coerce", utc=True).dt.tz_localize(None)


def apply_filter(frame: pd.DataFrame, column: str, op: str, value: Any = None, value2: Any = None) -> pd.DataFrame:
    if column not in frame.columns:
        raise ValueError(f"Unknown column: {column}")
    series = frame[column]
    text = series.astype(str).where(series.notna(), "")
    op = str(op or "contains")
    if op == "is_empty":
        return frame[text.str.strip() == ""]
    if op == "not_empty":
        return frame[text.str.strip() != ""]
    if op == "contains":
        return frame[text.str.contains(str(value or ""), case=False, regex=False)]
    if op == "starts_with":
        return frame[text.str.lower().str.startswith(str(value or "").lower())]
    if op in {"equals", "not_equals"}:
        mask = text.str.strip().str.casefold() == str(value or "").strip().casefold()
        return frame[mask if op == "equals" else ~mask]
    if op in {"gt", "gte", "lt", "lte", "between"}:
        numeric = _as_number(series)
        use_dates = numeric.notna().sum() == 0 or isinstance(value, (date, datetime))
        left = _as_date(series) if use_dates else numeric
        conv = (lambda v: pd.Timestamp(v)) if use_dates else (lambda v: float(v))
        try:
            a = conv(value) if value not in (None, "") else None
            b = conv(value2) if value2 not in (None, "") else None
        except Exception as exc:
            raise ValueError(f"Filter value for {pretty(column)} is not a valid {'date' if use_dates else 'number'}.") from exc
        if op == "between":
            mask = pd.Series(True, index=frame.index)
            if a is not None:
                mask &= left >= a
            if b is not None:
                mask &= left <= b
            return frame[mask.fillna(False)]
        if a is None:
            return frame
        mask = {"gt": left > a, "gte": left >= a, "lt": left < a, "lte": left <= a}[op]
        return frame[mask.fillna(False)]
    raise ValueError(f"Unknown filter operator: {op}")


@dataclass
class ReportSpec:
    dataset: str
    columns: list[str] = field(default_factory=list)
    filters: list[dict] = field(default_factory=list)
    group_by: list[str] = field(default_factory=list)
    measures: list[dict] = field(default_factory=list)  # [{"column": "quantity", "agg": "sum"}]
    sort_by: str = ""
    descending: bool = False
    limit: int = 0
    title: str = ""
    chart: str = ""  # "", "bar", "line", "pie"

    def to_json(self) -> str:
        return json.dumps(self.__dict__, ensure_ascii=False, default=str)

    @classmethod
    def from_json(cls, text: str) -> "ReportSpec":
        data = json.loads(text)
        return cls(**{k: v for k, v in data.items() if k in cls.__dataclass_fields__})


def build_report(frame: pd.DataFrame, spec: ReportSpec) -> pd.DataFrame:
    out = frame.copy()
    for f in spec.filters:
        if f.get("column"):
            out = apply_filter(out, f["column"], f.get("op", "contains"), f.get("value"), f.get("value2"))
    if spec.group_by:
        missing = [c for c in spec.group_by if c not in out.columns]
        if missing:
            raise ValueError("Unknown group column: " + ", ".join(missing))
        keys = [out[c].astype(str).where(out[c].notna(), "(blank)") for c in spec.group_by]
        grouped = out.groupby(keys, dropna=False)
        result = grouped.size().rename("Records").to_frame()
        for m in spec.measures or []:
            col, agg = m.get("column"), m.get("agg", "sum")
            if not col or col not in out.columns or agg not in AGGREGATES:
                continue
            series = out[col] if agg in {"count", "nunique"} else _as_number(out[col])
            name = f"{AGGREGATES[agg]} of {pretty(col)}"
            result[name] = series.groupby(keys, dropna=False).agg(agg)
        result = result.reset_index()
        result.columns = [pretty(c) if c in spec.group_by else c for c in result.columns]
    else:
        cols = [c for c in (spec.columns or list(out.columns)) if c in out.columns]
        result = out[cols].copy()
        result.columns = [pretty(c) for c in cols]
    if spec.sort_by:
        key = spec.sort_by if spec.sort_by in result.columns else pretty(spec.sort_by)
        if key in result.columns:
            numeric = _as_number(result[key])
            if numeric.notna().sum() >= max(1, len(result) // 2):
                result = result.assign(_k=numeric).sort_values("_k", ascending=not spec.descending, na_position="last").drop(columns="_k")
            else:
                result = result.sort_values(key, ascending=not spec.descending, na_position="last", key=lambda s: s.astype(str).str.lower())
    if spec.limit and spec.limit > 0:
        result = result.head(int(spec.limit))
    return result.reset_index(drop=True)


def excel_bytes(sheets: Sequence[tuple[str, pd.DataFrame]], *, title: str = "QCMS Report", subtitle: str = "") -> bytes:
    """Clean printable workbook: title rows, bold grey header band, borders, filter, frozen header."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="7F7F7F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="E6E9E9")
    buffer = BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets or [("Report", pd.DataFrame())]:
            sheet = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
            base, n = sheet, 2
            while sheet in used:
                sheet = f"{base[:28]} {n}"; n += 1
            used.add(sheet)
            frame = frame.copy()
            for column in frame.columns:
                if isinstance(frame[column].dtype, pd.DatetimeTZDtype):
                    frame[column] = frame[column].dt.tz_localize(None)
                if frame[column].map(lambda v: isinstance(v, (list, dict, tuple, set))).any():
                    frame[column] = frame[column].map(flatten)
            frame.to_excel(writer, index=False, sheet_name=sheet, startrow=3)
            ws = writer.sheets[sheet]
            ws["A1"] = f"FOUR STAR INDUSTRIES PVT. LTD. · {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = f"{name} · {len(frame):,} record(s) · generated {datetime.now().strftime('%d-%m-%Y %H:%M')}" + (f" · {subtitle}" if subtitle else "")
            ws["A2"].font = Font(size=10, color="555555")
            ncols = max(1, len(frame.columns))
            for col in range(1, ncols + 1):
                cell = ws.cell(row=4, column=col)
                cell.font = Font(bold=True, size=11)
                cell.fill = head_fill
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=5, max_row=4 + len(frame), max_col=ncols):
                for cell in row:
                    cell.border = border
                    cell.font = Font(size=11)
                    cell.alignment = Alignment(vertical="center", wrap_text=False)
            ws.row_dimensions[4].height = 24
            for r in range(5, 5 + len(frame)):
                ws.row_dimensions[r].height = 18
            for idx, column in enumerate(frame.columns, start=1):
                values = [str(column)] + [str(v) for v in frame[column].head(300).tolist()]
                ws.column_dimensions[get_column_letter(idx)].width = min(48, max(10, max(len(v) for v in values) + 2))
            ws.freeze_panes = "A5"
            if len(frame.columns):
                ws.auto_filter.ref = f"A4:{get_column_letter(ncols)}{4 + len(frame)}"
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = "4:4"
    return buffer.getvalue()


# ----------------------------------------------------------------------------- saved reports
def saved_reports(repo: Any) -> list[dict]:
    try:
        rows = repo.select("master_value_catalog", eq={"field_key": SAVED_KEY, "status": "ACTIVE"}, order_by="last_used_at", desc=True, limit=500)
    except Exception:
        return []
    out = []
    for row in rows or []:
        try:
            payload = json.loads(row.get("value_text") or "{}")
            out.append({"id": row.get("id"), "name": payload.get("name") or "Report", "owner": payload.get("owner") or "", "spec": ReportSpec.from_json(json.dumps(payload.get("spec") or {})), "created_by": row.get("created_by")})
        except Exception:
            continue
    return out


def save_report(repo: Any, name: str, spec: ReportSpec, owner: str, existing_id: str | None = None) -> dict:
    name = str(name or "").strip()
    if not name:
        raise ValueError("Enter a report name.")
    now = datetime.now(timezone.utc).isoformat()
    payload = {"field_key": SAVED_KEY, "value_text": json.dumps({"name": name, "owner": owner, "spec": json.loads(spec.to_json())}, ensure_ascii=False),
               "normalized_value": name.casefold()[:200], "status": "ACTIVE", "last_used_at": now, "updated_at": now}
    if existing_id:
        return repo.update("master_value_catalog", existing_id, payload)
    return repo.insert("master_value_catalog", payload)


def delete_report(repo: Any, report_id: str) -> None:
    repo.update("master_value_catalog", report_id, {"status": "INACTIVE", "updated_at": datetime.now(timezone.utc).isoformat()})
